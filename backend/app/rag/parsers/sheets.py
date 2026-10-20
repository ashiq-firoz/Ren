"""Spreadsheet parsers (XLSX via openpyxl read-only mode, CSV/TSV via the csv module).

Each sheet becomes a section; rows become markdown tables split into row blocks of ~900 tokens (each block
repeats the header row and is titled "Rows a–b" with the spreadsheet's row numbers), so every leaf is a
self-describing table fragment. Very large sheets are truncated with an explicit note.
"""
from __future__ import annotations

import csv
import io
import sys
from pathlib import Path

from app.rag.parsers import Block, ParsedDocument, decode_bytes, format_cell, table_to_markdown
from app.rag.textutil import estimate_tokens

MAX_ROWS_PER_SHEET = 20000
MAX_COLS = 40
MAX_CELL_CHARS = 300
BLOCK_TOKENS = 900

try:  # allow long text fields in CSV exports
    csv.field_size_limit(min(sys.maxsize, 2**31 - 1))
except OverflowError:  # pragma: no cover
    pass


def _row_values(row) -> list[str]:
    vals = [format_cell(v) for v in list(row)[:MAX_COLS + 20]]
    while vals and not vals[-1].strip():
        vals.pop()
    return vals


def emit_table_section(blocks: list[Block], name: str | None, rows: list[tuple[int, list[str]]], *,
                       level: int = 1, truncated: int = 0) -> None:
    """Append heading + row-block tables for one sheet. `rows` = [(spreadsheet row number, values)]."""
    rows = [(n, r) for n, r in rows if any(c.strip() for c in r)]
    if name:
        blocks.append(Block("heading", name, level))
    if not rows:
        return
    (_, header), data = rows[0], rows[1:]
    if not data:
        blocks.append(Block("table", table_to_markdown([header], max_cols=MAX_COLS, max_cell=MAX_CELL_CHARS)))
        return
    width = max(len(r) for _, r in rows)
    header = (header + [""] * width)[:width]
    head_tok = estimate_tokens(" | ".join(header))
    groups: list[list[tuple[int, list[str]]]] = []
    cur: list[tuple[int, list[str]]] = []
    cur_tok = head_tok
    for n, r in data:
        t = estimate_tokens(" | ".join(r)) + 2
        if cur and cur_tok + t > BLOCK_TOKENS:
            groups.append(cur)
            cur, cur_tok = [], head_tok
        cur.append((n, r))
        cur_tok += t
    if cur:
        groups.append(cur)
    for grp in groups:
        md = table_to_markdown([header] + [r for _, r in grp], max_rows=len(grp) + 1, max_cols=MAX_COLS,
                               max_cell=MAX_CELL_CHARS)
        if len(groups) > 1:
            blocks.append(Block("heading", f"Rows {grp[0][0]}–{grp[-1][0]}", level + 1 if name else 1))
        blocks.append(Block("table", md))
    if truncated:
        blocks.append(Block("text", f"(Truncated: {truncated} further rows were not indexed.)"))


def parse_xlsx(path: Path) -> ParsedDocument:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    blocks: list[Block] = []
    try:
        title = (wb.properties.title or "").strip() if wb.properties else ""
    except Exception:  # noqa: BLE001
        title = ""
    try:
        for ws in wb.worksheets:
            rows: list[tuple[int, list[str]]] = []
            extra = 0
            for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                vals = _row_values(row)
                if not any(v.strip() for v in vals):
                    continue
                if len(rows) >= MAX_ROWS_PER_SHEET:
                    extra += 1
                    continue
                rows.append((idx, vals))
            if rows:
                emit_table_section(blocks, ws.title, rows, level=1, truncated=extra)
    finally:
        wb.close()
    return ParsedDocument(title=title, blocks=blocks, format="xlsx")


def parse_csv(path: Path, *, delimiter: str | None = None) -> ParsedDocument:
    text = decode_bytes(path.read_bytes())
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    rows: list[tuple[int, list[str]]] = []
    extra = 0
    for idx, row in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter), start=1):
        vals = _row_values(row)
        if not any(v.strip() for v in vals):
            continue
        if len(rows) >= MAX_ROWS_PER_SHEET:
            extra += 1
            continue
        rows.append((idx, vals))
    blocks: list[Block] = []
    emit_table_section(blocks, None, rows, level=1, truncated=extra)
    return ParsedDocument(title="", blocks=blocks, format="csv")
