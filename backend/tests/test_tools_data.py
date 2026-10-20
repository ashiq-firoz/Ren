"""Data tools: sheets (xlsx/csv) and analytics (profile, DuckDB query, reconcile, charts, forecasting…)."""
from __future__ import annotations

import pytest

from app.tools.base import ToolError, ToolResult


def _c(res):
    return res.content if isinstance(res, ToolResult) else res


def _abs(ctx, vpath):
    return ctx.workspace.resolve(vpath, must_exist=True)


SALES = [
    {"date": "2026-01-05", "region": "North", "product": "Widget", "units": 10, "amount": 1200.50},
    {"date": "2026-01-18", "region": "South", "product": "Widget", "units": 4, "amount": 480.00},
    {"date": "2026-02-02", "region": "North", "product": "Gadget", "units": 7, "amount": 2100.00},
    {"date": "2026-02-20", "region": "East", "product": "Widget", "units": 12, "amount": 1440.00},
    {"date": "2026-03-03", "region": "South", "product": "Gadget", "units": 3, "amount": 900.00},
    {"date": "2026-03-15", "region": "North", "product": "Widget", "units": 9, "amount": 1080.00},
]


def _write_csv(ctx, vpath: str, text: str) -> str:
    p = ctx.workspace.resolve(vpath, for_write=True)
    p.write_text(text, encoding="utf-8")
    return vpath


# ----------------------------------------------------------------------------------------- sheets
async def test_sheet_create_write_read_calculate(tool_ctx):
    from openpyxl import load_workbook

    from app.tools.builtin import sheets as S

    res = await S.create(tool_ctx, path="output/tests/sales.xlsx", rows=SALES, sheet_name="Sales", overwrite=True)
    info = _c(res)
    assert info["path"] == "output/tests/sales.xlsx" and res.artifacts
    wb = load_workbook(str(_abs(tool_ctx, info["path"])))
    ws = wb["Sales"]
    assert ws["A1"].value == "date" and ws["A1"].font.b and ws.freeze_panes == "A2"
    assert ws["E2"].number_format == "#,##0.00"
    assert ws["A2"].value.year == 2026  # ISO strings became real dates

    read = await S.read(tool_ctx, path=info["path"])
    assert read["header"] == ["date", "region", "product", "units", "amount"]
    assert read["total_rows"] == 6 and read["rows"][0][1] == "North" and read["rows"][0][0].startswith("2026-01-05")
    part = await S.read(tool_ctx, path=info["path"], cell_range="Sales!B1:C3")
    assert part["header"] == ["region", "product"] and len(part["rows"]) == 2

    await S.write(tool_ctx, path=info["path"], cells={"G1": "Total", "G2": "=SUM(E2:E7)"}, sheet="Sales")
    await S.write(tool_ctx, path=info["path"], values=[["Note", "ok"]], start_cell="Notes!A1")
    wb = load_workbook(str(_abs(tool_ctx, info["path"])))
    assert wb["Sales"]["G2"].value == "=SUM(E2:E7)" and "Notes" in wb.sheetnames
    after = await S.read(tool_ctx, path=info["path"], cell_range="G1:G2", header="yes")
    assert after["rows"][0][0] == "=SUM(E2:E7)" and "never been computed" in after["note"]

    calc = await S.calculate(tool_ctx, path=info["path"], expression="SUM(E2:E7)", sheet="Sales")
    assert calc["value"] == pytest.approx(7200.5)
    calc = await S.calculate(tool_ctx, path=info["path"], expression="=AVERAGE(Sales!D:D)")
    assert calc["value"] == pytest.approx(7.5)
    grouped = await S.calculate(tool_ctx, path=info["path"], column="amount", operation="sum", group_by="region",
                                sheet="Sales")
    assert grouped["rows"][0] == ["North", pytest.approx(4380.5)]
    with pytest.raises(ToolError):
        await S.calculate(tool_ctx, path=info["path"], expression="VLOOKUP(A1, B:C, 2)")

    f = _c(await S.formula(tool_ctx, path=info["path"], cell="H2", formula="D2*2", sheet="Sales", fill_down_to=7))
    assert f["cells"] == 6
    wb = load_workbook(str(_abs(tool_ctx, info["path"])))
    assert wb["Sales"]["H7"].value == "=D7*2"

    await S.append(tool_ctx, path=info["path"], sheet="Sales",
                   rows=[{"date": "2026-04-01", "region": "West", "product": "Gizmo", "units": 1, "amount": 99,
                          "channel": "web"}])
    wb = load_workbook(str(_abs(tool_ctx, info["path"])))
    ws = wb["Sales"]
    headers = [c.value for c in ws[1]]
    assert "channel" in headers and ws.cell(row=8, column=headers.index("channel") + 1).value == "web"

    found = await S.find(tool_ctx, path=info["path"], query="gadget")
    assert found["total"] == 2 and {m["cell"] for m in found["matches"]} == {"C4", "C6"}


async def test_sheet_filter_sort_pivot_chart_format_export(tool_ctx):
    from openpyxl import load_workbook

    from app.tools.builtin import sheets as S

    path = _c(await S.create(tool_ctx, path="output/tests/sales2.xlsx", rows=SALES, overwrite=True))["path"]
    flt = _c(await S.filter_(tool_ctx, path=path, column="amount", operator=">=", value=1000))
    assert flt["matched"] == 4 and flt["new_sheet"] == "Sheet1 filtered"
    flt2 = _c(await S.filter_(tool_ctx, path=path, column="region", operator="in", value=["north", "east"],
                              new_sheet="NE"))
    assert flt2["matched"] == 4
    srt = _c(await S.sort(tool_ctx, path=path, by=["amount"], descending=True))
    assert srt["preview"][0][4] == 2100.0
    piv = _c(await S.pivot(tool_ctx, path=path, index=["region"], columns=["product"], values=["amount"]))
    assert piv["new_sheet"] == "Pivot" and "Total" in piv["columns"]
    north = next(r for r in piv["preview"] if r[0] == "North")
    assert north[piv["columns"].index("Widget")] == pytest.approx(2280.5)

    ch = _c(await S.chart(tool_ctx, path=path, chart_type="column", data_columns=["amount"], category_column="region",
                          title="Amount by deal", sheet="Sheet1"))
    assert ch["series"] == 1 and ch["points"] == 6
    await S.chart(tool_ctx, path=path, chart_type="pie", data_columns=["units"], category_column="product")
    await S.chart(tool_ctx, path=path, chart_type="line", sheet="Sheet1")

    fmt = _c(await S.format_(tool_ctx, path=path, number_formats={"amount": "$#,##0.00"},
                             column_widths={"A": 14}, banded_rows=True))
    assert "format amount" in fmt["applied"]
    wb = load_workbook(str(_abs(tool_ctx, path)))
    ws = wb["Sheet1"]
    assert len(ws._charts) == 3
    assert ws["E3"].number_format == "$#,##0.00" and ws.column_dimensions["A"].width == 14
    assert {"Sheet1 filtered", "NE", "Sheet1 sorted", "Pivot"} <= set(wb.sheetnames)

    await S.create_sheet(tool_ctx, path=path, name="Targets", rows=[["region", "target"], ["North", 5000]])
    await S.rename_sheet(tool_ctx, path=path, old_name="Targets", new_name="Goals")
    out = await S.delete_sheet(tool_ctx, path=path, name="NE")
    assert "Goals" in out["sheets"] and "NE" not in out["sheets"]
    with pytest.raises(ToolError):
        await S.rename_sheet(tool_ctx, path=path, old_name="Goals", new_name="Pivot")

    exp = _c(await S.export(tool_ctx, path=path, sheet="Pivot", output="output/tests/pivot.csv"))
    text = _abs(tool_ctx, exp["path"]).read_text(encoding="utf-8-sig")
    assert text.splitlines()[0].startswith("region,") and "Total" in text
    js = _c(await S.export(tool_ctx, path=path, sheet="Sheet1", format="json"))
    assert js["rows"] == 6


async def test_sheet_csv_support(tool_ctx):
    from app.tools.builtin import sheets as S

    path = _c(await S.create(tool_ctx, path="output/tests/people.csv", rows=[["name", "age"], ["Ana", 31], ["Bo", 25]],
                             overwrite=True))["path"]
    assert path.endswith(".csv")
    await S.append(tool_ctx, path=path, rows=[["Cy", 40]])
    read = await S.read(tool_ctx, path=path)
    assert read["header"] == ["name", "age"] and read["rows"][-1] == ["Cy", 40]
    flt = _c(await S.filter_(tool_ctx, path=path, column="age", operator="gt", value=30))
    assert flt["matched"] == 2 and flt["output"].endswith(".csv")
    with pytest.raises(ToolError, match="Excel workbook"):
        await S.chart(tool_ctx, path=path, chart_type="bar")
    await S.write(tool_ctx, path=path, cells={"B2": 32})
    assert (await S.read(tool_ctx, path=path))["rows"][0] == ["Ana", 32]


# ----------------------------------------------------------------------------------------- analytics
SALES_CSV = "date,region,product,units,amount\n" + "\n".join(
    f"{r['date']},{r['region']},{r['product']},{r['units']},{r['amount']}" for r in SALES) + "\n"


async def test_analytics_profile_and_duckdb_query(tool_ctx):
    from app.tools.builtin import analytics as A

    path = _write_csv(tool_ctx, "uploads/tests/sales_2026.csv", SALES_CSV)
    prof = await A.profile(tool_ctx, path=path)
    assert prof["rows"] == 6 and prof["columns"] == 5
    cols = {c["name"]: c for c in prof["columns_info"]}
    assert cols["amount"]["sum"] == pytest.approx(7200.5) and cols["amount"]["max"] == 2100
    assert cols["region"]["top"][0] == ["North", 3]
    assert cols["date"].get("looks_datetime") or "datetime" in cols["date"]["dtype"]

    q = await A.query(tool_ctx, sql="SELECT region, SUM(amount) AS total FROM sales_2026 GROUP BY 1 ORDER BY 2 DESC",
                      files=[path])
    assert q["columns"] == ["region", "total"] and q["rows"][0] == ["North", pytest.approx(4380.5)]
    assert q["tables"] == {"sales_2026": path}

    targets = _write_csv(tool_ctx, "uploads/tests/targets.csv", "region,target\nNorth,5000\nSouth,1000\nEast,2000\n")
    j = await A.query(tool_ctx, files={"s": path, "t": targets},
                      sql="SELECT t.region, t.target, SUM(s.amount) AS actual FROM s JOIN t USING (region) "
                          "GROUP BY ALL ORDER BY 1")
    assert [r[0] for r in j["rows"]] == ["East", "North", "South"]
    desc = await A.query(tool_ctx, sql="DESCRIBE s", files={"s": path})
    assert desc["returned"] == 5

    saved = _c(await A.query(tool_ctx, sql="SELECT * FROM s WHERE amount > 1000", files={"s": path},
                             save_as="output/tests/big.csv"))
    assert saved["row_count"] == 4 and saved["path"] == "output/tests/big.csv"

    # sandbox: no filesystem access, no writes, one statement only
    secret = tool_ctx.workspace.resolve("uploads/tests/sales_2026.csv")
    with pytest.raises(ToolError):
        await A.query(tool_ctx, sql=f"SELECT * FROM read_csv_auto('{secret.as_posix()}')", files={"s": path})
    with pytest.raises(ToolError, match="read-only"):
        await A.query(tool_ctx, sql="DROP TABLE s", files={"s": path})
    with pytest.raises(ToolError, match="one SQL statement"):
        await A.query(tool_ctx, sql="SELECT 1; SELECT 2", files={"s": path})
    with pytest.raises(ToolError):
        await A.query(tool_ctx, sql="COPY s TO 'leak.csv'", files={"s": path})
    with pytest.raises(ToolError):
        await A.query(tool_ctx, sql="SELECT * FROM missing_table", files={"s": path})


async def test_analytics_reconcile(tool_ctx):
    from openpyxl import load_workbook

    from app.tools.builtin import analytics as A

    bank = _write_csv(tool_ctx, "uploads/tests/bank.csv",
                      "ref,amount,memo\nINV-001,\"$1,000.00\",a\ninv-002,250.00,b\nINV-003,99.99,c\n"
                      "INV-004,500,d\nINV-004,500,d2\nINV-009,10,e\n")
    ledger = _write_csv(tool_ctx, "uploads/tests/ledger.csv",
                        "invoice,total\nINV-001,1000\nINV-002,240\nINV-003,100.00\nINV-004,1000\nINV-010,77\n")
    res = await A.reconcile(tool_ctx, left=bank, right=ledger, left_key="ref", right_key="invoice",
                            left_amount="amount", right_amount="total", tolerance=0.05,
                            left_label="Bank", right_label="Ledger", output="output/tests/recon.xlsx")
    info = _c(res)
    s = info["summary"]
    assert s["matched"] == 3  # 001, 003 (within tolerance), 004 (duplicates summed: 500+500)
    assert s["amount_differences"] == 1  # 002: 250 vs 240
    assert s["only_in_Bank"] == 1 and s["only_in_Ledger"] == 1
    assert s["duplicate_key_rows_Bank"] == 2
    assert s["total_Bank"] == pytest.approx(2359.99) and s["total_Ledger"] == pytest.approx(2417.0)
    assert info["top_differences"]["rows"][0][0] == "inv-002"
    assert info["top_differences"]["rows"][0][3] == pytest.approx(10.0)
    wb = load_workbook(str(_abs(tool_ctx, info["path"])))
    assert wb.sheetnames == ["Summary", "Differences", "Only in Bank", "Only in Ledger", "Matched"]
    assert wb["Only in Bank"]["A2"].value == "INV-009"
    assert res.artifacts

    # presence-only reconciliation
    pres = _c(await A.reconcile(tool_ctx, left=bank, right=ledger, left_key="ref", right_key="invoice"))
    assert pres["summary"]["matched"] == 4


async def test_analytics_visualize_png_artifact(tool_ctx):
    from PIL import Image

    from app.tools.builtin import analytics as A

    path = _write_csv(tool_ctx, "uploads/tests/sales_viz.csv", SALES_CSV)
    res = await A.visualize(tool_ctx, chart_type="bar", path=path, x="region", y="amount", title="Revenue by region",
                            subtitle="Q1 2026")
    info = _c(res)
    assert info["path"].startswith("output/charts/") and info["path"].endswith(".png")
    assert res.artifacts and res.artifacts[0]["path"] == info["path"]
    assert info["points"] == 3  # aggregated by region
    with Image.open(_abs(tool_ctx, info["path"])) as im:
        assert im.format == "PNG" and im.size[0] >= 1200
        assert im.getpixel((5, 5))[:3] == (0xF8, 0xF6, 0xF1)  # f69 ivory background

    data = [{"month": f"2026-{m:02d}-01", "revenue": 100 + 10 * m, "cost": 80 + 5 * m} for m in range(1, 13)]
    for ct in ("line", "area", "horizontal_bar", "pie", "donut", "scatter", "histogram"):
        kw = {"x": "month", "y": ["revenue", "cost"]} if ct in ("line", "area") else \
             {"x": "month", "y": "revenue"} if ct != "scatter" else {"x": "cost", "y": "revenue"}
        if ct == "histogram":
            kw = {"y": "revenue"}
        out = _c(await A.visualize(tool_ctx, chart_type=ct, data=data, title=f"{ct} chart", **kw))
        assert _abs(tool_ctx, out["path"]).stat().st_size > 5000, ct
    stacked = _c(await A.visualize(tool_ctx, chart_type="bar", path=path, x="region", y="amount", series="product",
                                   stacked=True))
    assert stacked["series"] == ["Gadget", "Widget"]
    with pytest.raises(ToolError):
        await A.visualize(tool_ctx, chart_type="bar")


async def test_analytics_aggregate_join_stats_forecast(tool_ctx):
    from app.tools.builtin import analytics as A

    path = _write_csv(tool_ctx, "uploads/tests/sales_more.csv", SALES_CSV)
    agg = await A.aggregate(tool_ctx, path=path, group_by=["region"], metrics=["sum:amount", "count", "mean:units"])
    assert agg["columns"] == ["region", "sum_amount", "count", "mean_units"]
    assert agg["rows"][0][:3] == ["North", pytest.approx(4380.5), 3]
    total = await A.aggregate(tool_ctx, path=path, metrics=["amount"])
    assert total["rows"][0][0] == pytest.approx(7200.5)

    regions = _write_csv(tool_ctx, "uploads/tests/regions.csv", "region,manager\nNorth,Ana\nSouth,Bo\n")
    j = _c(await A.join(tool_ctx, left=path, right=regions, on=["region"], how="left"))
    assert j["rows"] == 6 and j["matches"] == {"both": 5, "left_only": 1, "right_only": 0}
    assert "manager" in j["columns"]

    corr = await A.correlate(tool_ctx, path=path)
    assert corr["top_pairs"][0]["a"] in ("units", "amount")
    spread = _write_csv(tool_ctx, "uploads/tests/vals_spread.csv",
                        "v\n" + "\n".join(["8", "9", "10", "11", "12", "10", "9", "11", "10", "12", "500"]) + "\n")
    outl = await A.outliers(tool_ctx, path=spread, column="v")
    assert outl["method"] == "iqr" and outl["outliers"] == 1
    assert outl["rows"][0][1] == 500 and outl["rows"][0][0] == 12  # spreadsheet row of the outlier
    flat = _write_csv(tool_ctx, "uploads/tests/vals_flat.csv",
                      "v\n" + "\n".join(["10"] * 8 + ["11", "9", "500"]) + "\n")
    outl = await A.outliers(tool_ctx, path=flat, column="v")  # IQR = 0 → falls back to z-scores
    assert outl["method"] == "zscore" and outl["outliers"] == 1 and "z-score" in outl["note"]

    lines = ["month,revenue"]
    for i in range(36):  # trend + yearly seasonality
        y, m = 2023 + i // 12, i % 12 + 1
        lines.append(f"{y}-{m:02d}-01,{1000 + 20 * i + (150 if m in (11, 12) else 0)}")
    ts = _write_csv(tool_ctx, "uploads/tests/monthly.csv", "\n".join(lines) + "\n")
    fc = _c(await A.forecast(tool_ctx, path=ts, value_column="revenue", date_column="month", periods=6, chart=True))
    assert fc["frequency"] == "MS" and len(fc["forecast"]) == 6
    assert fc["forecast"][0]["period"] == "2026-01-01" and fc["method"] in ("linear", "holt", "holt_winters")
    assert fc["path"].endswith(".png")
    hw = _c(await A.forecast(tool_ctx, path=ts, value_column="revenue", date_column="month", periods=12,
                             method="holt_winters"))
    nov_dec = [r["forecast"] for r in hw["forecast"] if r["period"][5:7] in ("11", "12")]
    others = [r["forecast"] for r in hw["forecast"] if r["period"][5:7] in ("09", "10")]
    assert min(nov_dec) > max(others)  # seasonality captured
    lin = _c(await A.forecast(tool_ctx, path=ts, value_column="revenue", periods=3, method="linear"))
    assert lin["forecast"][0]["period"] == 37
    with pytest.raises(ToolError):
        await A.forecast(tool_ctx, path=regions, value_column="manager", periods=3)
