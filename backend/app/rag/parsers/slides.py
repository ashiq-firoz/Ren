"""PPTX parser: one section per slide, titled by the slide title (fallback "Slide n"); bullet text keeps its
indentation, tables become markdown, speaker notes are appended. Page numbers = slide numbers."""
from __future__ import annotations

from pathlib import Path

from app.rag.parsers import Block, ParsedDocument, table_to_markdown
from app.rag.textutil import clean_inline


def _shape_texts(shape, out: list[str], title_shape_id: int | None) -> None:
    try:
        if getattr(shape, "shape_id", None) == title_shape_id:
            return
        if shape.shape_type is not None and int(shape.shape_type) == 6:  # GROUP
            for s in sorted(shape.shapes, key=lambda s: ((s.top or 0), (s.left or 0))):
                _shape_texts(s, out, title_shape_id)
            return
    except Exception:  # noqa: BLE001
        pass
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = [[clean_inline(c.text) for c in row.cells] for row in shape.table.rows]
        md = table_to_markdown(rows)
        if md:
            out.append(md)
        return
    if getattr(shape, "has_chart", False) and shape.has_chart:
        try:
            chart = shape.chart
            if chart.has_title and chart.chart_title.has_text_frame:
                t = clean_inline(chart.chart_title.text_frame.text)
                if t:
                    out.append(f"Chart: {t}")
        except Exception:  # noqa: BLE001
            pass
        return
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        lines = []
        for para in shape.text_frame.paragraphs:
            t = clean_inline(getattr(para, "text", "") or "".join(r.text for r in para.runs))
            if t:
                lines.append("  " * int(para.level or 0) + ("- " if len(shape.text_frame.paragraphs) > 1 else "") + t)
        if lines:
            out.append("\n".join(lines))


def parse_pptx(path: Path) -> ParsedDocument:
    from pptx import Presentation

    prs = Presentation(str(path))
    title = (prs.core_properties.title or "").strip()
    blocks: list[Block] = []
    n = 0
    for n, slide in enumerate(prs.slides, start=1):
        stitle = ""
        title_id = None
        try:
            ts = slide.shapes.title
            if ts is not None and ts.has_text_frame:
                stitle = clean_inline(ts.text_frame.text)
                title_id = ts.shape_id
        except Exception:  # noqa: BLE001
            pass
        texts: list[str] = []
        shapes = sorted(slide.shapes, key=lambda s: ((s.top or 0), (s.left or 0)))
        for shape in shapes:
            _shape_texts(shape, texts, title_id)
        notes = ""
        try:
            if slide.has_notes_slide:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
        except Exception:  # noqa: BLE001
            notes = ""
        blocks.append(Block("heading", stitle or f"Slide {n}", 1, n))
        body = "\n\n".join(texts)
        if notes:
            body = (body + "\n\n" if body else "") + f"Speaker notes: {notes}"
        if body.strip():
            blocks.append(Block("text", body, 0, n))
    if not title and blocks:
        title = blocks[0].text if blocks[0].text != "Slide 1" else ""
    return ParsedDocument(title=title, blocks=blocks, page_count=n, format="pptx")
