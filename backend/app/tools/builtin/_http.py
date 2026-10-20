"""Shared plumbing for the SaaS integration tool groups (email, calendar, drive, chat, github, pm, kb, crm,
cloud, notify, secrets).

* `api()` / `RestClient` — authenticated JSON calls with retries (429 + idempotent 5xx, honouring Retry-After),
  one transparent re-auth on 401, and friendly `ToolError`s ("credentials invalid or expired — reconnect X").
* `pick_connector()` — choose which connected account serves a call (priority order, or an `account` argument).
* Workspace helpers — save downloads under `downloads/` without clobbering, read workspace files for uploads.
* `extract_text()` — text from pdf / docx / xlsx / pptx / html / eml / plain-text payloads.
* `preview()` — markdown approval previews that show outbound content verbatim.
* `llm_complete()` / `llm_structured()` — model calls for summaries and extraction inside tools.

Not a tool module (leading underscore → not auto-imported); the group modules import it.
"""
from __future__ import annotations

import asyncio
import io
import json as jsonlib
import logging
import mimetypes
import os
import re
import time
from collections.abc import Awaitable, Callable, Iterable, Sequence
from datetime import datetime, timezone
from email.utils import parsedate_to_datetime
from pathlib import Path
from typing import Any

import httpx

from app.tools import net
from app.tools.base import ToolContext, ToolError

log = logging.getLogger("ren.tools.integrations")

IDEMPOTENT = frozenset({"GET", "HEAD", "OPTIONS", "PUT", "DELETE"})
MAX_RETRY_WAIT = 30.0  # never sleep longer than this inside a tool call
DEFAULT_MAX_DOWNLOAD = 100 * 1024 * 1024


def client() -> httpx.AsyncClient:
    """HTTP client used by every integration adapter (tests monkeypatch `app.tools.net.http`)."""
    return net.http()


async def _sleep(seconds: float) -> None:  # indirection so tests can skip back-off waits
    await asyncio.sleep(seconds)


class ApiError(ToolError):
    """A failed API call. Subclass of ToolError so unhandled ones surface to the agent as friendly text."""

    def __init__(self, message: str, status: int | None = None, data: Any = None):
        super().__init__(message)
        self.status = status
        self.data = data


# --------------------------------------------------------------------------- low-level request
def clean_params(params: Any) -> Any:
    """Drop None values, render booleans as true/false; lists stay lists (repeated query keys)."""
    if not isinstance(params, dict):
        return params
    out: dict[str, Any] = {}
    for k, v in params.items():
        if v is None:
            continue
        if isinstance(v, bool):
            v = "true" if v else "false"
        out[k] = v
    return out


def parse_body(resp: httpx.Response) -> Any:
    if resp.status_code == 204 or not resp.content:
        return None
    ctype = resp.headers.get("content-type", "")
    if "json" in ctype:
        try:
            return resp.json()
        except ValueError:
            return resp.text
    text = resp.text
    if text.lstrip()[:1] in ("{", "["):
        try:
            return jsonlib.loads(text)
        except ValueError:
            pass
    return text


def retry_after(resp: httpx.Response) -> float | None:
    v = resp.headers.get("retry-after")
    if v:
        try:
            return max(0.0, float(v))
        except ValueError:
            try:
                return max(0.0, (parsedate_to_datetime(v) - datetime.now(timezone.utc)).total_seconds())
            except Exception:  # noqa: BLE001
                return None
    if resp.headers.get("x-ratelimit-remaining") == "0" and resp.headers.get("x-ratelimit-reset"):
        try:
            return max(0.0, float(resp.headers["x-ratelimit-reset"]) - time.time())
        except ValueError:
            return None
    return None


def _extract_error(data: Any) -> str:
    if isinstance(data, list):
        return "; ".join(p for p in (_extract_error(x) for x in data[:3]) if p)
    if isinstance(data, str):
        return data
    if not isinstance(data, dict):
        return ""
    if isinstance(data.get("error_summary"), str) and data["error_summary"]:  # Dropbox
        return data["error_summary"].rstrip("/.")
    err = data.get("error")
    if isinstance(err, dict):  # Google / Microsoft Graph style
        msg = err.get("message") or ""
        code = err.get("code") or err.get(".tag") or ""
        if not msg:
            msg = str(code)
        elif isinstance(code, str) and code and code not in msg and not code.isdigit():
            msg = f"{msg} [{code}]"
        return str(msg)
    pieces: list[str] = []
    for key in ("message", "error_description", "error_summary", "errorMessage", "detail", "title", "msg"):
        val = data.get(key)
        if isinstance(val, str) and val:
            pieces.append(val)
            break
    if isinstance(err, str) and err and not pieces:
        pieces.append(err)
    if isinstance(data.get("errorMessages"), list):  # Jira
        pieces += [str(x) for x in data["errorMessages"][:3]]
    errs = data.get("errors")
    if isinstance(errs, dict):  # Jira field errors
        pieces += [f"{k}: {v}" for k, v in list(errs.items())[:5]]
    elif isinstance(errs, list):  # GitHub / Asana / GraphQL
        for e in errs[:3]:
            if isinstance(e, dict):
                if e.get("message"):
                    pieces.append(str(e["message"]))
                elif e.get("field"):
                    pieces.append(f"{e.get('resource', '')}.{e['field']} {e.get('code', '')}".strip(". "))
                elif e.get("code"):
                    pieces.append(str(e["code"]))
            else:
                pieces.append(str(e))
    return "; ".join(dict.fromkeys(p for p in pieces if p))


def error_message(resp: httpx.Response) -> str:
    """Concise error text from a failed response (never includes request headers / credentials)."""
    try:
        data = resp.json()
    except Exception:  # noqa: BLE001
        data = None
    msg = _extract_error(data) if data is not None else ""
    if not msg:
        text = (resp.text or "").strip()
        if text.startswith("<"):
            m = re.search(r"<Message>(.*?)</Message>", text, re.S) or re.search(r"<title>(.*?)</title>", text, re.S | re.I)
            msg = m.group(1) if m else ""
        else:
            msg = text
    return re.sub(r"\s+", " ", msg).strip()[:400]


def _rate_limited(resp: httpx.Response) -> bool:
    if resp.headers.get("x-ratelimit-remaining") == "0":
        return True
    return "rate limit" in (resp.text or "")[:500].lower()


def http_error(service: str, reconnect: str, resp: httpx.Response) -> ApiError:
    status = resp.status_code
    detail = error_message(resp)
    tail = f": {detail}" if detail else ""
    if status == 401:
        msg = (f"{service} credentials are invalid or expired — reconnect {reconnect} in Connectors"
               + (f" ({detail})." if detail else "."))
    elif status == 403 and _rate_limited(resp):
        msg = f"{service} rate limit reached — wait a few minutes and try again."
    elif status == 403:
        msg = (f"{service} denied access (403){tail}. The connected account or app may lack the required "
               "permission or scope.")
    elif status == 404:
        msg = f"{service}: not found (404){tail}. Check the id/name — or the connected account may not have access."
    elif status == 429:
        msg = f"{service} rate limit reached{tail} — wait a minute and try again."
    elif status >= 500:
        msg = f"{service} is unavailable or failed (HTTP {status}){tail}. Try again later."
    else:
        msg = f"{service} rejected the request (HTTP {status}){tail}."
    try:
        data = resp.json()
    except Exception:  # noqa: BLE001
        data = None
    return ApiError(msg, status=status, data=data)


async def api(
    method: str,
    url: str,
    *,
    service: str,
    reconnect: str | None = None,
    headers: dict[str, str] | None = None,
    params: Any = None,
    json: Any = None,
    data: Any = None,
    content: bytes | str | None = None,
    files: Any = None,
    auth: Any = None,
    timeout: float | None = None,
    raw: bool = False,
    allow_status: Sequence[int] = (),
    retries: int = 3,
) -> Any:
    """One HTTP call with retries and friendly errors. Returns parsed JSON/text (or the Response if raw)."""
    method = method.upper()
    kwargs: dict[str, Any] = {}
    if headers:
        kwargs["headers"] = headers
    if params:
        kwargs["params"] = clean_params(params)
    if json is not None:
        kwargs["json"] = json
    if data is not None:
        kwargs["data"] = data
    if content is not None:
        kwargs["content"] = content
    if files is not None:
        kwargs["files"] = files
    if auth is not None:
        kwargs["auth"] = auth
    if timeout is not None:
        kwargs["timeout"] = timeout
    delay = 1.0
    for attempt in range(retries + 1):
        last = attempt >= retries
        try:
            resp = await client().request(method, url, **kwargs)
        except (httpx.UnsupportedProtocol, httpx.InvalidURL) as e:
            raise ApiError(f"The {service} URL looks invalid — check the connector settings ({e}).") from e
        except (httpx.ConnectError, httpx.ConnectTimeout) as e:  # never reached the server → safe to retry
            if not last:
                await _sleep(delay)
                delay = min(delay * 2, 8)
                continue
            raise ApiError(f"Could not connect to {service} ({type(e).__name__}). Check the connector's URL "
                           "and network access.") from e
        except httpx.TimeoutException as e:
            if method in IDEMPOTENT and not last:
                await _sleep(delay)
                delay = min(delay * 2, 8)
                continue
            raise ApiError(f"{service} did not respond in time. Try again shortly.") from e
        except httpx.TransportError as e:
            if method in IDEMPOTENT and not last:
                await _sleep(delay)
                delay = min(delay * 2, 8)
                continue
            raise ApiError(f"Network error while calling {service} ({type(e).__name__}).") from e
        status = resp.status_code
        if status < 400 or status in allow_status:
            return resp if raw else parse_body(resp)
        retryable = status == 429 or (status in (500, 502, 503, 504) and method in IDEMPOTENT)
        if retryable and not last:
            wait = retry_after(resp)
            wait = delay if wait is None else wait
            if wait <= MAX_RETRY_WAIT:
                await _sleep(wait)
                delay = min(delay * 2, 8)
                continue
        raise http_error(service, reconnect or service, resp)
    raise ApiError(f"{service} request failed.")  # pragma: no cover


class RestClient:
    """Authenticated JSON client for one service: base URL + auth headers + friendly errors + one re-auth on 401."""

    def __init__(
        self,
        service: str,
        base_url: str = "",
        *,
        auth: Callable[[], Awaitable[dict[str, str]]] | None = None,
        headers: dict[str, str] | None = None,
        basic: tuple[str, str] | None = None,
        reconnect: str | None = None,
        on_unauthorized: Callable[[], None] | None = None,
    ):
        self.service = service
        self.base_url = base_url.rstrip("/")
        self.auth = auth
        self.headers = dict(headers or {})
        self.basic = basic
        self.reconnect = reconnect or service
        self.on_unauthorized = on_unauthorized

    def url(self, path: str) -> str:
        if path.startswith(("http://", "https://")):
            return path
        return f"{self.base_url}{path}"

    async def _headers(self, extra: dict[str, str] | None) -> dict[str, str]:
        hdrs = dict(self.headers)
        if self.auth is not None:
            hdrs.update(await self.auth())
        if extra:
            hdrs.update(extra)
        return hdrs

    async def request(self, method: str, path: str, *, headers: dict[str, str] | None = None, **kw: Any) -> Any:
        if self.basic is not None and "auth" not in kw:
            kw["auth"] = self.basic
        hdrs = await self._headers(headers)
        try:
            return await api(method, self.url(path), service=self.service, reconnect=self.reconnect, headers=hdrs, **kw)
        except ApiError as e:
            if e.status == 401 and self.on_unauthorized is not None and self.auth is not None:
                self.on_unauthorized()  # drop the cached token and retry once with a fresh one
                hdrs = await self._headers(headers)
                return await api(method, self.url(path), service=self.service, reconnect=self.reconnect,
                                 headers=hdrs, **kw)
            raise

    async def get(self, path: str, **kw: Any) -> Any:
        return await self.request("GET", path, **kw)

    async def post(self, path: str, **kw: Any) -> Any:
        return await self.request("POST", path, **kw)

    async def put(self, path: str, **kw: Any) -> Any:
        return await self.request("PUT", path, **kw)

    async def patch(self, path: str, **kw: Any) -> Any:
        return await self.request("PATCH", path, **kw)

    async def delete(self, path: str, **kw: Any) -> Any:
        return await self.request("DELETE", path, **kw)

    async def fetch(self, path: str, *, params: Any = None, headers: dict[str, str] | None = None,
                    max_bytes: int = DEFAULT_MAX_DOWNLOAD, method: str = "GET", json: Any = None) -> tuple[bytes, httpx.Headers]:
        """Download a (possibly large) body with a size cap. Redirects to pre-signed URLs are followed
        (httpx drops the Authorization header on cross-origin redirects)."""
        hdrs = await self._headers(headers)
        try:
            return await fetch_bytes(self.url(path), service=self.service, reconnect=self.reconnect, headers=hdrs,
                                     params=params, max_bytes=max_bytes, method=method, json=json,
                                     auth=self.basic)
        except ApiError as e:
            if e.status == 401 and self.on_unauthorized is not None and self.auth is not None:
                self.on_unauthorized()
                hdrs = await self._headers(headers)
                return await fetch_bytes(self.url(path), service=self.service, reconnect=self.reconnect,
                                         headers=hdrs, params=params, max_bytes=max_bytes, method=method, json=json,
                                         auth=self.basic)
            raise


async def fetch_bytes(url: str, *, service: str, reconnect: str | None = None, headers: dict[str, str] | None = None,
                      params: Any = None, max_bytes: int = DEFAULT_MAX_DOWNLOAD, method: str = "GET",
                      json: Any = None, auth: Any = None) -> tuple[bytes, httpx.Headers]:
    kwargs: dict[str, Any] = {}
    if headers:
        kwargs["headers"] = headers
    if params:
        kwargs["params"] = clean_params(params)
    if json is not None:
        kwargs["json"] = json
    if auth is not None:
        kwargs["auth"] = auth
    limit_mb = max_bytes / 1_000_000
    try:
        async with client().stream(method, url, **kwargs) as resp:
            if resp.status_code >= 400:
                await resp.aread()
                raise http_error(service, reconnect or service, resp)
            length = resp.headers.get("content-length")
            if length and length.isdigit() and int(length) > max_bytes:
                raise ToolError(f"The file is too large ({int(length) / 1_000_000:.1f} MB; limit {limit_mb:.0f} MB).")
            buf = bytearray()
            async for chunk in resp.aiter_bytes():
                buf.extend(chunk)
                if len(buf) > max_bytes:
                    raise ToolError(f"The file is too large (over {limit_mb:.0f} MB).")
            return bytes(buf), resp.headers
    except (httpx.UnsupportedProtocol, httpx.InvalidURL) as e:
        raise ApiError(f"The {service} URL looks invalid ({e}).") from e
    except httpx.TransportError as e:
        raise ApiError(f"Network error while downloading from {service} ({type(e).__name__}).") from e


async def gather_limited(coros: Iterable[Awaitable[Any]], limit: int = 6) -> list[Any]:
    """Run awaitables concurrently (at most `limit` at a time); results keep input order."""
    sem = asyncio.Semaphore(limit)

    async def run(c: Awaitable[Any]) -> Any:
        async with sem:
            return await c

    return list(await asyncio.gather(*(run(c) for c in coros)))


# --------------------------------------------------------------------------- connector selection
KIND_LABELS = {
    "gmail": "Gmail", "microsoft365": "Microsoft 365", "imap_smtp": "IMAP/SMTP email",
    "google_calendar": "Google Calendar", "google_drive": "Google Drive", "dropbox": "Dropbox", "box": "Box",
    "slack": "Slack", "teams": "Microsoft Teams", "github": "GitHub", "jira": "Jira", "linear": "Linear",
    "asana": "Asana", "notion": "Notion", "confluence": "Confluence", "hubspot": "HubSpot",
    "salesforce": "Salesforce", "aws": "AWS", "azure": "Azure", "gcp": "Google Cloud", "webhook": "Webhook",
    "twilio": "Twilio",
}


def label(kind: str) -> str:
    return KIND_LABELS.get(kind, kind)


async def connected(ctx: ToolContext, kinds: Sequence[str]) -> list[tuple[str, str]]:
    """(kind, connector name) pairs available for `kinds`, in priority order of `kinds`."""
    overrides = ctx.extra.get("connectors") or {}
    found: list[tuple[str, str]] = [
        (k, str((overrides.get(k) or {}).get("_connector_name") or k)) for k in kinds if k in overrides
    ]
    rest = [k for k in kinds if k not in overrides]
    if rest:
        try:
            from sqlalchemy import select

            from app.models import Connector

            async with ctx.session() as db:
                rows = (await db.execute(
                    select(Connector.kind, Connector.name)
                    .where(Connector.org_id == ctx.org_id, Connector.kind.in_(rest))
                    .order_by(Connector.created_at)
                )).all()
            found += [(r[0], r[1]) for r in rows]
        except Exception:  # noqa: BLE001
            log.exception("connector lookup failed")
    order = {k: i for i, k in enumerate(kinds)}
    found.sort(key=lambda kn: order.get(kn[0], 99))  # stable: keeps created_at order within a kind
    return found


async def load_connector(ctx: ToolContext, kind: str, name: str | None = None) -> dict[str, Any]:
    overrides = ctx.extra.get("connectors") or {}
    if kind in overrides:
        return dict(overrides[kind])
    return await ctx.connector(kind, name=name)


async def pick_connector(ctx: ToolContext, kinds: Sequence[str], account: str | None = None, *,
                         what: str) -> tuple[str, dict[str, Any]]:
    """Return (kind, config) of the connector that should serve this call.

    With `account`, match a connector name or kind (case-insensitive); otherwise the first connected kind in
    priority order. Raises a friendly ToolError when nothing suitable is connected."""
    avail = await connected(ctx, kinds)
    if account and account.strip():
        a = account.strip().lower()
        for k, n in avail:
            if a in (k.lower(), n.lower(), label(k).lower()):
                return k, await load_connector(ctx, k, n)
        names = ", ".join(f"{n} ({label(k)})" for k, n in avail) or "none"
        raise ToolError(f"No {what} account named '{account}' is connected. Connected: {names}.")
    if not avail:
        raise ToolError(f"{what} is not connected — ask an admin to connect "
                        f"{' or '.join(label(k) for k in kinds)} in Connectors → Integrations.")
    k, n = avail[0]
    return k, await load_connector(ctx, k, n)


def need(cfg: dict[str, Any], keys: Sequence[str], service: str) -> None:
    missing = [k for k in keys if not str(cfg.get(k) or "").strip()]
    if missing:
        raise ToolError(f"The {service} connector is missing: {', '.join(missing)}. Update it in Connectors.")


def truthy(value: Any, default: bool = False) -> bool:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in ("1", "true", "yes", "y", "on")


def connector_test(fn: Callable[[dict[str, Any]], Awaitable[tuple[bool, str]]]):
    """Wrap a connector `test` coroutine: never raises, bounded time, friendly message."""

    async def wrapper(config: dict[str, Any]) -> tuple[bool, str]:
        try:
            return await asyncio.wait_for(fn(dict(config or {})), timeout=30)
        except asyncio.TimeoutError:
            return False, "Connection test timed out after 30 seconds."
        except ToolError as e:
            return False, str(e)
        except Exception as e:  # noqa: BLE001
            log.warning("connector test failed: %s", type(e).__name__)
            return False, f"Connection test failed ({type(e).__name__})."

    wrapper.__name__ = getattr(fn, "__name__", "test")
    wrapper.__doc__ = fn.__doc__
    return wrapper


# --------------------------------------------------------------------------- small formatting helpers
def clip(text: Any, limit: int) -> str:
    s = "" if text is None else str(text)
    return s if len(s) <= limit else s[: max(0, limit - 1)].rstrip() + "…"


def compact(d: dict[str, Any]) -> dict[str, Any]:
    return {k: v for k, v in d.items() if v is not None and v != "" and v != [] and v != {}}


def as_list(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, str):
        return [p.strip() for p in re.split(r"[,;\n]", value) if p.strip()]
    if isinstance(value, (list, tuple, set)):
        return [str(v).strip() for v in value if str(v).strip()]
    return [str(value)]


def iso(value: datetime | None) -> str | None:
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")


def parse_datetime(value: Any) -> datetime | None:
    """ISO 8601, RFC 2822 or epoch (s/ms) → aware UTC datetime (None if unparseable)."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    if isinstance(value, (int, float)) or (isinstance(value, str) and value.isdigit()):
        n = float(value)
        if n > 1e11:
            n /= 1000.0
        return datetime.fromtimestamp(n, tz=timezone.utc)
    s = str(value).strip()
    try:
        d = datetime.fromisoformat(s.replace("Z", "+00:00"))
        return d if d.tzinfo else d.replace(tzinfo=timezone.utc)
    except ValueError:
        pass
    m = re.match(r"^(\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2})(\.\d+)?(Z|[+-]\d{2}:?\d{2})?$", s)
    if m:  # e.g. Graph's 7-digit fractional seconds
        try:
            d = datetime.fromisoformat(m.group(1) + ((m.group(3) or "+00:00").replace("Z", "+00:00")))
            return d
        except ValueError:
            pass
    try:
        d = parsedate_to_datetime(s)
        return d if d.tzinfo else d.replace(tzinfo=timezone.utc)
    except Exception:  # noqa: BLE001
        return None


def iso_or_raw(value: Any) -> str | None:
    d = parse_datetime(value)
    return iso(d) if d else (str(value) if value else None)


# --------------------------------------------------------------------------- workspace files
_UNSAFE = re.compile(r'[<>:"/\\|?*\x00-\x1f]+')
_RESERVED = {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}


def safe_filename(name: str | None, default: str = "download") -> str:
    name = _UNSAFE.sub("_", (name or "").strip()).strip(" .")
    if not name:
        name = default
    stem, ext = os.path.splitext(name)
    if stem.upper() in _RESERVED:
        stem = "_" + stem
    return stem[:150] + ext[:20]


def ws_path(ctx: ToolContext, vpath: str, **kw: Any) -> Path:
    from app.services.workspace import WorkspaceError

    try:
        return ctx.workspace.resolve(vpath, **kw)
    except WorkspaceError as e:
        raise ToolError(str(e)) from e


def unique_path(ctx: ToolContext, folder: str, filename: str) -> Path:
    name = safe_filename(filename)
    stem, ext = os.path.splitext(name)
    folder = folder.strip("/") or "downloads"
    for i in range(1000):
        candidate = name if i == 0 else f"{stem} ({i}){ext}"
        p = ws_path(ctx, f"{folder}/{candidate}", for_write=True)
        if not p.exists():
            return p
    raise ToolError(f"Too many files named '{name}' in {folder}/.")


async def save_bytes(ctx: ToolContext, data: bytes, filename: str, folder: str = "downloads") -> dict[str, Any]:
    """Save bytes into the user's workspace (never overwrites: adds ' (1)', ' (2)'...)."""
    if not ctx.can_write:
        raise ToolError("Your role is read-only, so files can't be saved to the workspace.")
    p = unique_path(ctx, folder, filename)
    await asyncio.to_thread(p.write_bytes, data)
    return {"path": ctx.workspace.virtual(p), "name": p.name, "size": len(data),
            "mime": mimetypes.guess_type(p.name)[0] or "application/octet-stream"}


async def read_workspace_file(ctx: ToolContext, path: str, max_bytes: int = DEFAULT_MAX_DOWNLOAD) -> tuple[bytes, str, str]:
    """Read a workspace file for upload/attachment: (bytes, file name, mime)."""
    p = ws_path(ctx, path, must_exist=True)
    if not p.is_file():
        raise ToolError(f"'{path}' is a folder, not a file.")
    size = p.stat().st_size
    if size > max_bytes:
        raise ToolError(f"'{path}' is too large ({size / 1_000_000:.1f} MB; limit {max_bytes / 1_000_000:.0f} MB).")
    data = await asyncio.to_thread(p.read_bytes)
    return data, p.name, mimetypes.guess_type(p.name)[0] or "application/octet-stream"


# --------------------------------------------------------------------------- text extraction
TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".jsonl", ".xml", ".yaml", ".yml", ".log", ".ini", ".cfg",
    ".toml", ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".go", ".rs", ".rb", ".php", ".c", ".h", ".cpp", ".cs",
    ".sql", ".sh", ".ps1", ".bat", ".css", ".scss", ".vtt", ".srt", ".tex", ".rst", ".env.example",
}
_BLOCK_TAGS = {"p", "div", "section", "article", "header", "footer", "li", "tr", "table", "ul", "ol", "blockquote",
               "pre", "h1", "h2", "h3", "h4", "h5", "h6", "hr", "dt", "dd", "br", "main", "nav", "aside", "figure"}


def decode_text(data: bytes) -> str:
    if data.startswith(b"\xff\xfe") or data.startswith(b"\xfe\xff"):
        return data.decode("utf-16", "replace")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", "replace")


def html_to_text(html: str) -> str:
    """Readable text from HTML: block elements become line breaks, list items get bullets."""
    if not html:
        return ""
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "html.parser")
    for t in soup(["script", "style", "head", "noscript", "template"]):
        t.decompose()
    for tag in soup.find_all(True):
        if tag.name == "br":
            tag.replace_with("\n")
            continue
        if tag.name in _BLOCK_TAGS:
            tag.insert_before("\n")
            if tag.name == "li":
                tag.insert_before("- ")
            tag.insert_after("\n")
        elif tag.name in ("td", "th"):
            tag.insert_after(" | ")
        elif tag.name == "a" and tag.get("href", "").startswith("http"):
            text = tag.get_text(strip=True)
            if text and text != tag["href"] and len(tag["href"]) < 200:
                tag.insert_after(f" ({tag['href']})")
    text = soup.get_text()
    lines = [re.sub(r"[ \t ]+", " ", ln).strip() for ln in text.splitlines()]
    out = "\n".join(lines)
    return re.sub(r"\n{3,}", "\n\n", out).strip()


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as e:  # noqa: BLE001
                raise ToolError("This PDF is password-protected.") from e
        parts = []
        for i, page in enumerate(reader.pages[:500], 1):
            t = (page.extract_text() or "").strip()
            if t:
                parts.append(f"--- page {i} ---\n{t}")
        return "\n\n".join(parts)
    except ToolError:
        raise
    except Exception as e:  # noqa: BLE001
        raise ToolError(f"Could not read the PDF ({type(e).__name__}).") from e


def _docx_text(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    out: list[str] = []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading"):
            level = re.sub(r"\D", "", style) or "1"
            out.append("#" * min(int(level), 6) + " " + t)
        elif "list" in style:
            out.append("- " + t)
        else:
            out.append(t)
    for ti, table in enumerate(d.tables, 1):
        out.append(f"\n[table {ti}]")
        for row in table.rows[:500]:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(out)


def _xlsx_text(data: bytes, max_rows: int = 2000) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: list[str] = []
    try:
        for ws in wb.worksheets:
            out.append(f"## Sheet: {ws.title}")
            n = 0
            for row in ws.iter_rows(values_only=True):
                cells = ["" if v is None else str(v) for v in row[:60]]
                if not any(cells):
                    continue
                out.append("\t".join(cells).rstrip())
                n += 1
                if n >= max_rows:
                    out.append(f"… [more rows in '{ws.title}' not shown]")
                    break
    finally:
        wb.close()
    return "\n".join(out)


def _pptx_text(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    out: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"## Slide {i}")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                out.append(shape.text_frame.text.strip())
            elif getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    out.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append(f"[notes] {notes}")
    return "\n".join(out)


def _eml_text(data: bytes) -> str:
    import email
    import email.policy

    msg = email.message_from_bytes(data, policy=email.policy.default)
    head = "\n".join(f"{h}: {msg.get(h)}" for h in ("From", "To", "Cc", "Date", "Subject") if msg.get(h))
    body = msg.get_body(preferencelist=("plain", "html"))
    text = ""
    if body is not None:
        try:
            text = body.get_content()
        except Exception:  # noqa: BLE001
            text = (body.get_payload(decode=True) or b"").decode("utf-8", "replace")
        if body.get_content_type() == "text/html":
            text = html_to_text(text)
    return f"{head}\n\n{text}".strip()


def _looks_textual(data: bytes) -> bool:
    sample = data[:4096]
    if not sample or b"\x00" in sample:
        return False
    try:
        sample.decode("utf-8")
        return True
    except UnicodeDecodeError:
        return False


def extract_text(data: bytes, filename: str = "", mime: str | None = None) -> str:
    """Extract readable text from common document formats (sync — call via asyncio.to_thread)."""
    ext = Path(filename or "").suffix.lower()
    mime = (mime or mimetypes.guess_type(filename or "")[0] or "").lower()
    if ext == ".pdf" or mime == "application/pdf":
        return _pdf_text(data)
    if ext == ".docx" or "wordprocessingml" in mime:
        return _docx_text(data)
    if ext in (".xlsx", ".xlsm") or "spreadsheetml" in mime:
        return _xlsx_text(data)
    if ext == ".pptx" or "presentationml" in mime:
        return _pptx_text(data)
    if ext in (".html", ".htm", ".xhtml") or "html" in mime:
        return html_to_text(decode_text(data))
    if ext == ".eml" or mime == "message/rfc822":
        return _eml_text(data)
    if ext in TEXT_EXTS or mime.startswith("text/") or mime in ("application/json", "application/xml",
                                                                  "application/x-yaml", "application/csv"):
        return decode_text(data)
    if ext in (".doc", ".xls", ".ppt"):
        raise ToolError(f"Legacy {ext} files aren't supported for text extraction — download the file instead "
                        f"(or convert it to {ext}x).")
    if _looks_textual(data):
        return decode_text(data)
    raise ToolError(f"Can't extract text from this file type ({ext or mime or 'unknown'}). Use the download tool "
                    "to save it to the workspace instead.")


async def extract_text_async(data: bytes, filename: str = "", mime: str | None = None) -> str:
    return await asyncio.to_thread(extract_text, data, filename, mime)


def window(text: str, max_chars: int, offset: int = 0) -> dict[str, Any]:
    """Slice long text for the model: {text, chars, offset, truncated, next_offset}."""
    total = len(text)
    offset = max(0, min(offset, total))
    chunk = text[offset: offset + max(1, max_chars)]
    end = offset + len(chunk)
    out: dict[str, Any] = {"chars": total, "offset": offset, "text": chunk}
    if end < total:
        out["truncated"] = True
        out["next_offset"] = end
        out["note"] = f"Showing characters {offset}-{end} of {total}; call again with offset={end} for more."
    return out


# --------------------------------------------------------------------------- approval previews
def _fence(text: str) -> str:
    fence = "```"
    while fence in text:
        fence += "`"
    return fence


def preview(title: str, rows: Iterable[tuple[str, Any]], body: Any = None, *, body_label: str = "Message",
            note: str | None = None, max_body: int = 20000) -> str:
    """Markdown preview for approvers: key facts as a list, then the outbound content verbatim in a code block."""
    lines = [f"**{title}**", ""]
    for name, value in rows:
        if value is None or value == "" or value == [] or value == ():
            continue
        if isinstance(value, bool):
            value = "yes" if value else "no"
        elif isinstance(value, (list, tuple, set)):
            value = ", ".join(str(v) for v in value)
        elif isinstance(value, dict):
            value = jsonlib.dumps(value, ensure_ascii=False)
        lines.append(f"- **{name}:** {clip(str(value).replace(chr(10), ' '), 1500)}")
    if note:
        lines += ["", f"_{note}_"]
    if body is not None:
        text = str(body)
        if len(text) > max_body:
            text = text[:max_body] + f"\n… [{len(text) - max_body} more characters not shown]"
        fence = _fence(text)
        lines += ["", f"**{body_label}:**", fence, text if text else "(empty)", fence]
    return "\n".join(lines)


# --------------------------------------------------------------------------- LLM helpers
async def _record_usage(ctx: ToolContext, model: str, usage: dict | None) -> None:
    try:
        from app.services.usage import record_usage

        await record_usage(ctx.org_id, model, usage, purpose="agent", agent_id=ctx.agent_id, task_id=ctx.task_id)
    except Exception:  # noqa: BLE001
        log.debug("usage recording skipped", exc_info=True)


async def llm_complete(ctx: ToolContext, system: str, prompt: str, *, max_tokens: int = 1000) -> str:
    """Free-text model call for summaries/briefings (uses the org's utility model)."""
    try:
        llm = await ctx.llm()
        res = await llm.chat([{"role": "system", "content": system}, {"role": "user", "content": prompt}],
                             temperature=0.2, max_tokens=max_tokens)
    except ToolError:
        raise
    except Exception as e:  # noqa: BLE001
        raise ToolError(f"The language model step failed: {e}") from e
    await _record_usage(ctx, res.model, res.usage)
    return (res.content or "").strip()


async def llm_structured(ctx: ToolContext, system: str, prompt: str, schema: dict, *, max_tokens: int = 1500) -> Any:
    """Schema-constrained JSON model call for extraction."""
    try:
        llm = await ctx.llm()
        return await llm.json([{"role": "system", "content": system}, {"role": "user", "content": prompt}], schema,
                              max_tokens=max_tokens)
    except ToolError:
        raise
    except Exception as e:  # noqa: BLE001
        raise ToolError(f"The language model step failed: {e}") from e
