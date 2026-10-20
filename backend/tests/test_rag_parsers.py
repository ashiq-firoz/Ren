"""Parsers: every supported format → ParsedDocument (headings + text + pages), generated fixtures."""
import io

import pytest

from app.rag.parsers import ParseError, parse_file


def _headings(doc):
    return [(b.level, b.text) for b in doc.blocks if b.kind == "heading"]


def _text(doc):
    return "\n".join(b.text for b in doc.blocks if b.kind != "heading")


def test_markdown(tmp_path):
    p = tmp_path / "guide.md"
    p.write_text(
        "---\ntitle: Onboarding Guide\n---\n# Onboarding\n\nWelcome aboard.\n\n## Accounts\n\nYou get an email "
        "account.\n\n```python\n# not a heading\nprint('x')\n```\n\n## Equipment\n\n| Item | Qty |\n|---|---|\n"
        "| Laptop | 1 |\n", encoding="utf-8")
    doc = parse_file(p)
    assert doc.title == "Onboarding Guide"
    assert _headings(doc) == [(1, "Onboarding"), (2, "Accounts"), (2, "Equipment")]
    assert "# not a heading" in _text(doc)  # code fence content is not a heading
    assert "| Laptop | 1 |" in _text(doc)
    assert doc.structure_hint == "headings"


def test_html_strips_chrome(tmp_path):
    html = """<html><head><title>Remote Work Guide</title><script>var tracking = 1;</script>
    <style>.x{color:red}</style></head><body>
    <nav><a href="/">Home</a> <a href="/about">About us</a></nav>
    <header><div>ACME Corp</div></header>
    <main><h1>Remote Work Guide</h1><p>Everyone may work remotely <b>two days</b> per week.</p>
    <h2>Equipment</h2><p>Laptops are provided.</p><ul><li>Monitor</li><li>Headset</li></ul>
    <h2>Stipend</h2><table><tr><th>Item</th><th>Amount</th></tr><tr><td>Internet</td><td>50 USD</td></tr></table>
    </main><footer>Copyright 2026</footer><div class="cookie-banner">We use cookies</div></body></html>"""
    p = tmp_path / "remote.html"
    p.write_text(html, encoding="utf-8")
    doc = parse_file(p)
    text = _text(doc)
    assert doc.title == "Remote Work Guide"
    assert _headings(doc) == [(1, "Remote Work Guide"), (2, "Equipment"), (2, "Stipend")]
    assert "two days per week" in text
    assert "- Monitor" in text and "- Headset" in text
    assert "| Item | Amount |" in text and "| Internet | 50 USD |" in text
    for junk in ("tracking", "About us", "Copyright", "cookies", "color:red"):
        assert junk not in text


def test_txt_numbered_headings(tmp_path):
    p = tmp_path / "conduct.txt"
    p.write_text(
        "EMPLOYEE CODE OF CONDUCT\n\n1. Purpose\nThis code sets out the standards of behaviour expected\n"
        "from all employees.\n\n2. Conflicts of Interest\nEmployees must disclose outside business interests.\n\n"
        "2.1 Gifts\nGifts above 50 euros must be declared to compliance.\n\n2.2 Outside Employment\n"
        "Secondary employment requires written approval.\n\n3. Reporting\nConcerns can be reported anonymously.\n",
        encoding="utf-8")
    doc = parse_file(p)
    heads = _headings(doc)
    assert (1, "EMPLOYEE CODE OF CONDUCT") in heads
    assert (1, "1. Purpose") in heads and (2, "2.1 Gifts") in heads and (1, "3. Reporting") in heads
    # hard-wrapped lines are unwrapped into one paragraph
    assert "standards of behaviour expected from all employees." in _text(doc)


def test_txt_flat_and_lists_are_not_headings(tmp_path):
    p = tmp_path / "notes.txt"
    p.write_text("Shopping notes for the offsite.\n\n1. Buy water\n2. Buy snacks\n3. Book the bus\n\n"
                 "Remember to keep the receipts for reimbursement.\n", encoding="utf-8")
    doc = parse_file(p)
    assert _headings(doc) == []
    assert doc.structure_hint == "flat"


def test_docx_headings_tables_lists(tmp_path):
    import docx

    d = docx.Document()
    d.core_properties.title = "Travel Policy"
    d.add_heading("Introduction", level=1)
    d.add_paragraph("This policy covers business travel for all staff.")
    d.add_heading("Booking", level=1)
    d.add_heading("Flights", level=2)
    d.add_paragraph("Book economy class for flights under six hours.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Region", "Daily rate"
    t.cell(1, 0).text, t.cell(1, 1).text = "EMEA", "180 EUR"
    d.add_heading("Hotels", level=2)
    d.add_paragraph("Use the travel portal.")
    d.add_paragraph("Keep receipts", style="List Bullet")
    d.add_paragraph("Submit within 30 days", style="List Bullet")
    p = tmp_path / "travel.docx"
    d.save(p)
    doc = parse_file(p)
    assert doc.title == "Travel Policy"
    assert _headings(doc) == [(1, "Introduction"), (1, "Booking"), (2, "Flights"), (2, "Hotels")]
    text = _text(doc)
    assert "| Region | Daily rate |" in text and "| EMEA | 180 EUR |" in text
    assert "- Keep receipts\n- Submit within 30 days" in text


def test_pptx_slides(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Q3 Strategy"
    s1.placeholders[1].text = "Board update"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Revenue Growth"
    tf = s2.placeholders[1].text_frame
    tf.text = "Revenue grew 18 percent"
    para = tf.add_paragraph()
    para.text = "Driven by EMEA"
    para.level = 1
    s2.notes_slide.notes_text_frame.text = "Mention the two new contracts"
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Regional Breakdown"
    tbl = s3.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text, tbl.cell(0, 1).text = "Region", "Revenue"
    tbl.cell(1, 0).text, tbl.cell(1, 1).text = "APAC", "4.2M"
    p = tmp_path / "deck.pptx"
    prs.save(p)
    doc = parse_file(p)
    assert doc.page_count == 3
    assert [(b.text, b.page) for b in doc.blocks if b.kind == "heading"] == [
        ("Q3 Strategy", 1), ("Revenue Growth", 2), ("Regional Breakdown", 3)]
    text = _text(doc)
    assert "Revenue grew 18 percent" in text and "  - Driven by EMEA" in text
    assert "Speaker notes: Mention the two new contracts" in text
    assert "| APAC | 4.2M |" in text


def test_xlsx_sheets(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Revenue"])
    ws.append(["EMEA", 1200.5])
    ws.append(["APAC", 980])
    ws2 = wb.create_sheet("Headcount")
    ws2.append(["Team", "People"])
    ws2.append(["Engineering", 42])
    p = tmp_path / "metrics.xlsx"
    wb.save(p)
    doc = parse_file(p)
    assert _headings(doc) == [(1, "Sales"), (1, "Headcount")]
    text = _text(doc)
    assert "| Region | Revenue |" in text and "| EMEA | 1200.5 |" in text and "| APAC | 980 |" in text
    assert "| Engineering | 42 |" in text


def test_large_sheet_is_split_into_row_blocks(tmp_path):
    p = tmp_path / "big.csv"
    rows = ["id,customer,notes"] + [f"{i},Customer {i},Some fairly long note about account {i} and its renewal"
                                    for i in range(1, 400)]
    p.write_text("\n".join(rows), encoding="utf-8")
    doc = parse_file(p)
    heads = _headings(doc)
    assert len(heads) > 2 and heads[0][1].startswith("Rows 2–")
    tables = [b.text for b in doc.blocks if b.kind == "table"]
    assert all(t.startswith("| id | customer | notes |") for t in tables)  # header repeated in each block
    assert "| 399 | Customer 399 |" in "\n".join(tables)


def test_csv(tmp_path):
    p = tmp_path / "people.csv"
    p.write_text("name;dept;salary\nAlice;Eng;100\nBob;Ops;90\n", encoding="utf-8")
    doc = parse_file(p)
    assert "| name | dept | salary |" in _text(doc) and "| Bob | Ops | 90 |" in _text(doc)
    assert doc.title == "people"


def _pdf_bytes(sections):
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(True, 15)
    for title, size, body, new_page in sections:
        if new_page:
            pdf.add_page()
        pdf.set_font("Helvetica", "B", size)
        pdf.cell(0, 10, title, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 11)
        pdf.multi_cell(0, 6, body)
        pdf.ln(4)
    return bytes(pdf.output())


SECURITY_SECTIONS = [
    ("1 Introduction", 16, "This security policy applies to every employee and contractor. " * 3, True),
    ("2 Security Controls", 16, "The controls below protect company systems and customer data. " * 2, True),
    ("2.1 Passwords", 13, "Passwords must be at least 14 characters long and rotated every year. " * 3, False),
    ("2.2 Devices", 13, "Laptops must use full disk encryption and automatic screen locking. " * 3, True),
]


def test_pdf_native_outline(tmp_path):
    from pypdf import PdfReader, PdfWriter

    data = _pdf_bytes(SECURITY_SECTIONS)
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(data)))
    writer.add_outline_item("1 Introduction", 0)
    parent = writer.add_outline_item("2 Security Controls", 1)
    writer.add_outline_item("2.1 Passwords", 1, parent=parent)
    writer.add_outline_item("2.2 Devices", 2, parent=parent)
    p = tmp_path / "security.pdf"
    with open(p, "wb") as f:
        writer.write(f)
    doc = parse_file(p)
    assert doc.structure_hint == "native" and doc.page_count == 3
    assert [(e.level, e.title, e.page) for e in doc.native_toc] == [
        (1, "1 Introduction", 1), (1, "2 Security Controls", 2), (2, "2.1 Passwords", 2), (2, "2.2 Devices", 3)]
    assert _headings(doc) == [(1, "1 Introduction"), (1, "2 Security Controls"), (2, "2.1 Passwords"),
                              (2, "2.2 Devices")]
    # text between outline entries is mapped to the right section, with page numbers
    sections, cur = {}, None
    for b in doc.blocks:
        if b.kind == "heading":
            cur = b.text
            sections[cur] = []
        elif cur:
            sections[cur].append(b)
    pw = " ".join(b.text for b in sections["2.1 Passwords"])
    assert "14 characters" in pw and "disk encryption" not in pw
    assert "disk encryption" in " ".join(b.text for b in sections["2.2 Devices"])
    assert all(b.page == 3 for b in sections["2.2 Devices"])


def test_pdf_font_heading_inference(tmp_path):
    sections = [
        ("1. Chapter One Overview", 18, "This is body text for the first chapter of the guide. " * 12, True),
        ("1.1 Scope Details", 14, "Scope text that explains who the guide is meant for. " * 8, False),
        ("2. Chapter Two Procedures", 18, "Procedures describe the steps for every request type. " * 12, True),
        ("2.1 Approval Steps", 14, "Approvals are granted by the team lead within two days. " * 8, False),
    ]
    p = tmp_path / "guide.pdf"
    p.write_bytes(_pdf_bytes(sections))
    doc = parse_file(p)
    assert doc.native_toc is None and doc.structure_hint == "headings"
    heads = _headings(doc)
    assert (1, "1. Chapter One Overview") in heads and (2, "1.1 Scope Details") in heads
    assert (1, "2. Chapter Two Procedures") in heads and (2, "2.1 Approval Steps") in heads
    assert "within two days" in _text(doc)


def test_pdf_fast_layout_path_for_large_files(tmp_path, monkeypatch):
    import app.rag.parsers.pdf as pdfmod

    monkeypatch.setattr(pdfmod, "PLUMBER_MAX_PAGES", 0)  # force the pypdf text-visitor layout pass
    sections = [
        ("1. Chapter One Overview", 18, "This is body text for the first chapter of the guide. " * 12, True),
        ("1.1 Scope Details", 14, "Scope text that explains who the guide is meant for. " * 8, False),
        ("2. Chapter Two Procedures", 18, "Procedures describe the steps for every request type. " * 12, True),
    ]
    p = tmp_path / "large.pdf"
    p.write_bytes(_pdf_bytes(sections))
    doc = parse_file(p)
    assert doc.meta["parser"] == "pypdf-layout"
    assert _headings(doc) == [(1, "1. Chapter One Overview"), (2, "1.1 Scope Details"), (1, "2. Chapter Two Procedures")]
    assert [b.page for b in doc.blocks if b.kind == "heading"] == [1, 1, 2]
    assert "who the guide is meant for" in _text(doc)


def test_unsupported_and_corrupt_files(tmp_path):
    legacy = tmp_path / "old.doc"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0legacy")
    with pytest.raises(ParseError, match="Word 97-2003"):
        parse_file(legacy)
    bad = tmp_path / "broken.pdf"
    bad.write_bytes(b"%PDF-1.4\nthis is not really a pdf")
    with pytest.raises(ParseError):
        parse_file(bad)
    empty = tmp_path / "empty.txt"
    empty.write_bytes(b"")
    with pytest.raises(ParseError, match="empty"):
        parse_file(empty)
    baddocx = tmp_path / "fake.docx"
    baddocx.write_bytes(b"not a zip file")
    with pytest.raises(ParseError, match="DOCX"):
        parse_file(baddocx)
