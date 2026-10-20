"""Artifacts group: one-call creation of deliverables in any common format (md, html, json, csv, txt, docx, pdf,
xlsx, pptx), conversion between formats, validation, and listing what has been produced in output/."""
from __future__ import annotations

import asyncio
import json
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Literal

from app.tools.base import Risk, ToolContext, ToolError, ToolGroup, ToolResult
from app.tools.builtin import _office as O

artifacts = ToolGroup(
    "artifacts", "Artifacts",
    "Create finished deliverables in any common format (Markdown, HTML, JSON, CSV, text, Word, PDF, Excel, "
    "PowerPoint) from content, convert files between formats, validate that files open correctly, and list "
    "produced artifacts.",
    prefix="artifact", icon="package", category="productivity",
)

Format = Literal["md", "html", "json", "csv", "txt", "docx", "pdf", "xlsx", "pptx"]
_HTML_DOC = re.compile(r"^\s*(<!doctype\s+html|<html[\s>])", re.I)


def _as_text(content: Any) -> str:
    if isinstance(content, str):
        return content
    return json.dumps(content, ensure_ascii=False, indent=2, default=str)


def _write_any(fmt: str, content: Any, target: Path, *, title: str | None, resolver) -> dict:
    """Render `content` into `target` as `fmt`; returns a few stats."""
    if fmt in ("md", "txt"):
        text = _as_text(content)
        if fmt == "md" and title and not text.lstrip().startswith("#"):
            text = f"# {title}\n\n{text}"
        target.write_text(text, encoding="utf-8")
        return {"chars": len(text)}
    if fmt == "html":
        text = _as_text(content)
        html = text if _HTML_DOC.match(text) else O.markdown_to_html(text, title)
        target.write_text(html, encoding="utf-8")
        return {"chars": len(html)}
    if fmt == "json":
        if isinstance(content, str):
            try:
                content = json.loads(content)
            except json.JSONDecodeError as e:
                raise ToolError(f"Content is not valid JSON (line {e.lineno}, column {e.colno}): {e.msg}") from e
        target.write_text(json.dumps(content, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        return {"type": type(content).__name__}
    if fmt == "csv":
        hdr, rows = O.table_from_any(content)
        if not rows and not hdr:
            raise ToolError("No table data found in content (use rows, records or CSV text).")
        O.write_csv(target, hdr, rows)
        return {"rows": len(rows), "columns": max([len(hdr or [])] + [len(r) for r in rows])}
    if fmt == "xlsx":
        sheets: dict[str, tuple] = {}
        if isinstance(content, dict) and content and all(isinstance(v, (list, str)) for v in content.values()) \
                and not {"rows", "data", "columns", "header"} & set(content):
            for name, data in content.items():
                sheets[str(name)] = O.table_from_any(data)
        else:
            sheets[title or "Sheet1"] = O.table_from_any(content)
        if not any(r or h for h, r in sheets.values()):
            raise ToolError("No table data found in content.")
        O.write_xlsx(target, sheets)
        return {"sheets": {k: len(v[1]) for k, v in sheets.items()}}
    if fmt == "docx":
        doc = O.markdown_to_docx(_as_text(content), target, title=title, image_resolver=resolver)
        return {"paragraphs": len(doc.paragraphs), "tables": len(doc.tables)}
    if fmt == "pdf":
        O.markdown_to_pdf(_as_text(content), target, title=title, image_resolver=resolver)
        return {"pages": len(O.pdf_reader(target).pages)}
    if fmt == "pptx":
        if isinstance(content, list):
            prs = O.build_deck(target, title, content, image_resolver=resolver)
        else:
            spec = O.markdown_to_slide_specs(_as_text(content))
            if not spec["slides"] and not (title or spec["title"]):
                raise ToolError("No slides found — use '## Slide title' headings with bullet lists.")
            prs = O.build_deck(target, title or spec["title"], spec["slides"], subtitle=spec["subtitle"],
                               image_resolver=resolver)
        return {"slides": len(prs.slides)}
    raise ToolError(f"Unsupported format '{fmt}'.")


@artifacts.tool(risk=Risk.MEDIUM, scope="write", timeout=180)
async def create(ctx: ToolContext, content: Any, format: Format, name: str | None = None, title: str | None = None,
                 overwrite: bool = False) -> ToolResult:
    """Create a deliverable file in output/ from content, in one call:
    md/txt (text), html (markdown → styled page, or a full HTML document as-is), json (object or JSON text),
    csv/xlsx (rows, records, CSV or markdown-table text; xlsx also takes {"Sheet": rows, ...}), docx/pdf
    (markdown → formatted document), pptx (markdown with "## Slide" headings + bullets, or a list of
    {title, bullets, notes} slides).

    Args:
        content: The content (text/markdown, or structured data for json/csv/xlsx/pptx).
        format: Output format.
        name: File name or path, e.g. "q3-summary" or "output/reports/q3.docx" (default: from the title).
        title: Document/deck title (also used for the default file name).
        overwrite: Replace an existing file instead of saving under a new name.
    """
    if content is None or (isinstance(content, str) and not content.strip()):
        raise ToolError("`content` is empty.")
    target = O.ws_output(ctx, name, default_name=title or name or "artifact", ext="." + format, overwrite=overwrite)
    resolver = O.workspace_image_resolver(ctx)
    stats = await asyncio.to_thread(_write_any, format, content, target, title=title, resolver=resolver)
    return await O.produced(ctx, target, {"format": format, **stats}, display=f"Created {O.vpath(ctx, target)}")


# ----------------------------------------------------------------------------------------- export
_CONVERSIONS = {
    "md": {"docx", "pdf", "html", "txt", "pptx"}, "markdown": {"docx", "pdf", "html", "txt", "pptx"},
    "txt": {"docx", "pdf", "html", "md"}, "docx": {"pdf", "md", "html", "txt"},
    "html": {"md", "txt", "pdf", "docx"}, "htm": {"md", "txt", "pdf", "docx"},
    "csv": {"xlsx", "json", "md", "tsv"}, "tsv": {"xlsx", "json", "md", "csv"},
    "xlsx": {"csv", "json", "md", "tsv"}, "xlsm": {"csv", "json", "md"},
    "json": {"csv", "xlsx", "md"}, "pptx": {"md", "pdf", "txt"}, "pdf": {"txt", "md"},
}


def _convert(src: Path, to: str, target: Path, *, sheet: str | None, resolver) -> dict:
    ext = src.suffix.lower().lstrip(".")
    if ext in ("md", "markdown", "txt"):
        md = O.read_text_file(src)
        title = None
        if to == "html":
            m = re.match(r"\s*#\s+(.+)", md)
            title = m.group(1).strip() if m else src.stem
            target.write_text(O.markdown_to_html(md, None if m else title), encoding="utf-8")
            return {}
        if to in ("docx", "pdf"):
            return _write_any(to, md, target, title=title, resolver=resolver)
        if to == "pptx":
            return _write_any("pptx", md, target, title=None, resolver=resolver)
        target.write_text(md, encoding="utf-8")
        return {}
    if ext == "docx":
        if to == "pdf":
            from app.tools.builtin.documents import _libreoffice_pdf

            if _libreoffice_pdf(src, target):
                return {"engine": "libreoffice"}
            md = O.docx_to_markdown(src)
            O.markdown_to_pdf(md, target, image_resolver=resolver)
            return {"engine": "builtin", "pages": len(O.pdf_reader(target).pages)}
        md = O.docx_to_markdown(src)
        if to == "html":
            target.write_text(O.markdown_to_html(md, src.stem), encoding="utf-8")
        else:
            target.write_text(md, encoding="utf-8")
        return {}
    if ext in ("html", "htm"):
        title, text = O.html_to_text(O.read_text_file(src))
        md = (f"# {title}\n\n" if title else "") + text
        if to in ("docx", "pdf"):
            return _write_any(to, md, target, title=None, resolver=resolver)
        target.write_text(md, encoding="utf-8")
        return {}
    if ext in ("csv", "tsv", "xlsx", "xlsm", "json"):
        if ext in ("csv", "tsv"):
            rows = O.read_csv_rows(src)
            hdr, body = (rows[0], rows[1:]) if rows else ([], [])
        elif ext == "json":
            hdr, body = O.table_from_any(O.read_text_file(src))
        else:
            from openpyxl import load_workbook

            wb = load_workbook(str(src), read_only=True, data_only=True)
            try:
                ws = wb[sheet] if sheet else wb.worksheets[0]
                rows = [list(r) for r in ws.iter_rows(values_only=True)]
            except KeyError as e:
                raise ToolError(f"Sheet '{sheet}' not found. Sheets: {', '.join(wb.sheetnames)}.") from e
            finally:
                wb.close()
            while rows and all(v in (None, "") for v in rows[-1]):
                rows.pop()
            hdr, body = (rows[0], rows[1:]) if rows else ([], [])
        hdr = [("" if h is None else str(h)) for h in (hdr or [])]
        if to == "xlsx":
            O.write_xlsx(target, {src.stem[:31]: (hdr or None, body)})
        elif to in ("csv", "tsv"):
            O.write_csv(target, hdr or None, body, delimiter="\t" if to == "tsv" else ",")
        elif to == "json":
            recs = [{(hdr[i] if i < len(hdr) and hdr[i] else f"col{i + 1}"): O.jsonable(v) for i, v in enumerate(r)}
                    for r in body]
            target.write_text(json.dumps(recs, ensure_ascii=False, indent=2), encoding="utf-8")
        else:
            target.write_text(O.md_table(hdr, body), encoding="utf-8")
        return {"rows": len(body)}
    if ext == "pptx":
        md = O.pptx_to_markdown(src)
        if to == "pdf":
            O.markdown_to_pdf(md, target, image_resolver=resolver)
            return {"pages": len(O.pdf_reader(target).pages)}
        target.write_text(md, encoding="utf-8")
        return {}
    if ext == "pdf":
        n, pages = O.pdf_pages_text(src)
        sep = "\n\n" if to == "txt" else "\n\n---\n\n"
        target.write_text(sep.join(t for _, t in pages), encoding="utf-8")
        return {"pages": n}
    raise ToolError(f"Can't convert .{ext} files.")


@artifacts.tool(risk=Risk.MEDIUM, scope="write", timeout=240)
async def export(ctx: ToolContext, path: str, to: Literal["md", "docx", "pdf", "html", "txt", "csv", "tsv", "xlsx",
                                                           "json", "pptx"],
                 output: str | None = None, sheet: str | None = None) -> ToolResult:
    """Convert a workspace file to another format: md → docx/pdf/html/pptx, docx → pdf/md/html, html → md/pdf/docx,
    csv ↔ xlsx, xlsx/csv → json/md, json → csv/xlsx, pptx → md/pdf, pdf → txt/md.

    Args:
        path: Workspace path of the source file.
        to: Target format.
        output: Output path (default: output/<name>.<format>).
        sheet: Sheet to export when converting from Excel.
    """
    src = O.ws_input(ctx, path)
    ext = src.suffix.lower().lstrip(".")
    allowed = _CONVERSIONS.get(ext)
    if not allowed:
        raise ToolError(f"Converting .{ext or '?'} files isn't supported.")
    if to not in allowed:
        raise ToolError(f"Can't convert .{ext} to {to}. Supported targets: {', '.join(sorted(allowed))}.")
    target = O.ws_output(ctx, output, default_name=src.stem, ext="." + to, overwrite=bool(output))
    resolver = O.workspace_image_resolver(ctx)
    stats = await asyncio.to_thread(_convert, src, to, target, sheet=sheet, resolver=resolver)
    return await O.produced(ctx, target, {"from": O.vpath(ctx, src), "format": to, **stats},
                            display=f"Converted {src.name} → {target.name}")


# ----------------------------------------------------------------------------------------- validate / list
def _validate(p: Path) -> dict:
    ext = p.suffix.lower()
    issues: list[str] = []
    stats: dict[str, Any] = {}
    fmt = ext.lstrip(".") or "unknown"
    if p.stat().st_size == 0:
        return {"ok": False, "format": fmt, "issues": ["File is empty."], "stats": {}}
    try:
        if ext in (".docx", ".xlsx", ".xlsm", ".pptx"):
            with zipfile.ZipFile(p) as z:
                bad = z.testzip()
                if bad:
                    issues.append(f"Corrupt zip member: {bad}")
        if ext == ".docx":
            from docx import Document

            doc = Document(str(p))
            words = sum(len(par.text.split()) for par in doc.paragraphs)
            heads = sum(1 for par in doc.paragraphs if par.style is not None and
                        (par.style.name or "").startswith("Heading"))
            stats = {"paragraphs": len(doc.paragraphs), "tables": len(doc.tables), "words": words, "headings": heads,
                     "images": len(doc.inline_shapes)}
            if words == 0 and not doc.tables:
                issues.append("Document has no text.")
        elif ext in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook

            wb = load_workbook(str(p))
            stats = {"sheets": {ws.title: f"{ws.max_row}×{ws.max_column}" for ws in wb.worksheets},
                     "charts": sum(len(ws._charts) for ws in wb.worksheets)}
            errs = 0
            for ws in wb.worksheets:
                for row in ws.iter_rows(max_row=min(ws.max_row, 5000)):
                    for c in row:
                        if isinstance(c.value, str) and c.value.startswith("#") and c.value.rstrip("!?0/").upper() in \
                                ("#REF", "#DIV", "#VALUE", "#NAME", "#N/A", "#NUM", "#NULL"):
                            errs += 1
            if errs:
                issues.append(f"{errs} cell(s) contain Excel error values.")
        elif ext in (".csv", ".tsv"):
            text = O.read_text_file(p)
            rows = O.parse_csv_text(text, delimiter="\t" if ext == ".tsv" else None)
            widths = {len(r) for r in rows if r}
            stats = {"rows": max(0, len(rows) - 1), "columns": max(widths) if widths else 0}
            if len(widths) > 1:
                issues.append(f"Inconsistent column counts across rows: {sorted(widths)[:6]}.")
        elif ext == ".pptx":
            from pptx import Presentation

            prs = Presentation(str(p))
            titled = sum(1 for s in prs.slides if s.shapes.title is not None and s.shapes.title.text_frame.text.strip())
            stats = {"slides": len(prs.slides), "slides_with_titles": titled,
                     "size": f"{prs.slide_width / 914400:.2f}×{prs.slide_height / 914400:.2f} in"}
            if not len(prs.slides):
                issues.append("Deck has no slides.")
        elif ext == ".pdf":
            r = O.pdf_reader(p)
            sample = "".join((r.pages[i].extract_text() or "") for i in range(min(3, len(r.pages))))
            stats = {"pages": len(r.pages), "has_text_layer": len(sample.strip()) > 20}
        elif ext == ".json":
            data = json.loads(O.read_text_file(p))
            stats = {"type": type(data).__name__, "items": len(data) if isinstance(data, (list, dict)) else None}
        elif ext in (".html", ".htm"):
            from app.tools.builtin import _browser

            soup = _browser.make_soup(O.read_text_file(p))
            stats = {"title": _browser.page_title(soup), "links": len(soup.find_all("a")),
                     "images": len(soup.find_all("img")), "has_script": bool(soup.find("script"))}
        elif ext in (".md", ".markdown", ".txt"):
            text = O.read_text_file(p)
            stats = {"chars": len(text), "words": len(text.split()),
                     "headings": len(re.findall(r"^#{1,6}\s", text, re.M))}
        elif ext in (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff"):
            from PIL import Image

            with Image.open(p) as im:
                im.verify()
            with Image.open(p) as im:
                stats = {"width": im.size[0], "height": im.size[1], "mode": im.mode}
        else:
            issues.append(f"No validator for '{ext}' files; checked that the file exists and is non-empty.")
    except ToolError as e:
        issues.append(str(e))
    except Exception as e:  # noqa: BLE001
        issues.append(f"Failed to open as {fmt}: {e.__class__.__name__}: {str(e)[:200]}")
    blocking = [i for i in issues if not i.startswith("No validator")]
    return {"ok": not any(i.startswith(("Failed", "Corrupt", "File is empty", "Content")) for i in blocking)
            and not (ext == ".pdf" and not stats), "format": fmt, "stats": stats, "issues": issues}


@artifacts.tool(risk=Risk.LOW, scope="read", timeout=60)
async def validate(ctx: ToolContext, path: str) -> dict:
    """Check that a file opens and parses correctly (docx, xlsx, csv, pptx, pdf, json, html, md, images) and report
    basic stats and problems (corruption, empty content, inconsistent CSV rows, Excel error values).

    Args:
        path: Workspace path of the file to check.
    """
    p = O.ws_input(ctx, path)
    res = await asyncio.to_thread(_validate, p)
    return {"path": O.vpath(ctx, p), "size": O.human_size(p.stat().st_size), **res}


@artifacts.tool(name="list", risk=Risk.LOW, scope="read", timeout=30)
async def list_(ctx: ToolContext, path: str = "output", recursive: bool = True, extension: str | None = None,
                limit: int = 100) -> dict:
    """List produced files (default: everything under output/), newest first.

    Args:
        path: Folder to list.
        recursive: Include sub-folders.
        extension: Only files with this extension, e.g. "pdf".
        limit: Maximum files to return.
    """
    root = O.ws_path(ctx, path or "output")
    if not root.exists():
        return {"path": path, "count": 0, "files": []}
    if root.is_file():
        files = [root]
    else:
        it = root.rglob("*") if recursive else root.iterdir()
        files = [f for f in it if f.is_file()]
    if extension:
        ext = "." + extension.lower().lstrip(".")
        files = [f for f in files if f.suffix.lower() == ext]
    files.sort(key=lambda f: f.stat().st_mtime, reverse=True)
    n = max(1, min(int(limit or 100), 1000))
    out = []
    for f in files[:n]:
        st = f.stat()
        out.append({"path": O.vpath(ctx, f), "size": O.human_size(st.st_size),
                    "modified": datetime.fromtimestamp(st.st_mtime, tz=timezone.utc).strftime("%Y-%m-%d %H:%M UTC")})
    return {"path": O.vpath(ctx, root), "count": len(files), "files": out, "truncated": len(files) > n}

