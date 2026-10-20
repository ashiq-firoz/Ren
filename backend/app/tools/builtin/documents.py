"""Documents group: Word (.docx) deliverables — create from markdown, read back as markdown, append, find &
replace (across formatting runs), insert tables/images, comment, compare versions, export to PDF, restyle."""
from __future__ import annotations

import asyncio
import copy
import difflib
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any, Literal

from app.tools.base import Risk, ToolContext, ToolError, ToolGroup, ToolResult
from app.tools.builtin import _office as O

documents = ToolGroup(
    "documents", "Documents",
    "Create and edit Word documents (.docx): write professional documents from markdown, read them back as "
    "markdown, append sections, find & replace, insert tables and images, add review comments, compare two "
    "versions, export to PDF and restyle.",
    prefix="doc", icon="file-text", category="productivity",
)


def _open(path: Path):
    from docx import Document

    try:
        return Document(str(path))
    except Exception as e:  # noqa: BLE001
        raise ToolError(f"Could not open '{path.name}' as a Word document: {e}") from e


def _stats(doc) -> dict:
    words = 0
    headings = 0
    for p in doc.paragraphs:
        words += len(p.text.split())
        if p.style is not None and (p.style.name or "").startswith("Heading"):
            headings += 1
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                words += len(cell.text.split())
    return {"words": words, "paragraphs": len(doc.paragraphs), "tables": len(doc.tables), "headings": headings}


def _find_paragraph(doc, text: str, occurrence: int = 1):
    needle = text.lower().strip()
    seen = 0
    for p in O.iter_all_paragraphs(doc):
        if needle in p.text.lower():
            seen += 1
            if seen >= max(1, occurrence):
                return p
    return None


def _content_elements(doc) -> list:
    """Body children except the trailing section properties (new content is inserted before sectPr)."""
    return [el for el in doc.element.body if not el.tag.endswith("}sectPr")]


def _move_after(anchor_p, elements: list) -> None:
    ref = anchor_p._p
    for el in elements:
        ref.addnext(el)
        ref = el


# ----------------------------------------------------------------------------------------- tools
@documents.tool(risk=Risk.MEDIUM, scope="write", timeout=120)
async def create(ctx: ToolContext, title: str, markdown: str, path: str | None = None, subtitle: str | None = None,
                 author: str | None = None, title_page: bool = False,
                 page_size: Literal["letter", "a4"] = "letter", overwrite: bool = False) -> ToolResult:
    """Create a professionally formatted Word document (.docx) from markdown. Supports headings, paragraphs,
    bold/italic/strikethrough, inline code, links, bullet and numbered lists (nested), tables, code blocks,
    quotes, horizontal rules and workspace images (![caption](output/charts/x.png)).

    Args:
        title: Document title (shown at the top, or on a cover page).
        markdown: Body content in markdown. Don't repeat the title as a first-level heading.
        path: Output path, e.g. "output/q3-report.docx" (default: output/<title>.docx).
        subtitle: Optional subtitle under the title.
        author: Author shown on the cover page and in document properties.
        title_page: Put the title on its own cover page.
        page_size: Paper size.
        overwrite: Replace an existing file instead of saving under a new name.
    """
    target = O.ws_output(ctx, path, default_name=title or "document", ext=".docx", overwrite=overwrite)
    resolver = O.workspace_image_resolver(ctx)

    def build() -> dict:
        doc = O.markdown_to_docx(markdown or "", target, title=title, subtitle=subtitle, author=author,
                                 title_page=title_page, image_resolver=resolver, page_size=page_size)
        return _stats(doc)

    stats = await asyncio.to_thread(build)
    return await O.produced(ctx, target, stats, display=f"Created {O.vpath(ctx, target)}")


@documents.tool(risk=Risk.LOW, scope="read", timeout=60)
async def read(ctx: ToolContext, path: str, max_chars: int = 30000) -> dict:
    """Read a Word document as markdown (headings, lists, emphasis, links and tables preserved).

    Args:
        path: Workspace path of the .docx file.
        max_chars: Maximum characters of markdown to return.
    """
    p = O.ws_input(ctx, path, (".docx", ".md", ".txt"))
    if p.suffix.lower() != ".docx":
        text = O.read_text_file(p)
        body, truncated = O.clip(text, max(1000, int(max_chars or 30000)))
        return {"path": O.vpath(ctx, p), "markdown": body, "truncated": truncated}

    def run() -> tuple[str, dict, int]:
        doc = _open(p)
        md = O.docx_to_markdown(doc)
        comments = 0
        try:
            comments = len(list(doc.comments))
        except Exception:  # noqa: BLE001
            comments = 0
        return md, _stats(doc), comments

    md, stats, comments = await asyncio.to_thread(run)
    body, truncated = O.clip(md, max(1000, int(max_chars or 30000)))
    out: dict[str, Any] = {"path": O.vpath(ctx, p), **stats, "truncated": truncated, "markdown": body}
    if comments:
        out["comments"] = comments
    return out


@documents.tool(risk=Risk.MEDIUM, scope="write", timeout=120)
async def append(ctx: ToolContext, path: str, markdown: str, output: str | None = None) -> ToolResult:
    """Append markdown content (sections, lists, tables, images...) to the end of an existing Word document,
    keeping its styles.

    Args:
        path: Workspace path of the .docx file.
        markdown: Content to append, in markdown.
        output: Save to this path instead of modifying the original.
    """
    src = O.ws_input(ctx, path, (".docx",))
    target = O.ws_save_target(ctx, src, output, ext=".docx")
    resolver = O.workspace_image_resolver(ctx)

    def run() -> dict:
        doc = _open(src)
        O.markdown_to_docx(markdown or "", None, doc=doc, image_resolver=resolver)
        doc.save(str(target))
        return _stats(doc)

    stats = await asyncio.to_thread(run)
    return await O.produced(ctx, target, stats, display=f"Appended to {O.vpath(ctx, target)}")


@documents.tool(risk=Risk.MEDIUM, scope="write", timeout=120)
async def replace(ctx: ToolContext, path: str, find: str, replace: str, match_case: bool = True,
                  whole_word: bool = False, regex: bool = False, output: str | None = None) -> ToolResult:
    """Find and replace text everywhere in a Word document (body, tables, headers, footers), even when the text
    spans several formatting runs. Returns the number of replacements.

    Args:
        path: Workspace path of the .docx file.
        find: Text (or regular expression when regex=true) to find.
        replace: Replacement text (with regex, \\1 etc. refer to groups).
        match_case: Case-sensitive matching.
        whole_word: Only match whole words.
        regex: Treat `find` as a regular expression.
        output: Save to this path instead of modifying the original.
    """
    if not find:
        raise ToolError("`find` must not be empty.")
    src = O.ws_input(ctx, path, (".docx",))
    target = O.ws_save_target(ctx, src, output, ext=".docx")
    pattern_src = find if regex else re.escape(find)
    if whole_word:
        pattern_src = rf"(?<!\w)(?:{pattern_src})(?!\w)"
    try:
        pattern = re.compile(pattern_src, 0 if match_case else re.IGNORECASE)
    except re.error as e:
        raise ToolError(f"Invalid regular expression: {e}") from e
    repl = (lambda m: m.expand(replace)) if regex else (lambda m: replace)

    def run() -> int:
        doc = _open(src)
        count = sum(O.replace_in_paragraph(p, pattern, repl) for p in O.iter_all_paragraphs(doc))
        if count:
            doc.save(str(target))
        elif target != src:
            shutil.copyfile(src, target)
        return count

    count = await asyncio.to_thread(run)
    if count == 0:
        return ToolResult(content={"path": O.vpath(ctx, target), "replacements": 0,
                                   "hint": "No matches. Check spelling/case or use match_case=false."},
                          display="No matches found")
    return await O.produced(ctx, target, {"replacements": count},
                            display=f"{count} replacement(s) in {O.vpath(ctx, target)}")


@documents.tool(risk=Risk.MEDIUM, scope="write", timeout=120)
async def insert_table(ctx: ToolContext, path: str, rows: Any, header: bool = True, after: str | None = None,
                       caption: str | None = None, output: str | None = None) -> ToolResult:
    """Insert a styled table into a Word document — at the end, or right after the paragraph containing `after`.

    Args:
        path: Workspace path of the .docx file.
        rows: Table data: list of row lists (first row = header), list of objects, CSV text or a markdown table.
        header: Treat the first row as a header row.
        after: Insert after the first paragraph containing this text (default: end of document).
        caption: Optional caption shown above the table.
        output: Save to this path instead of modifying the original.
    """
    src = O.ws_input(ctx, path, (".docx",))
    target = O.ws_save_target(ctx, src, output, ext=".docx")
    hdr, body = O.table_from_any(rows, header=header)
    if not body and not hdr:
        raise ToolError("No table rows provided.")

    def run() -> dict:
        doc = _open(src)
        anchor = _find_paragraph(doc, after) if after else None
        if after and anchor is None:
            raise ToolError(f"Couldn't find a paragraph containing '{after}'.")
        writer = O.DocxWriter(doc)
        start = len(_content_elements(doc))
        if caption:
            cp = writer._para(doc, "Caption")
            cp.add_run(caption)
        writer.table_rows(hdr, body, doc)
        if anchor is not None:
            _move_after(anchor, _content_elements(doc)[start:])
        doc.save(str(target))
        return {"rows": len(body), "columns": max([len(hdr or [])] + [len(r) for r in body])}

    info = await asyncio.to_thread(run)
    return await O.produced(ctx, target, info, display=f"Inserted a {info['rows']}-row table")


@documents.tool(risk=Risk.MEDIUM, scope="write", timeout=120)
async def insert_image(ctx: ToolContext, path: str, image: str, width_inches: float | None = None,
                       caption: str | None = None, after: str | None = None,
                       output: str | None = None) -> ToolResult:
    """Insert a workspace image (PNG/JPG, e.g. a chart from analytics_visualize) into a Word document.

    Args:
        path: Workspace path of the .docx file.
        image: Workspace path of the image.
        width_inches: Display width (default: natural size, capped to the page width).
        caption: Optional caption under the image.
        after: Insert after the first paragraph containing this text (default: end of document).
        output: Save to this path instead of modifying the original.
    """
    src = O.ws_input(ctx, path, (".docx",))
    img = O.ws_input(ctx, image, (".png", ".jpg", ".jpeg", ".gif", ".bmp"))
    target = O.ws_save_target(ctx, src, output, ext=".docx")

    def run() -> None:
        doc = _open(src)
        anchor = _find_paragraph(doc, after) if after else None
        if after and anchor is None:
            raise ToolError(f"Couldn't find a paragraph containing '{after}'.")
        writer = O.DocxWriter(doc, image_resolver=lambda s: img)
        start = len(_content_elements(doc))
        writer.image(str(img), caption or "", doc, width_in=width_inches)
        if anchor is not None:
            _move_after(anchor, _content_elements(doc)[start:])
        doc.save(str(target))

    await asyncio.to_thread(run)
    return await O.produced(ctx, target, {"image": O.vpath(ctx, img)}, display=f"Inserted {img.name}")


@documents.tool(risk=Risk.MEDIUM, scope="write", timeout=120)
async def add_comment(ctx: ToolContext, path: str, anchor: str, comment: str, author: str = "Ren",
                      occurrence: int = 1, output: str | None = None) -> ToolResult:
    """Add a review comment (Word margin comment) anchored to specific text in a document.

    Args:
        path: Workspace path of the .docx file.
        anchor: Exact text the comment refers to (must appear within one paragraph).
        comment: The comment text.
        author: Comment author name.
        occurrence: Which occurrence of the anchor text to comment on (1 = first).
        output: Save to this path instead of modifying the original.
    """
    if not anchor or not comment:
        raise ToolError("Both `anchor` and `comment` are required.")
    src = O.ws_input(ctx, path, (".docx",))
    target = O.ws_save_target(ctx, src, output, ext=".docx")

    def run() -> str:
        from docx.enum.text import WD_COLOR_INDEX

        doc = _open(src)
        needle = anchor.lower()
        seen = 0
        for p in O.iter_all_paragraphs(doc):
            text = "".join(r.text for r in O.paragraph_runs(p))
            start = 0
            while True:
                idx = text.lower().find(needle, start)
                if idx < 0:
                    break
                seen += 1
                if seen == max(1, occurrence):
                    runs = O.isolate_runs(p, idx, idx + len(anchor))
                    if not runs:
                        raise ToolError("Could not anchor the comment to that text.")
                    initials = "".join(w[0] for w in author.split()[:3]).upper() or "R"
                    if hasattr(doc, "add_comment"):
                        try:
                            doc.add_comment(runs, text=comment, author=author, initials=initials)
                            doc.save(str(target))
                            return "comment"
                        except Exception:  # noqa: BLE001
                            pass
                    note = runs[-1]._r
                    from docx.text.run import Run

                    new_r = copy.deepcopy(note)
                    note.addnext(new_r)
                    nr = Run(new_r, runs[-1]._parent)
                    nr.text = f" [Comment ({author}): {comment}]"
                    nr.italic = True
                    nr.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    doc.save(str(target))
                    return "inline-note"
                start = idx + 1
        raise ToolError(f"Text '{anchor}' not found{'' if occurrence <= 1 else f' (occurrence {occurrence})'}.")

    kind = await asyncio.to_thread(run)
    return await O.produced(ctx, target, {"anchor": anchor, "type": kind},
                            display=f"Added comment on “{anchor[:40]}”")


def _doc_lines(path: Path) -> list[str]:
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = _open(path)
    lines: list[str] = []
    for block in doc.iter_inner_content():
        if isinstance(block, Paragraph):
            t = " ".join(block.text.split())
            if t:
                lines.append(t)
        elif isinstance(block, Table):
            for row in block.rows:
                cells = [" ".join(c.text.split()) for c in row.cells]
                if any(cells):
                    lines.append("| " + " | ".join(cells) + " |")
    return lines


def _word_diff(a: str, b: str) -> str:
    aw, bw = a.split(), b.split()
    out = []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, aw, bw, autojunk=False).get_opcodes():
        if op == "equal":
            out.append(" ".join(aw[i1:i2]))
        elif op == "delete":
            out.append("~~" + " ".join(aw[i1:i2]) + "~~")
        elif op == "insert":
            out.append("**" + " ".join(bw[j1:j2]) + "**")
        else:
            out.append("~~" + " ".join(aw[i1:i2]) + "~~ **" + " ".join(bw[j1:j2]) + "**")
    return " ".join(out)


def compare_docs(a: Path, b: Path) -> tuple[str, dict]:
    la, lb = _doc_lines(a), _doc_lines(b)
    sm = difflib.SequenceMatcher(None, la, lb, autojunk=False)
    added = removed = changed = 0
    sections: list[str] = []
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == "equal":
            continue
        if op == "replace":
            pairs = min(i2 - i1, j2 - j1)
            for k in range(pairs):
                changed += 1
                sections.append(f"### Changed (¶{i1 + k + 1} → ¶{j1 + k + 1})\n{_word_diff(la[i1 + k], lb[j1 + k])}")
            for k in range(i1 + pairs, i2):
                removed += 1
                sections.append(f"### Removed (¶{k + 1})\n> ~~{la[k]}~~")
            for k in range(j1 + pairs, j2):
                added += 1
                sections.append(f"### Added (¶{k + 1})\n> **{lb[k]}**")
        elif op == "delete":
            for k in range(i1, i2):
                removed += 1
                sections.append(f"### Removed (¶{k + 1})\n> ~~{la[k]}~~")
        elif op == "insert":
            for k in range(j1, j2):
                added += 1
                sections.append(f"### Added (¶{k + 1})\n> **{lb[k]}**")
    unchanged = sum(i2 - i1 for op, i1, i2, _, _ in sm.get_opcodes() if op == "equal")
    stats = {"added": added, "removed": removed, "changed": changed, "unchanged": unchanged,
             "similarity": round(sm.ratio() * 100, 1)}
    head = (f"# Comparison: {a.name} → {b.name}\n\n**Summary:** {added} added, {removed} removed, {changed} changed, "
            f"{unchanged} unchanged paragraphs · similarity {stats['similarity']}%.\n\n"
            "Legend: ~~removed~~ **added**\n")
    body = "\n\n".join(sections) if sections else "_The documents have identical text._"
    return head + "\n## Changes\n\n" + body + "\n", stats


@documents.tool(risk=Risk.LOW, scope="read", timeout=120,
                risk_fn=lambda a: Risk.MEDIUM if a.get("save_as") else Risk.LOW)
async def compare(ctx: ToolContext, path_a: str, path_b: str, save_as: str | None = None,
                  max_chars: int = 20000) -> dict | ToolResult:
    """Compare two Word documents paragraph by paragraph and report additions, deletions and word-level changes
    as markdown (~~removed~~ **added**).

    Args:
        path_a: The original / older version (.docx).
        path_b: The revised / newer version (.docx).
        save_as: Optionally save the full report as markdown (.md) or Word (.docx).
        max_chars: Maximum characters of the report to return.
    """
    a = O.ws_input(ctx, path_a, (".docx",))
    b = O.ws_input(ctx, path_b, (".docx",))
    report, stats = await asyncio.to_thread(compare_docs, a, b)
    text, truncated = O.clip(report, max(1000, int(max_chars or 20000)))
    out: dict[str, Any] = {**stats, "truncated": truncated, "report": text}
    if save_as:
        ext = ".docx" if str(save_as).lower().endswith(".docx") else ".md"
        target = O.ws_output(ctx, save_as, default_name=f"{a.stem}-vs-{b.stem}", ext=ext)
        if ext == ".docx":
            await asyncio.to_thread(O.markdown_to_docx, report, target)
        else:
            target.write_text(report, encoding="utf-8")
        return await O.produced(ctx, target, out, display=f"Compared {a.name} and {b.name}")
    return out


def _libreoffice_pdf(src: Path, target: Path) -> bool:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return False
    with tempfile.TemporaryDirectory(prefix="ren-lo-") as tmp:
        profile = Path(tmp, "profile").as_uri()
        try:
            subprocess.run([exe, f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf",
                            "--outdir", tmp, str(src)], capture_output=True, timeout=120, check=False)
        except (subprocess.TimeoutExpired, OSError):
            return False
        produced = Path(tmp, src.stem + ".pdf")
        if produced.exists() and produced.stat().st_size > 0:
            shutil.move(str(produced), str(target))
            return True
    return False


@documents.tool(risk=Risk.MEDIUM, scope="write", timeout=180)
async def export_pdf(ctx: ToolContext, path: str, output: str | None = None,
                     page_size: Literal["auto", "letter", "a4"] = "auto",
                     engine: Literal["auto", "builtin", "libreoffice"] = "auto") -> ToolResult:
    """Export a Word document (or a markdown file) to PDF. Uses LibreOffice for a pixel-faithful conversion when
    it is installed, otherwise Ren's built-in renderer (text, headings, lists, tables, code).

    Args:
        path: Workspace path of the .docx or .md file.
        output: Output PDF path (default: next to the source in output/).
        page_size: Paper size for the built-in renderer ("auto" follows the document).
        engine: Conversion engine.
    """
    src = O.ws_input(ctx, path, (".docx", ".md", ".markdown", ".txt"))
    target = O.ws_output(ctx, output, default_name=src.stem, ext=".pdf", overwrite=bool(output))
    resolver = O.workspace_image_resolver(ctx)

    def run() -> str:
        if src.suffix.lower() == ".docx" and engine in ("auto", "libreoffice"):
            if _libreoffice_pdf(src, target):
                return "libreoffice"
            if engine == "libreoffice":
                raise ToolError("LibreOffice (soffice) is not installed on this server; use engine='builtin'.")
        if src.suffix.lower() == ".docx":
            doc = _open(src)
            md = O.docx_to_markdown(doc)
            size = page_size
            if size == "auto":
                w = doc.sections[0].page_width
                size = "letter" if w and abs(w.inches - 8.5) < 0.2 else "a4"
            title = (doc.core_properties.title or "").strip() or None
            if title and md.lstrip().startswith("# "):
                title = None  # the markdown already starts with the title heading
            O.markdown_to_pdf(md, target, title=title, page_size=size, image_resolver=resolver)
        else:
            md = O.read_text_file(src)
            O.markdown_to_pdf(md, target, page_size="a4" if page_size == "auto" else page_size,
                              image_resolver=resolver)
        return "builtin"

    used = await asyncio.to_thread(run)
    pages = len(O.pdf_reader(target).pages)
    return await O.produced(ctx, target, {"pages": pages, "engine": used}, display=f"Exported {target.name}")


@documents.tool(risk=Risk.MEDIUM, scope="write", timeout=120)
async def apply_style(ctx: ToolContext, path: str, preset: Literal["ren", "none"] = "none",
                      heading_font: str | None = None, heading_color: str | None = None,
                      heading_size: float | None = None, body_font: str | None = None,
                      body_size: float | None = None, line_spacing: float | None = None,
                      output: str | None = None) -> ToolResult:
    """Restyle a Word document: heading font/size/colour, body font/size and line spacing (applied to the
    document's styles, so all matching paragraphs update). preset="ren" applies Ren's house style first.

    Args:
        path: Workspace path of the .docx file.
        preset: "ren" = f69 Labs house style (deep-green headings, green accents); "none" = only apply the overrides given.
        heading_font: Font for the title and headings, e.g. "Georgia".
        heading_color: Hex colour for headings, e.g. "#1F3864".
        heading_size: Point size of Heading 1 (lower levels scale down proportionally).
        body_font: Font for body text, e.g. "Calibri".
        body_size: Body text size in points.
        line_spacing: Line spacing multiple, e.g. 1.15.
        output: Save to this path instead of modifying the original.
    """
    src = O.ws_input(ctx, path, (".docx",))
    target = O.ws_save_target(ctx, src, output, ext=".docx")
    for c in (heading_color,):
        if c and not re.fullmatch(r"#?[0-9a-fA-F]{6}", c.strip()):
            raise ToolError(f"Invalid colour '{c}'. Use a hex value like #1F3864.")

    def run() -> list[str]:
        doc = _open(src)
        changed = []
        if preset == "ren":
            O.setup_docx_styles(doc)
            changed.append("Ren house style")
        normal = O.get_style(doc, "Normal")
        if normal is not None and (body_font or body_size or line_spacing):
            O.style_font(normal, body_font, body_size, None)
            if line_spacing:
                normal.paragraph_format.line_spacing = float(line_spacing)
            changed.append("body")
        if heading_font or heading_color or heading_size:
            color = heading_color.strip().lstrip("#") if heading_color else None
            scale = {1: 1.0, 2: 0.78, 3: 0.64, 4: 0.58, 5: 0.55, 6: 0.52}
            for name in ["Title"] + [f"Heading {i}" for i in range(1, 7)]:
                st = O.get_style(doc, name)
                if st is None:
                    continue
                size = None
                if heading_size:
                    lvl = int(name[-1]) if name[-1].isdigit() else 0
                    size = round(float(heading_size) * (1.3 if lvl == 0 else scale.get(lvl, 0.55)), 1)
                O.style_font(st, heading_font, size, color)
            changed.append("headings")
        if not changed:
            raise ToolError("Nothing to change — pass preset='ren' or at least one style option.")
        doc.save(str(target))
        return changed

    changed = await asyncio.to_thread(run)
    return await O.produced(ctx, target, {"updated": changed}, display=f"Restyled {target.name}")
