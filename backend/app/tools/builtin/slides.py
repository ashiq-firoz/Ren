"""Slides group: PowerPoint decks (.pptx) in the f69 Labs theme — create from an outline, read as markdown,
add/delete/duplicate/reorder slides, add text, images, tables and native (editable) charts, and switch between the
beige (light) and deep-green (dark) themes."""
from __future__ import annotations

import asyncio
import copy
from pathlib import Path
from typing import Any, Literal

from app.tools.base import Risk, ToolContext, ToolError, ToolGroup, ToolResult
from app.tools.builtin import _office as O

slides = ToolGroup(
    "slides", "Presentations",
    "Build and edit PowerPoint decks (.pptx) in the clean f69 Labs theme: create a deck from an outline, read it "
    "as markdown, add, delete, duplicate and reorder slides, add text, images, tables and native charts, and "
    "switch between light and dark themes.",
    prefix="ppt", icon="presentation", category="productivity",
)

Theme = Literal["light", "dark"]


def _open(path: Path):
    from pptx import Presentation

    try:
        return Presentation(str(path))
    except Exception as e:  # noqa: BLE001
        raise ToolError(f"Could not open '{path.name}' as a PowerPoint deck: {e}") from e


def _slide(prs, number: int):
    n = len(prs.slides)
    if not isinstance(number, int) or number < 1 or number > n:
        raise ToolError(f"Slide {number} doesn't exist (the deck has {n} slide{'s' if n != 1 else ''}).")
    return prs.slides[number - 1]


def _emu(inches: float | None, default: int) -> int:
    from pptx.util import Inches

    return int(Inches(float(inches))) if inches is not None else int(default)


def _is_decoration(shape) -> bool:
    return shape.name in ("Ren accent", "Ren slide number")


def _content_box(prs, slide, *, split: bool = True) -> tuple[int, int, int, int]:
    """Free content area for a new visual; if the slide already has body text, move it left and use the right."""
    g = O._Geo(prs)
    title = slide.shapes.title
    texts = []
    for shp in slide.shapes:
        if (title is not None and shp.shape_id == title.shape_id) or _is_decoration(shp):
            continue
        if shp.has_text_frame and shp.text_frame.text.strip() and (shp.top or 0) >= g.body_top - g.H * 0.05:
            texts.append(shp)
    if texts and split:
        text_w = int(g.inner_w * 0.42)
        for shp in texts:
            if (shp.width or 0) > text_w:
                shp.left, shp.width = g.mx, text_w
        gap = int(g.W * 0.03)
        return g.mx + text_w + gap, g.body_top, g.inner_w - text_w - gap, g.body_h
    return g.mx, g.body_top, g.inner_w, g.body_h


def _box(prs, slide, left, top, width, height) -> tuple[int, int, int, int]:
    x, y, w, h = _content_box(prs, slide, split=left is None and top is None)
    return _emu(left, x), _emu(top, y), _emu(width, w), _emu(height, h)


def duplicate_slide(prs, index: int):
    """Copy slide `index` (0-based) — shapes, pictures, charts (deep-copied), background and notes — right after it."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    src = prs.slides[index]
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    bg = src._element.cSld.bg
    if bg is not None:
        new._element.cSld.insert(0, copy.deepcopy(bg))
    rid_map: dict[str, str] = {}
    for rid, rel in list(src.part.rels.items()):
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        if rel.is_external:
            rid_map[rid] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif rel.reltype == RT.CHART:
            rid_map[rid] = new.part.relate_to(_clone_chart_part(rel.target_part), rel.reltype)
        else:
            rid_map[rid] = new.part.relate_to(rel.target_part, rel.reltype)
    r_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    for shp in src.shapes:
        el = copy.deepcopy(shp._element)
        for node in el.iter():
            for attr in list(node.attrib):
                if attr.startswith(r_ns) and node.attrib[attr] in rid_map:
                    node.attrib[attr] = rid_map[node.attrib[attr]]
        new.shapes._spTree.insert_element_before(el, "p:extLst")
    if src.has_notes_slide:
        text = src.notes_slide.notes_text_frame.text if src.notes_slide.notes_text_frame else ""
        if text:
            new.notes_slide.notes_text_frame.text = text
    O.move_slide(prs, len(prs.slides) - 1, index + 1)
    return new


def _clone_chart_part(chart_part):
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.opc.package import Part
    from pptx.parts.chart import ChartPart

    package = chart_part.package
    new_el = copy.deepcopy(chart_part._element)
    new_part = ChartPart(package.next_partname(ChartPart.partname_template), chart_part.content_type, package, new_el)
    r_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    rid_map = {}
    for rid, rel in list(chart_part.rels.items()):
        if rel.is_external:
            rid_map[rid] = new_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif rel.reltype == RT.PACKAGE:
            tgt = rel.target_part
            tmpl = "/ppt/embeddings/Microsoft_Excel_Sheet%d.xlsx"
            blob_part = Part(package.next_partname(tmpl), tgt.content_type, package, tgt.blob)
            rid_map[rid] = new_part.relate_to(blob_part, rel.reltype)
        else:
            rid_map[rid] = new_part.relate_to(rel.target_part, rel.reltype)
    for node in new_part._element.iter():
        for attr in list(node.attrib):
            if attr.startswith(r_ns) and node.attrib[attr] in rid_map:
                node.attrib[attr] = rid_map[node.attrib[attr]]
    return new_part


async def _save(ctx: ToolContext, prs, src: Path, output: str | None, summary: dict, display: str) -> ToolResult:
    target = O.ws_save_target(ctx, src, output, ext=".pptx")
    await asyncio.to_thread(prs.save, str(target))
    summary = {"slides": len(prs.slides), **summary}
    return await O.produced(ctx, target, summary, display=display)


# ----------------------------------------------------------------------------------------- tools
@slides.tool(risk=Risk.MEDIUM, scope="write", timeout=180)
async def create(ctx: ToolContext, title: str, slides: list[dict[str, Any]], subtitle: str | None = None,
                 theme: Theme = "light", path: str | None = None, overwrite: bool = False) -> ToolResult:
    """Create a 16:9 PowerPoint deck from an outline: a title slide plus one slide per entry. Each slide entry:
    {"title": "...", "bullets": ["point", ["sub-point"], "**bold** text"], "notes": "speaker notes"} and optionally
    "table": [[header...], [row...]], "chart": {"type": "column|bar|line|pie|doughnut|area", "categories": [...],
    "series": {"Revenue": [...]}}, or "image": "output/charts/x.png" (shown beside the bullets, or full width).
    An entry with only a title becomes a section divider.

    Args:
        title: Deck title (title slide).
        slides: Slide entries as described above.
        subtitle: Subtitle on the title slide (e.g. date, audience, author).
        theme: "light" (beige background, deep-green text) or "dark" (deep-green background, ivory text).
        path: Output path, e.g. "output/board-update.pptx" (default: output/<title>.pptx).
        overwrite: Replace an existing file instead of saving under a new name.
    """
    if not isinstance(slides, list):
        raise ToolError("`slides` must be a list of slide objects.")
    target = O.ws_output(ctx, path, default_name=title or "deck", ext=".pptx", overwrite=overwrite)
    resolver = O.workspace_image_resolver(ctx)

    def build() -> int:
        prs = O.build_deck(target, title, slides, subtitle=subtitle, theme=theme, image_resolver=resolver)
        return len(prs.slides)

    n = await asyncio.to_thread(build)
    return await O.produced(ctx, target, {"slides": n, "theme": theme}, display=f"Created a {n}-slide deck")


@slides.tool(risk=Risk.LOW, scope="read", timeout=60)
async def read(ctx: ToolContext, path: str, include_notes: bool = True, max_chars: int = 30000) -> dict:
    """Read a PowerPoint deck as markdown: each slide's title, bullets (with levels), tables, chart summaries,
    images and speaker notes.

    Args:
        path: Workspace path of the .pptx file.
        include_notes: Include speaker notes.
        max_chars: Maximum characters to return.
    """
    p = O.ws_input(ctx, path, O.DECK_EXTS)

    def run() -> tuple[str, int]:
        prs = _open(p)
        return O.pptx_to_markdown(prs, include_notes=include_notes), len(prs.slides)

    md, n = await asyncio.to_thread(run)
    body, truncated = O.clip(md, max(1000, int(max_chars or 30000)))
    return {"path": O.vpath(ctx, p), "slides": n, "truncated": truncated, "markdown": body}


@slides.tool(risk=Risk.MEDIUM, scope="write", timeout=120)
async def add_slide(ctx: ToolContext, path: str, title: str, bullets: list[Any] | None = None,
                    notes: str | None = None, position: int | None = None, table: list[list[Any]] | None = None,
                    chart: dict[str, Any] | None = None, image: str | None = None,
                    output: str | None = None) -> ToolResult:
    """Add a slide in the deck's style (title + bullets, optionally a table, chart or image).

    Args:
        path: Workspace path of the .pptx file.
        title: Slide title.
        bullets: Bullet points (nested lists = sub-bullets; **bold** supported).
        notes: Speaker notes.
        position: 1-based position for the new slide (default: end).
        table: Optional table rows (first row = header).
        chart: Optional chart {"type", "categories", "series", "title"}.
        image: Optional workspace image path.
        output: Save to this path instead of modifying the original.
    """
    p = O.ws_input(ctx, path, O.DECK_EXTS)
    resolver = O.workspace_image_resolver(ctx)

    def run():
        prs = _open(p)
        spec = {"title": title, "bullets": bullets, "notes": notes, "table": table, "chart": chart, "image": image}
        O.add_content_slide(prs, spec, image_resolver=resolver)
        n = len(prs.slides)
        pos = n
        if position is not None:
            pos = max(1, min(int(position), n))
            O.move_slide(prs, n - 1, pos - 1)
        return prs, pos

    prs, pos = await asyncio.to_thread(run)
    return await _save(ctx, prs, p, output, {"added_at": pos}, f"Added slide {pos}: {title[:50]}")


@slides.tool(risk=Risk.MEDIUM, scope="write", timeout=60)
async def delete_slide(ctx: ToolContext, path: str, slide: int, output: str | None = None) -> ToolResult:
    """Delete a slide.

    Args:
        path: Workspace path of the .pptx file.
        slide: 1-based slide number.
        output: Save to this path instead of modifying the original.
    """
    p = O.ws_input(ctx, path, O.DECK_EXTS)

    def run():
        prs = _open(p)
        s = _slide(prs, slide)
        title = s.shapes.title.text_frame.text if s.shapes.title is not None else ""
        if len(prs.slides) == 1:
            raise ToolError("Can't delete the only slide in the deck.")
        O.delete_slide(prs, slide - 1)
        return prs, title

    prs, title = await asyncio.to_thread(run)
    return await _save(ctx, prs, p, output, {"deleted": slide, "title": title}, f"Deleted slide {slide}")


@slides.tool(risk=Risk.MEDIUM, scope="write", timeout=60)
async def duplicate_slide_(ctx: ToolContext, path: str, slide: int, position: int | None = None,
                           output: str | None = None) -> ToolResult:
    """Duplicate a slide (inserted right after it, or at `position`).

    Args:
        path: Workspace path of the .pptx file.
        slide: 1-based number of the slide to copy.
        position: Optional 1-based position for the copy.
        output: Save to this path instead of modifying the original.
    """
    p = O.ws_input(ctx, path, O.DECK_EXTS)

    def run():
        prs = _open(p)
        _slide(prs, slide)
        duplicate_slide(prs, slide - 1)
        pos = slide + 1
        if position is not None:
            target = max(1, min(int(position), len(prs.slides)))
            O.move_slide(prs, slide, target - 1)
            pos = target
        return prs, pos

    prs, pos = await asyncio.to_thread(run)
    return await _save(ctx, prs, p, output, {"copy_at": pos}, f"Duplicated slide {slide}")


@slides.tool(risk=Risk.MEDIUM, scope="write", timeout=60)
async def reorder(ctx: ToolContext, path: str, order: list[int] | None = None, slide: int | None = None,
                  to: int | None = None, output: str | None = None) -> ToolResult:
    """Reorder slides: give the full new order (e.g. [1, 3, 2, 4]) or move one slide (`slide` → `to`).

    Args:
        path: Workspace path of the .pptx file.
        order: New order as a permutation of the current 1-based slide numbers.
        slide: Slide to move (with `to`).
        to: New 1-based position for `slide`.
        output: Save to this path instead of modifying the original.
    """
    p = O.ws_input(ctx, path, O.DECK_EXTS)

    def run():
        prs = _open(p)
        n = len(prs.slides)
        if order:
            if sorted(order) != list(range(1, n + 1)):
                raise ToolError(f"`order` must list every slide number 1–{n} exactly once.")
            lst = prs.slides._sldIdLst
            items = list(lst)
            for el in items:
                lst.remove(el)
            for k in order:
                lst.append(items[k - 1])
        elif slide is not None and to is not None:
            _slide(prs, slide)
            O.move_slide(prs, slide - 1, max(1, min(int(to), n)) - 1)
        else:
            raise ToolError("Provide `order`, or `slide` and `to`.")
        return prs

    prs = await asyncio.to_thread(run)
    titles = [(s.shapes.title.text_frame.text if s.shapes.title is not None else "")[:40] for s in prs.slides]
    return await _save(ctx, prs, p, output, {"order": titles}, "Reordered slides")


@slides.tool(risk=Risk.MEDIUM, scope="write", timeout=60)
async def add_text(ctx: ToolContext, path: str, slide: int, text: str, left: float | None = None,
                   top: float | None = None, width: float | None = None, height: float | None = None,
                   font_size: float = 18, bold: bool = False, color: str | None = None,
                   align: Literal["left", "center", "right"] = "left", bullets: bool = False,
                   output: str | None = None) -> ToolResult:
    """Add a text box to a slide (positions in inches; default: the free content area). Each line becomes a
    paragraph; with bullets=true lines become bullet points.

    Args:
        path: Workspace path of the .pptx file.
        slide: 1-based slide number.
        text: Text (newlines separate paragraphs; **bold** supported).
        left: Left offset in inches.
        top: Top offset in inches.
        width: Width in inches.
        height: Height in inches.
        font_size: Font size in points.
        bold: Bold text.
        color: Hex colour, e.g. "#1F5A45" (default: theme text colour).
        align: Paragraph alignment.
        bullets: Format lines as bullet points.
        output: Save to this path instead of modifying the original.
    """
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt

    p = O.ws_input(ctx, path, O.DECK_EXTS)

    def run():
        prs = _open(p)
        s = _slide(prs, slide)
        x, y, w, h = _box(prs, s, left, top, width, height)
        box = s.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        if bullets:
            O.fill_bullets(tf, text, explicit=True)
            for para in tf.paragraphs:
                for r in para.runs:
                    r.font.size = Pt(font_size)
        else:
            lines = text.split("\n")
            for i, line in enumerate(lines):
                para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
                for sp in O.inline_spans(line):
                    if sp.br or not sp.text:
                        continue
                    r = para.add_run()
                    r.text = sp.text
                    r.font.size = Pt(font_size)
                    r.font.bold = bold or sp.bold or None
                    if sp.italic:
                        r.font.italic = True
                    if color:
                        try:
                            r.font.color.rgb = RGBColor.from_string(color.strip().lstrip("#").upper())
                        except ValueError as e:
                            raise ToolError(f"Invalid colour '{color}'.") from e
        return prs

    prs = await asyncio.to_thread(run)
    return await _save(ctx, prs, p, output, {"slide": slide}, f"Added text to slide {slide}")


@slides.tool(risk=Risk.MEDIUM, scope="write", timeout=60)
async def add_image(ctx: ToolContext, path: str, slide: int, image: str, left: float | None = None,
                    top: float | None = None, width: float | None = None, height: float | None = None,
                    output: str | None = None) -> ToolResult:
    """Add a workspace image (PNG/JPG — e.g. a chart from analytics_visualize) to a slide, fitted into the free
    content area unless a position is given (inches).

    Args:
        path: Workspace path of the .pptx file.
        slide: 1-based slide number.
        image: Workspace path of the image.
        left: Left offset in inches.
        top: Top offset in inches.
        width: Box width in inches (image keeps its aspect ratio).
        height: Box height in inches.
        output: Save to this path instead of modifying the original.
    """
    p = O.ws_input(ctx, path, O.DECK_EXTS)
    img = O.ws_input(ctx, image, (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"))

    def run():
        prs = _open(p)
        s = _slide(prs, slide)
        x, y, w, h = _box(prs, s, left, top, width, height)
        O.add_picture_fit(s, img, x, y, w, h)
        return prs

    prs = await asyncio.to_thread(run)
    return await _save(ctx, prs, p, output, {"slide": slide, "image": O.vpath(ctx, img)},
                       f"Added {img.name} to slide {slide}")


@slides.tool(risk=Risk.MEDIUM, scope="write", timeout=60)
async def add_table(ctx: ToolContext, path: str, slide: int, rows: Any, header: bool = True,
                    left: float | None = None, top: float | None = None, width: float | None = None,
                    height: float | None = None, output: str | None = None) -> ToolResult:
    """Add a themed table to a slide (header row in the theme text colour, banded body rows, numbers right-aligned; long tables are
    truncated to fit).

    Args:
        path: Workspace path of the .pptx file.
        slide: 1-based slide number.
        rows: Table rows (first row = header), list of objects, CSV text or a markdown table.
        header: Style the first row as a header.
        left: Left offset in inches.
        top: Top offset in inches.
        width: Width in inches.
        height: Maximum height in inches.
        output: Save to this path instead of modifying the original.
    """
    p = O.ws_input(ctx, path, O.DECK_EXTS)
    hdr, body = O.table_from_any(rows, header=header)
    data = ([hdr] if hdr else []) + body
    if not data:
        raise ToolError("No table rows provided.")

    def run():
        prs = _open(p)
        s = _slide(prs, slide)
        x, y, w, h = _box(prs, s, left, top, width, height)
        _, truncated = O.add_table_shape(s, data, x, y, w, h, header=bool(hdr) or header)
        return prs, truncated

    prs, truncated = await asyncio.to_thread(run)
    summary: dict[str, Any] = {"slide": slide, "rows": len(body)}
    if truncated:
        summary["note"] = "Table truncated to fit the slide; consider splitting it across slides."
    return await _save(ctx, prs, p, output, summary, f"Added a table to slide {slide}")


@slides.tool(risk=Risk.MEDIUM, scope="write", timeout=60)
async def add_chart(ctx: ToolContext, path: str, slide: int, chart_type: str, categories: list[Any],
                    series: Any, title: str | None = None, number_format: str | None = None,
                    left: float | None = None, top: float | None = None, width: float | None = None,
                    height: float | None = None, output: str | None = None) -> ToolResult:
    """Add a native, editable PowerPoint chart (theme colours) to a slide.

    Args:
        path: Workspace path of the .pptx file.
        slide: 1-based slide number.
        chart_type: column, bar, line, pie, doughnut, area, stacked_column or stacked_bar.
        categories: Category labels (x-axis), e.g. ["Q1", "Q2", "Q3"].
        series: {"Revenue": [10, 12, 15], "Cost": [...]} or [{"name": ..., "values": [...]}].
        title: Optional chart title.
        number_format: Excel number format for values, e.g. "$#,##0" or "0%".
        left: Left offset in inches.
        top: Top offset in inches.
        width: Width in inches.
        height: Height in inches.
        output: Save to this path instead of modifying the original.
    """
    p = O.ws_input(ctx, path, O.DECK_EXTS)

    def run():
        prs = _open(p)
        s = _slide(prs, slide)
        x, y, w, h = _box(prs, s, left, top, width, height)
        O.add_chart_shape(s, chart_type, categories, series, x, y, w, h, title=title, number_format=number_format)
        return prs

    prs = await asyncio.to_thread(run)
    return await _save(ctx, prs, p, output, {"slide": slide, "chart_type": chart_type},
                       f"Added a {chart_type} chart to slide {slide}")


@slides.tool(risk=Risk.MEDIUM, scope="write", timeout=60)
async def apply_theme(ctx: ToolContext, path: str, theme: Theme = "light", reset_backgrounds: bool = True,
                      output: str | None = None) -> ToolResult:
    """Switch a deck to the f69 Labs beige theme (#EFE7DA background, #0F2D24 text, #1F5A45 accent) or the dark
    deep-green theme (#0F2D24 background, ivory text, mint accent). Theme colours drive backgrounds, text, tables and charts.

    Args:
        path: Workspace path of the .pptx file.
        theme: "light" or "dark".
        reset_backgrounds: Remove per-slide background overrides so every slide follows the theme.
        output: Save to this path instead of modifying the original.
    """
    p = O.ws_input(ctx, path, O.DECK_EXTS)

    def run():
        prs = _open(p)
        key = O.apply_deck_theme(prs, theme, reset_slide_backgrounds=reset_backgrounds)
        return prs, key

    prs, key = await asyncio.to_thread(run)
    return await _save(ctx, prs, p, output, {"theme": key}, f"Applied the {key} theme")
