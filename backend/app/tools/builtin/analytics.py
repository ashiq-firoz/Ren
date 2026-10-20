"""Analytics group: analyse tabular files (CSV/TSV, Excel, Parquet, JSON) — profiling, read-only SQL with DuckDB,
group-by aggregation, joins, two-way reconciliation (finance), correlation, outliers, forecasting (linear / Holt /
Holt-Winters, pure numpy) and clean matplotlib charts in the Ren palette."""
from __future__ import annotations

import asyncio
import csv
import math
import re
import warnings
from datetime import datetime
from pathlib import Path
from typing import Any, Literal

from app.tools.base import Risk, ToolContext, ToolError, ToolGroup, ToolResult
from app.tools.builtin import _office as O

analytics = ToolGroup(
    "analytics", "Data Analysis",
    "Analyse data files (CSV, Excel, Parquet, JSON): profile datasets, run read-only SQL with DuckDB, aggregate, "
    "join, reconcile two datasets (matched / unmatched / differences), find correlations and outliers, forecast "
    "time series and create presentation-ready charts.",
    prefix="analytics", icon="chart-line", category="data",
)

DATA_EXTS = (".csv", ".tsv", ".txt", ".xlsx", ".xlsm", ".parquet", ".pq", ".json", ".jsonl", ".ndjson")
# f69 Labs chart style: ivory canvas, deep-green ink, brand-green bars, bronze for forecasts (palette shared with Office charts)
IVORY, INK, ACCENT, BRONZE = "#F8F6F1", "#0F2D24", "#1F5A45", "#A8875A"
PALETTE = [f"#{c}" for c in O.CHART_PALETTE]
GRID, SOFT, MUTED = "#E3D8C6", "#2F4A40", f"#{O.MUTED}"


# ----------------------------------------------------------------------------------------- loading
def load_df(p: Path, sheet: str | None = None):
    import pandas as pd

    ext = p.suffix.lower()
    try:
        if ext in (".csv", ".tsv", ".txt"):
            raw = p.read_bytes()[:65536]
            enc = "utf-8-sig"
            try:
                sample = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                enc, sample = "cp1252", raw.decode("cp1252", errors="replace")
            if ext == ".tsv":
                sep = "\t"
            else:
                try:
                    sep = csv.Sniffer().sniff(sample[:8192], delimiters=",;\t|").delimiter
                except csv.Error:
                    sep = ","
            return pd.read_csv(p, sep=sep, encoding=enc, low_memory=False)
        if ext in (".xlsx", ".xlsm"):
            try:
                return pd.read_excel(p, sheet_name=sheet if sheet else 0, engine="openpyxl")
            except ValueError as e:
                from openpyxl import load_workbook

                wb = load_workbook(str(p), read_only=True)
                names = wb.sheetnames
                wb.close()
                raise ToolError(f"Sheet '{sheet}' not found. Sheets: {', '.join(names)}.") from e
        if ext in (".parquet", ".pq"):
            try:
                return pd.read_parquet(p)
            except ImportError:
                import duckdb

                con = duckdb.connect()
                try:
                    return con.execute("SELECT * FROM read_parquet(?)", [str(p)]).df()
                finally:
                    con.close()
        if ext in (".jsonl", ".ndjson"):
            return pd.read_json(p, lines=True)
        if ext == ".json":
            import json

            data = json.loads(O.read_text_file(p))
            if isinstance(data, dict):
                for key in ("data", "rows", "records", "items", "results"):
                    if isinstance(data.get(key), list):
                        data = data[key]
                        break
            if isinstance(data, list):
                return pd.json_normalize(data) if data and isinstance(data[0], dict) else pd.DataFrame(data)
            if isinstance(data, dict):
                return pd.DataFrame(data)
            raise ToolError("JSON must be a list of records or a table-like object.")
    except ToolError:
        raise
    except Exception as e:  # noqa: BLE001
        raise ToolError(f"Could not read '{p.name}': {e.__class__.__name__}: {str(e)[:200]}") from e
    raise ToolError(f"Unsupported data file type '{ext}'. Use CSV, TSV, XLSX, Parquet or JSON.")


async def _load(ctx: ToolContext, path: str, sheet: str | None = None):
    p = O.ws_input(ctx, path, DATA_EXTS)
    df = await asyncio.to_thread(load_df, p, sheet)
    df.columns = [str(c).strip() for c in df.columns]
    return p, df


def _col(df, name: str) -> str:
    if name in df.columns:
        return name
    low = {str(c).lower(): c for c in df.columns}
    key = str(name).strip().lower()
    if key in low:
        return low[key]
    raise ToolError(f"Column '{name}' not found. Columns: {', '.join(map(str, list(df.columns)[:50]))}.")


def _cols(df, names: list[str] | str | None) -> list[str]:
    if names is None:
        return []
    if isinstance(names, str):
        names = [n.strip() for n in names.split(",") if n.strip()]
    return [_col(df, n) for n in names]


def _records(df, limit: int) -> list[list[Any]]:
    return [[O.jsonable(v) for v in r] for r in df.head(limit).itertuples(index=False, name=None)]


def to_number(series):
    """Parse amounts like '1,234.50', '$99', '(120.00)', '12%' into floats."""
    import pandas as pd

    if pd.api.types.is_numeric_dtype(series) and not pd.api.types.is_bool_dtype(series):
        return series.astype(float)
    s = series.astype(str).str.strip()
    neg = s.str.match(r"^\(.*\)$") | s.str.endswith("-")
    s = s.str.replace(r"[,$€£¥₹\s%]", "", regex=True).str.replace(r"^\((.*)\)$", r"\1", regex=True).str.rstrip("-")
    num = pd.to_numeric(s, errors="coerce")
    return num.where(~neg, -num)


async def _save_table(ctx: ToolContext, df, save_as: str, default_name: str) -> tuple[Path, dict]:
    ext = Path(str(save_as)).suffix.lower()
    ext = ext if ext in (".csv", ".xlsx", ".json") else ".csv"
    target = O.ws_output(ctx, save_as, default_name=default_name, ext=ext, overwrite=True)

    def write() -> None:
        if ext == ".xlsx":
            O.write_xlsx(target, {"Data": ([str(c) for c in df.columns],
                                           [list(r) for r in df.itertuples(index=False, name=None)])})
        elif ext == ".json":
            target.write_text(df.to_json(orient="records", date_format="iso", indent=1), encoding="utf-8")
        else:
            df.to_csv(target, index=False, encoding="utf-8-sig")

    await asyncio.to_thread(write)
    return target, {"saved": O.vpath(ctx, target)}


def _save_risk(args: dict) -> Risk:
    return Risk.MEDIUM if args.get("save_as") or args.get("chart") else Risk.LOW


# ----------------------------------------------------------------------------------------- profile
def profile_df(df, *, max_columns: int = 60, top: int = 5) -> dict:
    import pandas as pd

    n = len(df)
    out_cols = []
    for c in list(df.columns)[:max_columns]:
        s = df[c]
        nulls = int(s.isna().sum())
        try:
            uniq = int(s.nunique(dropna=True))
        except TypeError:
            uniq = None
        info: dict[str, Any] = {"name": str(c), "dtype": str(s.dtype), "nulls": nulls,
                                "null_pct": round(100 * nulls / n, 1) if n else 0.0, "unique": uniq}
        if pd.api.types.is_bool_dtype(s):
            info["true"] = int(s.sum())
        elif pd.api.types.is_numeric_dtype(s):
            d = s.describe()
            info.update({k: O.jsonable(float(d[k2])) for k, k2 in (("mean", "mean"), ("std", "std"), ("min", "min"),
                                                                    ("p25", "25%"), ("median", "50%"),
                                                                    ("p75", "75%"), ("max", "max")) if k2 in d})
            info["sum"] = O.jsonable(float(s.sum()))
        elif pd.api.types.is_datetime64_any_dtype(s):
            info["min"], info["max"] = O.jsonable(s.min()), O.jsonable(s.max())
        else:
            sample = s.dropna().astype(str).head(1000)
            if len(sample):
                num = to_number(sample)
                if num.notna().mean() > 0.9:
                    info["looks_numeric"] = True
                    info["min"], info["max"] = O.jsonable(float(num.min())), O.jsonable(float(num.max()))
                else:
                    with warnings.catch_warnings():
                        warnings.simplefilter("ignore")
                        dt = pd.to_datetime(sample, errors="coerce", format="mixed")
                    if dt.notna().mean() > 0.9 and sample.str.len().median() >= 6:
                        info["looks_datetime"] = True
                        info["min"], info["max"] = O.jsonable(dt.min()), O.jsonable(dt.max())
                info["avg_length"] = round(float(sample.str.len().mean()), 1)
            try:
                vc = s.value_counts(dropna=True).head(top)
                info["top"] = [[O.jsonable(k), int(v)] for k, v in vc.items()]
            except TypeError:
                pass
        out_cols.append(info)
    try:
        dups = int(df.duplicated().sum())
    except TypeError:
        dups = None
    return {"rows": n, "columns": len(df.columns), "duplicate_rows": dups,
            "memory": O.human_size(int(df.memory_usage(deep=True).sum())), "columns_info": out_cols,
            "columns_omitted": max(0, len(df.columns) - max_columns), "sample": _records(df, 5)}


@analytics.tool(risk=Risk.LOW, scope="read", timeout=120)
async def profile(ctx: ToolContext, path: str, sheet: str | None = None) -> dict:
    """Profile a dataset: row/column counts, dtypes, nulls, unique counts, numeric statistics, top categories,
    date ranges, duplicate rows and a sample. Start here before analysing an unfamiliar file.

    Args:
        path: Workspace path of a CSV/TSV/XLSX/Parquet/JSON file.
        sheet: Sheet name for Excel files (default: first sheet).
    """
    p, df = await _load(ctx, path, sheet)
    res = await asyncio.to_thread(profile_df, df)
    return {"path": O.vpath(ctx, p), **res}


# ----------------------------------------------------------------------------------------- SQL (DuckDB)
def _view_name(stem: str, used: set[str]) -> str:
    base = re.sub(r"[^a-z0-9_]+", "_", stem.lower()).strip("_") or "t"
    if base[0].isdigit():
        base = "t_" + base
    name, i = base, 2
    while name in used:
        name, i = f"{base}_{i}", i + 1
    used.add(name)
    return name


def _register(con, name: str, p: Path, sheet: str | None) -> None:
    ext = p.suffix.lower()
    q = '"' + name.replace('"', '""') + '"'
    if ext in (".csv", ".tsv", ".txt"):
        try:
            con.execute(f"CREATE TABLE {q} AS SELECT * FROM read_csv_auto(?, header=true, sample_size=-1)", [str(p)])
            return
        except Exception:  # noqa: BLE001
            pass
    elif ext in (".parquet", ".pq"):
        con.execute(f"CREATE TABLE {q} AS SELECT * FROM read_parquet(?)", [str(p)])
        return
    elif ext in (".jsonl", ".ndjson", ".json"):
        try:
            con.execute(f"CREATE TABLE {q} AS SELECT * FROM read_json_auto(?)", [str(p)])
            return
        except Exception:  # noqa: BLE001
            pass
    df = load_df(p, sheet)
    df.columns = [str(c).strip() for c in df.columns]
    con.register("_ren_tmp", df)
    con.execute(f"CREATE TABLE {q} AS SELECT * FROM _ren_tmp")
    con.unregister("_ren_tmp")


def _run_sql(con, tables: dict[str, tuple[Path, str | None]], sql: str, max_rows: int, full: bool):
    import duckdb

    con.execute("SET threads TO 2")
    for name, (p, sheet) in tables.items():
        _register(con, name, p, sheet)
    # Sandbox: from here on the query can't touch the filesystem or network, and settings are frozen.
    con.execute("SET enable_external_access = false")
    con.execute("SET lock_configuration = true")
    try:
        stmts = con.extract_statements(sql)
    except duckdb.Error as e:
        raise ToolError(f"SQL error: {str(e).splitlines()[0][:300]}") from e
    if len(stmts) != 1:
        raise ToolError("Run exactly one SQL statement at a time.")
    if stmts[0].type not in (duckdb.StatementType.SELECT, duckdb.StatementType.EXPLAIN):
        raise ToolError("Only read-only queries are allowed (SELECT / WITH … SELECT / DESCRIBE / SUMMARIZE).")
    try:
        cur = con.execute(sql)
        cols = [d[0] for d in cur.description]
        if full:
            df = cur.df()
            return cols, df, False
        rows = cur.fetchmany(max_rows + 1)
    except duckdb.Error as e:
        raise ToolError(f"SQL error: {str(e).splitlines()[0][:400]}") from e
    return cols, rows[:max_rows], len(rows) > max_rows


@analytics.tool(risk=Risk.LOW, scope="read", timeout=180, risk_fn=_save_risk)
async def query(ctx: ToolContext, sql: str, files: Any, max_rows: int = 200, sheet: str | None = None,
                save_as: str | None = None) -> dict | ToolResult:
    """Run a read-only SQL query (DuckDB dialect) over workspace data files. Each file becomes a table named after
    the file (e.g. uploads/sales_2024.csv → sales_2024), or use {"alias": "path"} to pick names. Supports joins,
    window functions, CTEs, DESCRIBE/SUMMARIZE.

    Args:
        sql: One SELECT statement, e.g. "SELECT region, SUM(amount) AS total FROM sales GROUP BY 1 ORDER BY 2 DESC".
        files: Workspace file paths (list) or {"table_name": "path"} mapping.
        max_rows: Maximum rows to return.
        sheet: Sheet name for Excel files.
        save_as: Optionally save the full result (.csv, .xlsx or .json).
    """
    import duckdb

    if not (sql or "").strip():
        raise ToolError("`sql` is required.")
    if isinstance(files, str):
        files = [f.strip() for f in files.split(",") if f.strip()]
    if not files:
        raise ToolError("Provide `files`, e.g. ['uploads/sales.csv'] (queried as table `sales`).")
    used: set[str] = set()
    tables: dict[str, tuple[Path, str | None]] = {}
    items = files.items() if isinstance(files, dict) else [(None, f) for f in files]
    for alias, f in items:
        p = O.ws_input(ctx, str(f), DATA_EXTS)
        name = _view_name(alias or p.stem, used)
        tables[name] = (p, sheet)
    n = max(1, min(int(max_rows or 200), 5000))
    con = duckdb.connect(":memory:")
    try:
        cols, data, truncated = await asyncio.to_thread(_run_sql, con, tables, sql, n, bool(save_as))
    except asyncio.CancelledError:
        con.interrupt()
        raise
    finally:
        try:
            con.close()
        except Exception:  # noqa: BLE001
            pass
    table_map = {k: O.vpath(ctx, v[0]) for k, v in tables.items()}
    if save_as:
        df = data
        target, info = await _save_table(ctx, df, save_as, "query-result")
        out = {"tables": table_map, "columns": cols, "row_count": len(df), "rows": _records(df, min(n, 20)), **info}
        return await O.produced(ctx, target, out, display=f"Query returned {len(df)} rows")
    rows = [[O.jsonable(v) for v in r] for r in data]
    return {"tables": table_map, "columns": cols, "rows": rows, "row_count": len(rows) if not truncated else None,
            "returned": len(rows), "truncated": truncated}


# ----------------------------------------------------------------------------------------- aggregate / join
_AGG = {"sum": "sum", "total": "sum", "mean": "mean", "avg": "mean", "average": "mean", "median": "median",
        "min": "min", "max": "max", "count": "count", "nunique": "nunique", "distinct": "nunique", "std": "std",
        "first": "first", "last": "last"}


def _parse_metrics(df, metrics: list[str] | dict | None) -> list[tuple[str, str | None, str]]:
    specs: list[tuple[str, str | None, str]] = []
    if isinstance(metrics, dict):
        items = []
        for col, fns in metrics.items():
            for fn in (fns if isinstance(fns, list) else [fns]):
                items.append(f"{fn}:{col}")
        metrics = items
    for m in metrics or ["count"]:
        fn, _, col = str(m).partition(":")
        fn = fn.strip().lower()
        if fn not in _AGG:
            fn, col = "sum", str(m)  # bare column name → sum
        agg = _AGG[fn]
        column = _col(df, col.strip()) if col.strip() else None
        name = f"{agg}_{column}" if column else "count"
        specs.append((name, column, agg))
    return specs


@analytics.tool(risk=Risk.LOW, scope="read", timeout=120, risk_fn=_save_risk)
async def aggregate(ctx: ToolContext, path: str, metrics: list[str], group_by: list[str] | None = None,
                    sheet: str | None = None, sort_by: str | None = None, descending: bool = True, limit: int = 100,
                    save_as: str | None = None) -> dict | ToolResult:
    """Group rows and compute metrics, e.g. revenue by region and quarter.

    Args:
        path: Workspace path of the data file.
        metrics: Metric specs "fn:column" — fn in sum, mean, median, min, max, count, nunique, std, first, last —
            e.g. ["sum:amount", "count", "nunique:customer_id"]. A bare column name means sum.
        group_by: Columns to group by (omit for overall totals).
        sheet: Sheet name for Excel files.
        sort_by: Output column to sort by (default: first metric).
        descending: Sort descending.
        limit: Maximum rows to return.
        save_as: Optionally save the full result (.csv/.xlsx/.json).
    """
    import pandas as pd

    p, df = await _load(ctx, path, sheet)
    specs = _parse_metrics(df, metrics)
    groups = _cols(df, group_by)

    def run():
        work = df.copy()
        for _, col, agg in specs:
            if col and agg in ("sum", "mean", "median", "min", "max", "std") and not \
                    pd.api.types.is_datetime64_any_dtype(work[col]):
                num = to_number(work[col])
                if num.notna().sum() >= 0.5 * work[col].notna().sum():
                    work[col] = num
        if groups:
            g = work.groupby(groups, dropna=False)
            parts = {}
            for name, col, agg in specs:
                parts[name] = g.size() if col is None else g[col].agg(agg)
            res = pd.DataFrame(parts).reset_index()
        else:
            row = {}
            for name, col, agg in specs:
                row[name] = len(work) if col is None else work[col].agg(agg)
            res = pd.DataFrame([row])
        key = sort_by and next((c for c in res.columns if c.lower() == sort_by.lower()), None)
        key = key or (specs[0][0] if groups else None)
        if key:
            res = res.sort_values(key, ascending=not descending, kind="mergesort")
        return res

    res = await asyncio.to_thread(run)
    n = max(1, min(int(limit or 100), 5000))
    out: dict[str, Any] = {"path": O.vpath(ctx, p), "groups": len(res), "columns": [str(c) for c in res.columns],
                           "rows": _records(res, n), "truncated": len(res) > n}
    if save_as:
        target, info = await _save_table(ctx, res, save_as, f"{p.stem}-summary")
        return await O.produced(ctx, target, {**out, **info}, display=f"{len(res)} groups")
    return out


@analytics.tool(risk=Risk.MEDIUM, scope="write", timeout=180)
async def join(ctx: ToolContext, left: str, right: str, on: list[str] | None = None,
               left_on: list[str] | None = None, right_on: list[str] | None = None,
               how: Literal["inner", "left", "right", "outer"] = "inner", output: str | None = None,
               left_sheet: str | None = None, right_sheet: str | None = None) -> ToolResult:
    """Join two data files on key columns and save the result (CSV by default, or .xlsx) in output/.

    Args:
        left: Workspace path of the left file.
        right: Workspace path of the right file.
        on: Key column(s) present in both files.
        left_on: Key column(s) in the left file (when names differ).
        right_on: Key column(s) in the right file.
        how: Join type.
        output: Output path, e.g. "output/orders_with_customers.csv".
        left_sheet: Sheet for an Excel left file.
        right_sheet: Sheet for an Excel right file.
    """
    import pandas as pd

    lp, ldf = await _load(ctx, left, left_sheet)
    rp, rdf = await _load(ctx, right, right_sheet)
    if on:
        lk, rk = _cols(ldf, on), _cols(rdf, on)
    elif left_on and right_on:
        lk, rk = _cols(ldf, left_on), _cols(rdf, right_on)
    else:
        common = [c for c in ldf.columns if c in set(rdf.columns)]
        if not common:
            raise ToolError("No shared column names — pass `on` or `left_on`/`right_on`.")
        lk = rk = common[:1]
    if len(lk) != len(rk):
        raise ToolError("left_on and right_on must have the same number of columns.")

    def run():
        a, b = ldf.copy(), rdf.copy()
        for x, y in zip(lk, rk):
            if a[x].dtype != b[y].dtype:
                a[x] = _key_series(a[x])
                b[y] = _key_series(b[y])
        res = pd.merge(a, b, left_on=lk, right_on=rk, how=how, suffixes=("", f"_{O.slugify(rp.stem)[:20]}"),
                       indicator=True)
        stats = res["_merge"].value_counts().to_dict()
        res = res.drop(columns=["_merge"])
        return res, {"both": int(stats.get("both", 0)), "left_only": int(stats.get("left_only", 0)),
                     "right_only": int(stats.get("right_only", 0))}

    res, stats = await asyncio.to_thread(run)
    ext = ".xlsx" if output and str(output).lower().endswith(".xlsx") else ".csv"
    target = O.ws_output(ctx, output, default_name=f"{lp.stem}-{rp.stem}-joined", ext=ext, overwrite=bool(output))

    def save() -> None:
        if ext == ".xlsx":
            O.write_xlsx(target, {"Joined": ([str(c) for c in res.columns],
                                             [list(r) for r in res.itertuples(index=False, name=None)])})
        else:
            res.to_csv(target, index=False, encoding="utf-8-sig")

    await asyncio.to_thread(save)
    return await O.produced(ctx, target, {"rows": len(res), "columns": [str(c) for c in res.columns][:60],
                                          "matches": stats, "keys": {"left": lk, "right": rk},
                                          "preview": _records(res, 5)},
                            display=f"Joined → {len(res)} rows")


def _key_series(s):
    import pandas as pd

    if pd.api.types.is_float_dtype(s):
        return s.map(lambda v: "" if pd.isna(v) else (str(int(v)) if float(v).is_integer() else str(v)))
    return s.astype(str).str.strip().replace({"nan": "", "None": ""})


# ----------------------------------------------------------------------------------------- reconcile
def _norm_keys(s, normalize: bool):
    k = _key_series(s)
    if normalize:
        k = k.str.strip().str.upper().str.replace(r"\s+", " ", regex=True)
    return k


def reconcile_frames(a, b, *, key_a: str, key_b: str, amt_a: str | None, amt_b: str | None, tolerance: float,
                     normalize: bool, label_a: str, label_b: str) -> dict:
    import pandas as pd

    A, B = a.copy(), b.copy()
    A["_key"], B["_key"] = _norm_keys(A[key_a], normalize), _norm_keys(B[key_b], normalize)
    A = A[A["_key"] != ""]
    B = B[B["_key"] != ""]
    dup_a = int(A["_key"].duplicated(keep=False).sum())
    dup_b = int(B["_key"].duplicated(keep=False).sum())
    if amt_a and amt_b:
        A["_amt"], B["_amt"] = to_number(A[amt_a]).fillna(0.0), to_number(B[amt_b]).fillna(0.0)
        ga = A.groupby("_key").agg(amount_a=("_amt", "sum"), rows_a=("_amt", "size"), key_a=(key_a, "first"))
        gb = B.groupby("_key").agg(amount_b=("_amt", "sum"), rows_b=("_amt", "size"), key_b=(key_b, "first"))
    else:
        ga = A.groupby("_key").agg(rows_a=("_key", "size"), key_a=(key_a, "first"))
        gb = B.groupby("_key").agg(rows_b=("_key", "size"), key_b=(key_b, "first"))
    m = ga.join(gb, how="outer")
    in_a, in_b = m["rows_a"].notna(), m["rows_b"].notna()
    both = in_a & in_b
    summary: dict[str, Any] = {
        f"records_{label_a}": int(len(A)), f"records_{label_b}": int(len(B)),
        f"unique_keys_{label_a}": int(in_a.sum()), f"unique_keys_{label_b}": int(in_b.sum()),
        f"duplicate_key_rows_{label_a}": dup_a, f"duplicate_key_rows_{label_b}": dup_b,
        "keys_in_both": int(both.sum()), f"only_in_{label_a}": int((in_a & ~in_b).sum()),
        f"only_in_{label_b}": int((in_b & ~in_a).sum()),
    }
    out: dict[str, Any] = {"summary": summary}
    if amt_a and amt_b:
        m["difference"] = m["amount_a"].fillna(0) - m["amount_b"].fillna(0)
        matched = both & (m["difference"].abs() <= tolerance)
        diffs = both & ~matched
        total_a, total_b = float(A["_amt"].sum()), float(B["_amt"].sum())
        summary.update({
            "matched": int(matched.sum()), "amount_differences": int(diffs.sum()),
            f"total_{label_a}": round(total_a, 2), f"total_{label_b}": round(total_b, 2),
            "net_difference": round(total_a - total_b, 2),
            "sum_of_differences_on_matched_keys": round(float(m.loc[diffs, "difference"].sum()), 2),
            f"amount_only_in_{label_a}": round(float(m.loc[in_a & ~in_b, "amount_a"].sum()), 2),
            f"amount_only_in_{label_b}": round(float(m.loc[in_b & ~in_a, "amount_b"].sum()), 2),
            "match_rate_pct": round(100 * float(matched.sum()) / max(1, int((in_a | in_b).sum())), 1),
            "tolerance": tolerance,
        })
        d = m[diffs].copy()
        d["abs"] = d["difference"].abs()
        d = d.sort_values("abs", ascending=False)
        out["differences"] = pd.DataFrame({
            "key": d["key_a"].where(d["key_a"].notna(), d["key_b"]), f"amount_{label_a}": d["amount_a"],
            f"amount_{label_b}": d["amount_b"], "difference": d["difference"], f"rows_{label_a}": d["rows_a"],
            f"rows_{label_b}": d["rows_b"]})
        mm = m[matched]
        out["matched"] = pd.DataFrame({"key": mm["key_a"], f"amount_{label_a}": mm["amount_a"],
                                       f"amount_{label_b}": mm["amount_b"], "difference": mm["difference"]})
    else:
        summary["matched"] = int(both.sum())
        summary["match_rate_pct"] = round(100 * float(both.sum()) / max(1, int((in_a | in_b).sum())), 1)
        mm = m[both]
        out["matched"] = pd.DataFrame({"key": mm["key_a"], f"rows_{label_a}": mm["rows_a"],
                                       f"rows_{label_b}": mm["rows_b"]})
    only_a_keys = set(m.index[in_a & ~in_b])
    only_b_keys = set(m.index[in_b & ~in_a])
    out["only_a"] = A[A["_key"].isin(only_a_keys)].drop(columns=[c for c in ("_key", "_amt") if c in A.columns])
    out["only_b"] = B[B["_key"].isin(only_b_keys)].drop(columns=[c for c in ("_key", "_amt") if c in B.columns])
    return out


def _write_reconciliation(path: Path, res: dict, *, label_a: str, label_b: str, file_a: str, file_b: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Reconciliation summary"
    ws["A1"].font = Font(bold=True, size=14, color=O.INK)
    ws["A2"] = f"{label_a}: {file_a}   ·   {label_b}: {file_b}"
    ws["A3"] = f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')} by Ren"
    for c in ("A2", "A3"):
        ws[c].font = Font(italic=True, color=O.MUTED)
    ws.append([])
    ws.append(["Metric", "Value"])
    O.xlsx_header_style(ws, 5, 2)
    for k, v in res["summary"].items():
        ws.append([k.replace("_", " ").capitalize(), O.clean_cell(v)])
    for row in ws.iter_rows(min_row=6, min_col=2, max_col=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.00"
    O.xlsx_autofit(ws)
    ws.column_dimensions["A"].width = 38

    def sheet(name: str, df) -> None:
        if df is None:
            return
        w = wb.create_sheet(name)
        O.write_sheet(w, [str(c) for c in df.columns], [list(r) for r in df.itertuples(index=False, name=None)])

    sheet("Differences", res.get("differences"))
    sheet(f"Only in {label_a}"[:31], res.get("only_a"))
    sheet(f"Only in {label_b}"[:31], res.get("only_b"))
    sheet("Matched", res.get("matched"))
    wb.save(str(path))


@analytics.tool(risk=Risk.MEDIUM, scope="write", timeout=240)
async def reconcile(ctx: ToolContext, left: str, right: str, key: str | None = None, left_key: str | None = None,
                    right_key: str | None = None, amount: str | None = None, left_amount: str | None = None,
                    right_amount: str | None = None, tolerance: float = 0.01, normalize_keys: bool = True,
                    left_label: str = "A", right_label: str = "B", left_sheet: str | None = None,
                    right_sheet: str | None = None, output: str | None = None) -> ToolResult:
    """Reconcile two datasets on a key (e.g. bank statement vs. ledger, invoices vs. payments): finds matched
    records, amount differences beyond a tolerance, and records only in one side. Amounts like "$1,234.50" or
    "(120.00)" are parsed; duplicate keys are summed. Writes an Excel report (Summary, Differences, Only in A,
    Only in B, Matched) and returns the summary plus the largest differences.

    Args:
        left: Workspace path of dataset A.
        right: Workspace path of dataset B.
        key: Key column present in both (e.g. "invoice_id").
        left_key: Key column in A (when names differ).
        right_key: Key column in B.
        amount: Amount column present in both (omit for presence-only matching).
        left_amount: Amount column in A.
        right_amount: Amount column in B.
        tolerance: Absolute difference treated as a match (e.g. 0.01).
        normalize_keys: Trim/upper-case keys before matching.
        left_label: Short label for A in the report, e.g. "Bank".
        right_label: Short label for B, e.g. "Ledger".
        left_sheet: Sheet for an Excel A file.
        right_sheet: Sheet for an Excel B file.
        output: Report path (default: output/reconciliation-<A>-vs-<B>.xlsx).
    """
    lp, a = await _load(ctx, left, left_sheet)
    rp, b = await _load(ctx, right, right_sheet)
    ka = _col(a, left_key or key or "")
    kb = _col(b, right_key or key or "")
    amt_a = _col(a, left_amount or amount) if (left_amount or amount) else None
    amt_b = _col(b, right_amount or amount) if (right_amount or amount) else None
    if bool(amt_a) != bool(amt_b):
        raise ToolError("Give the amount column for both sides (amount, or left_amount + right_amount).")
    la = re.sub(r"[^A-Za-z0-9]+", "_", left_label or "A").strip("_") or "A"
    lb = re.sub(r"[^A-Za-z0-9]+", "_", right_label or "B").strip("_") or "B"
    if la.lower() == lb.lower():
        la, lb = la + "_1", lb + "_2"
    res = await asyncio.to_thread(reconcile_frames, a, b, key_a=ka, key_b=kb, amt_a=amt_a, amt_b=amt_b,
                                  tolerance=abs(float(tolerance or 0)), normalize=normalize_keys,
                                  label_a=la, label_b=lb)
    target = O.ws_output(ctx, output, default_name=f"reconciliation-{lp.stem}-vs-{rp.stem}", ext=".xlsx",
                         overwrite=bool(output))
    await asyncio.to_thread(_write_reconciliation, target, res, label_a=la, label_b=lb, file_a=O.vpath(ctx, lp),
                            file_b=O.vpath(ctx, rp))
    out: dict[str, Any] = {"summary": res["summary"]}
    if res.get("differences") is not None:
        d = res["differences"]
        out["top_differences"] = {"columns": [str(c) for c in d.columns], "rows": _records(d, 10)}
    out[f"sample_only_in_{la}"] = _records(res["only_a"], 5)
    out[f"sample_only_in_{lb}"] = _records(res["only_b"], 5)
    s = res["summary"]
    disp = f"{s.get('matched', 0)} matched · {s.get('amount_differences', 0)} differences · " \
           f"{s.get(f'only_in_{la}', 0)}/{s.get(f'only_in_{lb}', 0)} unmatched"
    return await O.produced(ctx, target, out, display=disp)


# ----------------------------------------------------------------------------------------- stats
@analytics.tool(risk=Risk.LOW, scope="read", timeout=120)
async def correlate(ctx: ToolContext, path: str, columns: list[str] | None = None,
                    method: Literal["pearson", "spearman", "kendall"] = "pearson", sheet: str | None = None,
                    top: int = 10) -> dict:
    """Correlation matrix between numeric columns, plus the strongest pairs.

    Args:
        path: Workspace path of the data file.
        columns: Columns to include (default: all numeric).
        method: Correlation method (spearman = rank-based, robust to outliers).
        sheet: Sheet name for Excel files.
        top: Number of strongest pairs to list.
    """
    import pandas as pd

    p, df = await _load(ctx, path, sheet)
    cols = _cols(df, columns) if columns else [c for c in df.columns if pd.api.types.is_numeric_dtype(df[c])
                                               and not pd.api.types.is_bool_dtype(df[c])]
    if len(cols) < 2:
        raise ToolError("Need at least two numeric columns.")
    num = pd.DataFrame({c: to_number(df[c]) for c in cols[:40]})
    corr = await asyncio.to_thread(num.corr, method=method)
    pairs = []
    names = list(corr.columns)
    for i in range(len(names)):
        for j in range(i + 1, len(names)):
            v = corr.iloc[i, j]
            if pd.notna(v):
                pairs.append({"a": names[i], "b": names[j], "r": round(float(v), 3),
                              "strength": "strong" if abs(v) >= 0.7 else "moderate" if abs(v) >= 0.4 else "weak"})
    pairs.sort(key=lambda x: -abs(x["r"]))
    matrix = {a: {b: (round(float(corr.loc[a, b]), 3) if pd.notna(corr.loc[a, b]) else None) for b in names}
              for a in names}
    return {"path": O.vpath(ctx, p), "method": method, "rows": len(df), "top_pairs": pairs[: max(1, int(top))],
            "matrix": matrix if len(names) <= 12 else "omitted (more than 12 columns)"}


@analytics.tool(risk=Risk.LOW, scope="read", timeout=120, risk_fn=_save_risk)
async def outliers(ctx: ToolContext, path: str, column: str, method: Literal["iqr", "zscore"] = "iqr",
                   threshold: float | None = None, sheet: str | None = None, max_rows: int = 20,
                   save_as: str | None = None) -> dict | ToolResult:
    """Find outliers in a numeric column with the IQR rule (default k=1.5) or z-scores (default |z| > 3).

    Args:
        path: Workspace path of the data file.
        column: Numeric column to test.
        method: "iqr" or "zscore".
        threshold: k for IQR, or the |z| cut-off.
        sheet: Sheet name for Excel files.
        max_rows: Maximum outlier rows to return.
        save_as: Optionally save all outlier rows (.csv/.xlsx).
    """
    p, df = await _load(ctx, path, sheet)
    col = _col(df, column)
    s = to_number(df[col])
    valid = s.dropna()
    if len(valid) < 4:
        raise ToolError(f"Column '{col}' has too few numeric values.")
    note = None
    q1, q3 = float(valid.quantile(0.25)), float(valid.quantile(0.75))
    if method == "iqr" and q3 - q1 == 0:
        # Most values are identical: the IQR rule would flag every other value. Use z-scores instead.
        method, threshold = "zscore", None
        note = "The IQR is zero (most values are identical), so the z-score rule (|z| > 3) was used instead."
    if method == "zscore":
        k = float(threshold or 3.0)
        mu, sd = float(valid.mean()), float(valid.std(ddof=0))
        if sd == 0:
            mask = s != s
        else:
            mask = ((s - mu) / sd).abs() > k
        bounds = {"mean": round(mu, 4), "std": round(sd, 4), "lower": round(mu - k * sd, 4),
                  "upper": round(mu + k * sd, 4)}
    else:
        k = float(threshold or 1.5)
        iqr = q3 - q1
        lo, hi = q1 - k * iqr, q3 + k * iqr
        mask = (s < lo) | (s > hi)
        bounds = {"q1": round(q1, 4), "q3": round(q3, 4), "iqr": round(iqr, 4), "lower": round(lo, 4),
                  "upper": round(hi, 4)}
    mask = mask.fillna(False)
    out_df = df[mask].copy()
    out_df.insert(0, "_row", out_df.index + 2)  # spreadsheet row number (header = row 1)
    order = (s[mask] - float(valid.median())).abs().sort_values(ascending=False).index
    out_df = out_df.loc[order]
    res: dict[str, Any] = {"path": O.vpath(ctx, p), "column": col, "method": method, "threshold": k,
                           "bounds": bounds, "outliers": int(mask.sum()),
                           "outlier_pct": round(100 * float(mask.sum()) / max(1, len(valid)), 2),
                           "columns": [str(c) for c in out_df.columns],
                           "rows": _records(out_df, max(1, int(max_rows or 20)))}
    if note:
        res["note"] = note
    if save_as:
        target, info = await _save_table(ctx, out_df, save_as, f"{p.stem}-outliers")
        return await O.produced(ctx, target, {**res, **info}, display=f"{res['outliers']} outliers")
    return res


# ----------------------------------------------------------------------------------------- forecasting
def _holt(y, alpha: float, beta: float, h: int):
    import numpy as np

    level, trend = y[0], (y[1] - y[0]) if len(y) > 1 else 0.0
    fitted = [y[0]]
    for t in range(1, len(y)):
        fitted.append(level + trend)
        prev = level
        level = alpha * y[t] + (1 - alpha) * (level + trend)
        trend = beta * (level - prev) + (1 - beta) * trend
    return np.array(fitted), np.array([level + (k + 1) * trend for k in range(h)]), trend


def _holt_winters(y, m: int, alpha: float, beta: float, gamma: float, h: int):
    import numpy as np

    level = float(np.mean(y[:m]))
    trend = float((np.mean(y[m:2 * m]) - np.mean(y[:m])) / m)
    season = [float(y[i] - level) for i in range(m)]
    fitted = []
    for t in range(len(y)):
        s = season[t % m]
        fitted.append(level + trend + s)
        prev = level
        level = alpha * (y[t] - s) + (1 - alpha) * (level + trend)
        trend = beta * (level - prev) + (1 - beta) * trend
        season[t % m] = gamma * (y[t] - level) + (1 - gamma) * s
    n = len(y)
    return np.array(fitted), np.array([level + (k + 1) * trend + season[(n + k) % m] for k in range(h)]), trend


def _fit(method: str, y, h: int, m: int | None):
    import numpy as np

    if method == "linear":
        t = np.arange(len(y))
        slope, icpt = np.polyfit(t, y, 1)
        fitted = icpt + slope * t
        fc = icpt + slope * np.arange(len(y), len(y) + h)
        return fitted, fc, {"slope": float(slope), "intercept": float(icpt)}, float(slope)
    best = None
    if method == "holt":
        for a in (0.1, 0.2, 0.3, 0.4, 0.5, 0.6, 0.7, 0.8, 0.9):
            for b in (0.01, 0.05, 0.1, 0.2, 0.3, 0.5):
                fitted, fc, tr = _holt(y, a, b, h)
                sse = float(np.sum((y[1:] - fitted[1:]) ** 2))
                if best is None or sse < best[0]:
                    best = (sse, fitted, fc, {"alpha": a, "beta": b}, tr)
    else:
        for a in (0.1, 0.2, 0.3, 0.5, 0.7, 0.9):
            for b in (0.01, 0.05, 0.1, 0.2, 0.4):
                for g in (0.05, 0.1, 0.2, 0.3, 0.5):
                    fitted, fc, tr = _holt_winters(y, int(m), a, b, g, h)
                    sse = float(np.sum((y[m:] - fitted[m:]) ** 2))
                    if best is None or sse < best[0]:
                        best = (sse, fitted, fc, {"alpha": a, "beta": b, "gamma": g, "season_length": m}, tr)
    _, fitted, fc, params, tr = best
    return fitted, fc, params, float(tr)


def _metrics(actual, pred) -> dict:
    import numpy as np

    actual, pred = np.asarray(actual, float), np.asarray(pred, float)
    rmse = float(np.sqrt(np.mean((actual - pred) ** 2)))
    nz = actual != 0
    mape = float(np.mean(np.abs((actual[nz] - pred[nz]) / actual[nz])) * 100) if nz.any() else None
    return {"rmse": round(rmse, 4), "mape_pct": round(mape, 2) if mape is not None else None}


_FREQ = {"d": "D", "day": "D", "daily": "D", "w": "W", "week": "W", "weekly": "W", "m": "MS", "month": "MS",
         "monthly": "MS", "q": "QS", "quarter": "QS", "quarterly": "QS", "y": "YS", "year": "YS", "yearly": "YS",
         "annual": "YS"}
_SEASON = {"D": 7, "W": 52, "MS": 12, "ME": 12, "M": 12, "QS": 4, "QE": 4, "Q": 4}


def forecast_series(df, value_col: str, date_col: str | None, periods: int, method: str, season_length: int | None,
                    freq: str | None, agg: str) -> dict:
    import numpy as np
    import pandas as pd

    data = df.copy()
    data[value_col] = to_number(data[value_col])
    freq_used = None
    if date_col:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            data[date_col] = pd.to_datetime(data[date_col], errors="coerce", format="mixed")
        data = data.dropna(subset=[date_col, value_col])
        if freq:
            f = _FREQ.get(freq.lower().strip(), freq)
            series = data.set_index(date_col)[value_col].resample(f).agg(agg)
            freq_used = f
        else:
            series = data.groupby(date_col)[value_col].agg(agg).sort_index()
            try:
                freq_used = pd.infer_freq(series.index) if len(series) >= 3 else None
            except (TypeError, ValueError):
                freq_used = None
            if freq_used is None and len(series) >= 3:
                gap = float(np.median(np.diff(series.index.values).astype("timedelta64[D]").astype(float)))
                freq_used = "D" if gap <= 1.5 else "W" if gap <= 8 else "MS" if gap <= 31.5 else \
                    "QS" if gap <= 92.5 else "YS"
        series = series.dropna()
    else:
        series = data[value_col].dropna().reset_index(drop=True)
    y = series.values.astype(float)
    if len(y) < 3:
        raise ToolError("Need at least 3 data points to forecast.")
    h = max(1, min(int(periods), 500))
    base = (freq_used or "").upper().split("-")[0]
    m = int(season_length) if season_length else _SEASON.get(base) if date_col else None
    candidates = ["linear"] + (["holt"] if len(y) >= 4 else [])
    if m and m >= 2 and len(y) >= 2 * m + 2:
        candidates.append("holt_winters")
    if method != "auto":
        if method == "holt_winters" and "holt_winters" not in candidates:
            raise ToolError(f"Holt-Winters needs a season length and at least {2 * (m or 2) + 2} points "
                            f"(have {len(y)}). Pass season_length or use 'holt'.")
        if method == "holt" and len(y) < 4:
            raise ToolError("Holt needs at least 4 points.")
        chosen, validation = method, None
    else:
        hold = max(2, min(len(y) // 5, h))
        scores = {}
        if len(y) - hold >= 3:
            for c in candidates:
                if c == "holt_winters" and len(y) - hold < 2 * m + 2:
                    continue
                if c == "holt" and len(y) - hold < 4:
                    continue
                _, fc, _, _ = _fit(c, y[:-hold], hold, m)
                scores[c] = _metrics(y[-hold:], fc)
        chosen = min(scores, key=lambda c: scores[c]["rmse"]) if scores else "linear"
        validation = {"holdout_points": hold, "scores": scores} if scores else None
    fitted, fc, params, trend = _fit(chosen, y, h, m)
    start = 1 if chosen == "holt" else (m if chosen == "holt_winters" else 0)
    resid = y[start:] - fitted[start:]
    sd = float(np.std(resid, ddof=1)) if len(resid) > 2 else 0.0
    if date_col and freq_used:
        future = pd.date_range(series.index[-1], periods=h + 1, freq=freq_used)[1:]
        labels = [d.date().isoformat() if d == d.normalize() else d.isoformat() for d in future]
        hist_labels = [d.date().isoformat() for d in series.index]
    else:
        labels = list(range(len(y) + 1, len(y) + h + 1))
        hist_labels = list(range(1, len(y) + 1))
    rows = [{"period": labels[k], "forecast": round(float(fc[k]), 4),
             "lower_95": round(float(fc[k] - 1.96 * sd * math.sqrt(k + 1)), 4),
             "upper_95": round(float(fc[k] + 1.96 * sd * math.sqrt(k + 1)), 4)} for k in range(h)]
    return {"method": chosen, "params": params, "history_points": len(y), "frequency": freq_used,
            "season_length": m if chosen == "holt_winters" else None, "trend_per_period": round(trend, 4),
            "fit": _metrics(y[start:], fitted[start:]), "validation": validation, "forecast": rows,
            "last_actual": {"period": hist_labels[-1], "value": round(float(y[-1]), 4)},
            "_history": (hist_labels, y.tolist()), "_fc": (labels, fc.tolist(), sd)}


@analytics.tool(risk=Risk.LOW, scope="read", timeout=180, risk_fn=_save_risk)
async def forecast(ctx: ToolContext, path: str, value_column: str, date_column: str | None = None, periods: int = 6,
                   method: Literal["auto", "linear", "holt", "holt_winters"] = "auto",
                   season_length: int | None = None, freq: str | None = None,
                   agg: Literal["sum", "mean", "last"] = "sum", sheet: str | None = None,
                   chart: bool = False) -> dict | ToolResult:
    """Forecast a time series: linear trend, Holt (level + trend) or Holt-Winters (additive seasonality), chosen
    automatically by holdout error when method="auto". Returns forecasts with approximate 95% intervals.

    Args:
        path: Workspace path of the data file.
        value_column: Numeric column to forecast (e.g. "revenue").
        date_column: Date column (rows are aggregated per date; frequency is inferred).
        periods: Number of future periods.
        method: Forecast model.
        season_length: Seasonal period (e.g. 12 for monthly data); inferred from the frequency when possible.
        freq: Resample to D, W, M, Q or Y before forecasting.
        agg: How to combine multiple rows per period.
        sheet: Sheet name for Excel files.
        chart: Also save a PNG chart of history + forecast (output/charts/).
    """
    p, df = await _load(ctx, path, sheet)
    vc = _col(df, value_column)
    dc = _col(df, date_column) if date_column else None
    res = await asyncio.to_thread(forecast_series, df, vc, dc, periods, method, season_length, freq, agg)
    hist, fcd = res.pop("_history"), res.pop("_fc")
    res = {"path": O.vpath(ctx, p), "value_column": vc, **res}
    if not chart:
        return res
    target = O.ws_output(ctx, None, default_name=f"{p.stem}-{vc}-forecast", ext=".png", folder="output/charts")
    await asyncio.to_thread(_forecast_chart, target, hist, fcd, f"{vc} forecast", bool(dc))
    return await O.produced(ctx, target, res, display=f"Forecast ({res['method']}) for {periods} periods")


# ----------------------------------------------------------------------------------------- charts
def _trim(x: float, decimals: int) -> str:
    return f"{x:.{decimals}f}".rstrip("0").rstrip(".") if decimals else f"{x:.0f}"


def _human(v: float) -> str:
    """Compact, exact-enough labels: 1,750,000 → 1.75M · 12,500 → 12.5K · 830,000 → 830K."""
    a = abs(v)
    if a >= 1e9:
        return _trim(v / 1e9, 2) + "B"
    if a >= 1e6:
        return _trim(v / 1e6, 2) + "M"
    if a >= 1e3:
        return _trim(v / 1e3, 1) + "K"
    if a == int(a):
        return f"{int(v)}"
    return _trim(v, 2)


def _fig(width: float = 10, height: float = 5.6):
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.figure import Figure

    fig = Figure(figsize=(width, height), dpi=150, facecolor=IVORY)
    FigureCanvasAgg(fig)
    ax = fig.add_subplot(111)
    ax.set_facecolor(IVORY)
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    ax.spines["bottom"].set_color(INK)
    ax.spines["bottom"].set_linewidth(0.8)
    ax.tick_params(colors=SOFT, labelsize=9, length=0, pad=6)
    ax.yaxis.grid(True, color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)
    return fig, ax


def _titles(fig, ax, title: str | None, subtitle: str | None, x_label: str | None, y_label: str | None,
            legend: bool = False) -> None:
    if title:
        fig.text(0.055, 0.955, title, fontsize=15, fontweight="bold", color=INK, ha="left", va="top")
    if subtitle:
        fig.text(0.055, 0.905 if title else 0.955, subtitle, fontsize=10, color=MUTED, ha="left", va="top")
    if x_label:
        ax.set_xlabel(x_label, color=SOFT, fontsize=9.5, labelpad=8)
    if y_label:
        ax.set_ylabel(y_label, color=SOFT, fontsize=9.5, labelpad=8)
    top = 0.86 if (title and subtitle) else 0.9 if (title or subtitle) else 0.95
    if legend:
        top -= 0.06  # room for the legend row above the plot
    fig.subplots_adjust(left=0.08, right=0.97, top=top, bottom=0.14)


def _save_png(fig, path: Path) -> None:
    fig.savefig(str(path), facecolor=IVORY, dpi=150)


def _forecast_chart(path: Path, hist, fcd, title: str, dates: bool) -> None:
    import numpy as np
    import pandas as pd
    from matplotlib.ticker import FuncFormatter

    hx, hy = hist
    fx, fy, sd = fcd
    fig, ax = _fig()
    if dates:
        hx = pd.to_datetime(hx)
        fx = pd.to_datetime(fx)
    ax.plot(hx, hy, color=INK, linewidth=2, label="Actual")
    fy = np.asarray(fy)
    band = np.array([1.96 * sd * math.sqrt(k + 1) for k in range(len(fy))])
    ax.plot([hx[-1]] + list(fx), [hy[-1]] + list(fy), color=BRONZE, linewidth=2.2, linestyle="--", label="Forecast")
    ax.fill_between(fx, fy - band, fy + band, color=BRONZE, alpha=0.2, linewidth=0, label="95% interval")
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: _human(v)))
    leg = ax.legend(frameon=False, fontsize=9, loc="upper left")
    for t in leg.get_texts():
        t.set_color(SOFT)
    if dates:
        import matplotlib.dates as mdates

        loc = mdates.AutoDateLocator()
        ax.xaxis.set_major_locator(loc)
        ax.xaxis.set_major_formatter(mdates.ConciseDateFormatter(loc))
    _titles(fig, ax, title, None, None, None)
    _save_png(fig, path)


def render_chart(df, path: Path, *, chart_type: str, x: str | None, ys: list[str], title: str | None,
                 subtitle: str | None, x_label: str | None, y_label: str | None, series_col: str | None,
                 agg: str, sort: str, limit: int | None, stacked: bool) -> dict:
    import numpy as np
    import pandas as pd
    from matplotlib.ticker import FuncFormatter

    data = df.copy()
    ct = chart_type
    is_date = False
    if ct == "histogram":
        col = ys[0] if ys else x
        vals = to_number(data[col]).dropna()
        if vals.empty:
            raise ToolError(f"Column '{col}' has no numeric values.")
        fig, ax = _fig()
        ax.hist(vals, bins=min(40, max(8, int(math.sqrt(len(vals))))), color=ACCENT, edgecolor=IVORY, linewidth=0.8)
        ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: _human(v)))
        ax.xaxis.set_major_formatter(FuncFormatter(lambda v, _: _human(v)))
        _titles(fig, ax, title or f"Distribution of {col}", subtitle, x_label or col, y_label or "Count")
        _save_png(fig, path)
        return {"points": int(len(vals)), "series": [col]}
    if not x:
        raise ToolError("`x` (category / date column) is required for this chart type.")
    if not ys:
        ys = [c for c in data.columns if c != x and pd.api.types.is_numeric_dtype(data[c])][:6]
        if not ys and ct != "pie":
            raise ToolError("No numeric columns to plot — pass `y`.")
    for y in ys:
        data[y] = to_number(data[y])
    if ct != "scatter":
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            parsed = pd.to_datetime(data[x], errors="coerce", format="mixed") \
                if not pd.api.types.is_numeric_dtype(data[x]) else None
        if parsed is not None and parsed.notna().mean() > 0.9 and data[x].astype(str).str.len().median() >= 6:
            data[x] = parsed
            is_date = True
    if series_col:
        if len(ys) != 1:
            raise ToolError("With `series`, pass exactly one `y` column.")
        piv = data.pivot_table(index=x, columns=series_col, values=ys[0], aggfunc=agg if agg != "none" else "sum")
        piv.columns = [str(c) for c in piv.columns]
        table = piv
    elif ct == "scatter":
        table = data[[x] + ys].dropna()
    else:
        if agg != "none" and data[x].duplicated().any():
            table = data.groupby(x)[ys].agg(agg) if agg != "count" else data.groupby(x)[ys].count()
        else:
            table = data.set_index(x)[ys]
    if ct != "scatter":
        if is_date or ct in ("line", "area"):
            table = table.sort_index()
        elif sort == "value":
            order = table.sum(axis=1, numeric_only=True).sort_values(ascending=False).index
            table = table.loc[order]
        elif sort == "x":
            table = table.sort_index()
        if limit and len(table) > limit and not is_date:
            if ct in ("pie", "donut"):
                head, rest = table.iloc[: limit - 1], table.iloc[limit - 1:]
                other = pd.DataFrame([rest.sum(numeric_only=True)], index=["Other"])
                table = pd.concat([head, other])
            else:
                table = table.iloc[:limit]
    fig, ax = _fig()
    series_names = [str(c) for c in (table.columns if ct != "scatter" else ys)]
    n_series = len(series_names)
    if ct in ("pie", "donut"):
        ax.yaxis.grid(False)
        ax.spines["bottom"].set_visible(False)
        vals = table.iloc[:, 0].clip(lower=0)
        colors = [PALETTE[i % len(PALETTE)] for i in range(len(vals))]
        wedges, _texts, autotexts = ax.pie(
            vals.values, labels=None, colors=colors, startangle=90, counterclock=False,
            autopct=lambda pct: f"{pct:.0f}%" if pct >= 4 else "", pctdistance=0.78 if ct == "donut" else 0.68,
            wedgeprops={"linewidth": 2, "edgecolor": IVORY, **({"width": 0.42} if ct == "donut" else {})},
            textprops={"color": "white", "fontsize": 9, "fontweight": "bold"})
        ax.axis("equal")
        leg = ax.legend(wedges, [f"{str(lbl)[:28]}  {_human(float(v))}" for lbl, v in zip(table.index, vals.values)],
                        loc="center left", bbox_to_anchor=(1.0, 0.5), frameon=False, fontsize=9)
        for t in leg.get_texts():
            t.set_color(SOFT)
        _titles(fig, ax, title, subtitle, None, None)
        fig.subplots_adjust(left=0.05, right=0.68)
        _save_png(fig, path)
        return {"points": int(len(vals)), "series": series_names[:1]}
    if ct == "scatter":
        xv = to_number(table[x]) if not pd.api.types.is_numeric_dtype(table[x]) else table[x]
        for i, y in enumerate(ys):
            ax.scatter(xv, table[y], s=26, alpha=0.75, color=PALETTE[i % len(PALETTE)], edgecolors="none", label=y)
        if len(ys) == 1 and len(table) >= 3:
            ok = xv.notna() & table[ys[0]].notna()
            if ok.sum() >= 3:
                k, b = np.polyfit(xv[ok].astype(float), table[ys[0]][ok].astype(float), 1)
                xs = np.linspace(float(xv[ok].min()), float(xv[ok].max()), 50)
                ax.plot(xs, k * xs + b, color=INK, linewidth=1.2, linestyle="--", alpha=0.7)
        ax.xaxis.set_major_formatter(FuncFormatter(lambda v, _: _human(v)))
        points = len(table)
    else:
        labels = table.index
        xs = np.arange(len(labels))
        if ct in ("bar", "horizontal_bar"):
            horizontal = ct == "horizontal_bar"
            width = 0.8 / (1 if stacked else n_series)
            bottom = np.zeros(len(labels))
            for i, col in enumerate(table.columns):
                vals = table[col].fillna(0).values
                color = PALETTE[i % len(PALETTE)]
                if stacked:
                    bars = (ax.barh(xs, vals, 0.65, left=bottom, color=color, label=str(col)) if horizontal else
                            ax.bar(xs, vals, 0.65, bottom=bottom, color=color, label=str(col)))
                    bottom = bottom + vals
                else:
                    offs = xs - 0.4 + width * (i + 0.5)
                    bars = (ax.barh(offs, vals, width * 0.92, color=color, label=str(col)) if horizontal else
                            ax.bar(offs, vals, width * 0.92, color=color, label=str(col)))
                if n_series == 1 and len(labels) <= 15:
                    ax.bar_label(bars, labels=[_human(float(v)) for v in vals], padding=3, fontsize=8.5, color=INK)
            tick_labels = [_date_label(v) if is_date else str(v)[:22] for v in labels]
            if horizontal:
                ax.set_yticks(xs, tick_labels)
                ax.invert_yaxis()
                ax.yaxis.grid(False)
                ax.xaxis.grid(True, color=GRID, linewidth=0.8)
                ax.spines["bottom"].set_visible(False)
                ax.spines["left"].set_visible(True)
                ax.spines["left"].set_color(INK)
                ax.xaxis.set_major_formatter(FuncFormatter(lambda v, _: _human(v)))
            else:
                ax.set_xticks(xs, tick_labels)
                if len(labels) > 8 or max((len(t) for t in tick_labels), default=0) > 10:
                    for lab in ax.get_xticklabels():
                        lab.set_rotation(35)
                        lab.set_ha("right")
        else:
            xvals = labels if is_date or pd.api.types.is_numeric_dtype(labels) else xs
            cum = np.zeros(len(labels))
            for i, col in enumerate(table.columns):
                vals = table[col].astype(float).values
                color = PALETTE[i % len(PALETTE)] if n_series > 1 else (INK if ct == "line" else ACCENT)
                if ct == "area":
                    base = cum if stacked else np.zeros(len(labels))
                    ax.fill_between(xvals, base, base + np.nan_to_num(vals), color=PALETTE[i % len(PALETTE)],
                                    alpha=0.35 if not stacked else 0.75, linewidth=0)
                    ax.plot(xvals, base + np.nan_to_num(vals), color=PALETTE[i % len(PALETTE)], linewidth=1.6,
                            label=str(col))
                    if stacked:
                        cum = cum + np.nan_to_num(vals)
                else:
                    ax.plot(xvals, vals, color=color, linewidth=2.2, label=str(col),
                            marker="o" if len(labels) <= 24 else None, markersize=4)
            if not (is_date or pd.api.types.is_numeric_dtype(labels)):
                ax.set_xticks(xs, [str(v)[:16] for v in labels])
                if len(labels) > 10:
                    for lab in ax.get_xticklabels():
                        lab.set_rotation(35)
                        lab.set_ha("right")
            if is_date:
                import matplotlib.dates as mdates

                loc = mdates.AutoDateLocator()
                ax.xaxis.set_major_locator(loc)
                ax.xaxis.set_major_formatter(mdates.ConciseDateFormatter(loc))
        points = len(labels)
    if ct != "horizontal_bar":
        ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: _human(v)))
    if ct in ("bar", "horizontal_bar"):
        totals = table.clip(lower=0).sum(axis=1) if stacked else table.max(axis=1)
        peak = float(totals.max()) if len(totals) else 0.0
        if peak > 0 and float(table.min().min()) >= 0:
            (ax.set_xlim if ct == "horizontal_bar" else ax.set_ylim)(0, peak * 1.1)
    show_legend = n_series > 1 or (ct == "scatter" and len(ys) > 1)
    if show_legend:
        leg = ax.legend(frameon=False, fontsize=9, ncol=min(5, max(1, n_series)), loc="lower left",
                        bbox_to_anchor=(0, 1.01), borderaxespad=0.2, handlelength=1.4)
        for t in leg.get_texts():
            t.set_color(SOFT)
    _titles(fig, ax, title, subtitle, x_label, y_label, legend=show_legend)
    if ct == "horizontal_bar":
        fig.subplots_adjust(left=0.22)
    _save_png(fig, path)
    return {"points": int(points), "series": series_names}


def _date_label(v) -> str:
    try:
        return v.strftime("%b %Y") if getattr(v, "day", 1) == 1 else v.strftime("%Y-%m-%d")
    except Exception:  # noqa: BLE001
        return str(v)


@analytics.tool(risk=Risk.MEDIUM, scope="write", timeout=180)
async def visualize(ctx: ToolContext, chart_type: Literal["bar", "horizontal_bar", "line", "area", "scatter", "pie",
                                                          "donut", "histogram"] = "bar",
                    path: str | None = None, data: list[dict[str, Any]] | None = None, x: str | None = None,
                    y: list[str] | str | None = None, series: str | None = None,
                    agg: Literal["sum", "mean", "median", "count", "max", "min", "none"] = "sum",
                    title: str | None = None, subtitle: str | None = None, x_label: str | None = None,
                    y_label: str | None = None, sort: Literal["auto", "value", "x", "none"] = "auto",
                    limit: int | None = None, stacked: bool = False, sheet: str | None = None,
                    output: str | None = None) -> ToolResult:
    """Create a presentation-ready chart (PNG, f69 ivory background, green/bronze palette) from a data file or
    inline records and save it under output/charts/. Rows with the same x are aggregated (agg). Insert the PNG
    into documents/decks with doc_insert_image / ppt_add_image.

    Args:
        chart_type: bar, horizontal_bar, line, area, scatter, pie, donut or histogram.
        path: Workspace data file (CSV/XLSX/Parquet/JSON) — or pass `data`.
        data: Inline records, e.g. [{"month": "Jan", "revenue": 120}, ...].
        x: Category / date column (x-axis; labels for pie).
        y: Value column(s) (for histogram: the column to bin).
        series: Column whose values become separate series (one y only), e.g. "region".
        agg: Aggregation for duplicate x values.
        title: Chart title.
        subtitle: Smaller line under the title (e.g. source or period).
        x_label: X-axis label.
        y_label: Y-axis label.
        sort: "value" (largest first), "x", "none"; auto = value for bars, x for lines/dates.
        limit: Keep the top N categories (pie groups the rest as "Other"; default 8 for pie).
        stacked: Stack series (bar/horizontal_bar/area).
        sheet: Sheet name for Excel files.
        output: PNG path (default: output/charts/<title>.png).
    """
    import pandas as pd

    if path:
        _, df = await _load(ctx, path, sheet)
    elif data:
        df = pd.DataFrame(data)
        df.columns = [str(c) for c in df.columns]
    else:
        raise ToolError("Provide `path` (a data file) or inline `data` records.")
    if df.empty:
        raise ToolError("No data to plot.")
    ys = [y] if isinstance(y, str) else list(y or [])
    xs = _col(df, x) if x else None
    ys = [_col(df, c) for c in ys]
    sc = _col(df, series) if series else None
    srt = sort if sort != "auto" else ("value" if chart_type in ("bar", "horizontal_bar", "pie", "donut") else "x")
    lim = limit or (8 if chart_type in ("pie", "donut") else 30 if chart_type in ("bar", "horizontal_bar") else None)
    target = O.ws_output(ctx, output, default_name=title or f"{chart_type}-chart", ext=".png", folder="output/charts",
                         overwrite=bool(output))
    info = await asyncio.to_thread(render_chart, df, target, chart_type=chart_type, x=xs, ys=ys, title=title,
                                   subtitle=subtitle, x_label=x_label, y_label=y_label, series_col=sc, agg=agg,
                                   sort=srt, limit=lim, stacked=stacked)
    return await O.produced(ctx, target, {"chart_type": chart_type, **info},
                            display=f"Chart saved: {O.vpath(ctx, target)}")

