"""Sheets group: Excel workbooks (.xlsx/.xlsm) and CSV files — create, read ranges, write, append, find, filter,
sort, formulas, calculations, sheet management, native Excel charts, pivots, formatting and CSV export.

Note on formulas: openpyxl stores formulas but does not compute them. Excel/LibreOffice compute them on open.
`sheet_calculate` evaluates simple aggregates (SUM/AVERAGE/MIN/MAX/COUNT/MEDIAN…) with pandas over the stored
values so agents can report numbers without a spreadsheet engine."""
from __future__ import annotations

import asyncio
import json
import re
from pathlib import Path
from typing import Any, Literal

from app.tools.base import Risk, ToolContext, ToolError, ToolGroup, ToolResult
from app.tools.builtin import _office as O

sheets = ToolGroup(
    "sheets", "Spreadsheets",
    "Work with Excel workbooks and CSV files: create styled sheets, read ranges, write and append rows, find "
    "values, filter and sort, set formulas, calculate aggregates, manage sheets, add native Excel charts, build "
    "pivot tables, format and export to CSV.",
    prefix="sheet", icon="sheet", category="data",
)

XLSX_EXTS = (".xlsx", ".xlsm")
CSV_EXTS = (".csv", ".tsv")
ALL_EXTS = XLSX_EXTS + CSV_EXTS


# ----------------------------------------------------------------------------------------- workbook wrapper
class Book:
    """openpyxl workbook for .xlsx, or a CSV file loaded into a one-sheet workbook (saved back as CSV)."""

    def __init__(self, path: Path, *, data_only: bool = False, new: bool = False):
        from openpyxl import Workbook, load_workbook

        self.path = path
        ext = path.suffix.lower()
        if ext == ".xls":
            raise ToolError("Legacy .xls files aren't supported — save as .xlsx first.")
        self.kind = "csv" if ext in CSV_EXTS else "xlsx"
        self.delimiter = "\t" if ext == ".tsv" else ","
        if new:
            self.wb = Workbook()
            if self.kind == "csv":
                self.wb.active.title = O.safe_sheet_name(path.stem)
            return
        if self.kind == "xlsx":
            try:
                self.wb = load_workbook(str(path), data_only=data_only, keep_vba=ext == ".xlsm")
            except Exception as e:  # noqa: BLE001
                raise ToolError(f"Could not open '{path.name}' as a workbook: {e}") from e
        else:
            self.wb = Workbook()
            ws = self.wb.active
            ws.title = O.safe_sheet_name(path.stem)
            for row in O.read_csv_rows(path):
                ws.append([O.clean_cell(v) for v in row])

    @property
    def names(self) -> list[str]:
        return [ws.title for ws in self.wb.worksheets]

    def sheet(self, name: str | None = None):
        if not name:
            return self.wb.active or self.wb.worksheets[0]
        for ws in self.wb.worksheets:
            if ws.title == name:
                return ws
        for ws in self.wb.worksheets:
            if ws.title.lower() == str(name).strip().lower():
                return ws
        raise ToolError(f"No sheet named '{name}'. Sheets: {', '.join(self.names)}.")

    def require_xlsx(self, what: str) -> None:
        if self.kind != "xlsx":
            raise ToolError(f"{what} needs an Excel workbook — CSV files hold a single plain table. "
                            "Convert with artifact_export (csv → xlsx) first.")

    def save(self, target: Path | None = None) -> None:
        target = target or self.path
        if target.suffix.lower() in CSV_EXTS:
            ws = self.wb.active
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            while rows and all(v is None or v == "" for v in rows[-1]):
                rows.pop()
            O.write_csv(target, None, rows, delimiter="\t" if target.suffix.lower() == ".tsv" else ",")
        else:
            self.wb.save(str(target))


def _split_sheet_ref(cell_range: str | None, sheet: str | None) -> tuple[str | None, str | None]:
    if cell_range and "!" in cell_range:
        s, _, r = cell_range.rpartition("!")
        return s.strip().strip("'"), r
    return sheet, cell_range


def _bounds(ws, cell_range: str | None) -> tuple[int, int, int, int]:
    """(min_col, min_row, max_col, max_row) of an A1 range (whole used area by default)."""
    max_c, max_r = max(1, ws.max_column or 1), max(1, ws.max_row or 1)
    if not cell_range:
        return 1, 1, max_c, max_r
    from openpyxl.utils.cell import range_boundaries

    try:
        c1, r1, c2, r2 = range_boundaries(cell_range.strip().replace("$", "").upper())
    except (ValueError, TypeError) as e:
        raise ToolError(f"Invalid range '{cell_range}'. Use A1 notation like 'A1:D20', 'B:B' or '2:10'.") from e
    return c1 or 1, r1 or 1, c2 or max_c, r2 or max_r


def _col_index(ws, ref: Any, header_row: int = 1) -> int:
    from openpyxl.utils import column_index_from_string

    s = str(ref).strip()
    headers = {}
    for c in ws[header_row]:
        if c.value is not None and str(c.value).strip():
            headers.setdefault(str(c.value).strip().lower(), c.column)
    if s.lower() in headers:
        return headers[s.lower()]
    if re.fullmatch(r"[A-Za-z]{1,3}", s):
        return column_index_from_string(s.upper())
    if s.isdigit() and int(s) >= 1:
        return int(s)
    raise ToolError(f"Column '{ref}' not found. Header columns: {', '.join(list(headers)[:40]) or '(none)'}.")


def _sheet_df(ws, header_row: int = 1):
    import pandas as pd

    rows = list(ws.iter_rows(min_row=header_row, values_only=True))
    if not rows:
        return pd.DataFrame()
    width = len(rows[0])
    while width > 0 and all((len(r) < width or r[width - 1] in (None, "")) for r in rows):
        width -= 1
    header, seen = [], {}
    for i, h in enumerate(rows[0][:width]):
        name = str(h).strip() if h not in (None, "") else f"Column{i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        header.append(name)
    body = [list(r[:width]) for r in rows[1:] if any(v not in (None, "") for v in r[:width])]
    return pd.DataFrame(body, columns=header)


def _resolve_cols(df, refs: list[str] | str | None) -> list[str]:
    if refs is None:
        return []
    if isinstance(refs, str):
        refs = [r.strip() for r in refs.split(",") if r.strip()]
    lower = {c.lower(): c for c in df.columns}
    out = []
    for r in refs:
        key = str(r).strip()
        if key in df.columns:
            out.append(key)
        elif key.lower() in lower:
            out.append(lower[key.lower()])
        elif re.fullmatch(r"[A-Za-z]{1,3}", key):
            from openpyxl.utils import column_index_from_string

            idx = column_index_from_string(key.upper()) - 1
            if 0 <= idx < len(df.columns):
                out.append(df.columns[idx])
            else:
                raise ToolError(f"Column {key} is outside the table.")
        else:
            raise ToolError(f"Column '{r}' not found. Columns: {', '.join(map(str, df.columns[:40]))}.")
    return out


def _df_to_sheet(ws, df, *, style: bool = True) -> None:
    header = [str(c) for c in df.columns]
    rows = [[O.clean_cell(v) for v in r] for r in df.itertuples(index=False, name=None)]
    O.write_sheet(ws, header, rows, style=style)


def _rows_payload(rows: list[list[Any]], width: int | None = None) -> list[list[Any]]:
    return [[O.jsonable(v) for v in (r if width is None else r[:width])] for r in rows]


def _new_or_existing(ctx: ToolContext, path: str) -> tuple[Path, bool]:
    p = O.ws_path(ctx, path)
    if p.exists():
        if p.suffix.lower() not in ALL_EXTS:
            raise ToolError(f"Expected an .xlsx or .csv file, got '{p.name}'.")
        return p, False
    O.require_write(ctx)
    if p.suffix.lower() not in ALL_EXTS:
        p = O.ws_output(ctx, path, default_name=Path(path).stem, ext=".xlsx")
    return O.ws_path(ctx, O.vpath(ctx, p), for_write=True), True


def _input(ctx: ToolContext, path: str) -> Path:
    return O.ws_input(ctx, path, ALL_EXTS)


# ----------------------------------------------------------------------------------------- create / read
@sheets.tool(risk=Risk.MEDIUM, scope="write", timeout=120)
async def create(ctx: ToolContext, path: str, rows: Any = None, sheet_name: str = "Sheet1",
                 sheets: dict[str, Any] | None = None, header: bool = True, overwrite: bool = False) -> ToolResult:
    """Create a new styled spreadsheet (.xlsx, or .csv) from rows. Header row is bold with a tinted fill, frozen
    and filterable; decimals get thousands separators; columns are auto-sized. Strings starting with "=" become
    formulas.

    Args:
        path: Output path, e.g. "output/budget.xlsx" (bare names go to output/; .csv makes a CSV file).
        rows: Table data — list of row lists (first row = header), list of objects, CSV text or a markdown table.
        sheet_name: Name of the sheet when `rows` is used.
        sheets: Several sheets at once: {"Summary": rows, "Detail": rows}.
        header: Treat the first row as a header.
        overwrite: Replace an existing file instead of saving under a new name.
    """
    ext = ".csv" if str(path).lower().endswith(".csv") else ".tsv" if str(path).lower().endswith(".tsv") else ".xlsx"
    target = O.ws_output(ctx, path, default_name=Path(str(path)).stem or "sheet", ext=ext, overwrite=overwrite)
    spec: dict[str, tuple] = {}
    if sheets:
        for name, data in sheets.items():
            spec[str(name)] = O.table_from_any(data, header=header)
    else:
        spec[sheet_name or "Sheet1"] = O.table_from_any(rows, header=header)

    def run() -> dict:
        if ext in CSV_EXTS:
            hdr, body = next(iter(spec.values()))
            O.write_csv(target, hdr, body, delimiter="\t" if ext == ".tsv" else ",")
        else:
            O.write_xlsx(target, spec)
        return {"sheets": {k: len(v[1]) for k, v in spec.items()}}

    info = await asyncio.to_thread(run)
    return await O.produced(ctx, target, info, display=f"Created {O.vpath(ctx, target)}")


def _detect_header(rows: list[list[Any]]) -> bool:
    if not rows:
        return False
    first = [v for v in rows[0] if v not in (None, "")]
    if not first or len(first) < 0.5 * len(rows[0]):
        return False
    if not all(isinstance(v, str) for v in first):
        return False
    if len(rows) == 1:
        return True
    return any(not isinstance(v, str) and v is not None for r in rows[1:6] for v in r) or len(set(first)) == len(first)


@sheets.tool(risk=Risk.LOW, scope="read", timeout=60)
async def read(ctx: ToolContext, path: str, sheet: str | None = None, cell_range: str | None = None,
               max_rows: int = 100, header: Literal["auto", "yes", "no"] = "auto") -> dict:
    """Read cells from a spreadsheet: a sheet (default: active) and optional A1 range, with header detection.
    Formula cells show their last computed value (or the formula if never computed).

    Args:
        path: Workspace path of the .xlsx/.csv file.
        sheet: Sheet name (also accepted inside the range, e.g. "Data!A1:F50").
        cell_range: A1 range like "A1:F50", "B:D" or "2:20" (default: the used area).
        max_rows: Maximum data rows to return.
        header: Whether the first row of the range is a header.
    """
    p = _input(ctx, path)
    sheet, cell_range = _split_sheet_ref(cell_range, sheet)
    n = max(1, min(int(max_rows or 100), 5000))

    def run() -> dict:
        book = Book(p)
        ws = book.sheet(sheet)
        c1, r1, c2, r2 = _bounds(ws, cell_range)
        r2_cap = min(r2, r1 + n + 1)
        rows = [list(r) for r in ws.iter_rows(min_row=r1, max_row=r2_cap, min_col=c1, max_col=c2, values_only=True)]
        formulas = any(isinstance(v, str) and v.startswith("=") for r in rows for v in r)
        note = None
        if formulas and book.kind == "xlsx":
            vals_ws = Book(p, data_only=True).sheet(ws.title)
            vrows = [list(r) for r in vals_ws.iter_rows(min_row=r1, max_row=r2_cap, min_col=c1, max_col=c2,
                                                         values_only=True)]
            uncomputed = 0
            for i, r in enumerate(rows):
                for j, v in enumerate(r):
                    if isinstance(v, str) and v.startswith("="):
                        cached = vrows[i][j] if i < len(vrows) and j < len(vrows[i]) else None
                        if cached is not None:
                            rows[i][j] = cached
                        else:
                            uncomputed += 1
            if uncomputed:
                note = (f"{uncomputed} formula cell(s) have never been computed (shown as formulas). Use "
                        "sheet_calculate to evaluate simple aggregates.")
        while rows and all(v in (None, "") for v in rows[-1]):
            rows.pop()
        width = max([len(r) for r in rows] + [0])
        while width > 0 and all((len(r) < width or r[width - 1] in (None, "")) for r in rows):
            width -= 1
        has_header = header == "yes" or (header == "auto" and _detect_header([r[:width] for r in rows]))
        hdr = [O.jsonable(v) for v in rows[0][:width]] if (has_header and rows) else None
        data = rows[1:] if has_header else rows
        truncated = len(data) > n or r2 > r2_cap
        data = data[:n]
        from openpyxl.utils import get_column_letter

        total_rows = (r2 - r1 + 1) - (1 if has_header else 0)
        out = {"path": O.vpath(ctx, p), "sheet": ws.title, "sheets": book.names,
               "range": f"{get_column_letter(c1)}{r1}:{get_column_letter(max(c1, c1 + width - 1))}{r2}",
               "total_rows": max(0, total_rows), "total_cols": width, "header": hdr,
               "rows": _rows_payload(data, width), "truncated": truncated}
        if note:
            out["note"] = note
        return out

    return await asyncio.to_thread(run)


# ----------------------------------------------------------------------------------------- write / append / clear
@sheets.tool(risk=Risk.MEDIUM, scope="write", timeout=120)
async def write(ctx: ToolContext, path: str, values: Any = None, start_cell: str = "A1",
                cells: dict[str, Any] | None = None, sheet: str | None = None) -> ToolResult:
    """Write values into a spreadsheet: a 2-D block starting at `start_cell`, and/or individual cells. Creates
    the file (and sheet) if missing. Strings starting with "=" are stored as formulas.

    Args:
        path: Workspace path of the .xlsx/.csv file.
        values: 2-D list of rows to write (a flat list writes one row).
        start_cell: Top-left cell for `values`, e.g. "B3" (may include a sheet: "Data!B3").
        cells: Individual cells, e.g. {"B2": 1200, "C10": "=SUM(C2:C9)"}.
        sheet: Sheet name (created if it doesn't exist).
    """
    if values is None and not cells:
        raise ToolError("Provide `values` and/or `cells` to write.")
    p, is_new = _new_or_existing(ctx, path)
    sheet, start_cell = _split_sheet_ref(start_cell, sheet)

    def run() -> dict:
        from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

        book = Book(p, new=is_new)
        if sheet and sheet not in book.names and sheet.lower() not in [n.lower() for n in book.names]:
            book.require_xlsx("Adding a sheet")
            ws = book.wb.active if (is_new and book.wb.active.max_row == 1 and book.wb.active["A1"].value is None) \
                else book.wb.create_sheet()
            ws.title = O.safe_sheet_name(sheet, [n for n in book.names if n != ws.title])
        else:
            ws = book.sheet(sheet)
        written = 0
        span = None
        if values is not None:
            block = values
            if isinstance(block, str):
                _, block = O.table_from_any(block, header=False)
            if block and not isinstance(block[0], (list, tuple, dict)):
                block = [block]
            if block and isinstance(block[0], dict):
                hdr, body = O.table_from_any(block)
                block = [hdr] + body
            try:
                col_letter, row0 = coordinate_from_string((start_cell or "A1").replace("$", "").upper())
            except ValueError as e:
                raise ToolError(f"Invalid start cell '{start_cell}'.") from e
            col0 = column_index_from_string(col_letter)
            for i, row in enumerate(block):
                for j, v in enumerate(row):
                    ws.cell(row=row0 + i, column=col0 + j, value=O.clean_cell(O.coerce_scalar(v)))
                    written += 1
            from openpyxl.utils import get_column_letter

            width = max([len(r) for r in block] + [1])
            span = f"{col_letter}{row0}:{get_column_letter(col0 + width - 1)}{row0 + max(1, len(block)) - 1}"
        for ref, v in (cells or {}).items():
            try:
                ws[str(ref).replace("$", "").upper()] = O.clean_cell(O.coerce_scalar(v))
            except (ValueError, KeyError) as e:
                raise ToolError(f"Invalid cell reference '{ref}'.") from e
            written += 1
        book.save()
        return {"sheet": ws.title, "cells_written": written, **({"range": span} if span else {})}

    info = await asyncio.to_thread(run)
    return await O.produced(ctx, p, info, display=f"Wrote {info['cells_written']} cell(s) to {p.name}")


@sheets.tool(risk=Risk.MEDIUM, scope="write", timeout=120)
async def append(ctx: ToolContext, path: str, rows: Any, sheet: str | None = None) -> ToolResult:
    """Append rows to the end of a sheet (creates the file if missing). Objects are matched to header columns by
    name; unknown keys add new columns.

    Args:
        path: Workspace path of the .xlsx/.csv file.
        rows: Rows to append — list of row lists, list of objects, or CSV text (without header).
        sheet: Sheet name (default: active sheet).
    """
    p, is_new = _new_or_existing(ctx, path)

    def run() -> dict:
        book = Book(p, new=is_new)
        ws = book.sheet(sheet) if (sheet and not is_new) else book.wb.active
        if is_new and sheet:
            ws.title = O.safe_sheet_name(sheet)
        data = rows
        if isinstance(data, str):
            s = data.strip()
            if not s:
                raise ToolError("No rows to append.")
            try:
                data = json.loads(s) if s[0] in "[{" else O.table_from_any(s, header=False)[1]
            except json.JSONDecodeError as e:
                raise ToolError(f"`rows` is not valid JSON: {e}") from e
        if isinstance(data, dict):
            data = [data]
        if not data:
            raise ToolError("No rows to append.")
        empty = ws.max_row == 1 and all(c.value is None for c in ws[1])
        if isinstance(data[0], dict):
            header = [] if empty else [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
            lower = [h.lower() for h in header]
            for rec in data:
                for k in rec:
                    if str(k).lower() not in lower:
                        header.append(str(k))
                        lower.append(str(k).lower())
            for j, h in enumerate(header, 1):
                ws.cell(row=1, column=j, value=h)
            if empty and book.kind == "xlsx":
                O.xlsx_header_style(ws, 1, len(header))
            start = 2 if empty else ws.max_row + 1
            for i, rec in enumerate(data):
                low = {str(k).lower(): v for k, v in rec.items()}
                for j, h in enumerate(header, 1):
                    ws.cell(row=start + i, column=j, value=O.clean_cell(O.coerce_scalar(low.get(h.lower()))))
        else:
            start = 1 if empty else ws.max_row + 1
            for i, r in enumerate(data):
                vals = r if isinstance(r, (list, tuple)) else [r]
                for j, v in enumerate(vals, 1):
                    ws.cell(row=start + i, column=j, value=O.clean_cell(O.coerce_scalar(v)))
        book.save()
        return {"sheet": ws.title, "appended": len(data), "total_rows": ws.max_row}

    info = await asyncio.to_thread(run)
    return await O.produced(ctx, p, info, display=f"Appended {info['appended']} row(s) to {p.name}")


@sheets.tool(risk=Risk.MEDIUM, scope="write", timeout=60)
async def clear(ctx: ToolContext, path: str, cell_range: str | None = None, sheet: str | None = None) -> dict:
    """Clear cell values in a range (formatting is kept), or empty the whole sheet when no range is given.

    Args:
        path: Workspace path of the .xlsx/.csv file.
        cell_range: A1 range to clear, e.g. "B2:D20" (default: entire sheet).
        sheet: Sheet name.
    """
    p = _input(ctx, path)
    O.require_write(ctx)
    sheet, cell_range = _split_sheet_ref(cell_range, sheet)

    def run() -> dict:
        book = Book(p)
        ws = book.sheet(sheet)
        if not cell_range:
            n = ws.max_row
            ws.delete_rows(1, n)
            book.save()
            return {"sheet": ws.title, "cleared": "all", "rows_removed": n}
        c1, r1, c2, r2 = _bounds(ws, cell_range)
        count = 0
        for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
            for cell in row:
                if cell.value is not None and type(cell).__name__ != "MergedCell":
                    cell.value = None
                    count += 1
        book.save()
        return {"sheet": ws.title, "range": cell_range, "cells_cleared": count}

    return {"path": O.vpath(ctx, p), **(await asyncio.to_thread(run))}


@sheets.tool(risk=Risk.LOW, scope="read", timeout=60)
async def find(ctx: ToolContext, path: str, query: str, sheet: str | None = None, match_case: bool = False,
               exact: bool = False, regex: bool = False, max_results: int = 50) -> dict:
    """Find cells whose value contains (or equals) a search term; returns sheet, cell reference and value.

    Args:
        path: Workspace path of the .xlsx/.csv file.
        query: Text or number to look for.
        sheet: Limit to one sheet (default: all sheets).
        match_case: Case-sensitive matching.
        exact: Whole-cell match instead of "contains".
        regex: Treat query as a regular expression.
        max_results: Maximum matches to return.
    """
    p = _input(ctx, path)
    flags = 0 if match_case else re.IGNORECASE
    try:
        if regex:
            pat = re.compile(query if not exact else f"^(?:{query})$", flags)
        else:
            pat = re.compile(("^" + re.escape(query) + "$") if exact else re.escape(query), flags)
    except re.error as e:
        raise ToolError(f"Invalid regular expression: {e}") from e
    limit = max(1, min(int(max_results or 50), 1000))

    def run() -> dict:
        book = Book(p, data_only=True)
        matches, total = [], 0
        for ws in ([book.sheet(sheet)] if sheet else book.wb.worksheets):
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if v is None:
                        continue
                    s = O._cell_text(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else str(v)
                    raw = str(v)
                    if pat.search(s) or pat.search(raw):
                        total += 1
                        if len(matches) < limit:
                            matches.append({"sheet": ws.title, "cell": cell.coordinate, "value": O.jsonable(v)})
        return {"path": O.vpath(ctx, p), "query": query, "total": total, "matches": matches,
                "truncated": total > len(matches)}

    return await asyncio.to_thread(run)


# ----------------------------------------------------------------------------------------- filter / sort
_OPS = {"=": "eq", "==": "eq", "equals": "eq", "eq": "eq", "!=": "ne", "<>": "ne", "not_equals": "ne", "ne": "ne",
        ">": "gt", "gt": "gt", ">=": "gte", "gte": "gte", "<": "lt", "lt": "lt", "<=": "lte", "lte": "lte",
        "contains": "contains", "not_contains": "not_contains", "starts_with": "starts_with",
        "startswith": "starts_with", "ends_with": "ends_with", "endswith": "ends_with", "in": "in",
        "not_in": "not_in", "between": "between", "empty": "empty", "is_empty": "empty", "not_empty": "not_empty",
        "is_not_empty": "not_empty"}


def apply_filter(df, column: str, operator: str, value: Any):
    import pandas as pd

    op = _OPS.get(str(operator).strip().lower())
    if op is None:
        raise ToolError(f"Unknown operator '{operator}'. Use one of: {', '.join(sorted(set(_OPS.values())))}.")
    s = df[column]
    num = pd.to_numeric(s, errors="coerce")
    text = s.astype(str).str.strip()

    def as_num(v: Any) -> float | None:
        try:
            return float(str(v).replace(",", ""))
        except (TypeError, ValueError):
            return None

    if op in ("gt", "gte", "lt", "lte"):
        v = as_num(value)
        if v is not None and num.notna().any():
            cmp = {"gt": num > v, "gte": num >= v, "lt": num < v, "lte": num <= v}[op]
        else:
            dt = pd.to_datetime(s, errors="coerce")
            dv = pd.to_datetime(value, errors="coerce")
            if pd.isna(dv):
                raise ToolError(f"'{value}' is not a number or date for operator '{operator}'.")
            cmp = {"gt": dt > dv, "gte": dt >= dv, "lt": dt < dv, "lte": dt <= dv}[op]
        return df[cmp.fillna(False)]
    if op == "between":
        if not isinstance(value, (list, tuple)) or len(value) != 2:
            raise ToolError("'between' needs value=[low, high].")
        lo, hi = as_num(value[0]), as_num(value[1])
        if lo is None or hi is None:
            dt = pd.to_datetime(s, errors="coerce")
            return df[((dt >= pd.to_datetime(value[0])) & (dt <= pd.to_datetime(value[1]))).fillna(False)]
        return df[((num >= lo) & (num <= hi)).fillna(False)]
    if op in ("eq", "ne"):
        v = as_num(value)
        if v is not None and num.notna().any():
            mask = (num == v) | (text.str.lower() == str(value).strip().lower())
        else:
            mask = text.str.lower() == str(value).strip().lower()
        return df[mask] if op == "eq" else df[~mask]
    if op in ("in", "not_in"):
        vals = value if isinstance(value, (list, tuple)) else [x.strip() for x in str(value).split(",")]
        lv = {str(x).strip().lower() for x in vals}
        mask = text.str.lower().isin(lv)
        return df[mask] if op == "in" else df[~mask]
    if op == "empty":
        return df[s.isna() | (text == "") | (text.str.lower() == "nan")]
    if op == "not_empty":
        return df[~(s.isna() | (text == "") | (text.str.lower() == "nan"))]
    lowv = str(value or "").lower()
    low = text.str.lower()
    mask = {"contains": low.str.contains(lowv, regex=False), "not_contains": ~low.str.contains(lowv, regex=False),
            "starts_with": low.str.startswith(lowv), "ends_with": low.str.endswith(lowv)}[op]
    return df[mask.fillna(False)]


async def _write_result(ctx: ToolContext, p: Path, book: Book, df, *, new_sheet: str | None, default_sheet: str,
                        output: str | None, default_stem: str) -> tuple[Path, dict]:
    """Put a derived table into a new sheet (xlsx) or a new file (csv source / explicit output)."""
    if output or book.kind == "csv":
        ext = ".xlsx" if (output and str(output).lower().endswith(".xlsx")) else ".csv"
        target = O.ws_output(ctx, output, default_name=default_stem, ext=ext, overwrite=bool(output))

        def save_file() -> None:
            header = [str(c) for c in df.columns]
            rows = [list(r) for r in df.itertuples(index=False, name=None)]
            if ext == ".xlsx":
                O.write_xlsx(target, {new_sheet or default_sheet: (header, rows)})
            else:
                O.write_csv(target, header, rows)

        await asyncio.to_thread(save_file)
        return target, {"output": O.vpath(ctx, target)}
    O.require_write(ctx)
    name = O.safe_sheet_name(new_sheet or default_sheet, book.names if not new_sheet else
                             [n for n in book.names if n.lower() != (new_sheet or "").lower()])

    def save_sheet() -> None:
        if name in book.names:
            del book.wb[name]
        ws = book.wb.create_sheet(name)
        _df_to_sheet(ws, df)
        book.save()

    await asyncio.to_thread(save_sheet)
    return p, {"new_sheet": name}


@sheets.tool(name="filter", risk=Risk.MEDIUM, scope="write", timeout=120)
async def filter_(ctx: ToolContext, path: str, column: str, operator: str, value: Any = None,
                  sheet: str | None = None, new_sheet: str | None = None, output: str | None = None) -> ToolResult:
    """Filter rows by a condition and write the matching rows into a new sheet (or a new file for CSV inputs).

    Args:
        path: Workspace path of the .xlsx/.csv file.
        column: Header name (or column letter) to test.
        operator: eq, ne, gt, gte, lt, lte, contains, not_contains, starts_with, ends_with, in, not_in, between,
            empty, not_empty (symbols like ">=" also work). Numbers and dates are compared as such.
        value: Value to compare with (a list for in/between).
        sheet: Source sheet (default: active).
        new_sheet: Name for the result sheet (default: "<sheet> filtered").
        output: Write the result to a separate .xlsx/.csv file instead.
    """
    p = _input(ctx, path)
    book = await asyncio.to_thread(Book, p, data_only=True)
    ws = book.sheet(sheet)
    df = await asyncio.to_thread(_sheet_df, ws)
    if df.empty:
        raise ToolError("The sheet has no data rows.")
    col = _resolve_cols(df, [column])[0]
    res = apply_filter(df, col, operator, value)
    if book.kind == "xlsx" and not output:
        book = await asyncio.to_thread(Book, p)  # keep formulas/formatting of the original sheets
    target, where = await _write_result(ctx, p, book, res, new_sheet=new_sheet, default_sheet=f"{ws.title} filtered",
                                        output=output, default_stem=f"{p.stem}-filtered")
    preview = _rows_payload([list(r) for r in res.head(5).itertuples(index=False, name=None)])
    return await O.produced(ctx, target, {"matched": len(res), "of": len(df), **where,
                                          "columns": [str(c) for c in res.columns], "preview": preview},
                            display=f"{len(res)} of {len(df)} rows match")


@sheets.tool(risk=Risk.MEDIUM, scope="write", timeout=120)
async def sort(ctx: ToolContext, path: str, by: list[str], descending: bool = False, sheet: str | None = None,
               new_sheet: str | None = None, in_place: bool = False, output: str | None = None) -> ToolResult:
    """Sort rows by one or more columns, writing the result to a new sheet (default), in place, or to a file.

    Args:
        path: Workspace path of the .xlsx/.csv file.
        by: Column names/letters to sort by, in priority order.
        descending: Sort largest/latest first.
        sheet: Source sheet (default: active).
        new_sheet: Name for the result sheet (default: "<sheet> sorted").
        in_place: Overwrite the source sheet's rows with the sorted order (values only).
        output: Write the result to a separate .xlsx/.csv file instead.
    """
    import pandas as pd

    p = _input(ctx, path)
    book = await asyncio.to_thread(Book, p, data_only=True)
    ws = book.sheet(sheet)
    df = await asyncio.to_thread(_sheet_df, ws)
    if df.empty:
        raise ToolError("The sheet has no data rows.")
    cols = _resolve_cols(df, by)
    keys = {}
    for c in cols:
        num = pd.to_numeric(df[c], errors="coerce")
        keys[c] = num if num.notna().sum() >= 0.8 * df[c].notna().sum() else df[c].astype(str).str.lower()
    order = pd.DataFrame(keys).sort_values(cols, ascending=not descending, kind="mergesort").index
    res = df.loc[order]
    if in_place and not output:
        O.require_write(ctx)

        def save() -> None:
            b = Book(p)
            w = b.sheet(ws.title)
            for i, row in enumerate(res.itertuples(index=False, name=None), start=2):
                for j, v in enumerate(row, start=1):
                    w.cell(row=i, column=j, value=O.clean_cell(v))
            b.save()

        await asyncio.to_thread(save)
        return await O.produced(ctx, p, {"sorted_rows": len(res), "by": cols, "sheet": ws.title},
                                display=f"Sorted {len(res)} rows in place")
    if book.kind == "xlsx" and not output:
        book = await asyncio.to_thread(Book, p)
    target, where = await _write_result(ctx, p, book, res, new_sheet=new_sheet, default_sheet=f"{ws.title} sorted",
                                        output=output, default_stem=f"{p.stem}-sorted")
    preview = _rows_payload([list(r) for r in res.head(5).itertuples(index=False, name=None)])
    return await O.produced(ctx, target, {"sorted_rows": len(res), "by": cols, **where, "preview": preview},
                            display=f"Sorted {len(res)} rows")


# ----------------------------------------------------------------------------------------- formulas / calculate
@sheets.tool(risk=Risk.MEDIUM, scope="write", timeout=60)
async def formula(ctx: ToolContext, path: str, cell: str, formula: str, sheet: str | None = None,
                  fill_down_to: int | None = None) -> ToolResult:
    """Set an Excel formula in a cell (optionally filled down a column with relative references adjusted).
    Excel computes it on open; use sheet_calculate to preview simple aggregates now.

    Args:
        path: Workspace path of the .xlsx file.
        cell: Target cell, e.g. "E2" (may include a sheet: "Data!E2").
        formula: Formula, e.g. "=C2*D2" or "=SUM(B2:B50)" (leading "=" optional).
        sheet: Sheet name.
        fill_down_to: Copy the formula down to this row number (e.g. 50).
    """
    p = _input(ctx, path)
    O.require_write(ctx)
    sheet, cell = _split_sheet_ref(cell, sheet)
    f = formula.strip()
    if not f.startswith("="):
        f = "=" + f

    def run() -> dict:
        from openpyxl.formula.translate import Translator
        from openpyxl.utils.cell import coordinate_from_string

        book = Book(p)
        book.require_xlsx("Formulas")
        ws = book.sheet(sheet)
        ref = cell.replace("$", "").upper()
        try:
            col, row = coordinate_from_string(ref)
        except ValueError as e:
            raise ToolError(f"Invalid cell '{cell}'.") from e
        ws[ref] = f
        filled = 1
        if fill_down_to and int(fill_down_to) > row:
            for r in range(row + 1, int(fill_down_to) + 1):
                ws[f"{col}{r}"] = Translator(f, origin=ref).translate_formula(f"{col}{r}")
                filled += 1
        book.save()
        return {"sheet": ws.title, "cell": ref, "formula": f, "cells": filled}

    info = await asyncio.to_thread(run)
    preview = None
    m = _AGG_RE.match(f)
    if m:
        try:
            preview = (await calculate(ctx, path=path, expression=f, sheet=info["sheet"]))["value"]
        except ToolError:
            preview = None
    if preview is not None:
        info["preview_value"] = preview
    info["note"] = "Excel computes formulas when the file is opened."
    return await O.produced(ctx, p, info, display=f"Set {info['cell']} = {f}")


_AGG_RE = re.compile(r"^\s*=?\s*(SUM|AVERAGE|AVG|MEAN|MIN|MAX|COUNT|COUNTA|MEDIAN|STDEV|STDEV\.S|PRODUCT)\s*\(\s*"
                     r"(?:'?([^'!]+)'?!)?\$?([A-Z]{1,3})\$?(\d+)?\s*(?::\s*\$?([A-Z]{1,3})\$?(\d+)?)?\s*\)\s*$", re.I)
_OPS_CALC = {"sum": "sum", "total": "sum", "average": "mean", "avg": "mean", "mean": "mean", "min": "min",
             "max": "max", "count": "count", "counta": "count", "median": "median", "std": "std", "stdev": "std",
             "stdev.s": "std", "nunique": "nunique", "distinct": "nunique", "product": "prod"}


def _aggregate(series, op: str):
    import pandas as pd

    if op in ("count", "nunique"):
        return int(series.notna().sum()) if op == "count" else int(series.nunique(dropna=True))
    num = pd.to_numeric(series, errors="coerce").dropna()
    if num.empty:
        return None
    v = getattr(num, op)()
    return O.jsonable(float(v))


@sheets.tool(risk=Risk.LOW, scope="read", timeout=60)
async def calculate(ctx: ToolContext, path: str, expression: str | None = None, column: str | None = None,
                    operation: str = "sum", sheet: str | None = None, group_by: str | None = None) -> dict:
    """Compute aggregates from a sheet's stored values with pandas (openpyxl can't evaluate formulas): either a
    simple Excel-style expression like "SUM(C2:C50)" / "AVERAGE(D:D)", or an operation over a column, optionally
    grouped by another column.

    Args:
        path: Workspace path of the .xlsx/.csv file.
        expression: SUM/AVERAGE/MIN/MAX/COUNT/COUNTA/MEDIAN/STDEV/PRODUCT over one A1 range.
        column: Column name/letter to aggregate (alternative to expression).
        operation: sum, mean/average, min, max, count, median, std, nunique.
        sheet: Sheet name.
        group_by: Column to group by (with `column`).
    """
    p = _input(ctx, path)

    def run() -> dict:
        import pandas as pd

        book = Book(p, data_only=True)
        if expression:
            m = _AGG_RE.match(expression)
            if not m:
                raise ToolError("Only single-range aggregates are supported, e.g. SUM(B2:B20), AVERAGE(C:C), "
                                "MAX(Data!D2:D99). Use analytics_query for anything more complex.")
            fn, sh, c1, r1, c2, r2 = m.groups()
            ws = book.sheet(sh or sheet)
            from openpyxl.utils import column_index_from_string

            ci1 = column_index_from_string(c1.upper())
            ci2 = column_index_from_string((c2 or c1).upper())
            ri1 = int(r1) if r1 else 1
            ri2 = int(r2) if r2 else (int(r1) if (r1 and not c2) else ws.max_row)
            vals = [v for row in ws.iter_rows(min_row=ri1, max_row=ri2, min_col=ci1, max_col=ci2, values_only=True)
                    for v in row]
            op = _OPS_CALC[fn.lower()]
            if fn.upper() == "COUNT":
                value = sum(1 for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool))
            elif fn.upper() == "COUNTA":
                value = sum(1 for v in vals if v not in (None, ""))
            else:
                nums = [float(v) for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool)]
                if not nums:
                    value = None
                else:
                    value = _aggregate(pd.Series(nums), op)
            formulas = 0
            if book.kind == "xlsx":
                fws = Book(p).sheet(ws.title)
                formulas = sum(1 for row in fws.iter_rows(min_row=ri1, max_row=ri2, min_col=ci1, max_col=ci2,
                                                          values_only=True)
                               for v in row if isinstance(v, str) and v.startswith("="))
            uncomputed = sum(1 for v in vals if v is None) if formulas else 0
            out = {"expression": expression.strip().lstrip("="), "sheet": ws.title, "value": value,
                   "cells": len(vals)}
            if formulas and uncomputed:
                out["note"] = (f"{min(formulas, uncomputed)} formula cell(s) in the range have no stored value and "
                               "were ignored (they are computed when opened in Excel).")
            return out
        if not column:
            raise ToolError("Provide an `expression` like SUM(B2:B20) or a `column` + `operation`.")
        op = _OPS_CALC.get(operation.strip().lower())
        if op is None:
            raise ToolError(f"Unknown operation '{operation}'. Use sum, mean, min, max, count, median, std, nunique.")
        ws = book.sheet(sheet)
        df = _sheet_df(ws)
        col = _resolve_cols(df, [column])[0]
        if group_by:
            g = _resolve_cols(df, [group_by])[0]
            if op not in ("count", "nunique"):
                df[col] = pd.to_numeric(df[col], errors="coerce")
            res = df.groupby(g, dropna=False)[col].agg("count" if op == "count" else op).reset_index()
            res = res.sort_values(col, ascending=False)
            return {"sheet": ws.title, "operation": op, "column": col, "group_by": g, "groups": len(res),
                    "rows": _rows_payload([list(r) for r in res.head(200).itertuples(index=False, name=None)])}
        return {"sheet": ws.title, "operation": op, "column": col, "value": _aggregate(df[col], op),
                "rows_considered": int(df[col].notna().sum())}

    return await asyncio.to_thread(run)


# ----------------------------------------------------------------------------------------- sheet management
@sheets.tool(risk=Risk.MEDIUM, scope="write", timeout=60)
async def create_sheet(ctx: ToolContext, path: str, name: str, rows: Any = None, index: int | None = None) -> dict:
    """Add a new sheet to a workbook, optionally filled with rows (styled header).

    Args:
        path: Workspace path of the .xlsx file.
        name: New sheet name.
        rows: Optional data (list of row lists, list of objects, CSV text).
        index: Position (0 = first; default: last).
    """
    p = _input(ctx, path)
    O.require_write(ctx)

    def run() -> dict:
        book = Book(p)
        book.require_xlsx("Multiple sheets")
        if name.lower() in [n.lower() for n in book.names]:
            raise ToolError(f"A sheet named '{name}' already exists.")
        ws = book.wb.create_sheet(O.safe_sheet_name(name, book.names), index)
        if rows is not None:
            hdr, body = O.table_from_any(rows)
            O.write_sheet(ws, hdr, body)
        book.save()
        return {"path": O.vpath(ctx, p), "sheet": ws.title, "sheets": book.names}

    return await asyncio.to_thread(run)


@sheets.tool(risk=Risk.MEDIUM, scope="write", timeout=60)
async def rename_sheet(ctx: ToolContext, path: str, old_name: str, new_name: str) -> dict:
    """Rename a sheet.

    Args:
        path: Workspace path of the .xlsx file.
        old_name: Current sheet name.
        new_name: New sheet name.
    """
    p = _input(ctx, path)
    O.require_write(ctx)

    def run() -> dict:
        book = Book(p)
        book.require_xlsx("Renaming sheets")
        ws = book.sheet(old_name)
        others = [n for n in book.names if n != ws.title]
        if new_name.lower() in [n.lower() for n in others]:
            raise ToolError(f"A sheet named '{new_name}' already exists.")
        ws.title = O.safe_sheet_name(new_name)
        book.save()
        return {"path": O.vpath(ctx, p), "sheets": book.names}

    return await asyncio.to_thread(run)


@sheets.tool(risk=Risk.MEDIUM, scope="write", timeout=60)
async def delete_sheet(ctx: ToolContext, path: str, name: str) -> dict:
    """Delete a sheet from a workbook (the last remaining sheet can't be deleted).

    Args:
        path: Workspace path of the .xlsx file.
        name: Sheet to delete.
    """
    p = _input(ctx, path)
    O.require_write(ctx)
    if O.vpath(ctx, p).startswith("shared") and not ctx.is_admin:
        raise ToolError("Only admins can delete data from shared files.")

    def run() -> dict:
        book = Book(p)
        book.require_xlsx("Deleting sheets")
        ws = book.sheet(name)
        if len(book.names) == 1:
            raise ToolError("Can't delete the only sheet in a workbook.")
        del book.wb[ws.title]
        book.save()
        return {"path": O.vpath(ctx, p), "deleted": ws.title, "sheets": book.names}

    return await asyncio.to_thread(run)


# ----------------------------------------------------------------------------------------- charts / pivot / format
@sheets.tool(risk=Risk.MEDIUM, scope="write", timeout=120)
async def chart(ctx: ToolContext, path: str, chart_type: Literal["column", "bar", "line", "pie", "area"] = "column",
                data_columns: list[str] | None = None, category_column: str | None = None, sheet: str | None = None,
                title: str | None = None, anchor: str | None = None, x_title: str | None = None,
                y_title: str | None = None, stacked: bool = False, last_row: int | None = None) -> ToolResult:
    """Add a native Excel chart (editable in Excel) built from sheet columns. Assumes a header in row 1.

    Args:
        path: Workspace path of the .xlsx file.
        chart_type: column (vertical bars), bar (horizontal), line, pie or area.
        data_columns: Value columns (header names or letters). Default: all numeric columns.
        category_column: Column with the labels/x-axis (default: first column).
        sheet: Sheet with the data (default: active).
        title: Chart title.
        anchor: Top-left cell for the chart, e.g. "H2" (default: right of the data).
        x_title: X-axis title.
        y_title: Y-axis title.
        stacked: Stack series (column/bar/area).
        last_row: Last data row to include (default: last used row).
    """
    p = _input(ctx, path)
    O.require_write(ctx)

    def run() -> dict:
        from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, Reference
        from openpyxl.chart.series import DataPoint
        from openpyxl.utils import get_column_letter

        book = Book(p)
        book.require_xlsx("Charts")
        ws = book.sheet(sheet)
        last = int(last_row) if last_row else ws.max_row
        while last > 2 and all(c.value in (None, "") for c in ws[last]):
            last -= 1
        if last < 2:
            raise ToolError("The sheet needs a header row and at least one data row.")
        cat = _col_index(ws, category_column) if category_column else 1
        if data_columns:
            cols = [_col_index(ws, c) for c in data_columns]
        else:
            cols = []
            for c in range(1, ws.max_column + 1):
                if c == cat:
                    continue
                vals = [ws.cell(row=r, column=c).value for r in range(2, min(last, 200) + 1)]
                nums = [v for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool)]
                if nums and len(nums) >= 0.6 * len([v for v in vals if v is not None]):
                    cols.append(c)
            if not cols:
                raise ToolError("No numeric columns found — pass data_columns explicitly.")
        if chart_type == "pie":
            cols = cols[:1]
            ch = PieChart()
        elif chart_type == "line":
            ch = LineChart()
        elif chart_type == "area":
            ch = AreaChart()
            if stacked:
                ch.grouping = "stacked"
        else:
            ch = BarChart()
            ch.type = "bar" if chart_type == "bar" else "col"
            ch.gapWidth = 80
            if stacked:
                ch.grouping = "stacked"
                ch.overlap = 100
        for c in cols:
            ch.add_data(Reference(ws, min_col=c, min_row=1, max_row=last), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=cat, min_row=2, max_row=last))
        ch.title = title or (str(ws.cell(row=1, column=cols[0]).value) if len(cols) == 1 else None)
        ch.height, ch.width = 9.5, 18
        pal = O.CHART_PALETTE
        for i, s in enumerate(ch.series):
            color = pal[i % len(pal)]
            if chart_type == "pie":
                for k in range(last - 1):
                    pt = DataPoint(idx=k)
                    pt.graphicalProperties.solidFill = pal[k % len(pal)]
                    s.dPt.append(pt)
            elif chart_type == "line":
                s.graphicalProperties.line.solidFill = color
                s.graphicalProperties.line.width = 28575
                s.smooth = False
            else:
                s.graphicalProperties.solidFill = color
                s.graphicalProperties.line.solidFill = color
        if chart_type != "pie":
            ch.x_axis.delete = False
            ch.y_axis.delete = False
            if x_title:
                ch.x_axis.title = x_title
            if y_title:
                ch.y_axis.title = y_title
            ch.y_axis.number_format = "#,##0"
            if len(cols) == 1:
                ch.legend = None
            else:
                ch.legend.position = "b"
        else:
            ch.legend.position = "r"
        where = anchor or f"{get_column_letter(ws.max_column + 2)}2"
        ws.add_chart(ch, where)
        book.save()
        return {"sheet": ws.title, "chart_type": chart_type, "anchor": where, "series": len(cols),
                "points": last - 1}

    info = await asyncio.to_thread(run)
    return await O.produced(ctx, p, info, display=f"Added a {chart_type} chart to {p.name}")


_AGGFUNCS = {"sum": "sum", "mean": "mean", "average": "mean", "avg": "mean", "count": "count", "min": "min",
             "max": "max", "median": "median", "nunique": "nunique", "distinct": "nunique"}


@sheets.tool(risk=Risk.MEDIUM, scope="write", timeout=120)
async def pivot(ctx: ToolContext, path: str, index: list[str], values: list[str], columns: list[str] | None = None,
                aggfunc: str = "sum", sheet: str | None = None, new_sheet: str | None = None, totals: bool = True,
                output: str | None = None) -> ToolResult:
    """Build a pivot table (pandas) and write it to a new sheet (or a new file for CSV inputs).

    Args:
        path: Workspace path of the .xlsx/.csv file.
        index: Row grouping column(s), e.g. ["Region"].
        values: Column(s) to aggregate, e.g. ["Revenue"].
        columns: Optional column grouping, e.g. ["Quarter"].
        aggfunc: sum, mean, count, min, max, median or nunique.
        sheet: Source sheet (default: active).
        new_sheet: Result sheet name (default: "Pivot").
        totals: Add a Total row/column.
        output: Write the result to a separate .xlsx/.csv file instead.
    """
    import pandas as pd

    p = _input(ctx, path)
    agg = _AGGFUNCS.get(aggfunc.strip().lower())
    if agg is None:
        raise ToolError(f"Unknown aggfunc '{aggfunc}'. Use sum, mean, count, min, max, median or nunique.")
    book = await asyncio.to_thread(Book, p, data_only=True)
    ws = book.sheet(sheet)
    df = await asyncio.to_thread(_sheet_df, ws)
    if df.empty:
        raise ToolError("The sheet has no data rows.")
    idx = _resolve_cols(df, index)
    vals = _resolve_cols(df, values)
    cols = _resolve_cols(df, columns) if columns else []
    if agg not in ("count", "nunique"):
        for v in vals:
            df[v] = pd.to_numeric(df[v], errors="coerce")

    def build():
        pt = pd.pivot_table(df, index=idx, columns=cols or None, values=vals, aggfunc=agg, fill_value=0,
                            margins=totals, margins_name="Total", observed=True)
        pt = pt.reset_index()
        if isinstance(pt.columns, pd.MultiIndex):
            names = []
            for tup in pt.columns:
                parts = [str(x) for x in tup if str(x) not in ("", "nan")]
                if len(vals) == 1 and len(parts) > 1 and parts[0] == vals[0]:
                    parts = parts[1:]
                names.append(" · ".join(parts) or "value")
            pt.columns = names
        return pt

    try:
        pt = await asyncio.to_thread(build)
    except (TypeError, ValueError) as e:
        raise ToolError(f"Pivot failed: {e}") from e
    if book.kind == "xlsx" and not output:
        book = await asyncio.to_thread(Book, p)
    target, where = await _write_result(ctx, p, book, pt, new_sheet=new_sheet, default_sheet="Pivot", output=output,
                                        default_stem=f"{p.stem}-pivot")
    preview = _rows_payload([list(r) for r in pt.head(15).itertuples(index=False, name=None)])
    return await O.produced(ctx, target, {"rows": len(pt), "columns": [str(c) for c in pt.columns], **where,
                                          "preview": preview}, display=f"Pivot with {len(pt)} rows")


@sheets.tool(name="format", risk=Risk.MEDIUM, scope="write", timeout=120)
async def format_(ctx: ToolContext, path: str, sheet: str | None = None, header: bool = True,
                  number_formats: dict[str, str] | None = None, column_widths: dict[str, float] | None = None,
                  autofit: bool = True, freeze: str | None = "A2", autofilter: bool = True, wrap_text: bool = False,
                  banded_rows: bool = False) -> ToolResult:
    """Format a sheet for presentation: bold tinted header, number formats per column, column widths (or
    auto-fit), freeze panes, auto-filter, text wrapping and banded rows.

    Args:
        path: Workspace path of the .xlsx file.
        sheet: Sheet name (default: active).
        header: Style row 1 as a header.
        number_formats: Column → Excel format, e.g. {"Revenue": "$#,##0", "Margin": "0.0%", "Date": "yyyy-mm-dd"}.
        column_widths: Column → width in characters, e.g. {"A": 30, "Notes": 60}.
        autofit: Auto-size columns not given in column_widths.
        freeze: Cell to freeze panes at ("A2" = keep header visible; null to unfreeze).
        autofilter: Add filter buttons to the header.
        wrap_text: Wrap long text in cells.
        banded_rows: Light zebra striping on data rows.
    """
    p = _input(ctx, path)
    O.require_write(ctx)

    def run() -> dict:
        from openpyxl.styles import Alignment, PatternFill
        from openpyxl.utils import get_column_letter

        book = Book(p)
        book.require_xlsx("Formatting")
        ws = book.sheet(sheet)
        done = []
        if header and ws.max_row >= 1:
            O.xlsx_header_style(ws, 1, ws.max_column)
            done.append("header")
        if autofit:
            O.xlsx_autofit(ws)
            done.append("autofit")
        for ref, fmt in (number_formats or {}).items():
            c = _col_index(ws, ref)
            for r in range(2 if header else 1, ws.max_row + 1):
                ws.cell(row=r, column=c).number_format = fmt
            done.append(f"format {ref}")
        for ref, width in (column_widths or {}).items():
            ws.column_dimensions[get_column_letter(_col_index(ws, ref))].width = float(width)
            done.append(f"width {ref}")
        ws.freeze_panes = freeze.upper() if freeze else None
        if autofilter and ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
            done.append("autofilter")
        if wrap_text or banded_rows:
            band = PatternFill("solid", fgColor=O.BAND)
            for r in range(2 if header else 1, ws.max_row + 1):
                for cell in ws[r]:
                    if wrap_text:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
                    if banded_rows and r % 2 == 1:
                        cell.fill = band
            done += (["wrap"] if wrap_text else []) + (["banded"] if banded_rows else [])
        book.save()
        return {"sheet": ws.title, "applied": done}

    info = await asyncio.to_thread(run)
    return await O.produced(ctx, p, info, display=f"Formatted {p.name}")


@sheets.tool(risk=Risk.MEDIUM, scope="write", timeout=120)
async def export(ctx: ToolContext, path: str, sheet: str | None = None, output: str | None = None,
                 format: Literal["csv", "tsv", "json", "md"] = "csv") -> ToolResult:
    """Export one sheet to CSV (default), TSV, JSON records or a markdown table.

    Args:
        path: Workspace path of the .xlsx/.csv file.
        sheet: Sheet to export (default: active).
        output: Output path (default: output/<file>-<sheet>.<ext>).
        format: Output format.
    """
    p = _input(ctx, path)
    ext = "." + format

    def load():
        book = Book(p, data_only=True)
        ws = book.sheet(sheet)
        return ws.title, [list(r) for r in ws.iter_rows(values_only=True)]

    title, rows = await asyncio.to_thread(load)
    while rows and all(v in (None, "") for v in rows[-1]):
        rows.pop()
    target = O.ws_output(ctx, output, default_name=f"{p.stem}-{title}", ext=ext, overwrite=bool(output))

    def save() -> None:
        if format in ("csv", "tsv"):
            O.write_csv(target, None, rows, delimiter="\t" if format == "tsv" else ",")
        elif format == "json":
            hdr = [str(h) if h is not None else f"col{i + 1}" for i, h in enumerate(rows[0])] if rows else []
            recs = [{h: O.jsonable(v) for h, v in zip(hdr, r)} for r in rows[1:]]
            target.write_text(json.dumps(recs, ensure_ascii=False, indent=2), encoding="utf-8")
        else:
            target.write_text(O.md_table([O.jsonable(v) for v in rows[0]], rows[1:]) if rows else "", encoding="utf-8")

    await asyncio.to_thread(save)
    return await O.produced(ctx, target, {"sheet": title, "rows": max(0, len(rows) - 1), "format": format},
                            display=f"Exported {title} → {target.name}")
