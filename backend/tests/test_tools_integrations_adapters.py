"""Adapter-level request/response shapes for the remaining integrations (Graph, OneDrive, Google Calendar/Drive,
Dropbox, Box, Notion, Confluence, Linear, Asana, Salesforce, Teams, Twilio, webhooks, GCP, Azure) plus the
model-assisted tools. Shares the mocked-HTTP plumbing of test_tools_integrations.py.
"""
from __future__ import annotations

import base64
import hashlib
import hmac
import io
import json
from typing import Any

import httpx
import pytest

from app.tools.base import ToolError, registry
from app.tools.builtin import _oauth
from app.tools.connectors import connector_kinds
from tests.test_tools_integrations import (  # noqa: F401  (fixtures are used by name)
    GOOGLE,
    _ORG,
    body_json,
    call,
    ctx,
    form,
    google_token_route,
    http,
    integ_org,
    query,
)

MS = {"tenant_id": "contoso.onmicrosoft.com", "client_id": "app-id", "client_secret": "ms-secret",
      "refresh_token": "ms-refresh"}


def ms_token_route(http, token: str = "ms-tok", refresh: str | None = None) -> None:
    payload: dict[str, Any] = {"access_token": token, "expires_in": 3600, "token_type": "Bearer"}
    if refresh:
        payload["refresh_token"] = refresh
    http.add("POST", r"login\.microsoftonline\.com/.+/oauth2/v2\.0/token", json=payload)


async def _add_connector(org_id: str, kind: str, name: str, config: dict, secrets: dict) -> str:
    from app.core.db import session_scope
    from app.core.security import encrypt_json
    from app.models import Connector

    async with session_scope() as db:
        row = Connector(org_id=org_id, kind=kind, name=name, config=config, secrets_enc=encrypt_json(secrets))
        db.add(row)
        await db.flush()
        return row.id


async def _drop_connectors(org_id: str) -> None:
    from sqlalchemy import delete

    from app.core.db import session_scope
    from app.models import Connector

    async with session_scope() as db:
        await db.execute(delete(Connector).where(Connector.org_id == org_id))


# =========================================================================== Microsoft Graph mail
async def test_graph_send_mail_and_rotated_refresh_token_is_persisted(ctx, http):
    from app.core.db import session_scope
    from app.core.security import decrypt_json
    from app.models import Connector

    cid = await _add_connector(ctx.org_id, "microsoft365", "Contoso 365",
                               {"tenant_id": MS["tenant_id"], "client_id": MS["client_id"]},
                               {"client_secret": MS["client_secret"], "refresh_token": MS["refresh_token"]})
    try:
        ms_token_route(http, refresh="ms-refresh-rotated")
        http.add("POST", r"graph\.microsoft\.com/v1\.0/me/sendMail", status=202)
        res = await call("email_send", ctx, to=["Jane Doe <jane@example.com>"], subject="Plan",
                         body="**Bold** plan", body_format="markdown")
        assert res["sent"] and res["account"] == "Microsoft 365"
        tok = form(http.one("POST", r"oauth2/v2\.0/token"))
        assert tok == {"grant_type": "refresh_token", "client_id": "app-id", "refresh_token": "ms-refresh",
                       "client_secret": "ms-secret", "scope": "https://graph.microsoft.com/.default offline_access"}
        send = http.one("POST", r"/me/sendMail")
        assert send.headers["Authorization"] == "Bearer ms-tok"
        payload = body_json(send)
        assert payload["saveToSentItems"] is True
        msg = payload["message"]
        assert msg["subject"] == "Plan" and msg["body"]["contentType"] == "HTML"
        assert "<strong>Bold</strong>" in msg["body"]["content"]
        assert msg["toRecipients"] == [{"emailAddress": {"address": "jane@example.com", "name": "Jane Doe"}}]
        async with session_scope() as db:  # the rotated refresh token was stored back (encrypted)
            row = await db.get(Connector, cid)
            assert decrypt_json(row.secrets_enc) == {"client_secret": "ms-secret", "refresh_token": "ms-refresh-rotated"}
    finally:
        await _drop_connectors(ctx.org_id)


async def test_graph_search_filter_rules_and_app_only_user_path(ctx, http):
    ctx.extra["connectors"]["microsoft365"] = dict(MS)
    ms_token_route(http)
    http.add("GET", r"/me/mailFolders/inbox/messages", json={"value": [{
        "id": "AAMk/1+x=", "conversationId": "conv-1", "subject": "Budget", "isRead": False,
        "from": {"emailAddress": {"name": "Ana", "address": "ana@client.com"}},
        "toRecipients": [{"emailAddress": {"address": "me@contoso.com"}}],
        "receivedDateTime": "2026-10-05T09:00:00Z", "bodyPreview": "Numbers attached", "hasAttachments": True}]})
    res = await call("email_search", ctx, unread_only=True, from_address="ana@client.com")
    q = query(http.one("GET", r"/me/mailFolders/inbox/messages"))
    assert q["$filter"] == ["receivedDateTime ge 1900-01-01T00:00:00Z and isRead eq false and "
                            "from/emailAddress/address eq 'ana@client.com'"]
    assert q["$orderby"] == ["receivedDateTime desc"]
    assert res["messages"][0] == {"id": "AAMk/1+x=", "thread_id": "conv-1", "from": "Ana <ana@client.com>",
                                  "to": "me@contoso.com", "date": "2026-10-05T09:00:00Z", "subject": "Budget",
                                  "snippet": "Numbers attached", "unread": True, "has_attachments": True}
    http.add("GET", r"/me/messages\?", json={"value": []})
    await call("email_search", ctx, query='budget "Q3"', since="2026-10-01")
    assert query(http.one("GET", r"/me/messages\?"))["$search"] == ['"budget \\"Q3\\" received>=2026-10-01"']
    # app-only (client credentials) addresses the configured mailbox instead of /me
    ctx.extra["connectors"]["microsoft365"] = {"tenant_id": "contoso", "client_id": "c", "client_secret": "s",
                                               "user": "jane@contoso.com"}
    _oauth.clear_cache()
    http.add("GET", r"/users/jane%40contoso\.com/mailFolders/inbox/messages", json={"value": []})
    await call("email_search", ctx)
    cc = form(http.find("POST", r"login\.microsoftonline\.com/contoso/oauth2")[0])
    assert cc["grant_type"] == "client_credentials" and cc["scope"] == "https://graph.microsoft.com/.default"


# =========================================================================== OneDrive
def _docx_bytes(text: str) -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Plan", level=1)
    d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


async def test_onedrive_list_and_read_docx_follows_redirect_without_token(ctx, http):
    ctx.extra["connectors"]["microsoft365"] = dict(MS)
    ms_token_route(http)
    http.add("GET", r"/me/drive/root:/Reports:/children", json={"value": [
        {"id": "ITEM1", "name": "Plan.docx", "size": 1234, "file": {"mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
         "lastModifiedDateTime": "2026-10-01T00:00:00Z", "webUrl": "https://contoso-my.sharepoint.com/Plan.docx",
         "parentReference": {"path": "/drive/root:/Reports"}},
        {"id": "FOLD", "name": "Archive", "folder": {"childCount": 3}, "parentReference": {"path": "/drive/root:/Reports"}}]})
    listed = await call("cloud_list", ctx, folder="/Reports")
    assert listed["drive"] == "OneDrive / SharePoint"
    assert listed["items"][0]["path"] == "/Reports/Plan.docx" and listed["items"][1] == {
        "id": "FOLD", "name": "Archive", "type": "folder", "path": "/Reports/Archive", "children": 3}
    http.add("GET", r"/me/drive/items/ITEM1\?", json={"id": "ITEM1", "name": "Plan.docx", "file": {
        "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}})
    http.add("GET", r"/me/drive/items/ITEM1/content", status=302,
             headers={"Location": "https://dl.files.example.net/x/Plan.docx"})
    http.add("GET", r"dl\.files\.example\.net", content=_docx_bytes("Launch in November."))
    doc = await call("cloud_read", ctx, file_id="ITEM1")
    assert "# Plan" in doc["text"] and "Launch in November." in doc["text"]
    redirected = http.one("GET", r"dl\.files\.example\.net")
    assert "Authorization" not in redirected.headers  # the token never leaves graph.microsoft.com


# =========================================================================== Google Calendar
EVENT = {"id": "ev1", "summary": "Kickoff", "start": {"dateTime": "2026-10-07T15:00:00+02:00"},
         "end": {"dateTime": "2026-10-07T15:30:00+02:00"}, "htmlLink": "https://calendar.google.com/ev1",
         "hangoutLink": "https://meet.google.com/abc-defg-hij", "organizer": {"email": "me@acme.com", "self": True},
         "attendees": [{"email": "ana@client.com", "responseStatus": "accepted"},
                       {"email": "me@acme.com", "self": True, "organizer": True, "responseStatus": "accepted"}]}


async def test_google_calendar_create_invites_with_meet_and_update_guard(ctx, http):
    ctx.extra["connectors"]["google_calendar"] = dict(GOOGLE)
    google_token_route(http)
    http.add("POST", r"/calendar/v3/calendars/primary/events\?", json=EVENT)
    created = await call("calendar_create", ctx, title="Kickoff", start="2026-10-07T15:00", duration_minutes=30,
                         attendees=["Ana <ana@client.com>"], add_meeting_link=True, timezone="Europe/Berlin")
    req = http.one("POST", r"/calendars/primary/events")
    assert query(req) == {"sendUpdates": ["all"], "conferenceDataVersion": ["1"]}
    body = body_json(req)
    assert body["summary"] == "Kickoff" and body["attendees"] == [{"email": "ana@client.com"}]
    assert body["start"] == {"dateTime": "2026-10-07T15:00:00+02:00", "timeZone": "Europe/Berlin"}
    assert body["end"] == {"dateTime": "2026-10-07T15:30:00+02:00", "timeZone": "Europe/Berlin"}
    assert body["conferenceData"]["createRequest"]["conferenceSolutionKey"] == {"type": "hangoutsMeet"}
    assert created["meeting_link"] == "https://meet.google.com/abc-defg-hij"
    # updating a meeting with other attendees requires notify_attendees=true (→ HIGH risk approval)
    http.add("GET", r"/calendars/primary/events/ev1$", json=EVENT)
    http.add("GET", r"/calendars/primary$", json={"id": "me@acme.com", "summary": "me@acme.com"})
    with pytest.raises(ToolError, match="other attendee"):
        await call("calendar_update", ctx, event_id="ev1", description="New agenda")
    http.add("PATCH", r"/calendars/primary/events/ev1\?", json={**EVENT, "description": "New agenda"})
    upd = await call("calendar_update", ctx, event_id="ev1", description="New agenda", notify_attendees=True)
    patch = http.one("PATCH", r"/events/ev1")
    assert query(patch)["sendUpdates"] == ["all"] and body_json(patch) == {"description": "New agenda"}
    assert upd["updated"] and upd["description"] == "New agenda"


async def test_google_calendar_availability_merges_busy_and_suggests_slots(ctx, http):
    ctx.extra["connectors"]["google_calendar"] = dict(GOOGLE)
    google_token_route(http)
    http.add("POST", r"/calendar/v3/freeBusy", json={"calendars": {
        "primary": {"busy": [{"start": "2026-10-07T10:00:00Z", "end": "2026-10-07T11:00:00Z"}]},
        "ana@client.com": {"busy": [{"start": "2026-10-07T13:00:00Z", "end": "2026-10-07T14:30:00Z"}]},
        "ghost@nowhere.dev": {"errors": [{"domain": "global", "reason": "notFound"}]}}})
    res = await call("calendar_availability", ctx, start="2026-10-07T09:00:00Z", end="2026-10-07T17:00:00Z",
                     attendees=["ana@client.com", "ghost@nowhere.dev"], duration_minutes=60,
                     working_hours="09:00-17:00", timezone="UTC")
    fb = body_json(http.one("POST", r"/freeBusy"))
    assert fb["items"] == [{"id": "primary"}, {"id": "ana@client.com"}, {"id": "ghost@nowhere.dev"}]
    assert res["people"]["ghost@nowhere.dev"] == {"error": "notFound"}
    assert [(s["start"][11:16], s["end"][11:16], s["minutes"]) for s in res["suggested_slots"]] == [
        ("09:00", "10:00", 60), ("11:00", "13:00", 120), ("14:30", "17:00", 150)]


# =========================================================================== Google Drive
async def test_google_drive_read_exports_google_doc_from_url(ctx, http):
    ctx.extra["connectors"]["google_drive"] = dict(GOOGLE)
    google_token_route(http)
    http.add("GET", r"/drive/v3/files/DOCID1234567890\?", json={
        "id": "DOCID1234567890", "name": "Launch plan", "mimeType": "application/vnd.google-apps.document"})
    http.add("GET", r"/files/DOCID1234567890/export\?", text="Launch plan\n\nShip v2 in November.")
    res = await call("cloud_read", ctx, file_id="https://docs.google.com/document/d/DOCID1234567890/edit?usp=sharing")
    assert query(http.one("GET", r"/export"))["mimeType"] == ["text/plain"]
    assert res["name"] == "Launch plan.txt" and "Ship v2 in November." in res["text"]
    http.add("GET", r"/files/DOCID1234567890/export\?", content=b"PK fake docx")
    dl = await call("cloud_download", ctx, file_id="DOCID1234567890")
    assert dl["path"] == "downloads/Launch plan.docx" or dl["path"].startswith("downloads/Launch plan")
    with pytest.raises(ToolError, match="options"):
        await call("cloud_download", ctx, file_id="DOCID1234567890", export_format="mp3")


async def test_google_drive_share_with_people_and_public_link(ctx, http):
    ctx.extra["connectors"]["google_drive"] = dict(GOOGLE)
    google_token_route(http)
    http.add("GET", r"/drive/v3/files/FILE1234567890\?", json={
        "id": "FILE1234567890", "name": "Deck.pdf", "mimeType": "application/pdf",
        "webViewLink": "https://drive.google.com/file/d/FILE1234567890/view", "owners": [{"emailAddress": "me@acme.com"}]})
    http.add("POST", r"/files/FILE1234567890/permissions", json={"id": "perm"})
    res = await call("cloud_share", ctx, file_id="FILE1234567890", emails=["ana@client.com"], role="editor",
                     link_scope="anyone", message="Here you go")
    perms = http.find("POST", r"/permissions")
    assert body_json(perms[0]) == {"type": "user", "role": "writer", "emailAddress": "ana@client.com"}
    assert query(perms[0])["sendNotificationEmail"] == ["true"] and query(perms[0])["emailMessage"] == ["Here you go"]
    assert body_json(perms[1]) == {"type": "anyone", "role": "writer"}
    assert res["link"] == "https://drive.google.com/file/d/FILE1234567890/view" and res["with"] == ["ana@client.com"]
    preview = registry.get("cloud_share").render_preview({"file_id": "FILE1234567890", "link_scope": "anyone"})
    assert "ANYONE WITH THE LINK" in preview


# =========================================================================== Dropbox / Box
async def test_dropbox_refresh_list_and_download_with_ascii_header(ctx, http):
    ctx.extra["connectors"]["dropbox"] = {"app_key": "dbx-key", "app_secret": "dbx-secret", "refresh_token": "dbx-rt"}
    http.add("POST", r"api\.dropboxapi\.com/oauth2/token", json={"access_token": "dbx-tok", "expires_in": 14400})
    http.add("POST", r"api\.dropboxapi\.com/2/files/list_folder$", json={"entries": [
        {".tag": "file", "name": "Q3.pdf", "id": "id:q3", "path_display": "/Reports/Q3.pdf", "size": 10,
         "server_modified": "2026-10-01T00:00:00Z"},
        {".tag": "folder", "name": "Old", "id": "id:old", "path_display": "/Reports/Old"}]})
    res = await call("cloud_list", ctx, folder="Reports", account="dropbox")
    assert body_json(http.one("POST", r"/files/list_folder"))["path"] == "/Reports"
    assert form(http.one("POST", r"oauth2/token")) == {"grant_type": "refresh_token", "refresh_token": "dbx-rt",
                                                       "client_id": "dbx-key", "client_secret": "dbx-secret"}
    assert [i["type"] for i in res["items"]] == ["file", "folder"]
    http.add("POST", r"/2/files/get_metadata", json={".tag": "file", "name": "Résumé.txt",
                                                      "path_display": "/Reports/Résumé.txt"})
    http.add("POST", r"content\.dropboxapi\.com/2/files/download", content=b"cv text")
    dl = await call("cloud_download", ctx, file_id="/Reports/Résumé.txt")
    req = http.one("POST", r"files/download")
    assert req.headers["Dropbox-API-Arg"] == '{"path": "/Reports/R\\u00e9sum\\u00e9.txt"}'
    assert req.headers["Authorization"] == "Bearer dbx-tok" and dl["size"] == 7


async def test_box_list_and_company_link(ctx, http):
    ctx.extra["connectors"]["box"] = {"developer_token": "box-dev"}
    http.add("GET", r"api\.box\.com/2\.0/folders/0/items", json={"entries": [
        {"type": "file", "id": "123", "name": "Deck.pptx", "size": 5, "modified_at": "2026-10-01T00:00:00Z",
         "path_collection": {"entries": [{"name": "All Files"}, {"name": "Sales"}]}}]})
    listed = await call("cloud_list", ctx)
    assert listed["items"][0] == {"id": "file:123", "name": "Deck.pptx", "type": "file", "size": 5,
                                  "modified": "2026-10-01T00:00:00Z", "path": "/Sales/Deck.pptx",
                                  "url": "https://app.box.com/file/123"}
    http.add("PUT", r"api\.box\.com/2\.0/files/123\?", json={"shared_link": {"url": "https://app.box.com/s/xyz"}})
    res = await call("cloud_share", ctx, file_id="file:123", link_scope="organization")
    assert body_json(http.one("PUT", r"/files/123")) == {"shared_link": {"access": "company"}}
    assert res["link"] == "https://app.box.com/s/xyz"
    with pytest.raises(ToolError, match="Invalid Box item"):
        await call("cloud_get", ctx, file_id="Deck.pptx")


# =========================================================================== Notion / Confluence
async def test_notion_create_page_converts_markdown_to_blocks(ctx, http):
    ctx.extra["connectors"]["notion"] = {"token": "ntn_test"}
    http.add("POST", r"api\.notion\.com/v1/pages$", json={
        "object": "page", "id": "pg-1", "url": "https://www.notion.so/pg1",
        "properties": {"title": {"type": "title", "title": [{"plain_text": "Launch plan"}]}}})
    res = await call("kb_create", ctx, title="Launch plan",
                     content="# Goals\nShip **v2** by Q4\n- [x] design\n- [ ] build\n- docs\n```py\nprint(1)\n```",
                     parent_id="https://www.notion.so/acme/Plans-0123456789abcdef0123456789abcdef")
    req = http.one("POST", r"/v1/pages")
    assert req.headers["Notion-Version"] == "2022-06-28" and req.headers["Authorization"] == "Bearer ntn_test"
    body = body_json(req)
    assert body["parent"] == {"page_id": "0123456789abcdef0123456789abcdef"}
    assert body["properties"] == {"title": {"title": [{"type": "text", "text": {"content": "Launch plan"}}]}}
    t = lambda s, **a: {"type": "text", "text": {"content": s}, **({"annotations": a} if a else {})}  # noqa: E731
    assert body["children"] == [
        {"object": "block", "type": "heading_1", "heading_1": {"rich_text": [t("Goals")]}},
        {"object": "block", "type": "paragraph", "paragraph": {"rich_text": [t("Ship "), t("v2", bold=True), t(" by Q4")]}},
        {"object": "block", "type": "to_do", "to_do": {"rich_text": [t("design")], "checked": True}},
        {"object": "block", "type": "to_do", "to_do": {"rich_text": [t("build")], "checked": False}},
        {"object": "block", "type": "bulleted_list_item", "bulleted_list_item": {"rich_text": [t("docs")]}},
        {"object": "block", "type": "code", "code": {"rich_text": [t("print(1)")], "language": "python"}},
    ]
    assert res["created"] and res["id"] == "pg-1" and res["title"] == "Launch plan" and res["blocks"] == 6


async def test_notion_database_query_and_row_update_use_schema(ctx, http):
    ctx.extra["connectors"]["notion"] = {"token": "ntn_test"}
    db_id = "fedcba9876543210fedcba9876543210"
    http.add("GET", rf"/v1/databases/{db_id}$", json={"object": "database", "properties": {
        "Name": {"type": "title"}, "Status": {"type": "select"}, "Due": {"type": "date"},
        "Tags": {"type": "multi_select"}, "Done": {"type": "checkbox"}, "Points": {"type": "number"}}})
    http.add("POST", rf"/v1/databases/{db_id}/query", json={"has_more": False, "results": [{
        "id": "row-1", "url": "https://notion.so/row1", "properties": {
            "Name": {"type": "title", "title": [{"plain_text": "Launch"}]},
            "Status": {"type": "select", "select": {"name": "Doing"}},
            "Tags": {"type": "multi_select", "multi_select": [{"name": "q4"}]},
            "Done": {"type": "checkbox", "checkbox": False}}}]})
    rows = await call("kb_database_query", ctx, database_id=db_id, match={"Status": "Doing", "Done": False})
    assert body_json(http.one("POST", r"/query"))["filter"] == {"and": [
        {"property": "Status", "select": {"equals": "Doing"}}, {"property": "Done", "checkbox": {"equals": False}}]}
    assert rows["rows"] == [{"id": "row-1", "url": "https://notion.so/row1", "Name": "Launch", "Status": "Doing",
                             "Tags": ["q4"], "Done": False}]
    http.add("POST", r"/v1/pages$", json={"id": "row-2", "url": "https://notion.so/row2"})
    added = await call("kb_database_update", ctx, database_id=db_id, values={
        "Name": "Write brief", "status": "Todo", "Due": "2026-11-01", "Tags": ["q4", "launch"], "Done": "no", "Points": 3})
    assert body_json(http.one("POST", r"/v1/pages$")) == {"parent": {"database_id": db_id}, "properties": {
        "Name": {"title": [{"type": "text", "text": {"content": "Write brief"}}]}, "Status": {"select": {"name": "Todo"}},
        "Due": {"date": {"start": "2026-11-01"}}, "Tags": {"multi_select": [{"name": "q4"}, {"name": "launch"}]},
        "Done": {"checkbox": False}, "Points": {"number": 3.0}}}
    assert added["created"] and added["row_id"] == "row-2"
    with pytest.raises(ToolError, match="Unknown property 'Owner'"):
        await call("kb_database_update", ctx, database_id=db_id, values={"Owner": "x"})


async def test_confluence_create_with_code_macro_and_read_as_markdown(ctx, http):
    ctx.extra["connectors"]["confluence"] = {"site_url": "https://acme.atlassian.net", "email": "me@acme.com",
                                             "api_token": "conf-tok", "default_space": "ENG"}
    http.add("GET", r"/wiki/api/v2/spaces\?", json={"results": [{"id": "98304", "key": "ENG"}]})
    http.add("POST", r"/wiki/api/v2/pages$", json={"id": "555", "title": "Runbook",
                                                   "_links": {"webui": "/spaces/ENG/pages/555/Runbook"}})
    res = await call("kb_create", ctx, title="Runbook", content="# Deploy\n\n```bash\nmake deploy\n```")
    req = http.one("POST", r"/wiki/api/v2/pages$")
    assert req.headers["Authorization"] == "Basic " + base64.b64encode(b"me@acme.com:conf-tok").decode()
    body = body_json(req)
    assert body["spaceId"] == "98304" and body["status"] == "current" and body["body"]["representation"] == "storage"
    assert "<h1>Deploy</h1>" in body["body"]["value"]
    assert ('<ac:structured-macro ac:name="code"><ac:parameter ac:name="language">bash</ac:parameter>'
            "<ac:plain-text-body><![CDATA[make deploy]]></ac:plain-text-body></ac:structured-macro>") in body["body"]["value"]
    assert res["url"] == "https://acme.atlassian.net/wiki/spaces/ENG/pages/555/Runbook"
    http.add("GET", r"/wiki/api/v2/pages/555\?", json={
        "id": "555", "title": "Runbook", "spaceId": "98304", "version": {"number": 3}, "_links": {"webui": "/x"},
        "body": {"storage": {"value": '<h2>Steps</h2><ul><li>Build</li><li>Ship <strong>fast</strong></li></ul>'
                                      '<ac:structured-macro ac:name="code"><ac:plain-text-body><![CDATA[make all]]>'
                                      "</ac:plain-text-body></ac:structured-macro>"}}})
    page = await call("kb_get", ctx, page_id="https://acme.atlassian.net/wiki/spaces/ENG/pages/555/Runbook")
    assert page["version"] == 3
    for part in ("# Runbook", "## Steps", "- Build", "- Ship **fast**", "```\nmake all\n```"):
        assert part in page["text"], part
    with pytest.raises(ToolError, match="Notion"):
        await call("kb_database_query", ctx, database_id="123")


# =========================================================================== Linear / Asana
async def test_linear_create_task_graphql_variables(ctx, http):
    ctx.extra["connectors"]["linear"] = {"api_key": "lin_api_123"}

    def gql(request: httpx.Request) -> httpx.Response:
        q = body_json(request)["query"]
        if "issueCreate" in q:
            return httpx.Response(200, json={"data": {"issueCreate": {"success": True, "issue": {
                "id": "uuid-7", "identifier": "ENG-7", "title": "Fix crash", "url": "https://linear.app/acme/issue/ENG-7",
                "priorityLabel": "High", "state": {"name": "Todo", "type": "unstarted"}, "team": {"key": "ENG"}}}}})
        return httpx.Response(200, json={"data": {"teams": {"nodes": [{"id": "team-1"}]}}})

    http.add("POST", r"api\.linear\.app/graphql", gql)
    res = await call("pm_create_task", ctx, title="Fix crash", project="ENG", description="Steps to reproduce",
                     priority="high", due="2026-10-30")
    calls = http.find("POST", r"linear\.app/graphql")
    assert calls[0].headers["Authorization"] == "lin_api_123"  # personal keys are sent without "Bearer"
    assert body_json(calls[0])["variables"] == {"k": "ENG"}
    assert body_json(calls[1])["variables"] == {"i": {"teamId": "team-1", "title": "Fix crash",
                                                      "description": "Steps to reproduce", "priority": 2,
                                                      "dueDate": "2026-10-30"}}
    assert res == {"tool": "Linear", "created": True, "id": "ENG-7", "title": "Fix crash", "status": "Todo",
                   "priority": "High", "team": "ENG", "url": "https://linear.app/acme/issue/ENG-7"}
    http.add("POST", r"api\.linear\.app/graphql", json={"errors": [{"message": "Authentication required"}]})
    with pytest.raises(ToolError, match="reconnect Linear"):
        await call("pm_get_task", ctx, task_id="ENG-7")


async def test_asana_search_falls_back_to_typeahead_and_create_task(ctx, http):
    ctx.extra["connectors"]["asana"] = {"token": "asana-pat", "workspace_gid": "111"}
    http.add("GET", r"/workspaces/111/tasks/search", status=402,
             json={"errors": [{"message": "Search is only available to premium users."}]})
    http.add("GET", r"/workspaces/111/typeahead", json={"data": [
        {"gid": "t1", "name": "Launch email", "completed": False, "permalink_url": "https://app.asana.com/0/1/t1"}]})
    res = await call("pm_search", ctx, query="launch")
    assert res["tasks"] == [{"id": "t1", "title": "Launch email", "status": "open", "url": "https://app.asana.com/0/1/t1"}]
    assert "Premium" in res["note"]
    http.add("GET", r"app\.asana\.com/api/1\.0/projects\?", json={"data": [{"gid": "p1", "name": "Marketing"}]})
    http.add("POST", r"app\.asana\.com/api/1\.0/tasks\?", json={"data": {
        "gid": "t2", "name": "Write brief", "permalink_url": "https://app.asana.com/0/p1/t2", "due_on": "2026-10-20"}})
    created = await call("pm_create_task", ctx, title="Write brief", project="marketing", assignee="me",
                         due="2026-10-20")
    post = http.one("POST", r"/api/1\.0/tasks")
    assert post.headers["Authorization"] == "Bearer asana-pat"
    assert body_json(post) == {"data": {"name": "Write brief", "assignee": "me", "due_on": "2026-10-20",
                                        "projects": ["p1"]}}
    assert created["id"] == "t2" and created["due"] == "2026-10-20"


# =========================================================================== Salesforce
async def test_salesforce_refresh_soql_escaping_and_pipeline(ctx, http):
    ctx.extra["connectors"]["salesforce"] = {"instance_url": "https://acme.my.salesforce.com", "client_id": "sf-cid",
                                             "client_secret": "sf-cs", "refresh_token": "sf-rt"}
    http.add("POST", r"login\.salesforce\.com/services/oauth2/token",
             json={"access_token": "sf-tok", "instance_url": "https://acme.my.salesforce.com"})

    def soql(request: httpx.Request) -> httpx.Response:
        q = query(request)["q"][0]
        if "FROM OpportunityStage" in q:
            return httpx.Response(200, json={"records": [
                {"MasterLabel": "Prospecting", "SortOrder": 1, "IsClosed": False},
                {"MasterLabel": "Negotiation/Review", "SortOrder": 2, "IsClosed": False},
                {"MasterLabel": "Closed Won", "SortOrder": 3, "IsClosed": True}]})
        if "GROUP BY StageName" in q:
            return httpx.Response(200, json={"records": [
                {"StageName": "Negotiation/Review", "deals": 2, "amount": 75000, "weighted": 67500},
                {"StageName": "Prospecting", "deals": 1, "amount": 10000, "weighted": 1000}]})
        return httpx.Response(200, json={"records": [{
            "attributes": {"type": "Opportunity"}, "Id": "006A0000001", "Name": "Acme's renewal", "Amount": 50000,
            "StageName": "Negotiation/Review", "CloseDate": "2026-11-30", "Probability": 90,
            "Account": {"Name": "Acme"}, "Owner": {"Name": "Sam"}}]})

    http.add("GET", r"acme\.my\.salesforce\.com/services/data/v61\.0/query", soql)
    res = await call("crm_list_deals", ctx, stage="Negotiation/Review", query="Acme's")
    assert form(http.one("POST", r"oauth2/token")) == {"grant_type": "refresh_token", "client_id": "sf-cid",
                                                       "refresh_token": "sf-rt", "client_secret": "sf-cs"}
    q = query(http.find("GET", r"/query")[0])["q"][0]
    assert q == ("SELECT Id, Name, Amount, StageName, CloseDate, Probability, Account.Name, Owner.Name FROM Opportunity "
                 "WHERE IsClosed = false AND StageName = 'Negotiation/Review' AND Name LIKE '%Acme\\'s%' "
                 "ORDER BY CloseDate ASC LIMIT 25")
    assert res["deals"] == [{"id": "006A0000001", "name": "Acme's renewal", "amount": 50000.0,
                             "stage": "Negotiation/Review", "close_date": "2026-11-30", "probability": 90.0,
                             "company": "Acme", "owner": "Sam", "url": "https://acme.my.salesforce.com/006A0000001"}]
    assert http.find("GET", r"/query")[0].headers["Authorization"] == "Bearer sf-tok"
    summary = await call("crm_pipeline_summary", ctx)
    assert [s["stage"] for s in summary["stages"]] == ["Prospecting", "Negotiation/Review"]
    assert summary["totals"] == {"deals": 3, "amount": 85000.0, "weighted": 68500.0}


# =========================================================================== Teams / Twilio / webhooks
async def test_teams_webhook_posts_adaptive_card(ctx, http):
    ctx.extra["connectors"]["teams"] = {"webhook_url": "https://acme.webhook.office.com/webhookb2/xyz"}
    http.add("POST", r"acme\.webhook\.office\.com/webhookb2/xyz", status=200, text="1")
    res = await call("chat_send", ctx, channel="Sales / General", text="Deal closed 🎉")
    card = body_json(http.one("POST", r"webhook\.office\.com"))
    assert card["type"] == "message"
    att = card["attachments"][0]
    assert att["contentType"] == "application/vnd.microsoft.card.adaptive"
    assert att["content"]["body"] == [{"type": "TextBlock", "text": "Deal closed 🎉", "wrap": True}]
    assert res == {"workspace": "Microsoft Teams", "sent": True, "via": "webhook"}
    with pytest.raises(ToolError, match="Graph credentials"):
        await call("chat_list_channels", ctx)


async def test_twilio_sms_and_whatsapp(ctx, http):
    ctx.extra["connectors"]["twilio"] = {"account_sid": "AC123", "auth_token": "tw-tok", "from_number": "+15550001111",
                                         "whatsapp_from": "+14155238886"}
    http.add("POST", r"api\.twilio\.com/2010-04-01/Accounts/AC123/Messages\.json", status=201,
             json={"sid": "SM1", "status": "queued", "num_segments": "1"})
    res = await call("notify_sms", ctx, to="+1 (415) 555-0123", body="Your order shipped")
    req = http.one("POST", r"/Messages\.json")
    assert req.headers["Authorization"] == "Basic " + base64.b64encode(b"AC123:tw-tok").decode()
    assert form(req) == {"To": "+14155550123", "Body": "Your order shipped", "From": "+15550001111"}
    assert res == {"sent": True, "sid": "SM1", "status": "queued", "to": "+14155550123", "segments": "1"}
    await call("notify_whatsapp", ctx, to="+14155550123", body="Hi!")
    wa = form(http.find("POST", r"/Messages\.json")[-1])
    assert (wa["To"], wa["From"]) == ("whatsapp:+14155550123", "whatsapp:+14155238886")
    with pytest.raises(ToolError, match="E.164"):
        await call("notify_sms", ctx, to="555-0123", body="x")
    http.add("POST", r"/Messages\.json", status=400, json={"code": 63016, "message": "Failed to send freeform message"})
    with pytest.raises(ToolError, match="24 hours"):
        await call("notify_whatsapp", ctx, to="+14155550123", body="Hi again")


async def test_webhook_notification_is_signed_and_never_uses_agent_urls(ctx, http):
    ctx.extra["connectors"]["webhook"] = {"url": "https://hooks.example.com/ren", "signing_secret": "whsec",
                                          "header_name": "X-Api-Key", "header_value": "k-123"}
    http.add("POST", r"hooks\.example\.com/ren", status=200, text="ok")
    res = await call("notify_webhook", ctx, message="Invoice paid", data_json='{"invoice": "INV-42", "amount": 120}',
                     event="invoice.paid")
    req = http.one("POST", r"hooks\.example\.com")
    ts = req.headers["X-Ren-Timestamp"]
    expected = hmac.new(b"whsec", ts.encode() + b"." + req.content, hashlib.sha256).hexdigest()
    assert req.headers["X-Ren-Signature"] == f"sha256={expected}" and req.headers["X-Api-Key"] == "k-123"
    payload = body_json(req)
    assert payload["event"] == "invoice.paid" and payload["message"] == "Invoice paid"
    assert payload["data"] == {"invoice": "INV-42", "amount": 120} and payload["source"] == "ren"
    assert res["sent"] and res["status"] == 200
    assert "url" not in registry.get("notify_webhook").parameters["properties"]  # no agent-supplied destinations
    with pytest.raises(ToolError, match="valid JSON"):
        await call("notify_webhook", ctx, message="x", data_json="{not json")
    # notify_slack falls back to the webhook connector when no Slack/Teams is connected
    sent = await call("notify_slack", ctx, text="Build failed on main")
    assert body_json(http.find("POST", r"hooks\.example\.com")[-1]) == {"text": "Build failed on main"}
    assert sent["via"] == "webhook" and sent["sent"]


# =========================================================================== GCP / Azure
def _service_account() -> tuple[dict, Any]:
    from cryptography.hazmat.primitives import serialization
    from cryptography.hazmat.primitives.asymmetric import rsa

    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    pem = key.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8,
                            serialization.NoEncryption()).decode()
    return {"type": "service_account", "project_id": "proj-1", "private_key_id": "kid-1", "private_key": pem,
            "client_email": "ren@proj-1.iam.gserviceaccount.com",
            "token_uri": "https://oauth2.googleapis.com/token"}, key.public_key()


def _b64d(s: str) -> bytes:
    return base64.urlsafe_b64decode(s + "=" * (-len(s) % 4))


async def test_gcp_signs_service_account_jwt_and_lists_buckets(ctx, http):
    from cryptography.hazmat.primitives import hashes
    from cryptography.hazmat.primitives.asymmetric import padding

    sa, public_key = _service_account()
    ctx.extra["connectors"]["gcp"] = {"service_account_json": json.dumps(sa)}
    http.add("POST", r"oauth2\.googleapis\.com/token", json={"access_token": "gcp-tok", "expires_in": 3599})
    http.add("GET", r"storage\.googleapis\.com/storage/v1/b\?", json={"items": [
        {"name": "proj-1-reports", "location": "EU", "storageClass": "STANDARD", "timeCreated": "2026-01-01T00:00:00Z"}]})
    res = await call("gcp_storage_list_buckets", ctx)
    tok = form(http.one("POST", r"oauth2\.googleapis\.com/token"))
    assert tok["grant_type"] == "urn:ietf:params:oauth:grant-type:jwt-bearer"
    header_b64, claims_b64, sig_b64 = tok["assertion"].split(".")
    assert json.loads(_b64d(header_b64)) == {"alg": "RS256", "typ": "JWT", "kid": "kid-1"}
    claims = json.loads(_b64d(claims_b64))
    assert claims["iss"] == sa["client_email"] and claims["aud"] == "https://oauth2.googleapis.com/token"
    assert claims["scope"] == "https://www.googleapis.com/auth/cloud-platform.read-only"
    assert claims["exp"] - claims["iat"] == 3600 and "sub" not in claims
    public_key.verify(_b64d(sig_b64), f"{header_b64}.{claims_b64}".encode(), padding.PKCS1v15(), hashes.SHA256())
    listing = http.one("GET", r"/storage/v1/b\?")
    assert query(listing)["project"] == ["proj-1"] and listing.headers["Authorization"] == "Bearer gcp-tok"
    assert res["buckets"] == [{"name": "proj-1-reports", "location": "EU", "class": "STANDARD",
                               "created": "2026-01-01T00:00:00Z"}]
    await call("gcp_storage_list_buckets", ctx)
    assert len(http.find("POST", r"oauth2\.googleapis\.com/token")) == 1  # signed token reused from cache


async def test_azure_client_credentials_resource_groups_and_blob_xml(ctx, http):
    ctx.extra["connectors"]["azure"] = {"tenant_id": "t-1", "client_id": "c-1", "client_secret": "s-1",
                                        "subscription_id": "sub-1"}

    def token(request: httpx.Request) -> httpx.Response:
        scope = form(request)["scope"]
        return httpx.Response(200, json={"access_token": "arm-tok" if "management" in scope else "blob-tok",
                                         "expires_in": 3600})

    http.add("POST", r"login\.microsoftonline\.com/t-1/oauth2/v2\.0/token", token)
    http.add("GET", r"management\.azure\.com/subscriptions/sub-1/resourcegroups\?", json={"value": [
        {"name": "rg-prod", "location": "westeurope", "properties": {"provisioningState": "Succeeded"},
         "tags": {"env": "prod"}}]})
    res = await call("azure_list_resource_groups", ctx)
    assert res["resource_groups"] == [{"name": "rg-prod", "location": "westeurope", "state": "Succeeded",
                                       "tags": {"env": "prod"}}]
    arm = http.one("GET", r"/resourcegroups")
    assert query(arm)["api-version"] == ["2021-04-01"] and arm.headers["Authorization"] == "Bearer arm-tok"
    tok = form(http.find("POST", r"oauth2/v2\.0/token")[0])
    assert tok == {"grant_type": "client_credentials", "client_id": "c-1", "client_secret": "s-1",
                   "scope": "https://management.azure.com/.default"}
    xml = ('﻿<?xml version="1.0" encoding="utf-8"?><EnumerationResults ContainerName="reports"><Prefix>2026/</Prefix>'
           "<Blobs><Blob><Name>2026/q3.pdf</Name><Properties><Content-Length>2048</Content-Length>"
           "<Last-Modified>Mon, 05 Oct 2026 10:00:00 GMT</Last-Modified><Content-Type>application/pdf</Content-Type>"
           "</Properties></Blob><BlobPrefix><Name>2026/archive/</Name></BlobPrefix></Blobs><NextMarker /></EnumerationResults>")
    http.add("GET", r"acmedata\.blob\.core\.windows\.net/reports\?", content=xml.encode("utf-8"),
             headers={"Content-Type": "application/xml"})
    blobs = await call("azure_storage_list_blobs", ctx, account="acmedata", container="reports", prefix="2026/")
    req = http.one("GET", r"blob\.core\.windows\.net")
    assert req.headers["Authorization"] == "Bearer blob-tok" and req.headers["x-ms-version"] == "2023-11-03"
    assert query(req)["comp"] == ["list"] and query(req)["prefix"] == ["2026/"]
    assert blobs["blobs"] == [{"name": "2026/q3.pdf", "size": 2048, "modified": "Mon, 05 Oct 2026 10:00:00 GMT",
                               "type": "application/pdf"}]
    assert blobs["folders"] == ["2026/archive/"]
    with pytest.raises(ToolError, match="3-24 lowercase"):
        await call("azure_storage_list_blobs", ctx, account="evil.example.com/x", container="reports")


# =========================================================================== model-assisted tools
def _gmail_msg(mid: str, sender: str, text: str, date_ms: str) -> dict:
    return {"id": mid, "threadId": "t1", "labelIds": ["INBOX"], "internalDate": date_ms, "snippet": text[:40],
            "payload": {"mimeType": "text/plain", "headers": [{"name": "From", "value": sender},
                                                              {"name": "Subject", "value": "Launch date"}],
                        "body": {"data": base64.urlsafe_b64encode(text.encode()).decode().rstrip("=")}}}


async def test_email_summarize_and_extract_use_the_org_model(ctx, http):
    ctx.extra["connectors"]["gmail"] = dict(GOOGLE)
    google_token_route(http)
    http.add("GET", r"/users/me/threads/t1\?", json={"id": "t1", "messages": [
        _gmail_msg("m1", "Ana <ana@client.com>", "Can we move the launch to Friday?", "1759654800000"),
        _gmail_msg("m2", "me@acme.com", "Friday works. I'll send the deck by Thursday.", "1759658400000")]})
    res = await call("email_summarize_thread", ctx, thread_id="t1")
    assert res["messages"] == 2 and res["subject"] == "Launch date"
    assert res["participants"] == ["Ana <ana@client.com>", "me@acme.com"]
    assert isinstance(res["summary"], str) and "Launch date" in res["summary"]  # mock model echoes its prompt
    http.add("GET", r"/users/me/profile", json={"emailAddress": "me@acme.com"})
    items = await call("email_extract_action_items", ctx, thread_id="t1")
    assert items["messages_analyzed"] == 2 and items["action_items"] == []  # mock model returns an empty list


async def test_slack_summarize_channel_reads_history(ctx, http):
    ctx.extra["connectors"]["slack"] = {"bot_token": "xoxb-1"}
    http.add("GET", r"conversations\.list", json={"ok": True, "channels": [{"id": "C0SALES001", "name": "sales"}],
                                                   "response_metadata": {"next_cursor": ""}})
    http.add("GET", r"conversations\.history", json={"ok": True, "has_more": False, "messages": [
        {"type": "message", "user": "U0ANA00001", "text": "We closed <@U0BOB00001> deal with <https://acme.com|Acme>",
         "ts": "1759658400.000200"},
        {"type": "message", "user": "U0BOB00001", "text": "Great &amp; thanks!", "ts": "1759654800.000100"}]})
    http.add("GET", r"users\.info", lambda req: httpx.Response(200, json={"ok": True, "user": {
        "name": query(req)["user"][0], "profile": {"display_name": {"U0ANA00001": "ana", "U0BOB00001": "bob"}[query(req)["user"][0]]}}}))
    res = await call("chat_summarize_channel", ctx, channel="#sales", hours=48)
    hist = query(http.one("GET", r"conversations\.history"))
    assert hist["channel"] == ["C0SALES001"] and float(hist["oldest"][0]) > 0
    assert res["messages"] == 2 and isinstance(res["summary"], str) and res["summary"]
    assert "bob: Great & thanks!" in res["summary"] and "@bob deal with Acme (https://acme.com)" in res["summary"]


# =========================================================================== GitHub PRs and Actions
async def test_github_get_pr_diff_stats_reviews_and_checks(ctx, http):
    ctx.extra["connectors"]["github"] = {"token": "t"}
    http.add("GET", r"/repos/o/r/pulls/5$", json={
        "number": 5, "title": "Add SSO", "state": "open", "user": {"login": "ana"}, "head": {"ref": "sso", "sha": "abc123def456789"},
        "base": {"ref": "main"}, "additions": 120, "deletions": 8, "changed_files": 3, "commits": 2, "mergeable": True,
        "mergeable_state": "clean", "comments": 1, "review_comments": 2, "body": "Implements SAML."})
    http.add("GET", r"/repos/o/r/pulls/5/files", json=[{"filename": "auth.py", "status": "modified", "additions": 100,
                                                        "deletions": 5, "patch": "@@ -1 +1 @@"}])
    http.add("GET", r"/repos/o/r/pulls/5/reviews", json=[{"user": {"login": "lee"}, "state": "CHANGES_REQUESTED"},
                                                         {"user": {"login": "lee"}, "state": "APPROVED"}])
    http.add("GET", r"/commits/abc123def456789/check-runs", json={"check_runs": [
        {"name": "ci", "conclusion": "failure"}, {"name": "lint", "conclusion": "success"}]})
    pr = await call("github_get_pr", ctx, repo="o/r", number=5)
    assert (pr["additions"], pr["deletions"], pr["changed_files"], pr["comments"]) == (120, 8, 3, 3)
    assert pr["files"] == [{"file": "auth.py", "status": "modified", "additions": 100, "deletions": 5}]
    assert pr["reviews"] == {"lee": "APPROVED"}
    assert pr["checks"] == {"summary": {"failure": 1, "success": 1}, "failing": ["ci"]}


async def test_github_workflow_logs_are_trimmed_around_errors(ctx, http):
    ctx.extra["connectors"]["github"] = {"token": "t"}
    http.add("GET", r"/repos/o/r/actions/runs/77/jobs", json={"jobs": [
        {"id": 9, "name": "test", "conclusion": "failure"}, {"id": 10, "name": "lint", "conclusion": "success"}]})
    noise = "\n".join(f"2026-10-05T10:00:{i % 60:02d}.0000000Z step {i}" for i in range(400))
    log = noise + "\n2026-10-05T10:07:00.0000000Z FAILED tests/test_api.py::test_login\n" \
                  "2026-10-05T10:07:01.0000000Z ##[error]Process completed with exit code 1."
    http.add("GET", r"/repos/o/r/actions/jobs/9/logs", status=302,
             headers={"Location": "https://pipelines.actions.githubusercontent.example/logs/9"})
    http.add("GET", r"pipelines\.actions\.githubusercontent\.example/logs/9", text=log)
    res = await call("github_workflow_logs", ctx, repo="o/r", run_id=77, max_chars=3000)
    assert [j["job_id"] for j in res["jobs"]] == [9]
    excerpt = res["jobs"][0]["log"]
    assert "##[error]Process completed with exit code 1." in excerpt and "FAILED tests/test_api.py" in excerpt
    assert "step 0\n" not in excerpt and "2026-10-05T10:07:01" not in excerpt  # trimmed + timestamps stripped
    assert len(excerpt) <= 3001
    assert "Authorization" not in http.one("GET", r"githubusercontent\.example").headers


# =========================================================================== Gmail drafts and labels
async def test_gmail_labels_archive_and_draft_lifecycle(ctx, http):
    ctx.extra["connectors"]["gmail"] = dict(GOOGLE)
    google_token_route(http)
    http.add("GET", r"/users/me/labels$", json={"labels": [{"id": "INBOX", "name": "INBOX"},
                                                           {"id": "Label_1", "name": "Clients"}]})
    http.add("POST", r"/users/me/labels$", json={"id": "Label_2", "name": "Follow up"})
    http.add("POST", r"/messages/batchModify", status=204)
    res = await call("email_label", ctx, message_ids=["m1", "m2"], add=["clients", "Follow up"])
    assert body_json(http.find("POST", r"/batchModify")[0]) == {"ids": ["m1", "m2"], "addLabelIds": ["Label_1", "Label_2"]}
    assert body_json(http.one("POST", r"/users/me/labels$"))["name"] == "Follow up" and res["updated"] == 2
    await call("email_archive", ctx, message_ids=["m1"])
    assert body_json(http.find("POST", r"/batchModify")[-1]) == {"ids": ["m1"], "removeLabelIds": ["INBOX"]}
    await call("email_mark_read", ctx, message_ids=["m2"])
    assert body_json(http.find("POST", r"/batchModify")[-1]) == {"ids": ["m2"], "removeLabelIds": ["UNREAD"]}
    http.add("POST", r"/users/me/drafts$", json={"id": "r-1", "message": {"id": "m9", "threadId": "t9"}})
    draft = await call("email_create_draft", ctx, to=["ana@client.com"], subject="Proposal", body="Draft v1")
    assert draft["draft_id"] == "r-1" and draft["saved"] and "nothing was sent" in draft["note"]
    http.add("GET", r"/users/me/drafts/m9\?", status=404, json={"error": {"message": "Requested entity was not found."}})
    http.add("GET", r"/users/me/drafts\?", json={"drafts": [{"id": "r-1", "message": {"id": "m9"}}]})
    http.add("DELETE", r"/users/me/drafts/r-1$", status=204)
    gone = await call("email_delete_draft", ctx, draft_id="m9")  # message id resolved to its draft id
    assert gone == {"draft_id": "r-1", "deleted": True, "account": "Gmail"}


# =========================================================================== connection tests
async def test_connector_kind_tests_report_identity(http):
    http.add("GET", r"slack\.com/api/auth\.test", json={"ok": True, "team": "Acme", "user": "renbot"})
    ok, msg = await connector_kinds["slack"].test({"bot_token": "xoxb-1"})
    assert ok and "Acme" in msg and "renbot" in msg
    http.add("GET", r"/account-info/v3/details", status=403, json={"message": "missing scope"})
    http.add("GET", r"/crm/v3/objects/contacts\?", json={"results": []})
    ok, msg = await connector_kinds["hubspot"].test({"token": "pat"})
    assert ok and "contacts readable" in msg
    http.add("GET", r"api\.twilio\.com/2010-04-01/Accounts/AC1\.json", json={"friendly_name": "Acme", "status": "active"})
    assert (await connector_kinds["twilio"].test({"account_sid": "AC1", "auth_token": "t"}))[0]
    http.add("POST", r"hooks\.example\.com/ping", status=204)
    ok, msg = await connector_kinds["webhook"].test({"url": "https://hooks.example.com/ping"})
    ping = body_json(http.one("POST", r"hooks\.example\.com/ping"))
    assert ok and "204" in msg and ping["event"] == "ren.test" and ping["test"] is True
    http.add("GET", r"api\.notion\.com/v1/users/me", status=401, json={"code": "unauthorized", "message": "API token is invalid."})
    ok, msg = await connector_kinds["notion"].test({"token": "bad"})
    assert not ok and "reconnect Notion" in msg
    ok, msg = await connector_kinds["teams"].test({"webhook_url": "https://evil.example.com/hook"})
    assert not ok
    ok, msg = await connector_kinds["jira"].test({"site_url": "acme.atlassian.net"})
    assert not ok and "email" in msg and "api_token" in msg  # missing fields named, no exception
