"""Shared helpers for the document / data tool groups (leading underscore → not auto-imported as a tool module).

Contents
- Workspace I/O: resolve inputs/outputs inside the sandbox, publish artifacts, compact ToolResults.
- Markdown → blocks (markdown-it-py) → DOCX (python-docx) · PDF (fpdf2, Unicode DejaVu fonts) · HTML · PPTX outline.
- DOCX / PPTX / XLSX / PDF / CSV → markdown or text extraction.
- PPTX deck builder with the f69 Labs beige / deep-green themes (theme colours live in the deck's theme part, so they can
  be swapped).
- Tabular helpers: coerce any table-ish input into header + rows, styled XLSX writing.

Theme: f69 Labs palette — deep green #0F2D24 text, brand green #1F5A45 accents, sand #E8DCC8 / beige fills.
"""
from __future__ import annotations

import copy
import csv
import importlib.util
import io
import json
import math
import re
import unicodedata
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, datetime, time as dtime
from decimal import Decimal
from pathlib import Path
from typing import Any

from app.services.workspace import WorkspaceError
from app.tools.base import ToolContext, ToolError, ToolResult

# --------------------------------------------------------------------------------------------- theme
# f69 Labs brand palette (F69-labs-logo.png): Deep Green 0F2D24 · Green 1F5A45 · Mint CFE7D9 · Beige D9CBB8 ·
# Sand E8DCC8 · Ivory F8F6F1. Documents print on white, so text is deep green and fills are light sand / beige.
INK = "0F2D24"  # deep green: body text and headings
INK_SOFT = "2F4A40"  # softer deep green: minor headings, code
MUTED = "5E6E65"  # subtitles, captions, metadata
BEIGE = "F3EDE2"  # light beige: code blocks, inline code, cover page
SAND = "E8DCC8"  # brand sand: table header fill
BAND = "FAF7F0"  # zebra rows in tables
BORDER = "D9CBB8"  # brand beige: table and code borders
ACCENT = "1F5A45"  # brand green: rules, bullets, quote bars, header underlines
LINK = "1F5A45"  # brand green (underlined) — readable on white paper
FONT = "Calibri"  # humanist sans available in every Office install (Inter-like proportions)
MONO = "Consolas"
# Native chart series (PPTX / XLSX) and matplotlib charts: green, bronze, slate, terracotta, sage, deep green, sand, plum
CHART_PALETTE = ["1F5A45", "A8875A", "4F6D8A", "B0654F", "7FB59A", "0F2D24", "C9B48F", "7D6E8F"]
MONO_FONTS = {"consolas", "courier new", "courier", "menlo", "monaco", "source code pro", "fira code",
              "jetbrains mono", "cascadia code", "cascadia mono", "dejavu sans mono", "lucida console"}

DOC_EXTS = (".docx",)
SHEET_EXTS = (".xlsx", ".xlsm", ".csv", ".tsv")
DECK_EXTS = (".pptx",)


def rgb(hex_color: str) -> tuple[int, int, int]:
    h = hex_color.lstrip("#")
    return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)


# --------------------------------------------------------------------------------------------- workspace I/O
def slugify(text: str, max_len: int = 60) -> str:
    s = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    s = re.sub(r"[^\w\s-]", "", s).strip().lower()
    s = re.sub(r"[\s_-]+", "-", s).strip("-")
    return s[:max_len].strip("-") or "untitled"


def safe_filename(name: str, max_len: int = 120) -> str:
    s = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", str(name or "")).strip().strip(".")
    return s[:max_len].strip()


def human_size(n: int | float) -> str:
    n = float(n or 0)
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024
    return f"{n:.1f} GB"


def ws_path(ctx: ToolContext, path: str | None, *, must_exist: bool = False, for_write: bool = False) -> Path:
    try:
        return ctx.workspace.resolve(path, must_exist=must_exist, for_write=for_write)
    except WorkspaceError as e:
        raise ToolError(str(e)) from e


def vpath(ctx: ToolContext, abs_path: Path) -> str:
    return ctx.workspace.virtual(abs_path)


def ws_input(ctx: ToolContext, path: str | None, exts: tuple[str, ...] | None = None) -> Path:
    """Resolve an existing workspace file (optionally restricted to extensions)."""
    if not path or not str(path).strip():
        raise ToolError("A workspace file path is required (e.g. 'uploads/report.docx').")
    p = ws_path(ctx, str(path).strip(), must_exist=True)
    if p.is_dir():
        raise ToolError(f"'{path}' is a folder, not a file.")
    if exts and p.suffix.lower() not in exts:
        raise ToolError(f"Expected a {' / '.join(exts)} file but got '{p.name}'.")
    return p


def require_write(ctx: ToolContext) -> None:
    if not ctx.can_write:
        raise ToolError("Your role is read-only: you can't create or modify files.")


def _with_ext(name: str, ext: str) -> str:
    stem, dot, suf = name.rpartition(".")
    if dot and stem and 1 <= len(suf) <= 5 and suf.isalnum():
        return name if "." + suf.lower() == ext else stem + ext
    return name + ext


def unique_path(p: Path) -> Path:
    if not p.exists():
        return p
    for i in range(2, 1000):
        cand = p.with_name(f"{p.stem}-{i}{p.suffix}")
        if not cand.exists():
            return cand
    raise ToolError(f"Too many files named like '{p.name}'.")


def ws_output(ctx: ToolContext, path: str | None, *, default_name: str, ext: str, overwrite: bool = False,
              folder: str = "output") -> Path:
    """Resolve where to write a new file. Bare names go to `output/`; the extension is enforced; existing files
    get a numeric suffix unless `overwrite`."""
    require_write(ctx)
    ext = "." + ext.lstrip(".").lower()
    if path and str(path).strip():
        raw = str(path).strip().replace("\\", "/")
        directory, _, name = raw.rpartition("/")
        name = safe_filename(name) or slugify(default_name)
        directory = directory.strip("/") or folder
        raw = f"{directory}/{_with_ext(name, ext)}"
    else:
        raw = f"{folder}/{slugify(default_name)}{ext}"
    target = ws_path(ctx, raw, for_write=True)
    if target.is_dir():
        raise ToolError(f"'{vpath(ctx, target)}' is a folder.")
    if target.exists() and not overwrite:
        target = unique_path(target)
    return target


def ws_save_target(ctx: ToolContext, src: Path, output: str | None, *, ext: str | None = None) -> Path:
    """Target for tools that modify a file: in place by default, or `output` (always overwritten)."""
    require_write(ctx)
    if output and str(output).strip():
        return ws_output(ctx, output, default_name=src.stem, ext=ext or src.suffix, overwrite=True)
    return src


async def publish(ctx: ToolContext, abs_path: Path, *, name: str | None = None) -> dict:
    """Register a produced file as an artifact (once per path per tool context, so repeated edits of the same
    deliverable during one task don't create duplicate artifacts); returns the artifact record."""
    vp = vpath(ctx, abs_path)
    seen = ctx.extra.setdefault("_ren_artifacts", {})
    if vp in seen:
        rec = dict(seen[vp])
        rec["size"] = abs_path.stat().st_size
        return rec
    art = await ctx.add_artifact(vp, name=name)
    seen[vp] = art
    return art


async def produced(ctx: ToolContext, abs_path: Path, summary: dict | None = None, *, display: str | None = None,
                   name: str | None = None, artifact: bool = True) -> ToolResult:
    """Standard compact result for tools that create/modify a deliverable."""
    vp = vpath(ctx, abs_path)
    size = abs_path.stat().st_size
    content: dict[str, Any] = {"path": vp, "size": human_size(size)}
    arts: list[dict] = []
    if artifact:
        art = await publish(ctx, abs_path, name=name)
        arts.append(art)
    if summary:
        summary = dict(summary)
        src = summary.pop("path", None)
        if src and src != vp:
            content["source"] = src
        content.update(summary)
    return ToolResult(content=content, artifacts=arts, display=display or f"Saved {vp} ({human_size(size)})")


def clip(text: str, max_chars: int) -> tuple[str, bool]:
    text = text or ""
    if max_chars and len(text) > max_chars:
        return text[:max_chars].rstrip() + f"\n…[truncated — {len(text) - max_chars:,} more characters]", True
    return text, False


def jsonable(v: Any) -> Any:
    """Make a cell / scalar JSON-friendly (dates → ISO, NaN → None, numpy → python)."""
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        if math.isnan(v) or math.isinf(v):
            return None
        return round(v, 6)
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, (datetime, date, dtime)):
        return v.isoformat()
    if isinstance(v, str):
        return v
    try:  # numpy / pandas scalars
        import numpy as np

        if isinstance(v, np.bool_):
            return bool(v)
        if isinstance(v, np.integer):
            return int(v)
        if isinstance(v, np.floating):
            return jsonable(float(v))
        if isinstance(v, np.datetime64):
            import pandas as pd

            ts = pd.Timestamp(v)
            return None if pd.isna(ts) else ts.isoformat()
    except Exception:  # noqa: BLE001
        pass
    try:
        import pandas as pd

        if isinstance(v, pd.Timestamp):
            return None if pd.isna(v) else v.isoformat()
        if v is pd.NaT or (not isinstance(v, (list, dict, tuple)) and pd.isna(v)):
            return None
    except Exception:  # noqa: BLE001
        pass
    if isinstance(v, (list, tuple)):
        return [jsonable(x) for x in v]
    if isinstance(v, dict):
        return {str(k): jsonable(x) for k, x in v.items()}
    return str(v)


# --------------------------------------------------------------------------------------------- markdown → blocks
@dataclass
class Span:
    text: str = ""
    bold: bool = False
    italic: bool = False
    code: bool = False
    strike: bool = False
    link: str | None = None
    image: str | None = None  # image src (text = alt)
    br: bool = False


@dataclass
class Block:
    kind: str  # heading | para | list | code | quote | hr | table | image
    level: int = 0
    spans: list[Span] = field(default_factory=list)
    ordered: bool = False
    start: int = 1
    items: list[list[Block]] = field(default_factory=list)
    children: list[Block] = field(default_factory=list)
    text: str = ""
    lang: str = ""
    src: str = ""
    header: list[list[Span]] = field(default_factory=list)
    rows: list[list[list[Span]]] = field(default_factory=list)
    aligns: list[str | None] = field(default_factory=list)


_MD_IR = None
_MD_HTML = None


def _md_ir():
    global _MD_IR
    if _MD_IR is None:
        from markdown_it import MarkdownIt

        _MD_IR = MarkdownIt("commonmark", {"html": True}).enable(["table", "strikethrough"])
    return _MD_IR


def _md_html():
    global _MD_HTML
    if _MD_HTML is None:
        from markdown_it import MarkdownIt

        # raw HTML disabled: agent-written markdown can't smuggle scripts into HTML deliverables
        _MD_HTML = MarkdownIt("commonmark", {"html": False, "linkify": False}).enable(["table", "strikethrough"])
    return _MD_HTML


_TAG_RE = re.compile(r"<[^>]+>")
_BR_RE = re.compile(r"<br\s*/?>", re.I)


def _strip_html(s: str) -> str:
    import html as _html

    return _html.unescape(_TAG_RE.sub("", _BR_RE.sub("\n", s or ""))).strip()


def _inline(nodes, st: dict) -> list[Span]:
    out: list[Span] = []
    for ch in nodes or []:
        t = ch.type
        if t == "text":
            out.append(Span(ch.content, **st))
        elif t == "softbreak":
            out.append(Span(" ", **st))
        elif t == "hardbreak":
            out.append(Span("", br=True))
        elif t == "code_inline":
            out.append(Span(ch.content, code=True, link=st.get("link")))
        elif t == "strong":
            out += _inline(ch.children, {**st, "bold": True})
        elif t == "em":
            out += _inline(ch.children, {**st, "italic": True})
        elif t == "s":
            out += _inline(ch.children, {**st, "strike": True})
        elif t == "link":
            out += _inline(ch.children, {**st, "link": str(ch.attrs.get("href") or "") or None})
        elif t == "image":
            alt = "".join(c.content for c in (ch.children or []) if c.type == "text") or ch.content or ""
            out.append(Span(alt, image=str(ch.attrs.get("src") or "")))
        elif t == "html_inline":
            if _BR_RE.fullmatch(ch.content.strip()):
                out.append(Span("", br=True))
        elif ch.children:
            out += _inline(ch.children, st)
        elif getattr(ch, "content", ""):
            out.append(Span(ch.content, **st))
    return out


def _node_spans(node) -> list[Span]:
    for ch in node.children or []:
        if ch.type == "inline":
            return _inline(ch.children, {})
    return []


def _blocks(nodes) -> list[Block]:
    out: list[Block] = []
    for n in nodes or []:
        t = n.type
        if t == "heading":
            out.append(Block("heading", level=int(n.tag[1]), spans=_node_spans(n)))
        elif t == "paragraph":
            spans = _node_spans(n)
            real = [s for s in spans if s.image is not None or s.br or s.text.strip()]
            if len(real) == 1 and real[0].image is not None:
                out.append(Block("image", text=real[0].text, src=real[0].image or ""))
            elif spans:
                out.append(Block("para", spans=spans))
        elif t in ("bullet_list", "ordered_list"):
            items = [_blocks(li.children) for li in n.children]
            start = 1
            if t == "ordered_list":
                try:
                    start = int(n.attrs.get("start", 1))
                except (TypeError, ValueError):
                    start = 1
            out.append(Block("list", ordered=t == "ordered_list", start=start, items=items))
        elif t in ("fence", "code_block"):
            out.append(Block("code", text=(n.content or "").rstrip("\n"), lang=(n.info or "").strip()))
        elif t == "blockquote":
            out.append(Block("quote", children=_blocks(n.children)))
        elif t == "hr":
            out.append(Block("hr"))
        elif t == "table":
            out.append(_table_block(n))
        elif t == "html_block":
            text = _strip_html(n.content)
            if text:
                out.append(Block("para", spans=[Span(text)]))
        elif n.children:
            out.extend(_blocks(n.children))
    return out


def _table_block(n) -> Block:
    header: list[list[Span]] = []
    rows: list[list[list[Span]]] = []
    aligns: list[str | None] = []
    for sec in n.children:
        for tr in sec.children:
            cells = []
            for cell in tr.children:
                cells.append(_node_spans(cell))
                if sec.type == "thead":
                    style = str(cell.attrs.get("style", "")) if cell.attrs else ""
                    aligns.append("right" if "right" in style else "center" if "center" in style else None)
            if sec.type == "thead":
                header = cells
            else:
                rows.append(cells)
    return Block("table", header=header, rows=rows, aligns=aligns)


def parse_markdown(md: str) -> list[Block]:
    from markdown_it.tree import SyntaxTreeNode

    return _blocks(SyntaxTreeNode(_md_ir().parse(md or "")).children)


def inline_spans(text: str) -> list[Span]:
    """Inline markdown (bold/italic/code/links) of a single line, without block interpretation."""
    from markdown_it.tree import SyntaxTreeNode

    root = SyntaxTreeNode(_md_ir().parseInline(text or ""))
    for ch in root.children:
        if ch.type == "inline":
            return _inline(ch.children, {}) or [Span(text or "")]
    return [Span(text or "")]


def spans_text(spans: list[Span]) -> str:
    return "".join("\n" if s.br else s.text for s in spans).strip()


def spans_md(spans: list[Span]) -> str:
    out = []
    for s in spans:
        if s.br:
            out.append("  \n")
            continue
        if s.image is not None:
            out.append(f"![{s.text}]({s.image})")
            continue
        t = s.text
        if not t:
            continue
        if s.code:
            t = f"`{t}`"
        else:
            core = t.strip()
            if core and (s.bold or s.italic or s.strike):
                if s.strike:
                    core = f"~~{core}~~"
                if s.bold and s.italic:
                    core = f"***{core}***"
                elif s.bold:
                    core = f"**{core}**"
                elif s.italic:
                    core = f"*{core}*"
                t = t[: len(t) - len(t.lstrip())] + core + t[len(t.rstrip()):]
        if s.link:
            t = f"[{t}]({s.link})"
        out.append(t)
    return "".join(out).strip()


def md_table(header: list[Any] | None, rows: list[list[Any]], max_cell: int = 200) -> str:
    def fmt(v: Any) -> str:
        if v is None:
            return ""
        if isinstance(v, float):
            s = f"{v:,.4f}".rstrip("0").rstrip(".") if not v.is_integer() else f"{int(v)}"
        else:
            s = str(v)
        s = s.replace("|", "\\|").replace("\r", " ").replace("\n", " ")
        return s[:max_cell] + "…" if len(s) > max_cell else s

    ncols = max([len(header or [])] + [len(r) for r in rows] + [1])
    hdr = list(header or [f"col{i + 1}" for i in range(ncols)]) + [""] * (ncols - len(header or []))
    lines = ["| " + " | ".join(fmt(h) for h in hdr) + " |", "|" + "|".join(["---"] * ncols) + "|"]
    for r in rows:
        r = list(r) + [""] * (ncols - len(r))
        lines.append("| " + " | ".join(fmt(v) for v in r) + " |")
    return "\n".join(lines)


_NUM_RE = re.compile(r"^[\s$€£¥₹(+-]*\d[\d,]*(\.\d+)?\s*%?\)?\s*$")


def is_numeric_text(s: Any) -> bool:
    return isinstance(s, (int, float)) and not isinstance(s, bool) or bool(_NUM_RE.match(str(s or "")))


def numeric_columns(rows: list[list[Any]], ncols: int) -> list[bool]:
    out = []
    for j in range(ncols):
        vals = [r[j] for r in rows if j < len(r) and r[j] not in (None, "")]
        out.append(bool(vals) and sum(1 for v in vals if is_numeric_text(v)) >= 0.8 * len(vals))
    return out


# --------------------------------------------------------------------------------------------- DOCX writing
PPR_AFTER_SHD = ("w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct",
                 "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
                 "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection",
                 "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr",
                 "w:sectPr", "w:pPrChange")
PPR_AFTER_PBDR = ("w:shd",) + PPR_AFTER_SHD
RPR_AFTER_SHD = ("w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout",
                 "w:specVanish", "w:oMath")
RPR_AFTER_HIGHLIGHT = ("w:u", "w:effect", "w:bdr", "w:shd") + RPR_AFTER_SHD
TCPR_AFTER_SHD = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark", "w:headers",
                  "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange")
TBLPR_AFTER_BORDERS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription",
                       "w:tblPrChange")
_BORDER_ORDER = ("top", "left", "bottom", "right", "between", "bar")


def _wel(tag: str, **attrs: Any):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(f"w:{k}"), str(v))
    return e


def _set_child(parent, new, successors: tuple[str, ...]):
    old = parent.find(new.tag)
    if old is not None:
        parent.remove(old)
    parent.insert_element_before(new, *successors)
    return new


def ppr_border(pPr, sides: dict[str, tuple[str, int, int]]) -> None:
    pbdr = _wel("w:pBdr")
    for side in _BORDER_ORDER:
        if side in sides:
            color, sz, space = sides[side]
            pbdr.append(_wel(f"w:{side}", val="single", sz=sz, space=space, color=color))
    _set_child(pPr, pbdr, PPR_AFTER_PBDR)


def ppr_shading(pPr, fill: str) -> None:
    _set_child(pPr, _wel("w:shd", val="clear", color="auto", fill=fill), PPR_AFTER_SHD)


def run_shading(run, fill: str) -> None:
    _set_child(run._r.get_or_add_rPr(), _wel("w:shd", val="clear", color="auto", fill=fill), RPR_AFTER_SHD)


def cell_shading(cell, fill: str) -> None:
    _set_child(cell._tc.get_or_add_tcPr(), _wel("w:shd", val="clear", color="auto", fill=fill), TCPR_AFTER_SHD)


def table_borders(table, color: str = BORDER, sz: int = 4, vertical: bool = False) -> None:
    borders = _wel("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in ("left", "right") or (edge == "insideV" and not vertical):
            borders.append(_wel(f"w:{edge}", val="nil"))
        else:
            borders.append(_wel(f"w:{edge}", val="single", sz=sz, space=0, color=color))
    _set_child(table._tbl.tblPr, borders, TBLPR_AFTER_BORDERS)


def repeat_header(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(_wel("w:tblHeader").tag) is None:
        trPr.append(_wel("w:tblHeader", val="true"))


def style_font(style, name: str | None = None, size: float | None = None, color: str | None = None,
               bold: bool | None = None, italic: bool | None = None) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    f = style.font
    if name:
        f.name = name
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is not None:
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                rfonts.attrib.pop(qn(a), None)
            rfonts.set(qn("w:cs"), name)
            rfonts.set(qn("w:eastAsia"), name)
    if size:
        f.size = Pt(size)
    if color:
        f.color.rgb = RGBColor.from_string(color.lstrip("#").upper())
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic


def get_style(doc, name: str):
    try:
        return doc.styles[name]
    except KeyError:
        return None


def ensure_code_style(doc):
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Pt

    st = get_style(doc, "Code Block")
    if st is not None:
        return st
    st = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    normal = get_style(doc, "Normal")
    if normal is not None:
        st.base_style = normal
    style_font(st, MONO, 9, INK, bold=False, italic=False)
    pf = st.paragraph_format
    pf.space_before = Pt(4)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.0
    pPr = st.element.get_or_add_pPr()
    ppr_shading(pPr, BEIGE)
    ppr_border(pPr, {s: (BORDER, 4, 4) for s in ("top", "left", "bottom", "right")})
    return st


def setup_docx_styles(doc, *, body_font: str = FONT, heading_font: str = FONT, body_size: float = 10.5,
                      heading_color: str = INK, accent: str = ACCENT) -> None:
    """Professional defaults: calm sans, deep-green headings, brand-green accents, compact lists and tables."""
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    normal = get_style(doc, "Normal")
    if normal is not None:
        style_font(normal, body_font, body_size, INK)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
    sizes = {1: 20, 2: 15, 3: 12.5, 4: 11.5, 5: 11, 6: 10.5}
    for lvl, size in sizes.items():
        st = get_style(doc, f"Heading {lvl}")
        if st is None:
            continue
        style_font(st, heading_font, size, heading_color if lvl <= 3 else INK_SOFT, bold=True, italic=False)
        st.paragraph_format.space_before = Pt({1: 20, 2: 16, 3: 12}.get(lvl, 10))
        st.paragraph_format.space_after = Pt(6 if lvl <= 2 else 4)
        st.paragraph_format.keep_with_next = True
    h1 = get_style(doc, "Heading 1")
    if h1 is not None:
        ppr_border(h1.element.get_or_add_pPr(), {"bottom": (accent, 8, 4)})
    title = get_style(doc, "Title")
    if title is not None:
        pPr = title.element.get_or_add_pPr()
        old = pPr.find(qn("w:pBdr"))
        if old is not None:
            pPr.remove(old)
        style_font(title, heading_font, 26, heading_color, bold=True, italic=False)
        title.font.all_caps = False
        title.paragraph_format.space_before = Pt(0)
        title.paragraph_format.space_after = Pt(4)
    sub = get_style(doc, "Subtitle")
    if sub is not None:
        style_font(sub, body_font, 13, MUTED, bold=False, italic=False)
        sub.paragraph_format.space_after = Pt(18)
    quote = get_style(doc, "Quote")
    if quote is not None:
        style_font(quote, None, None, INK_SOFT, italic=True)
        quote.paragraph_format.left_indent = Inches(0.25)
        ppr_border(quote.element.get_or_add_pPr(), {"left": (accent, 12, 8)})
    for name in ("List Bullet", "List Bullet 2", "List Bullet 3", "List Number", "List Number 2", "List Number 3",
                 "List Continue", "List Continue 2", "List Continue 3"):
        st = get_style(doc, name)
        if st is not None:
            st.paragraph_format.space_after = Pt(3)
    cap = get_style(doc, "Caption")
    if cap is not None:
        style_font(cap, None, 9, MUTED, bold=False, italic=True)
    footer = get_style(doc, "Footer")
    if footer is not None:
        style_font(footer, body_font, 8.5, MUTED)
    ensure_code_style(doc)


def new_document(page_size: str = "letter"):
    from docx import Document
    from docx.shared import Inches, Mm

    doc = Document()
    sec = doc.sections[0]
    if str(page_size).lower() == "a4":
        sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin = sec.bottom_margin = Inches(1)
    setup_docx_styles(doc)
    return doc


def _add_field(paragraph, instr: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


def page_number_footer(doc, section=None) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    section = section or doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    if get_style(doc, "Footer") is not None:
        p.style = doc.styles["Footer"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_field(p, "PAGE")
    p.add_run(" / ")
    _add_field(p, "NUMPAGES")


def restart_numbering(doc, paragraph, style_name: str, start: int = 1) -> int | None:
    """Give an ordered list its own numbering instance so it starts at `start`."""
    try:
        numbering = doc.part.numbering_part.element
        st = doc.styles[style_name]
        numPr = st.element.pPr.numPr
        ilvl = numPr.ilvl.val if numPr.ilvl is not None else 0
        num = numbering.num_having_numId(numPr.numId.val)
        new_num = numbering.add_num(num.abstractNumId.val)
        new_num.add_lvlOverride(ilvl=ilvl).add_startOverride(start)
        apply_numbering(paragraph, new_num.numId, ilvl)
        return new_num.numId
    except Exception:  # noqa: BLE001
        return None


def apply_numbering(paragraph, num_id: int, ilvl: int = 0) -> None:
    numPr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    numPr.get_or_add_ilvl().val = ilvl
    numPr.get_or_add_numId().val = num_id


def _style_ilvl(doc, style_name: str) -> int:
    try:
        numPr = doc.styles[style_name].element.pPr.numPr
        return numPr.ilvl.val if numPr.ilvl is not None else 0
    except Exception:  # noqa: BLE001
        return 0


def add_hyperlink_run(paragraph, run, url: str) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)
    run._r.addprevious(hl)
    hl.append(run._r)


def _data_uri_image(src: str) -> io.BytesIO | None:
    m = re.match(r"data:image/(png|jpe?g|gif|bmp);base64,(.+)$", src or "", re.S | re.I)
    if not m:
        return None
    import base64

    try:
        return io.BytesIO(base64.b64decode(m.group(2)))
    except Exception:  # noqa: BLE001
        return None


def _image_width_in(img_src: Path | io.BytesIO, max_w: float) -> float:
    try:
        from PIL import Image

        with Image.open(img_src) as im:
            dpi = im.info.get("dpi", (96, 96))[0] or 96
            w = im.size[0] / float(dpi if dpi > 30 else 96)
        if isinstance(img_src, io.BytesIO):
            img_src.seek(0)
        return max(0.5, min(max_w, w))
    except Exception:  # noqa: BLE001
        return max_w


ImageResolver = Callable[[str], Path | None]


class DocxWriter:
    """Render markdown blocks into a python-docx Document (or table cell)."""

    def __init__(self, doc, *, image_resolver: ImageResolver | None = None):
        self.doc = doc
        self.resolve_image = image_resolver
        sec = doc.sections[-1]
        self.max_width = max(2.0, (sec.page_width - sec.left_margin - sec.right_margin) / 914400)
        ensure_code_style(doc)

    # -- helpers
    def _para(self, c, style: str | None = None):
        if style and get_style(self.doc, style) is not None:
            return c.add_paragraph(style=style)
        return c.add_paragraph()

    def _image_source(self, src: str) -> Path | io.BytesIO | None:
        data = _data_uri_image(src)
        if data is not None:
            return data
        if self.resolve_image and src and not re.match(r"^[a-z]+://", src, re.I):
            try:
                return self.resolve_image(src)
            except Exception:  # noqa: BLE001
                return None
        return None

    def runs(self, p, spans: list[Span], *, size: float | None = None, bold: bool = False) -> None:
        from docx.shared import Inches, Pt, RGBColor

        last = None
        for s in spans:
            if s.br:
                (last or p.add_run()).add_break()
                continue
            if s.image is not None:
                img = self._image_source(s.image)
                if img is not None:
                    try:
                        p.add_run().add_picture(img if isinstance(img, io.BytesIO) else str(img),
                                                width=Inches(_image_width_in(img, min(self.max_width, 3.0))))
                        continue
                    except Exception:  # noqa: BLE001
                        pass
                r = p.add_run(f"[{s.text or 'image'}]")
                r.italic = True
                continue
            if not s.text:
                continue
            r = p.add_run(s.text)
            if s.bold or bold:
                r.bold = True
            if s.italic:
                r.italic = True
            if s.strike:
                r.font.strike = True
            if size:
                r.font.size = Pt(size)
            if s.code:
                r.font.name = MONO
                r.font.size = Pt((size or 10.5) - 1)
                run_shading(r, BEIGE)
            if s.link and re.match(r"^(https?://|mailto:)", s.link, re.I):
                r.font.color.rgb = RGBColor.from_string(LINK)
                r.font.underline = True
                add_hyperlink_run(p, r, s.link)
            last = r

    # -- blocks
    def write(self, blocks: list[Block], container=None) -> None:
        c = container if container is not None else self.doc
        for b in blocks:
            self.block(b, c)

    def block(self, b: Block, c) -> None:
        k = b.kind
        if k == "heading":
            self.heading(b, c)
        elif k == "para":
            self.runs(self._para(c), b.spans)
        elif k == "list":
            self.list_(b, c, 0)
        elif k == "code":
            self.code(b, c)
        elif k == "quote":
            for ch in b.children:
                if ch.kind == "para":
                    self.runs(self._para(c, "Quote"), ch.spans)
                else:
                    self.block(ch, c)
        elif k == "hr":
            from docx.shared import Pt

            p = self._para(c)
            ppr_border(p._p.get_or_add_pPr(), {"bottom": (BORDER, 6, 1)})
            p.paragraph_format.space_after = Pt(8)
        elif k == "table":
            self.table(b, c)
        elif k == "image":
            self.image(b.src, b.text, c)

    def heading(self, b: Block, c) -> None:
        from docx.shared import Pt

        lvl = max(1, min(6, b.level))
        name = f"Heading {lvl}"
        if get_style(self.doc, name) is not None:
            self.runs(c.add_paragraph(style=name), b.spans)
        else:
            p = c.add_paragraph()
            self.runs(p, b.spans, size={1: 20, 2: 15, 3: 12.5}.get(lvl, 11.5), bold=True)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.keep_with_next = True

    def list_(self, b: Block, c, level: int) -> None:
        from docx.shared import Inches

        base = "List Number" if b.ordered else "List Bullet"
        style = base if level == 0 else f"{base} {min(level + 1, 3)}"
        has_style = get_style(self.doc, style) is not None
        num_id = None
        ilvl = _style_ilvl(self.doc, style) if has_style else 0
        for idx, item in enumerate(b.items):
            first = True
            for blk in item:
                if blk.kind == "para" and first:
                    p = self._para(c, style if has_style else None)
                    if not has_style:
                        p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
                        p.add_run(f"{b.start + idx}. " if b.ordered else "• ")
                    elif b.ordered:
                        if num_id is None:
                            num_id = restart_numbering(self.doc, p, style, b.start)
                        else:
                            apply_numbering(p, num_id, ilvl)
                    self.runs(p, blk.spans)
                    first = False
                elif blk.kind == "list":
                    self.list_(blk, c, level + 1)
                elif blk.kind == "para":
                    cont = "List Continue" if level == 0 else f"List Continue {min(level + 1, 3)}"
                    p = self._para(c, cont)
                    self.runs(p, blk.spans)
                else:
                    self.block(blk, c)
                    first = False

    def code(self, b: Block, c) -> None:
        p = self._para(c, "Code Block")
        lines = (b.text or "").split("\n")
        for i, line in enumerate(lines):
            r = p.add_run(line)
            if i < len(lines) - 1:
                r.add_break()

    def table(self, b: Block, c) -> None:
        header = [spans_text(h) for h in b.header]
        rows_txt = [[spans_text(x) for x in r] for r in b.rows]
        self.table_rows(header if b.header else None, b.rows, c, aligns=b.aligns, spans_mode=True,
                        text_rows=rows_txt, header_spans=b.header)

    def table_rows(self, header: list[str] | None, rows: list[list[Any]], c, *, aligns: list[str | None] | None = None,
                   spans_mode: bool = False, text_rows: list[list[str]] | None = None,
                   header_spans: list[list[Span]] | None = None):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt

        text_rows = text_rows if text_rows is not None else [["" if v is None else str(v) for v in r] for r in rows]
        ncols = max([len(header or [])] + [len(r) for r in text_rows] + [0])
        if ncols == 0:
            return None
        nrows = len(text_rows) + (1 if header else 0)
        t = c.add_table(rows=nrows, cols=ncols)
        table_borders(t)
        numeric = numeric_columns(text_rows, ncols)
        lens = []
        for j in range(ncols):
            cells = ([header[j]] if header and j < len(header) else []) + [r[j] for r in text_rows[:200] if j < len(r)]
            lens.append(min(40, max([len(x) for x in cells] + [3])))
        total = sum(lens) or 1
        widths = [Inches(max(0.6, self.max_width * w / total)) for w in lens]
        align_map = {"right": WD_ALIGN_PARAGRAPH.RIGHT, "center": WD_ALIGN_PARAGRAPH.CENTER}
        font_size = 9.5 if ncols <= 6 else 8.5
        r0 = 0
        if header:
            for j in range(ncols):
                cell = t.cell(0, j)
                p = cell.paragraphs[0]
                if header_spans is not None and j < len(header_spans):
                    self.runs(p, header_spans[j], size=font_size, bold=True)
                else:
                    self.runs(p, [Span(header[j] if j < len(header) else "")], size=font_size, bold=True)
                cell_shading(cell, SAND)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                a = (aligns[j] if aligns and j < len(aligns) else None) or ("right" if numeric[j] else None)
                if a in align_map:
                    p.alignment = align_map[a]
            repeat_header(t.rows[0])
            r0 = 1
        for i, row in enumerate(rows):
            for j in range(ncols):
                cell = t.cell(r0 + i, j)
                p = cell.paragraphs[0]
                if spans_mode:
                    spans = row[j] if j < len(row) else []
                    self.runs(p, spans, size=font_size)
                else:
                    v = row[j] if j < len(row) else None
                    self.runs(p, [Span(_cell_text(v))], size=font_size)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                a = (aligns[j] if aligns and j < len(aligns) else None) or ("right" if numeric[j] else None)
                if a in align_map:
                    p.alignment = align_map[a]
        for j in range(ncols):
            for row in t.rows:
                row.cells[j].width = widths[j]
        spacer = self._para(c)
        spacer.paragraph_format.space_after = Pt(4)
        return t

    def image(self, src: str, alt: str, c, width_in: float | None = None) -> bool:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches

        img = self._image_source(src)
        if img is None:
            p = self._para(c)
            r = p.add_run(f"[Image: {alt or src}]")
            r.italic = True
            return False
        p = self._para(c)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        w = width_in or _image_width_in(img, self.max_width)
        p.add_run().add_picture(img if isinstance(img, io.BytesIO) else str(img), width=Inches(min(w, self.max_width)))
        if alt:
            cap = self._para(c, "Caption")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.add_run(alt)
        return True


def _cell_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, int) and not isinstance(v, bool) and abs(v) >= 10000:
        return f"{v:,}"
    if isinstance(v, float):
        return f"{v:,.2f}" if not v.is_integer() else f"{int(v):,}"
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return str(v)


def _today_long() -> str:
    d = datetime.now()
    return f"{d.strftime('%B')} {d.day}, {d.year}"


def docx_cover(doc, title: str, subtitle: str | None = None, author: str | None = None) -> None:
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt, RGBColor

    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(150)
    bar = doc.add_paragraph()
    bar.paragraph_format.space_after = Pt(10)
    br = bar.add_run("\u2003" * 10)
    br.font.size = Pt(5)
    run_shading(br, ACCENT)
    p = doc.add_paragraph(style="Title") if get_style(doc, "Title") is not None else doc.add_paragraph()
    r = p.add_run(title)
    r.font.size = Pt(34)
    if subtitle:
        sp = doc.add_paragraph(style="Subtitle") if get_style(doc, "Subtitle") is not None else doc.add_paragraph()
        sp.add_run(subtitle)
    meta = " · ".join(x for x in (author, _today_long()) if x)
    mp = doc.add_paragraph()
    mp.paragraph_format.space_before = Pt(36)
    mr = mp.add_run(meta)
    mr.font.size = Pt(10)
    mr.font.color.rgb = RGBColor.from_string(MUTED)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _norm_title(s: str) -> str:
    return re.sub(r"\W+", " ", (s or "").lower()).strip()


def markdown_to_docx(markdown: str, path: Path | str | None = None, *, title: str | None = None,
                     subtitle: str | None = None, author: str | None = None, title_page: bool = False,
                     image_resolver: ImageResolver | None = None, doc=None, page_size: str = "letter",
                     footer: bool = True):
    """Render markdown into a styled .docx (new document, or appended to `doc`)."""
    fresh = doc is None
    if fresh:
        doc = new_document(page_size)
    writer = DocxWriter(doc, image_resolver=image_resolver)
    blocks = parse_markdown(markdown)
    if title:
        if blocks and blocks[0].kind == "heading" and blocks[0].level == 1 and \
                _norm_title(spans_text(blocks[0].spans)) == _norm_title(title):
            blocks = blocks[1:]
        if title_page:
            docx_cover(doc, title, subtitle, author)
        else:
            p = doc.add_paragraph(style="Title") if get_style(doc, "Title") is not None else doc.add_paragraph()
            p.add_run(title)
            if subtitle:
                sp = doc.add_paragraph(style="Subtitle") if get_style(doc, "Subtitle") is not None \
                    else doc.add_paragraph()
                sp.add_run(subtitle)
    writer.write(blocks)
    if fresh:
        if footer:
            page_number_footer(doc)
        cp = doc.core_properties
        if title:
            cp.title = title
        if author:
            cp.author = author
        else:
            cp.author = "Ren"
    if path:
        doc.save(str(path))
    return doc


# --------------------------------------------------------------------------------------------- DOCX reading
def _is_mono_run(run) -> bool:
    try:
        name = (run.font.name or "").lower()
        if name in MONO_FONTS:
            return True
        st = run.style.name.lower() if run.style is not None else ""
        return st in ("html code", "code", "verbatim char")
    except Exception:  # noqa: BLE001
        return False


def _runs_md(runs) -> str:
    segs: list[tuple[str, tuple]] = []
    for r in runs:
        t = r.text
        if not t:
            if r._r.xpath(".//*[local-name()='drawing' or local-name()='pict']"):
                segs.append(("[image]", (False, False, False, False)))
            continue
        flags = (bool(r.bold), bool(r.italic), bool(r.font.strike), _is_mono_run(r))
        if segs and segs[-1][1] == flags:
            segs[-1] = (segs[-1][0] + t, flags)
        else:
            segs.append((t, flags))
    out = []
    for text, (b, i, s, code) in segs:
        if code and text.strip():
            out.append(f"`{text}`")
            continue
        core = text.strip()
        if not core or not (b or i or s):
            out.append(text)
            continue
        if s:
            core = f"~~{core}~~"
        if b and i:
            core = f"***{core}***"
        elif b:
            core = f"**{core}**"
        elif i:
            core = f"*{core}*"
        out.append(text[: len(text) - len(text.lstrip())] + core + text[len(text.rstrip()):])
    return "".join(out)


def paragraph_md(p) -> str:
    parts: list[str] = []
    pending = []
    try:
        items = list(p.iter_inner_content())
    except Exception:  # noqa: BLE001
        items = list(p.runs)
    for item in items:
        if hasattr(item, "address") and hasattr(item, "runs"):  # Hyperlink
            if pending:
                parts.append(_runs_md(pending))
                pending = []
            text = _runs_md(item.runs).strip()
            addr = getattr(item, "url", None) or item.address
            parts.append(f"[{text}]({addr})" if addr and text else text)
        else:
            pending.append(item)
    if pending:
        parts.append(_runs_md(pending))
    return "".join(parts)


def _num_format(numbering_el, num_id: int, ilvl: int) -> str | None:
    try:
        abs_ids = numbering_el.xpath(f'./w:num[@w:numId="{num_id}"]/w:abstractNumId/@w:val')
        if not abs_ids:
            return None
        fmts = numbering_el.xpath(
            f'./w:abstractNum[@w:abstractNumId="{abs_ids[0]}"]/w:lvl[@w:ilvl="{ilvl}"]/w:numFmt/@w:val')
        return fmts[0] if fmts else None
    except Exception:  # noqa: BLE001
        return None


def _list_info(p, numbering_el) -> tuple[str, int] | None:
    style = (p.style.name if p.style is not None else "") or ""
    m = re.match(r"List (Bullet|Number)(?: (\d))?$", style)
    if m:
        return ("ol" if m.group(1) == "Number" else "ul"), (int(m.group(2)) - 1 if m.group(2) else 0)
    pPr = p._p.pPr
    numPr = pPr.numPr if pPr is not None else None
    if numPr is None or numPr.numId is None or numPr.numId.val == 0:
        return None
    ilvl = numPr.ilvl.val if numPr.ilvl is not None else 0
    fmt = _num_format(numbering_el, numPr.numId.val, ilvl) if numbering_el is not None else None
    return ("ul" if fmt in (None, "bullet", "none") else "ol"), ilvl


def table_md(t) -> str:
    rows = []
    for row in t.rows:
        cells = []
        for cell in row.cells:
            cells.append(" ".join(par.text.strip() for par in cell.paragraphs if par.text.strip()))
        rows.append(cells)
    if not rows:
        return ""
    return md_table(rows[0], rows[1:])


def docx_to_markdown(src) -> str:
    """Convert a .docx (path or Document) to markdown: headings, lists, emphasis, links, tables, code."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(src)) if isinstance(src, (str, Path)) else src
    try:
        numbering_el = doc.part.numbering_part.element
    except Exception:  # noqa: BLE001
        numbering_el = None
    out: list[str] = []
    prev_kind = None
    prev_list_kind: str | None = None
    code_buf: list[str] = []
    counters: dict[int, int] = {}

    def flush_code():
        nonlocal code_buf
        if code_buf:
            out.append("```\n" + "\n".join(code_buf) + "\n```")
            code_buf = []

    for block in doc.iter_inner_content():
        if isinstance(block, Table):
            flush_code()
            md = table_md(block)
            if md:
                out.append(md)
            prev_kind = "table"
            continue
        if not isinstance(block, Paragraph):
            continue
        p = block
        style = (p.style.name if p.style is not None else "") or ""
        if style == "Code Block" or (p.runs and all(_is_mono_run(r) for r in p.runs if r.text.strip())
                                     and any(r.text.strip() for r in p.runs) and style not in ("Title",)
                                     and not style.startswith("Heading")):
            code_buf.extend(p.text.split("\n"))
            prev_kind = "code"
            continue
        flush_code()
        text = paragraph_md(p).strip()
        if not text:
            prev_kind = None
            continue
        if style == "Title":
            out.append("# " + p.text.strip())
            kind = "title"
        elif style.startswith("Heading"):
            m = re.search(r"(\d)", style)
            lvl = int(m.group(1)) if m else 2
            out.append("#" * min(lvl, 6) + " " + p.text.strip())
            kind = "heading"
        elif style == "Subtitle":
            out.append(f"*{p.text.strip()}*")
            kind = "para"
        elif style in ("Quote", "Intense Quote"):
            out.append("> " + text)
            kind = "quote"
        else:
            li = _list_info(p, numbering_el)
            if li:
                lkind, lvl = li
                if lkind == "ol":
                    counters[lvl] = counters.get(lvl, 0) + 1 if prev_kind == "list" else 1
                    for deeper in [k for k in counters if k > lvl]:
                        counters.pop(deeper, None)
                    marker = f"{counters[lvl]}."
                else:
                    marker = "-"
                line = "   " * lvl + f"{marker} {text}"
                if prev_kind == "list" and out and not (lvl == 0 and prev_list_kind not in (None, lkind)):
                    out[-1] = out[-1] + "\n" + line
                else:
                    if lvl == 0 and lkind == "ol" and prev_kind == "list":
                        counters[0] = 1
                        line = f"1. {text}"
                    out.append(line)
                prev_kind = "list"
                if lvl == 0:
                    prev_list_kind = lkind
                continue
            out.append(text)
            kind = "para"
        prev_kind = kind
    flush_code()
    return "\n\n".join(out).strip() + "\n"


def iter_all_paragraphs(doc):
    """Body, table-cell (recursive), header and footer paragraphs."""
    def from_container(container):
        for p in container.paragraphs:
            yield p
        for t in getattr(container, "tables", []):
            for row in t.rows:
                seen = set()
                for cell in row.cells:
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    yield from from_container(cell)

    yield from from_container(doc)
    for sec in doc.sections:
        for hf in (sec.header, sec.footer):
            try:
                if not hf.is_linked_to_previous or sec is doc.sections[0]:
                    yield from from_container(hf)
            except Exception:  # noqa: BLE001
                continue


def paragraph_runs(p) -> list:
    """All text runs of a paragraph in document order, including runs inside hyperlinks."""
    from docx.text.run import Run

    return [Run(r, p) for r in p._p.xpath("./w:r | ./w:hyperlink/w:r | ./w:ins/w:r")]


def _run_at(bounds: list[int], pos: int) -> int:
    lo, hi = 0, len(bounds) - 1
    while lo < hi:
        mid = (lo + hi + 1) // 2
        if bounds[mid] <= pos:
            lo = mid
        else:
            hi = mid - 1
    return lo


def replace_in_paragraph(p, pattern: re.Pattern, repl: Callable[[re.Match], str]) -> int:
    """Replace matches across run boundaries, keeping the formatting of the run where each match starts."""
    runs = paragraph_runs(p)
    if not runs:
        return 0
    texts = [r.text for r in runs]
    full = "".join(texts)
    matches = [m for m in pattern.finditer(full) if m.end() > m.start()]
    if not matches:
        return 0
    bounds, pos = [], 0
    for t in texts:
        bounds.append(pos)
        pos += len(t)
    cur = list(texts)
    for m in reversed(matches):
        s, e = m.span()
        rs, re_ = _run_at(bounds, s), _run_at(bounds, e - 1)
        new_text = repl(m)
        off_s = s - bounds[rs]
        off_e = e - bounds[re_]
        if rs == re_:
            cur[rs] = cur[rs][:off_s] + new_text + cur[rs][off_e:]
        else:
            cur[rs] = cur[rs][:off_s] + new_text
            for k in range(rs + 1, re_):
                cur[k] = ""
            cur[re_] = cur[re_][off_e:]
    for r, old, new in zip(runs, texts, cur):
        if old != new:
            r.text = new
    return len(matches)


def split_run(run, offset: int):
    """Split a run at `offset`; returns (left, right) runs sharing the original formatting."""
    from docx.text.run import Run

    text = run.text
    new_r = copy.deepcopy(run._r)
    run._r.addnext(new_r)
    right = Run(new_r, run._parent)
    run.text = text[:offset]
    right.text = text[offset:]
    return run, right


def isolate_runs(p, start: int, end: int) -> list:
    """Split runs so that characters [start, end) of the paragraph text are covered by whole runs."""
    runs = paragraph_runs(p)
    pos = 0
    for r in runs:
        t = r.text
        r_start, r_end = pos, pos + len(t)
        if r_start < end and r_end > start:
            if start > r_start:
                split_run(r, start - r_start)
                return isolate_runs(p, start, end)
            if end < r_end:
                split_run(r, end - r_start)
                return isolate_runs(p, start, end)
        pos = r_end
    out, pos = [], 0
    for r in paragraph_runs(p):
        t = r.text
        if pos >= start and pos + len(t) <= end and t:
            out.append(r)
        pos += len(t)
    return out


# --------------------------------------------------------------------------------------------- PDF (fpdf2)
_LATIN1_MAP = {"\u2018": "'", "\u2019": "'", "\u201a": ",", "\u201c": '"', "\u201d": '"', "\u2013": "-",
               "\u2014": "-", "\u2026": "...", "\u2022": "-", "\u00a0": " ", "\u2192": "->", "\u2190": "<-",
               "\u2713": "v", "\u2714": "v", "\u2717": "x", "\u20ac": "EUR", "\u2264": "<=", "\u2265": ">=",
               "\u2212": "-", "\u2009": " ", "\u200b": "", "\u2003": " "}


def latin1(s: str) -> str:
    s = "".join(_LATIN1_MAP.get(ch, ch) for ch in s or "")
    return s.encode("latin-1", "replace").decode("latin-1")


def dejavu_dir() -> Path | None:
    spec = importlib.util.find_spec("matplotlib")
    if spec and spec.submodule_search_locations:
        for loc in spec.submodule_search_locations:
            d = Path(loc) / "mpl-data" / "fonts" / "ttf"
            if (d / "DejaVuSans.ttf").exists() and (d / "DejaVuSans-Bold.ttf").exists():
                return d
    return None


_PDF_CLASS = None


def _pdf_class():
    global _PDF_CLASS
    if _PDF_CLASS is not None:
        return _PDF_CLASS
    from fpdf import FPDF

    class RenPDF(FPDF):
        doc_title = ""
        unicode_ok = False
        sans = "helvetica"
        mono = "courier"
        skip_pages: set = set()
        first_content_page = 1

        def fix(self, s: str) -> str:
            return s if self.unicode_ok else latin1(s)

        def header(self):
            if self.page_no() in self.skip_pages or self.page_no() <= self.first_content_page or not self.doc_title:
                return
            self.set_font(self.sans, "", 7.5)
            self.set_text_color(*rgb(MUTED))
            self.set_xy(self.l_margin, 9)
            self.cell(0, 4, self.fix(self.doc_title[:100]), align="L")
            self.set_draw_color(*rgb(BORDER))
            self.set_line_width(0.2)
            self.line(self.l_margin, 14.5, self.w - self.r_margin, 14.5)
            self.set_xy(self.l_margin, self.t_margin)

        def footer(self):
            if self.page_no() in self.skip_pages:
                return
            self.set_y(-12)
            self.set_font(self.sans, "", 7.5)
            self.set_text_color(*rgb(MUTED))
            self.cell(0, 5, f"{self.page_no()} / {{nb}}", align="C")

    _PDF_CLASS = RenPDF
    return RenPDF


def register_fonts(pdf, *, mono: bool = True, styles: tuple[str, ...] = ("", "B", "I", "BI")) -> None:
    """Register the Unicode DejaVu family (or fall back to latin-1 core fonts)."""
    d = dejavu_dir()
    pdf.unicode_ok, pdf.sans, pdf.mono = False, "helvetica", "courier"
    if d is None:
        return
    files = {"": "DejaVuSans.ttf", "B": "DejaVuSans-Bold.ttf", "I": "DejaVuSans-Oblique.ttf",
             "BI": "DejaVuSans-BoldOblique.ttf"}
    try:
        for style in styles:
            pdf.add_font("RenSans", style, str(d / files[style]))
        if mono:
            pdf.add_font("RenMono", "", str(d / "DejaVuSansMono.ttf"))  # code is always set upright/regular
        pdf.unicode_ok, pdf.sans, pdf.mono = True, "RenSans", ("RenMono" if mono else "courier")
    except Exception:  # noqa: BLE001
        pdf.unicode_ok, pdf.sans, pdf.mono = False, "helvetica", "courier"


def new_pdf(title: str | None = None, *, page_size: str = "A4", orientation: str = "portrait"):
    """FPDF instance with Unicode fonts (DejaVu, shipped with matplotlib) or latin-1 core-font fallback."""
    cls = _pdf_class()
    fmt = "letter" if str(page_size).lower() == "letter" else ("A3" if str(page_size).upper() == "A3" else "A4")
    pdf = cls(orientation="L" if str(orientation).lower().startswith("l") else "P", unit="mm", format=fmt)
    pdf.skip_pages = set()
    pdf.doc_title = title or ""
    pdf.set_margins(20, 20, 20)
    pdf.set_auto_page_break(True, margin=18)
    pdf.alias_nb_pages()
    register_fonts(pdf)
    if title:
        pdf.set_title(pdf.fix(title))
    pdf.set_creator("Ren by f69 Labs")
    return pdf


class PdfRenderer:
    def __init__(self, pdf, *, image_resolver: ImageResolver | None = None, base_size: float = 10.0):
        self.pdf = pdf
        self.resolve_image = image_resolver
        self.size = base_size
        self.base_l = pdf.l_margin

    @property
    def lh(self) -> float:
        return self.size * 0.5

    def fix(self, s: str) -> str:
        return self.pdf.fix(s)

    def _margin(self, indent: float) -> None:
        self.pdf.set_left_margin(self.base_l + indent)
        if self.pdf.get_x() < self.pdf.l_margin:
            self.pdf.set_x(self.pdf.l_margin)

    def _ensure_space(self, h: float) -> None:
        if self.pdf.get_y() + h > self.pdf.page_break_trigger:
            self.pdf.add_page()

    def render(self, blocks: list[Block], indent: float = 0.0) -> None:
        for b in blocks:
            self.block(b, indent)
        self._margin(0)

    def block(self, b: Block, indent: float) -> None:
        k = b.kind
        if k == "heading":
            self.heading(b, indent)
        elif k == "para":
            self.para(b.spans, indent)
        elif k == "list":
            self.list_(b, indent)
        elif k == "code":
            self.code(b, indent)
        elif k == "quote":
            self.quote(b, indent)
        elif k == "hr":
            self.hr(indent)
        elif k == "table":
            self.table(b, indent)
        elif k == "image":
            self.image(b.src, b.text, indent)

    def inline(self, spans: list[Span], *, size: float, color: str = INK, italic: bool = False,
               bold: bool = False) -> None:
        pdf = self.pdf
        lh = size * 0.5
        for s in spans:
            if s.br:
                pdf.ln(lh)
                continue
            text = s.text if s.image is None else f"[{s.text or 'image'}]"
            text = self.fix(text)
            if not text:
                continue
            style = ("B" if (s.bold or bold) else "") + ("I" if (s.italic or italic or s.image is not None) else "")
            if s.code:
                style = ""
            link = s.link if (s.link and re.match(r"^(https?://|mailto:)", s.link, re.I)) else ""
            if link:
                style += "U"
            if s.strike:
                style += "S"
            pdf.set_font(self.pdf.mono if s.code else self.pdf.sans, style, size - 1 if s.code else size)
            pdf.set_text_color(*rgb(LINK if link else (INK_SOFT if s.code else color)))
            pdf.write(lh, text, link=link)

    def para(self, spans: list[Span], indent: float, *, color: str = INK, italic: bool = False,
             after: float = 2.2) -> None:
        pdf = self.pdf
        self._margin(indent)
        self._ensure_space(self.lh * 2)
        pdf.set_x(pdf.l_margin)
        self.inline(spans, size=self.size, color=color, italic=italic)
        pdf.ln(self.lh)
        pdf.ln(after)

    def heading(self, b: Block, indent: float) -> None:
        pdf = self.pdf
        lvl = max(1, min(6, b.level))
        size = {1: 17, 2: 13.5, 3: 11.5, 4: 10.5}.get(lvl, 10)
        if pdf.get_y() > pdf.page_break_trigger - 28:
            pdf.add_page()
        elif pdf.get_y() > pdf.t_margin + 2:
            pdf.ln({1: 6, 2: 4.5, 3: 3}.get(lvl, 2.5))
        self._margin(indent)
        pdf.set_x(pdf.l_margin)
        pdf.set_font(pdf.sans, "B", size)
        pdf.set_text_color(*rgb(INK if lvl <= 2 else INK_SOFT))
        pdf.multi_cell(0, size * 0.5, self.fix(spans_text(b.spans)), align="L", new_x="LMARGIN", new_y="NEXT")
        if lvl == 1:
            y = pdf.get_y() + 1.2
            pdf.set_draw_color(*rgb(ACCENT))
            pdf.set_line_width(0.7)
            pdf.line(pdf.l_margin, y, pdf.l_margin + 16, y)
            pdf.ln(4)
        else:
            pdf.ln(1.6)

    def list_(self, b: Block, indent: float) -> None:
        pdf = self.pdf
        for i, item in enumerate(b.items):
            marker = f"{b.start + i}." if b.ordered else ("\u2022" if pdf.unicode_ok else "-")
            self._margin(indent)
            self._ensure_space(self.lh * 1.5)
            y0 = pdf.get_y()
            pdf.set_xy(pdf.l_margin + (0 if b.ordered else 1.0), y0)
            pdf.set_font(pdf.sans, "B" if b.ordered else "", self.size)
            pdf.set_text_color(*rgb(INK_SOFT if b.ordered else ACCENT))
            pdf.cell(6, self.lh, marker)
            first = True
            for blk in item:
                if blk.kind == "para" and first:
                    self._margin(indent + 6.5)
                    pdf.set_xy(pdf.l_margin, y0)
                    self.inline(blk.spans, size=self.size)
                    pdf.ln(self.lh)
                    pdf.ln(0.9)
                else:
                    if first:
                        pdf.ln(self.lh)
                    self.block(blk, indent + 6.5)
                first = False
            if first:
                pdf.ln(self.lh)
        self._margin(indent)
        pdf.ln(1.3)

    def code(self, b: Block, indent: float) -> None:
        from fpdf.enums import WrapMode

        pdf = self.pdf
        self._margin(indent)
        pdf.ln(0.8)
        pdf.set_x(pdf.l_margin)
        pdf.set_font(pdf.mono, "", 8.3)
        pdf.set_text_color(*rgb(INK))
        pdf.set_fill_color(*rgb(BEIGE))
        pdf.set_draw_color(*rgb(BORDER))
        pdf.set_line_width(0.2)
        text = self.fix((b.text or " ").expandtabs(4)) or " "
        pdf.multi_cell(0, 4.1, text, border=1, fill=True, align="L", new_x="LMARGIN", new_y="NEXT",
                       wrapmode=WrapMode.CHAR, padding=(2.2, 3, 2.2, 3))
        pdf.ln(2.6)

    def quote(self, b: Block, indent: float) -> None:
        pdf = self.pdf
        page0, y0 = pdf.page, pdf.get_y()
        for ch in b.children:
            if ch.kind == "para":
                self.para(ch.spans, indent + 5, color=INK_SOFT, italic=True)
            else:
                self.block(ch, indent + 5)
        if pdf.page == page0:
            x = self.base_l + indent + 1.5
            pdf.set_draw_color(*rgb(ACCENT))
            pdf.set_line_width(0.8)
            pdf.line(x, y0 + 0.4, x, max(y0 + 2, pdf.get_y() - 2.2))
        self._margin(indent)

    def hr(self, indent: float) -> None:
        pdf = self.pdf
        self._margin(indent)
        pdf.ln(1.5)
        y = pdf.get_y()
        pdf.set_draw_color(*rgb(BORDER))
        pdf.set_line_width(0.3)
        pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
        pdf.ln(4)

    def table(self, b: Block, indent: float) -> None:
        header = [spans_text(h) for h in b.header] if b.header else None
        body = [[spans_text(c) for c in r] for r in b.rows]
        self.table_rows(header, body, indent, aligns=b.aligns)

    def table_rows(self, header: list[str] | None, body: list[list[Any]], indent: float = 0.0,
                   aligns: list[str | None] | None = None) -> None:
        from fpdf.fonts import FontFace

        pdf = self.pdf
        body = [[_cell_text(v) for v in r] for r in body]
        ncols = max([len(header or [])] + [len(r) for r in body] + [0])
        if ncols == 0:
            return
        hdr = [self.fix(str(h)) for h in (header or [])] + [""] * (ncols - len(header or [])) if header else None
        rows = [[self.fix(str(v)) for v in (list(r) + [""] * (ncols - len(r)))[:ncols]] for r in body]
        lens = [max([len(hdr[j]) if hdr else 0] + [len(r[j]) for r in rows[:300]] + [3]) for j in range(ncols)]
        lens = [min(max(n, 4), 42) for n in lens]
        numeric = numeric_columns(rows, ncols)
        al = []
        for j in range(ncols):
            a = aligns[j] if aligns and j < len(aligns) else None
            if a == "right" or (a is None and numeric[j]):
                al.append("RIGHT")
            else:
                al.append("CENTER" if a == "center" else "LEFT")
        size = 8.5 if ncols <= 5 else 7.6 if ncols <= 8 else 6.6
        self._margin(indent)
        pdf.ln(0.8)
        pdf.set_x(pdf.l_margin)
        pdf.set_font(pdf.sans, "", size)
        pdf.set_text_color(*rgb(INK))
        pdf.set_draw_color(*rgb(BORDER))
        pdf.set_line_width(0.2)
        width = pdf.w - pdf.l_margin - pdf.r_margin
        pdf.set_fill_color(255, 255, 255)
        with pdf.table(col_widths=lens, width=width, align="LEFT", text_align=tuple(al), line_height=size * 0.52,
                       borders_layout="HORIZONTAL_LINES", first_row_as_headings=bool(hdr),
                       headings_style=FontFace(emphasis="BOLD", color=rgb(INK), fill_color=rgb(SAND)),
                       cell_fill_color=rgb(BAND), cell_fill_mode="ROWS", padding=(1.3, 1.8)) as t:
            if hdr:
                row = t.row()
                for h in hdr:
                    row.cell(h)
            for r in rows:
                row = t.row()
                for v in r:
                    row.cell(v)
        pdf.ln(3)

    def image(self, src: str, alt: str, indent: float) -> None:
        pdf = self.pdf
        img = _data_uri_image(src)
        if img is None and self.resolve_image and src and not re.match(r"^[a-z]+://", src, re.I):
            try:
                img = self.resolve_image(src)
            except Exception:  # noqa: BLE001
                img = None
        if img is None:
            self.para([Span(f"[Image: {alt or src}]", italic=True)], indent, color=MUTED)
            return
        try:
            from PIL import Image

            with Image.open(img) as im:
                wpx, hpx = im.size
            if isinstance(img, io.BytesIO):
                img.seek(0)
            self._margin(indent)
            max_w = pdf.w - pdf.l_margin - pdf.r_margin
            w = min(max_w, wpx * 25.4 / 110)
            h = w * hpx / max(1, wpx)
            max_h = pdf.page_break_trigger - pdf.t_margin - 12
            if h > max_h:
                h, w = max_h, max_h * wpx / max(1, hpx)
            self._ensure_space(h + 6)
            y = pdf.get_y()
            x = pdf.l_margin + (max_w - w) / 2
            pdf.image(img if isinstance(img, io.BytesIO) else str(img), x=x, y=y, w=w, h=h)
            pdf.set_y(y + h + 1.5)
            if alt:
                pdf.set_font(pdf.sans, "I", 8)
                pdf.set_text_color(*rgb(MUTED))
                pdf.multi_cell(0, 4, self.fix(alt), align="C", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2.5)
        except Exception:  # noqa: BLE001
            self.para([Span(f"[Image: {alt or src}]", italic=True)], indent, color=MUTED)


def pdf_cover(pdf, title: str, subtitle: str | None = None, author: str | None = None) -> None:
    pdf.add_page()
    pdf.skip_pages.add(pdf.page_no())
    pdf.set_fill_color(*rgb(BEIGE))
    pdf.rect(0, 0, pdf.w, pdf.h, style="F")
    x = 24.0
    y = pdf.h * 0.36
    pdf.set_fill_color(*rgb(ACCENT))
    pdf.rect(x, y - 9, 22, 1.6, style="F")
    pdf.set_xy(x, y)
    pdf.set_font(pdf.sans, "B", 28)
    pdf.set_text_color(*rgb(INK))
    pdf.multi_cell(pdf.w - 2 * x, 12, pdf.fix(title), align="L", new_x="LEFT", new_y="NEXT")
    if subtitle:
        pdf.ln(3)
        pdf.set_x(x)
        pdf.set_font(pdf.sans, "", 13)
        pdf.set_text_color(*rgb(MUTED))
        pdf.multi_cell(pdf.w - 2 * x, 6.5, pdf.fix(subtitle), align="L", new_x="LEFT", new_y="NEXT")
    meta = " · ".join(v for v in (author, _today_long()) if v)
    pdf.set_xy(x, pdf.h - 36)
    pdf.set_font(pdf.sans, "", 9.5)
    pdf.set_text_color(*rgb(INK_SOFT))
    pdf.cell(pdf.w - 2 * x, 5, pdf.fix(meta))


def markdown_to_pdf(markdown: str, path: Path | str | None = None, *, title: str | None = None,
                    subtitle: str | None = None, author: str | None = None, title_page: bool = False,
                    page_size: str = "A4", image_resolver: ImageResolver | None = None) -> bytes:
    """Render markdown into a clean, paginated PDF (headings, lists, tables, code, quotes, links, images)."""
    pdf = new_pdf(title, page_size=page_size)
    blocks = parse_markdown(markdown)
    if title and blocks and blocks[0].kind == "heading" and blocks[0].level == 1 and \
            _norm_title(spans_text(blocks[0].spans)) == _norm_title(title):
        blocks = blocks[1:]
    if title and title_page:
        pdf_cover(pdf, title, subtitle, author)
    pdf.add_page()
    pdf.first_content_page = pdf.page_no()
    if title and not title_page:
        pdf.set_font(pdf.sans, "B", 21)
        pdf.set_text_color(*rgb(INK))
        pdf.multi_cell(0, 9.5, pdf.fix(title), align="L", new_x="LMARGIN", new_y="NEXT")
        if subtitle:
            pdf.ln(1)
            pdf.set_font(pdf.sans, "", 11.5)
            pdf.set_text_color(*rgb(MUTED))
            pdf.multi_cell(0, 6, pdf.fix(subtitle), align="L", new_x="LMARGIN", new_y="NEXT")
        y = pdf.get_y() + 2.5
        pdf.set_fill_color(*rgb(ACCENT))
        pdf.rect(pdf.l_margin, y, 18, 1.1, style="F")
        pdf.set_y(y + 7)
    if author:
        pdf.set_author(pdf.fix(author))
    PdfRenderer(pdf, image_resolver=image_resolver).render(blocks)
    data = bytes(pdf.output())
    if path:
        Path(path).write_bytes(data)
    return data


# --------------------------------------------------------------------------------------------- PDF reading
def parse_page_ranges(spec: str | int | None, n: int) -> list[int]:
    """'1-3,5,9-' → zero-based page indices (validated against n pages)."""
    if spec is None or str(spec).strip().lower() in ("", "all", "*"):
        return list(range(n))
    out: list[int] = []
    for part in re.split(r"[,\s;]+", str(spec).strip().lower()):
        if not part:
            continue
        part = part.replace("last", str(n))
        m = re.fullmatch(r"(\d+)?-(\d+)?", part)
        if m:
            a = int(m.group(1)) if m.group(1) else 1
            b = int(m.group(2)) if m.group(2) else n
        elif part.isdigit():
            a = b = int(part)
        else:
            raise ToolError(f"Invalid page range '{part}'. Use forms like '1-3,5,8-'.")
        if a < 1 or b > n or a > b:
            raise ToolError(f"Page range '{part}' is outside 1–{n}.")
        for i in range(a - 1, b):
            if i not in out:
                out.append(i)
    return out


def pdf_reader(path: Path):
    from pypdf import PdfReader

    try:
        r = PdfReader(str(path))
    except Exception as e:  # noqa: BLE001
        raise ToolError(f"Could not open '{path.name}' as a PDF: {e}") from e
    if r.is_encrypted:
        ok = 0
        try:
            ok = r.decrypt("")
        except Exception:  # noqa: BLE001
            ok = 0
        if not ok:
            raise ToolError(f"'{path.name}' is password-protected.")
    return r


def pdf_pages_text(path: Path, pages: str | None = None) -> tuple[int, list[tuple[int, str]]]:
    r = pdf_reader(path)
    n = len(r.pages)
    out = []
    for i in parse_page_ranges(pages, n):
        try:
            t = r.pages[i].extract_text() or ""
        except Exception:  # noqa: BLE001
            t = ""
        out.append((i + 1, re.sub(r"[ \t]+\n", "\n", t).strip()))
    return n, out


# --------------------------------------------------------------------------------------------- HTML
HTML_CSS = """
:root{--bg:#EFE7DA;--paper:#F8F6F1;--ink:#0F2D24;--soft:#2F4A40;--muted:#5E6E65;--border:#DCCFBB;--accent:#1F5A45;--link:#1F5A45;--code:#F1EADF;--head:#E8DCC8;--band:#F3EEE5}
*{box-sizing:border-box}
html{background:var(--bg)}
body{margin:0;font-family:Inter,-apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,Helvetica,Arial,sans-serif;
  color:var(--ink);line-height:1.62;font-size:16px;-webkit-font-smoothing:antialiased}
main.doc{max-width:820px;margin:40px auto;background:var(--paper);padding:56px 64px;border-radius:14px;
  box-shadow:0 1px 2px rgba(15,45,36,.06),0 8px 24px rgba(15,45,36,.07)}
h1,h2,h3,h4,h5,h6{line-height:1.25;margin:1.6em 0 .5em;font-weight:650;letter-spacing:-.01em}
h1{font-size:1.9rem;padding-bottom:.35em;border-bottom:2px solid var(--accent)}
h2{font-size:1.4rem}h3{font-size:1.15rem;color:var(--soft)}h4,h5,h6{font-size:1rem;color:var(--soft)}
p{margin:.65em 0}a{color:var(--link);text-underline-offset:2px}
ul,ol{padding-left:1.4em}li{margin:.25em 0}li::marker{color:var(--accent)}
blockquote{margin:1em 0;padding:.2em 1.1em;border-left:3px solid var(--accent);color:var(--soft);font-style:italic}
code{font-family:"JetBrains Mono",Consolas,SFMono-Regular,Menlo,monospace;font-size:.88em;background:var(--code);
  padding:.12em .35em;border-radius:4px}
pre{background:var(--code);border:1px solid var(--border);border-radius:8px;padding:14px 16px;overflow:auto}
pre code{background:none;padding:0;font-size:.85em}
table{border-collapse:collapse;width:100%;margin:1.2em 0;font-size:.92em}
th,td{padding:8px 10px;border-bottom:1px solid var(--border);text-align:left;vertical-align:top}
th{background:var(--head);font-weight:600}tr:nth-child(even) td{background:var(--band)}
hr{border:none;border-top:1px solid var(--border);margin:2em 0}
img{max-width:100%;height:auto;border-radius:6px}
.doc-title{font-size:2.3rem;margin:0 0 .15em;border:none;padding:0}
.doc-subtitle{color:var(--muted);margin:0 0 1.6em;font-size:1.1rem}
.doc-rule{width:56px;height:3px;background:var(--accent);border-radius:2px;margin:0 0 2em}
@media print{html{background:#fff}main.doc{box-shadow:none;margin:0;max-width:none;padding:0}}
@media (max-width:640px){main.doc{margin:0;border-radius:0;padding:28px 20px}}
"""


def markdown_to_html(markdown: str, title: str | None = None, *, subtitle: str | None = None,
                     standalone: bool = True) -> str:
    import html as _html

    body = _md_html().render(markdown or "")
    if not standalone:
        return body
    head_title = _html.escape(title or "Document")
    hero = ""
    if title:
        hero = f'<h1 class="doc-title">{_html.escape(title)}</h1>'
        if subtitle:
            hero += f'<p class="doc-subtitle">{_html.escape(subtitle)}</p>'
        hero += '<div class="doc-rule"></div>'
    return ("<!doctype html>\n<html lang=\"en\">\n<head>\n<meta charset=\"utf-8\">\n"
            "<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">\n"
            f"<title>{head_title}</title>\n<style>{HTML_CSS}</style>\n</head>\n"
            f"<body>\n<main class=\"doc\">\n{hero}\n{body}</main>\n</body>\n</html>\n")


def html_to_text(html: str) -> tuple[str, str]:
    """(title, readable text) for an HTML document (delegates to the web readability extractor)."""
    from app.tools.builtin import _browser

    r = _browser.readable(html, "")
    return r.get("title") or "", r.get("text") or ""


# --------------------------------------------------------------------------------------------- PPTX
THEMES: dict[str, dict[str, Any]] = {
    # f69 beige: deep-green text on beige, brand-green accent
    "light": {"bg": "EFE7DA", "text": "0F2D24", "dk2": "56665D", "lt2": "E8DCC8", "hlink": "1F5A45",
              "accents": ["1F5A45", "A8875A", "4F6D8A", "B0654F", "7FB59A", "0F2D24"]},
    # f69 dark: ivory text on deep green, mint accent
    "dark": {"bg": "0F2D24", "text": "F8F6F1", "dk2": "B2AB9D", "lt2": "15392D", "hlink": "CFE7D9",
             "accents": ["8FCBAA", "C9AD7C", "9DB4CC", "D9A08A", "CFE7D9", "E8DCC8"]},
}
_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _theme_key(name: str | None) -> str:
    n = (name or "light").lower()
    return "dark" if n in ("dark", "ink", "dark-ink", "dark ink", "black", "green", "deep-green", "deep green") else "light"


def apply_deck_theme(prs, theme: str = "light", *, reset_slide_backgrounds: bool = False) -> str:
    """Swap the deck's theme colours (and master background) to the f69 beige (light) or deep-green (dark) palette."""
    from lxml import etree
    from pptx.dml.color import RGBColor
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    key = _theme_key(theme)
    th = THEMES[key]
    for master in prs.slide_masters:
        try:
            part = master.part.part_related_by(RT.THEME)
            root = etree.fromstring(part.blob)
            cs = root.find(f".//{{{_A_NS}}}clrScheme")
            if cs is not None:
                cs.set("name", "f69 Beige" if key == "light" else "f69 Dark")
                values = {"dk1": th["text"], "lt1": th["bg"], "dk2": th["dk2"], "lt2": th["lt2"],
                          "hlink": th["hlink"], "folHlink": th["dk2"]}
                for i, c in enumerate(th["accents"], 1):
                    values[f"accent{i}"] = c
                for tag, hexval in values.items():
                    el = cs.find(f"{{{_A_NS}}}{tag}")
                    if el is None:
                        continue
                    for child in list(el):
                        el.remove(child)
                    etree.SubElement(el, f"{{{_A_NS}}}srgbClr").set("val", hexval)
            for font_tag in ("majorFont", "minorFont"):
                latin = root.find(f".//{{{_A_NS}}}{font_tag}/{{{_A_NS}}}latin")
                if latin is not None:
                    latin.set("typeface", FONT)
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        except Exception:  # noqa: BLE001
            pass
        try:
            fill = master.background.fill
            fill.solid()
            fill.fore_color.rgb = RGBColor.from_string(th["bg"])
        except Exception:  # noqa: BLE001
            pass
    if reset_slide_backgrounds:
        for slide in prs.slides:
            bg = slide._element.cSld.bg
            if bg is not None:
                slide._element.cSld.remove(bg)
    return key


def _layout(prs, name: str, fallback_idx: int):
    for lay in prs.slide_layouts:
        if lay.name == name:
            return lay
    layouts = list(prs.slide_layouts)
    return layouts[min(fallback_idx, len(layouts) - 1)]


class _Geo:
    def __init__(self, prs):
        self.W, self.H = int(prs.slide_width), int(prs.slide_height)
        self.mx = int(self.W * 0.045)
        self.title_top = int(self.H * 0.06)
        self.title_h = int(self.H * 0.13)
        self.rule_top = int(self.H * 0.195)
        self.body_top = int(self.H * 0.235)
        self.body_h = int(self.H * 0.67)
        self.inner_w = self.W - 2 * self.mx


def _place(shape, x, y, w, h) -> None:
    shape.left, shape.top, shape.width, shape.height = int(x), int(y), int(w), int(h)


def _remove_shape(shape) -> None:
    el = shape._element
    el.getparent().remove(el)


def _placeholder(slide, idx: int):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            return ph
    return None


def _set_text(shape, text: str, *, size: float, bold: bool = False, color_theme=None, anchor=None,
              align=None) -> None:
    from pptx.util import Pt

    tf = shape.text_frame
    tf.word_wrap = True
    if anchor is not None:
        tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    for extra in list(tf.paragraphs[1:]):
        extra._p.getparent().remove(extra._p)
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    if align is not None:
        p.alignment = align
    for s in inline_spans(text):
        if s.br or not s.text:
            continue
        r = p.add_run()
        r.text = s.text
        r.font.size = Pt(size)
        r.font.bold = bold or s.bold or None
        if s.italic:
            r.font.italic = True
        if color_theme is not None:
            r.font.color.theme_color = color_theme


def _accent_bar(slide, x, y, w, h):
    from pptx.enum.dml import MSO_THEME_COLOR
    from pptx.enum.shapes import MSO_SHAPE

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(x), int(y), int(w), int(h))
    bar.fill.solid()
    bar.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
    bar.line.fill.background()
    bar.shadow.inherit = False
    bar.name = "Ren accent"
    return bar


def _slide_number(slide, g: _Geo) -> None:
    """Bottom-right slide-number *field* (stays correct after reordering)."""
    from lxml import etree
    from pptx.enum.text import PP_ALIGN

    box = slide.shapes.add_textbox(g.W - g.mx - int(g.W * 0.08), int(g.H * 0.925), int(g.W * 0.08), int(g.H * 0.05))
    box.name = "Ren slide number"
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    endp = p._p.find(f"{{{_A_NS}}}endParaRPr")
    if endp is not None:
        p._p.remove(endp)
    fld = etree.SubElement(p._p, f"{{{_A_NS}}}fld")
    fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}")
    fld.set("type", "slidenum")
    rpr = etree.SubElement(fld, f"{{{_A_NS}}}rPr")
    rpr.set("lang", "en-US")
    rpr.set("sz", "1000")
    solid = etree.SubElement(rpr, f"{{{_A_NS}}}solidFill")
    etree.SubElement(solid, f"{{{_A_NS}}}schemeClr").set("val", "tx2")
    etree.SubElement(fld, f"{{{_A_NS}}}t").text = "‹#›"


def normalize_bullets(bullets: Any) -> list[tuple[str, int]]:
    """Accepts strings (leading spaces / '-' markers give levels), nested lists, or {text, level} dicts."""
    out: list[tuple[str, int]] = []

    def walk(items: Any, level: int) -> None:
        if items is None:
            return
        if isinstance(items, str):
            for line in items.splitlines():
                if line.strip():
                    walk_line(line, level)
            return
        for it in items:
            if isinstance(it, (list, tuple)):
                walk(list(it), level + 1)
            elif isinstance(it, dict):
                text = str(it.get("text") or it.get("title") or "")
                lvl = int(it.get("level", level) or 0)
                if text:
                    out.append((text, max(0, min(4, lvl))))
                if it.get("children") or it.get("bullets"):
                    walk(it.get("children") or it.get("bullets"), lvl + 1)
            elif it is not None:
                walk_line(str(it), level)

    def walk_line(line: str, level: int) -> None:
        stripped = line.lstrip(" \t")
        indent = len(line) - len(stripped)
        lvl = level + indent // 2
        m = re.match(r"^([-*•+]|\d+[.)])\s+(.*)$", stripped)
        text = m.group(2) if m else stripped
        if text.strip():
            out.append((text.strip(), max(0, min(4, lvl))))

    walk(bullets, 0)
    return out


def _bullet_sizes(items: list[tuple[str, int]], compact: bool = False) -> float:
    n = len(items)
    total = sum(len(t) for t, _ in items)
    if n <= 5 and total < 380:
        base = 20.0
    elif n <= 8 and total < 650:
        base = 18.0
    elif n <= 11 and total < 950:
        base = 16.0
    else:
        base = 14.0
    return base - (2 if compact else 0)


def fill_bullets(tf, bullets: Any, *, explicit: bool = False, compact: bool = False) -> int:
    """Write bullet paragraphs into a text frame. `explicit` adds bullet glyphs (for plain text boxes)."""
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt

    if isinstance(bullets, list) and bullets and all(
            isinstance(b, tuple) and len(b) == 2 and isinstance(b[0], str) and isinstance(b[1], int) for b in bullets):
        items = list(bullets)  # already normalized
    else:
        items = normalize_bullets(bullets)
    tf.word_wrap = True
    base = _bullet_sizes(items, compact)
    first = tf.paragraphs[0]
    for extra in list(tf.paragraphs[1:]):
        extra._p.getparent().remove(extra._p)
    for r in list(first.runs):
        r._r.getparent().remove(r._r)
    for i, (text, level) in enumerate(items):
        p = first if i == 0 else tf.add_paragraph()
        p.level = min(level, 4)
        size = max(10.0, base - 2 * level)
        for s in inline_spans(text):
            if s.br or not s.text:
                continue
            r = p.add_run()
            r.text = s.text
            r.font.size = Pt(size)
            if s.bold:
                r.font.bold = True
            if s.italic:
                r.font.italic = True
            if s.code:
                r.font.name = MONO
            if s.link and re.match(r"^https?://", s.link):
                r.hyperlink.address = s.link
        p.space_after = Pt(9 if level == 0 else 4)
        pPr = p._p.get_or_add_pPr()
        if explicit:
            mar = Emu(int(285750 * (level + 1)))
            pPr.set("marL", str(int(mar)))
            pPr.set("indent", str(-285750))
        clr = pPr.makeelement(qn("a:buClr"), {})
        sc = clr.makeelement(qn("a:schemeClr"), {"val": "accent1"})
        clr.append(sc)
        pPr.insert_element_before(clr, "a:buSzTx", "a:buSzPct", "a:buSzPts", "a:buFontTx", "a:buFont", "a:buNone",
                                  "a:buAutoNum", "a:buChar", "a:buBlip", "a:tabLst", "a:defRPr", "a:extLst")
        if explicit:
            font = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
            char = pPr.makeelement(qn("a:buChar"), {"char": "•" if level == 0 else "–"})
            pPr.insert_element_before(font, "a:buNone", "a:buAutoNum", "a:buChar", "a:buBlip", "a:tabLst",
                                      "a:defRPr", "a:extLst")
            pPr.insert_element_before(char, "a:buBlip", "a:tabLst", "a:defRPr", "a:extLst")
    return len(items)


def add_table_shape(slide, rows: list[list[Any]], x, y, cx, cy, *, header: bool = True, max_rows: int = 14):
    from pptx.enum.dml import MSO_THEME_COLOR
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Pt

    rows = [list(r) for r in rows if r is not None]
    if not rows:
        raise ToolError("Table has no rows.")
    truncated = len(rows) > max_rows + (1 if header else 0)
    if truncated:
        rows = rows[: max_rows + (1 if header else 0)]
    ncols = max(len(r) for r in rows)
    rows = [r + [""] * (ncols - len(r)) for r in rows]
    nrows = len(rows)
    font = 14.0 if nrows <= 6 and ncols <= 4 else 12.0 if nrows <= 10 and ncols <= 6 else 10.0
    row_h = Emu(int(Pt(font * 2.2)))
    height = min(int(cy), row_h * nrows)
    shape = slide.shapes.add_table(nrows, ncols, int(x), int(y), int(cx), int(height))
    table = shape.table
    table.first_row = header
    table.horz_banding = False
    texts = [[_cell_text(v) for v in r] for r in rows]
    lens = [min(36, max([len(texts[i][j]) for i in range(nrows)] + [3])) for j in range(ncols)]
    total = sum(lens) or 1
    for j in range(ncols):
        table.columns[j].width = Emu(int(int(cx) * lens[j] / total))
    numeric = numeric_columns(texts[1:] if header else texts, ncols)
    for i in range(nrows):
        table.rows[i].height = row_h
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.text = texts[i][j]
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Emu(int(Pt(6)))
            cell.margin_top = cell.margin_bottom = Emu(int(Pt(3)))
            cell.fill.solid()
            is_head = header and i == 0
            if is_head:
                cell.fill.fore_color.theme_color = MSO_THEME_COLOR.TEXT_1
            elif (i % 2 == 0 and header) or (i % 2 == 1 and not header):
                cell.fill.fore_color.theme_color = MSO_THEME_COLOR.BACKGROUND_2
            else:
                cell.fill.fore_color.theme_color = MSO_THEME_COLOR.BACKGROUND_1
            for p in cell.text_frame.paragraphs:
                if numeric[j] and not is_head:
                    p.alignment = PP_ALIGN.RIGHT
                for r in p.runs:
                    r.font.size = Pt(font)
                    r.font.bold = True if is_head else None
                    r.font.color.theme_color = MSO_THEME_COLOR.BACKGROUND_1 if is_head else MSO_THEME_COLOR.TEXT_1
    return shape, truncated


CHART_TYPES = ("bar", "column", "line", "pie", "doughnut", "area", "stacked_bar", "stacked_column")


def add_chart_shape(slide, chart_type: str, categories: list[Any], series: Any, x, y, cx, cy, *,
                    title: str | None = None, number_format: str | None = None):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Pt

    types = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
             "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT,
             "area": XL_CHART_TYPE.AREA, "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
             "stacked_column": XL_CHART_TYPE.COLUMN_STACKED}
    ct = (chart_type or "column").lower()
    if ct not in types:
        raise ToolError(f"Unsupported chart type '{chart_type}'. Use one of: {', '.join(CHART_TYPES)}.")
    pairs: list[tuple[str, list]] = []
    if isinstance(series, dict):
        pairs = [(str(k), list(v or [])) for k, v in series.items()]
    elif isinstance(series, list):
        for i, s in enumerate(series):
            if isinstance(s, dict):
                pairs.append((str(s.get("name") or f"Series {i + 1}"), list(s.get("values") or [])))
            elif isinstance(s, (list, tuple)):
                pairs.append((f"Series {i + 1}", list(s)))
    if not pairs or not categories:
        raise ToolError("A chart needs categories and at least one series of values.")
    if ct in ("pie", "doughnut"):
        pairs = pairs[:1]
    data = CategoryChartData()
    data.categories = [str(c) for c in categories]
    parsed: list[list[float | None]] = []
    for name, values in pairs:
        vals: list[float | None] = []
        for v in values[: len(categories)]:
            try:
                vals.append(None if v is None or v == "" else float(str(v).replace(",", "")))
            except ValueError:
                vals.append(None)
        vals += [None] * (len(categories) - len(vals))
        parsed.append(vals)
        data.add_series(name, vals)
    gf = slide.shapes.add_chart(types[ct], int(x), int(y), int(cx), int(cy), data)
    chart = gf.chart
    chart.font.size = Pt(12)
    chart.has_legend = len(pairs) > 1 or ct in ("pie", "doughnut")
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    if title:
        chart.has_title = True
        chart.chart_title.text_frame.text = title
        for p in chart.chart_title.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(14)
                r.font.bold = True
    else:
        chart.has_title = False
    plot = chart.plots[0]
    if ct in ("pie", "doughnut"):
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.show_percentage = True
        dl.show_value = False
        dl.number_format = "0%"
        dl.number_format_is_linked = False
        dl.font.size = Pt(11)
    else:
        try:
            plot.gap_width = 70
        except Exception:  # noqa: BLE001
            pass
        if len(pairs) == 1 and ct in ("bar", "column") and len(categories) <= 12:
            vals = [v for v in parsed[0] if v is not None]
            fmt = number_format or ("#,##0" if all(float(v).is_integer() for v in vals) else "#,##0.0#")
            plot.has_data_labels = True
            plot.data_labels.font.size = Pt(10)
            plot.data_labels.number_format = fmt
            plot.data_labels.number_format_is_linked = False
        try:
            va = chart.value_axis
            va.has_major_gridlines = True
            va.major_gridlines.format.line.width = Pt(0.5)
            if number_format:
                va.tick_labels.number_format = number_format
                va.tick_labels.number_format_is_linked = False
            va.format.line.fill.background()
        except Exception:  # noqa: BLE001
            pass
    return gf


def add_picture_fit(slide, image_path: Path, x, y, cx, cy):
    from PIL import Image

    with Image.open(image_path) as im:
        w, h = im.size
    scale = min(int(cx) / max(1, w), int(cy) / max(1, h))
    pw, ph = int(w * scale), int(h * scale)
    return slide.shapes.add_picture(str(image_path), int(x) + (int(cx) - pw) // 2, int(y) + (int(cy) - ph) // 2,
                                    pw, ph)


def move_slide(prs, old_index: int, new_index: int) -> None:
    lst = prs.slides._sldIdLst
    items = list(lst)
    el = items[old_index]
    lst.remove(el)
    lst.insert(max(0, min(new_index, len(items) - 1)), el)


def delete_slide(prs, index: int) -> None:
    lst = prs.slides._sldIdLst
    sld = list(lst)[index]
    prs.part.drop_rel(sld.rId)
    lst.remove(sld)


def add_title_slide(prs, title: str, subtitle: str | None = None, notes: str | None = None):
    from pptx.enum.dml import MSO_THEME_COLOR
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

    g = _Geo(prs)
    slide = prs.slides.add_slide(_layout(prs, "Title Slide", 0))
    x = int(g.mx * 1.6)
    w = g.W - 2 * x
    ts = slide.shapes.title
    if ts is None:
        ts = slide.shapes.add_textbox(x, int(g.H * 0.30), w, int(g.H * 0.26))
    _place(ts, x, int(g.H * 0.28), w, int(g.H * 0.27))
    _set_text(ts, title or "Untitled", size=40 if len(title or "") <= 48 else 32, bold=True,
              anchor=MSO_ANCHOR.BOTTOM, align=PP_ALIGN.LEFT)
    _accent_bar(slide, x, int(g.H * 0.585), int(g.W * 0.07), int(g.H * 0.009))
    sub = _placeholder(slide, 1)
    if subtitle:
        if sub is None:
            sub = slide.shapes.add_textbox(x, int(g.H * 0.62), w, int(g.H * 0.15))
        _place(sub, x, int(g.H * 0.62), w, int(g.H * 0.16))
        _set_text(sub, subtitle, size=18, color_theme=MSO_THEME_COLOR.TEXT_2, anchor=MSO_ANCHOR.TOP,
                  align=PP_ALIGN.LEFT)
    elif sub is not None:
        _remove_shape(sub)
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx not in (0, 1):
            _remove_shape(ph)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def add_section_slide(prs, title: str, subtitle: str | None = None, notes: str | None = None):
    from pptx.enum.dml import MSO_THEME_COLOR
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

    g = _Geo(prs)
    slide = prs.slides.add_slide(_layout(prs, "Title Only", 5))
    x = int(g.mx * 1.6)
    ts = slide.shapes.title or slide.shapes.add_textbox(x, int(g.H * 0.36), g.W - 2 * x, int(g.H * 0.2))
    _place(ts, x, int(g.H * 0.33), g.W - 2 * x, int(g.H * 0.2))
    _set_text(ts, title or "", size=36, bold=True, anchor=MSO_ANCHOR.BOTTOM, align=PP_ALIGN.LEFT)
    _accent_bar(slide, x, int(g.H * 0.56), int(g.W * 0.05), int(g.H * 0.008))
    if subtitle:
        box = slide.shapes.add_textbox(x, int(g.H * 0.59), g.W - 2 * x, int(g.H * 0.15))
        _set_text(box, subtitle, size=18, color_theme=MSO_THEME_COLOR.TEXT_2, anchor=MSO_ANCHOR.TOP,
                  align=PP_ALIGN.LEFT)
    _slide_number(slide, g)
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx != 0:
            _remove_shape(ph)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def add_content_slide(prs, spec: dict, *, image_resolver: ImageResolver | None = None):
    """One slide from a spec: {title, bullets, notes, subtitle, table, chart, image}. Picks a layout:
    section (title only) · bullets · visual (table/chart/image) · bullets + visual side by side."""
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

    title = str(spec.get("title") or "").strip()
    bullets = spec.get("bullets") or spec.get("points") or spec.get("content")
    notes = spec.get("notes")
    table = spec.get("table")
    chart = spec.get("chart")
    image = spec.get("image")
    has_visual = bool(table or chart or image)
    items = normalize_bullets(bullets) if bullets else []
    if not items and not has_visual:
        return add_section_slide(prs, title, spec.get("subtitle"), notes)
    g = _Geo(prs)
    layout = _layout(prs, "Title and Content", 1) if (items and not has_visual) else _layout(prs, "Title Only", 5)
    slide = prs.slides.add_slide(layout)
    ts = slide.shapes.title or slide.shapes.add_textbox(g.mx, g.title_top, g.inner_w, g.title_h)
    _place(ts, g.mx, g.title_top, g.inner_w, g.title_h)
    _set_text(ts, title, size=30 if len(title) <= 55 else 24, bold=True, anchor=MSO_ANCHOR.BOTTOM,
              align=PP_ALIGN.LEFT)
    _accent_bar(slide, g.mx, g.rule_top, int(g.W * 0.06), int(g.H * 0.007))
    body = _placeholder(slide, 1)
    vis_x, vis_w = g.mx, g.inner_w
    if items:
        if has_visual:
            text_w = int(g.inner_w * 0.42)
            box = slide.shapes.add_textbox(g.mx, g.body_top, text_w, g.body_h)
            fill_bullets(box.text_frame, items, explicit=True, compact=True)
            vis_x, vis_w = g.mx + text_w + int(g.W * 0.03), g.inner_w - text_w - int(g.W * 0.03)
        else:
            if body is None:
                body = slide.shapes.add_textbox(g.mx, g.body_top, g.inner_w, g.body_h)
                fill_bullets(body.text_frame, items, explicit=True)
            else:
                _place(body, g.mx, g.body_top, g.inner_w, g.body_h)
                fill_bullets(body.text_frame, items)
    notes_extra = []
    if table:
        rows = table.get("rows") if isinstance(table, dict) else table
        if isinstance(table, dict) and table.get("header"):
            rows = [table["header"]] + list(rows or [])
        _, truncated = add_table_shape(slide, rows or [], vis_x, g.body_top, vis_w, g.body_h)
        if truncated:
            notes_extra.append("Table truncated to fit the slide.")
    elif chart:
        add_chart_shape(slide, chart.get("type", "column"), chart.get("categories") or [], chart.get("series"),
                        vis_x, g.body_top, vis_w, g.body_h, title=chart.get("title"),
                        number_format=chart.get("number_format"))
    elif image:
        path = image_resolver(str(image)) if image_resolver else None
        if path is None:
            raise ToolError(f"Image not found in the workspace: {image}")
        add_picture_fit(slide, path, vis_x, g.body_top, vis_w, g.body_h)
    for ph in list(slide.placeholders):
        idx = ph.placeholder_format.idx
        if idx == 0 or (idx == 1 and items and not has_visual):
            continue
        _remove_shape(ph)
    _slide_number(slide, g)
    note_text = "\n".join([n for n in [notes] + notes_extra if n])
    if note_text:
        slide.notes_slide.notes_text_frame.text = note_text
    return slide


def new_presentation(theme: str = "light", widescreen: bool = True):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    if widescreen:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    apply_deck_theme(prs, theme)
    return prs


def build_deck(path: Path | str | None, title: str | None, slides: list[dict], *, subtitle: str | None = None,
               theme: str = "light", image_resolver: ImageResolver | None = None, author: str | None = None):
    prs = new_presentation(theme)
    if title:
        add_title_slide(prs, title, subtitle)
    for spec in slides or []:
        if isinstance(spec, str):
            spec = {"title": spec}
        add_content_slide(prs, dict(spec), image_resolver=image_resolver)
    cp = prs.core_properties
    cp.title = title or ""
    cp.author = author or "Ren"
    if path:
        prs.save(str(path))
    return prs


def _iter_shapes(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shp in shapes:
        if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shp.shapes)
        else:
            yield shp


def slide_md(slide, index: int, include_notes: bool = True) -> str:
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

    title_shape = slide.shapes.title
    title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
    lines = [f"## Slide {index}: {title}" if title else f"## Slide {index}"]
    title_id = title_shape.shape_id if title_shape is not None else None
    shapes = [s for s in _iter_shapes(slide.shapes) if s.shape_id != title_id]
    shapes.sort(key=lambda s: ((s.top or 0) // 100000, s.left or 0))
    for shp in shapes:
        if getattr(shp, "has_table", False) and shp.has_table:
            rows = [[c.text.strip() for c in r.cells] for r in shp.table.rows]
            if rows:
                lines.append(md_table(rows[0], rows[1:]))
            continue
        if getattr(shp, "has_chart", False) and shp.has_chart:
            ch = shp.chart
            try:
                names = [s.name for s in ch.plots[0].series]
                cats = list(ch.plots[0].categories)
                kind = str(ch.chart_type).split(".")[-1].split(" ")[0].lower()
                lines.append(f"[Chart ({kind}): series {', '.join(names)} over {len(cats)} categories"
                             f"{' — ' + ', '.join(map(str, cats[:6])) if cats else ''}]")
            except Exception:  # noqa: BLE001
                lines.append("[Chart]")
            continue
        if shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
            lines.append(f"[Image: {shp.name}]")
            continue
        if not shp.has_text_frame:
            continue
        is_body = False
        try:
            if shp.is_placeholder and shp.placeholder_format.type in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT):
                is_body = True
        except Exception:  # noqa: BLE001
            pass
        for p in shp.text_frame.paragraphs:
            text = "".join(r.text for r in p.runs).strip()
            if not text:
                continue
            has_bullet = is_body or p._p.pPr is not None and p._p.pPr.find(
                "{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is not None
            lines.append(("  " * p.level + "- " + text) if has_bullet else text)
    if include_notes and slide.has_notes_slide:
        nt = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
        if nt:
            lines.append("> Notes: " + nt.replace("\n", " "))
    return "\n".join(lines)


def pptx_to_markdown(src, include_notes: bool = True) -> str:
    from pptx import Presentation

    prs = Presentation(str(src)) if isinstance(src, (str, Path)) else src
    return "\n\n".join(slide_md(s, i, include_notes) for i, s in enumerate(prs.slides, 1)) + "\n"


def markdown_to_slide_specs(md: str) -> dict:
    """Split markdown into a deck outline: '# ' = deck title (a paragraph right after it = subtitle),
    '## ' (or '---') = new slide, lists/paragraphs = bullets, tables = slide tables, 'Notes:' = speaker notes.
    Returns {"title", "subtitle", "slides"}."""
    deck_title: str | None = None
    subtitle: str | None = None
    slides: list[dict] = []
    cur: dict | None = None

    def ensure() -> dict:
        nonlocal cur
        if cur is None:
            cur = {"title": "", "bullets": []}
            slides.append(cur)
        return cur

    def flatten(b: Block, level: int) -> list[str]:
        out = []
        for item in b.items:
            for blk in item:
                if blk.kind == "para":
                    out.append("  " * level + "- " + spans_md(blk.spans))
                elif blk.kind == "list":
                    out.extend(flatten(blk, level + 1))
        return out

    for b in parse_markdown(md):
        if b.kind == "heading" and b.level == 1 and deck_title is None and not slides:
            deck_title = spans_text(b.spans)
            continue
        if b.kind == "heading" and b.level <= 2:
            cur = {"title": spans_text(b.spans), "bullets": []}
            slides.append(cur)
            continue
        if b.kind == "hr":
            cur = None
            continue
        if b.kind == "para" and deck_title is not None and subtitle is None and not slides and cur is None:
            subtitle = spans_text(b.spans)
            continue
        s = ensure()
        if b.kind == "list":
            s["bullets"].extend(flatten(b, 0))
        elif b.kind == "para":
            text = spans_md(b.spans)
            if re.match(r"^(\*\*)?notes?:", text, re.I):
                s["notes"] = re.sub(r"^(\*\*)?notes?:(\*\*)?\s*", "", text, flags=re.I)
            else:
                s["bullets"].append(text)
        elif b.kind == "heading":
            s["bullets"].append("**" + spans_text(b.spans) + "**")
        elif b.kind == "table":
            s["table"] = [[spans_text(c) for c in b.header]] + [[spans_text(c) for c in r] for r in b.rows]
        elif b.kind == "quote":
            for ch in b.children:
                if ch.kind == "para":
                    s["bullets"].append("*" + spans_text(ch.spans) + "*")
        elif b.kind == "image":
            s["image"] = b.src
        elif b.kind == "code":
            s["bullets"].extend(f"`{ln}`" for ln in b.text.splitlines()[:8] if ln.strip())
    return {"title": deck_title, "subtitle": subtitle,
            "slides": [s for s in slides if s.get("title") or s.get("bullets") or s.get("table") or s.get("image")]}


# --------------------------------------------------------------------------------------------- tables / XLSX
_INT_RE = re.compile(r"^-?(0|[1-9]\d{0,14})$")
_FLOAT_RE = re.compile(r"^-?(0|[1-9]\d*)?\.\d+$|^-?(0|[1-9]\d*)\.\d*$")
_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_DATETIME_RE = re.compile(r"^\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(:\d{2}(\.\d+)?)?$")


def coerce_scalar(v: Any) -> Any:
    """Turn clean numeric / ISO-date strings into numbers / dates (keeps codes like '00123' as text)."""
    if not isinstance(v, str):
        return v
    s = v.strip()
    if not s or s.startswith("="):
        return v
    if _INT_RE.match(s):
        return int(s)
    if _FLOAT_RE.match(s):
        try:
            return float(s)
        except ValueError:
            return v
    if _DATE_RE.match(s):
        try:
            return datetime.strptime(s, "%Y-%m-%d").date()
        except ValueError:
            return v
    if _DATETIME_RE.match(s):
        try:
            return datetime.fromisoformat(s.replace(" ", "T"))
        except ValueError:
            return v
    return v


def _looks_like_md_table(s: str) -> bool:
    lines = [ln.strip() for ln in s.strip().splitlines() if ln.strip()]
    return len(lines) >= 2 and lines[0].startswith("|") and re.match(r"^\|?\s*:?-{2,}", lines[1]) is not None


def _parse_md_table(s: str) -> tuple[list[str], list[list[Any]]]:
    lines = [ln.strip() for ln in s.strip().splitlines() if ln.strip().startswith("|")]

    def cells(line: str) -> list[str]:
        line = line.strip().strip("|")
        return [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", line)]

    header = cells(lines[0])
    rows = [[coerce_scalar(c) for c in cells(ln)] for ln in lines[2:]]
    return header, rows


def parse_csv_text(s: str, *, delimiter: str | None = None) -> list[list[str]]:
    sample = s[:4096]
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = "\t" if sample.count("\t") > sample.count(",") else ","
    return [row for row in csv.reader(io.StringIO(s), delimiter=delimiter)]


def table_from_any(data: Any, *, header: bool = True) -> tuple[list[str] | None, list[list[Any]]]:
    """Coerce JSON records / row lists / column dicts / CSV / markdown-table text into (header, rows)."""
    if data is None:
        return None, []
    if isinstance(data, str):
        s = data.strip()
        if not s:
            return None, []
        if s[0] in "[{":
            try:
                return table_from_any(json.loads(s), header=header)
            except json.JSONDecodeError:
                pass
        if _looks_like_md_table(s):
            return _parse_md_table(s)
        rows = parse_csv_text(s)
        if header and rows:
            return [str(h).strip() for h in rows[0]], [[coerce_scalar(c) for c in r] for r in rows[1:]]
        return None, [[coerce_scalar(c) for c in r] for r in rows]
    if isinstance(data, dict):
        if "rows" in data or "data" in data:
            rows = data.get("rows") if "rows" in data else data.get("data")
            hdr = data.get("columns") or data.get("header") or data.get("headers")
            h2, r2 = table_from_any(rows, header=hdr is None and header)
            return ([str(x) for x in hdr] if hdr else h2), r2
        if data and all(isinstance(v, (list, tuple)) for v in data.values()):
            cols = [str(c) for c in data]
            n = max(len(v) for v in data.values())
            return cols, [[(list(data[c])[i] if i < len(data[c]) else None) for c in data] for i in range(n)]
        return [str(k) for k in data], [list(data.values())]
    if isinstance(data, (list, tuple)):
        data = list(data)
        if not data:
            return None, []
        if all(isinstance(r, dict) for r in data):
            cols: list[str] = []
            for r in data:
                for k in r:
                    if str(k) not in cols:
                        cols.append(str(k))
            return cols, [[r.get(c, r.get(c)) for c in cols] for r in data]
        rows = [list(r) if isinstance(r, (list, tuple)) else [r] for r in data]
        if header and rows and rows[0] and all(isinstance(v, str) for v in rows[0]) and len(rows) > 1:
            return [str(v) for v in rows[0]], rows[1:]
        return None, rows
    raise ToolError("Unsupported table data: pass a list of row lists, a list of objects, or CSV text.")


def clean_cell(v: Any) -> Any:
    """Value safe for openpyxl (strips illegal XML control chars; serialises containers)."""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    if v is None or isinstance(v, bool):
        return v
    if isinstance(v, float):
        return None if (math.isnan(v) or math.isinf(v)) else v
    if hasattr(v, "to_pydatetime"):  # pandas Timestamp
        try:
            import pandas as pd

            if pd.isna(v):
                return None
        except Exception:  # noqa: BLE001
            pass
        v = v.to_pydatetime()
    if isinstance(v, datetime) and v.tzinfo is not None:
        return v.replace(tzinfo=None)  # Excel has no time zones
    if isinstance(v, (int, datetime, date, dtime, Decimal)):
        return v
    if isinstance(v, (dict, list, tuple)):
        v = json.dumps(v, ensure_ascii=False, default=str)
    else:
        jv = jsonable(v)
        if isinstance(jv, (int, float, bool)) or jv is None:
            return jv
        v = str(v)
    return ILLEGAL_CHARACTERS_RE.sub("", v)


def safe_sheet_name(name: str, existing: list[str] | None = None) -> str:
    s = re.sub(r"[\[\]:*?/\\]", "-", str(name or "Sheet")).strip().strip("'") or "Sheet"
    s = s[:31]
    if existing:
        base, i = s, 2
        low = {e.lower() for e in existing}
        while s.lower() in low:
            suffix = f" ({i})"
            s = base[: 31 - len(suffix)] + suffix
            i += 1
    return s


def xlsx_header_style(ws, row: int = 1, ncols: int | None = None, col_start: int = 1) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    ncols = ncols or ws.max_column
    for c in range(col_start, col_start + ncols):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color=INK)
        cell.fill = PatternFill("solid", fgColor=SAND)
        cell.border = Border(bottom=Side(style="medium", color=ACCENT))
        cell.alignment = Alignment(vertical="center")


def xlsx_autofit(ws, min_width: float = 6, max_width: float = 60, sample: int = 2000) -> None:
    from openpyxl.utils import get_column_letter

    widths: dict[int, int] = {}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, sample)):
        for cell in row:
            v = cell.value
            if v is None:
                continue
            if isinstance(v, (datetime, date)):
                s = "2024-01-01 00:00" if isinstance(v, datetime) else "2024-01-01"
            elif isinstance(v, float):
                s = f"{v:,.2f}"
            elif isinstance(v, int):
                s = f"{v:,}"
            else:
                s = str(v)
            n = max(len(x) for x in s.split("\n")) if s else 0
            if cell.font is not None and cell.font.b:
                n = int(n * 1.1) + 1
            col = cell.column if isinstance(cell.column, int) else None
            if col:
                widths[col] = max(widths.get(col, 0), n)
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(max_width, w + 2))


def xlsx_number_formats(ws, first_data_row: int = 2) -> None:
    """Thousands separators for decimals, ISO-ish date format for dates (integers left alone: IDs, years)."""
    for col in ws.iter_cols(min_row=first_data_row, max_row=min(ws.max_row, 5000)):
        for cell in col:
            v = cell.value
            if isinstance(v, float) and cell.number_format == "General":
                cell.number_format = "#,##0.00"
            elif isinstance(v, datetime) and cell.number_format == "General":
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif isinstance(v, date) and not isinstance(v, datetime) and cell.number_format == "General":
                cell.number_format = "yyyy-mm-dd"


def write_sheet(ws, header: list[Any] | None, rows: list[list[Any]], *, style: bool = True) -> None:
    if header:
        ws.append([clean_cell(h) for h in header])
    for r in rows:
        ws.append([clean_cell(coerce_scalar(v)) for v in r])
    if style:
        if header:
            xlsx_header_style(ws, 1, len(header))
            ws.freeze_panes = "A2"
            if ws.max_row > 1:
                ws.auto_filter.ref = ws.dimensions
        xlsx_number_formats(ws, 2 if header else 1)
        xlsx_autofit(ws)


def write_xlsx(path: Path | str, sheets: dict[str, tuple[list[Any] | None, list[list[Any]]]]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    names: list[str] = []
    for name, (header, rows) in sheets.items():
        nm = safe_sheet_name(name, names)
        names.append(nm)
        write_sheet(wb.create_sheet(nm), header, rows)
    if not names:
        wb.create_sheet("Sheet1")
    wb.save(str(path))


def write_csv(path: Path | str, header: list[Any] | None, rows: list[list[Any]], delimiter: str = ",") -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=delimiter)
        if header:
            w.writerow(["" if h is None else h for h in header])
        for r in rows:
            w.writerow(["" if v is None else (v.isoformat() if isinstance(v, (datetime, date)) else v) for v in r])


def read_csv_rows(path: Path, *, max_rows: int | None = None) -> list[list[Any]]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    delim = "\t" if path.suffix.lower() == ".tsv" else None
    rows = parse_csv_text(text, delimiter=delim)
    if max_rows is not None:
        rows = rows[:max_rows]
    return [[coerce_scalar(c) for c in r] for r in rows]


def xlsx_to_markdown(path: Path, max_rows: int = 50, sheets: list[str] | None = None) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    out = []
    try:
        for ws in wb.worksheets:
            if sheets and ws.title not in sheets:
                continue
            rows = []
            for row in ws.iter_rows(values_only=True):
                if any(v is not None and str(v).strip() != "" for v in row):
                    rows.append(list(row))
                if len(rows) > max_rows:
                    break
            dims = f"{ws.max_row or 0} rows × {ws.max_column or 0} cols"
            out.append(f"## Sheet: {ws.title} ({dims})")
            if rows:
                width = max(len(r) for r in rows)
                while width > 0 and all((len(r) < width or r[width - 1] in (None, "")) for r in rows):
                    width -= 1
                rows = [r[:width] for r in rows]
                out.append(md_table([("" if v is None else str(v)) for v in rows[0]], rows[1:max_rows]))
                if len(rows) > max_rows:
                    out.append(f"…(showing first {max_rows - 1} data rows)")
            else:
                out.append("(empty)")
    finally:
        wb.close()
    return "\n\n".join(out) + "\n"


# --------------------------------------------------------------------------------------------- generic extraction
TEXT_EXTS = (".md", ".markdown", ".txt", ".text", ".log", ".rst", ".yaml", ".yml", ".toml", ".ini", ".xml", ".py",
             ".js", ".ts", ".sql", ".sh", ".css")


def read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-16") if raw[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def extract_text(path: Path, *, max_rows: int = 200) -> str:
    """Best-effort text/markdown for any supported office/data/text file."""
    ext = path.suffix.lower()
    if ext == ".docx":
        return docx_to_markdown(path)
    if ext == ".pptx":
        return pptx_to_markdown(path)
    if ext in (".xlsx", ".xlsm"):
        return xlsx_to_markdown(path, max_rows=max_rows)
    if ext in (".csv", ".tsv"):
        rows = read_csv_rows(path, max_rows=max_rows + 1)
        return md_table(rows[0], rows[1:]) if rows else ""
    if ext == ".pdf":
        n, pages = pdf_pages_text(path)
        return "\n\n".join(f"--- Page {i} ---\n{t}" for i, t in pages)
    if ext in (".html", ".htm"):
        title, text = html_to_text(read_text_file(path))
        return (f"# {title}\n\n" if title else "") + text
    if ext == ".json":
        try:
            return json.dumps(json.loads(read_text_file(path)), indent=2, ensure_ascii=False)
        except json.JSONDecodeError:
            return read_text_file(path)
    if ext in TEXT_EXTS or ext == "":
        return read_text_file(path)
    raise ToolError(f"Can't extract text from '{path.name}' (unsupported type '{ext}').")


def workspace_image_resolver(ctx: ToolContext) -> ImageResolver:
    """Resolve markdown image sources to workspace files (never fetches remote URLs)."""
    def resolve(src: str) -> Path | None:
        if not src or re.match(r"^[a-z][a-z0-9+.-]*:", src, re.I):
            return None
        try:
            p = ctx.workspace.resolve(src.split("?")[0].split("#")[0], must_exist=True)
        except WorkspaceError:
            return None
        if p.is_file() and p.suffix.lower() in (".png", ".jpg", ".jpeg", ".gif", ".bmp"):
            return p
        return None

    return resolve
