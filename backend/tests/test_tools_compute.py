"""Compute & engineering tool groups: registration, filesystem, shell, code and git (local execution)."""
import asyncio
import json
import re
import shutil
import sys
import uuid
import zipfile

import pytest

import app.tools.builtin.code  # noqa: F401  (registers the groups)
import app.tools.builtin.containers  # noqa: F401
import app.tools.builtin.database  # noqa: F401
import app.tools.builtin.filesystem  # noqa: F401
import app.tools.builtin.git  # noqa: F401
import app.tools.builtin.shell  # noqa: F401
from app.core.config import settings
from app.tools.base import Risk, ToolContext, ToolError, registry
from app.tools.builtin import _exec

HAS_GIT = shutil.which("git") is not None
HAS_NODE = shutil.which("node") is not None


async def call(ctx, tool, /, **args):
    spec = registry.get(tool)
    assert spec is not None, f"tool {tool} is not registered"
    return await spec.handler(ctx, **spec.validate(args))


def uniq(prefix: str) -> str:
    return f"{prefix}-{uuid.uuid4().hex[:8]}"


# =========================================================================== registration
EXPECTED = {
    "filesystem": ("fs", "core", "folder", True, {
        "fs_list", "fs_find", "fs_search", "fs_read", "fs_write", "fs_append", "fs_edit", "fs_mkdir", "fs_move",
        "fs_copy", "fs_delete", "fs_metadata", "fs_diff", "fs_archive", "fs_extract", "fs_extract_text"}),
    "shell": ("shell", "engineering", "terminal", True, {
        "shell_exec", "shell_exec_background", "shell_poll", "shell_kill", "shell_processes", "shell_which",
        "shell_env", "shell_install_package"}),
    "code": ("code", "engineering", "code", True, {
        "code_python", "code_javascript", "code_test", "code_lint", "code_format"}),
    "git": ("git", "engineering", "git-branch", True, {
        "git_status", "git_log", "git_diff", "git_show", "git_branch", "git_checkout", "git_create_branch",
        "git_add", "git_commit", "git_pull", "git_push", "git_merge", "git_clone", "git_init"}),
    "docker": ("docker", "infrastructure", "container", False, {
        "docker_list", "docker_images", "docker_logs", "docker_inspect", "docker_build", "docker_run",
        "docker_stop"}),
    "kubernetes": ("k8s", "infrastructure", "server", False, {
        "k8s_get", "k8s_describe", "k8s_logs", "k8s_apply", "k8s_restart", "k8s_scale", "k8s_rollback"}),
    "database": ("db", "data", "database", True, {
        "db_list_connections", "db_list_tables", "db_describe_table", "db_search_schema", "db_query", "db_explain",
        "db_execute"}),
}


def test_groups_and_schemas_registered():
    for gid, (prefix, category, icon, default_enabled, names) in EXPECTED.items():
        group = registry.group(gid)
        assert group is not None, gid
        assert (group.prefix, group.category, group.icon, group.default_enabled) == (
            prefix, category, icon, default_enabled), gid
        assert {t.name for t in group.tools} == names, gid
        for spec in group.tools:
            assert re.match(r"^[a-zA-Z0-9_-]{1,64}$", spec.name)
            assert len(spec.description) > 20, spec.name
            assert spec.parameters["type"] == "object"
            for pname, prop in spec.parameters["properties"].items():
                assert prop.get("description"), f"{spec.name}.{pname} has no description"
            if spec.risk in (Risk.HIGH, Risk.CRITICAL):
                assert spec.preview is not None, f"{spec.name} needs an approval preview"
            assert spec.openai_schema()["function"]["name"] == spec.name
    assert registry.group("database").connector_kinds == ("postgres", "mysql", "sqlite")


def test_risk_and_scope_policy(monkeypatch):
    get = registry.get
    assert (get("fs_read").risk, get("fs_read").scope) == (Risk.LOW, "read")
    assert (get("fs_write").risk, get("fs_write").scope) == (Risk.MEDIUM, "write")
    assert (get("fs_delete").risk, get("fs_delete").scope) == (Risk.MEDIUM, "write")
    assert get("git_status").risk == Risk.LOW and get("git_commit").risk == Risk.MEDIUM
    assert get("git_clone").risk == Risk.MEDIUM and get("git_pull").risk == Risk.MEDIUM
    assert get("git_merge").risk == Risk.HIGH
    assert get("git_push").effective_risk({"branch": "feature/x"}) == Risk.HIGH
    assert get("git_push").effective_risk({"branch": "main"}) == Risk.CRITICAL
    assert get("git_push").effective_risk({"branch": "feature/x", "force": True}) == Risk.CRITICAL
    assert (get("docker_list").risk, get("docker_logs").risk, get("docker_inspect").risk) == (Risk.LOW,) * 3
    assert (get("docker_build").risk, get("docker_run").risk) == (Risk.HIGH, Risk.HIGH)
    assert (get("docker_stop").risk, get("docker_stop").scope) == (Risk.CRITICAL, "admin")
    for name in ("k8s_restart", "k8s_scale", "k8s_rollback"):
        assert (get(name).risk, get(name).scope) == (Risk.CRITICAL, "admin")
    assert get("k8s_apply").effective_risk({"dry_run": True}) == Risk.LOW
    assert get("k8s_apply").effective_risk({}) == Risk.CRITICAL
    assert get("db_query").risk == Risk.LOW and get("db_execute").risk == Risk.CRITICAL
    assert get("shell_install_package").effective_risk({}) == Risk.HIGH
    monkeypatch.setattr(settings, "sandbox_url", None)
    for name in ("shell_exec", "shell_exec_background", "code_python", "code_javascript", "code_test"):
        assert get(name).effective_risk({}) == Risk.HIGH, name
    monkeypatch.setattr(settings, "sandbox_url", "http://sandbox:8080")
    for name in ("shell_exec", "shell_exec_background", "code_python", "code_javascript", "code_test"):
        assert get(name).effective_risk({}) == Risk.MEDIUM, name
    assert get("shell_install_package").effective_risk({}) == Risk.HIGH
    preview = get("shell_exec").render_preview({"command": "rm -rf build", "cwd": "projects/app"})
    assert "rm -rf build" in preview and "sandbox" in preview


def test_connector_kinds_registered():
    from app.tools.connectors import connector_kinds

    for kind in ("postgres", "mysql", "sqlite"):
        spec = connector_kinds[kind]
        assert spec.test is not None and "database" in spec.tool_groups
        assert "allow_writes" in {f.key for f in spec.fields}
    assert connector_kinds["postgres"].secret_keys == {"password"}
    assert connector_kinds["mysql"].secret_keys == {"password"}
    assert "sslmode" in {f.key for f in connector_kinds["postgres"].fields}


# =========================================================================== filesystem
async def test_fs_write_read_edit_search_diff(tool_ctx):
    d = f"workspace/{uniq('fs')}"
    assert "Created" in await call(tool_ctx, "fs_write", path=f"{d}/notes.md", content="alpha\nbeta\ngamma\n")
    assert await call(tool_ctx, "fs_read", path=f"{d}/notes.md") == "alpha\nbeta\ngamma"
    part = await call(tool_ctx, "fs_read", path=f"{d}/notes.md", offset=2, limit=1, line_numbers=True)
    assert "2\tbeta" in part and "offset=3" in part

    assert "line 2" in await call(tool_ctx, "fs_edit", path=f"{d}/notes.md", old_text="beta", new_text="BETA")
    with pytest.raises(ToolError, match="not found"):
        await call(tool_ctx, "fs_edit", path=f"{d}/notes.md", old_text="zeta", new_text="x")
    await call(tool_ctx, "fs_write", path=f"{d}/dup.txt", content="x = 1\nx = 1\n")
    with pytest.raises(ToolError, match="Found 2 occurrences"):
        await call(tool_ctx, "fs_edit", path=f"{d}/dup.txt", old_text="x = 1", new_text="y = 2")
    await call(tool_ctx, "fs_edit", path=f"{d}/dup.txt", old_text="x = 1", new_text="y = 2", count=2)
    assert await call(tool_ctx, "fs_read", path=f"{d}/dup.txt") == "y = 2\ny = 2"

    await call(tool_ctx, "fs_append", path=f"{d}/notes.md", content="delta\n")
    found = await call(tool_ctx, "fs_search", query="beta", path=d)
    assert f"{d}/notes.md:2: BETA" in found
    assert f"{d}/notes.md:3: gamma" in await call(tool_ctx, "fs_search", query=r"^gam+a$", path=d, regex=True)
    assert "No matches" in await call(tool_ctx, "fs_search", query="beta", path=d, case_sensitive=True)

    await call(tool_ctx, "fs_write", path=f"{d}/v2.md", content="alpha\nBETA\nomega\ndelta\n")
    diff = await call(tool_ctx, "fs_diff", path_a=f"{d}/notes.md", path_b=f"{d}/v2.md")
    assert "-gamma" in diff and "+omega" in diff and "1 line(s) added, 1 removed" in diff
    assert await call(tool_ctx, "fs_diff", path_a=f"{d}/v2.md", path_b=f"{d}/v2.md") == "Files are identical."

    listing = await call(tool_ctx, "fs_list", path=d)
    assert f"{d}/notes.md" in listing and f"{d}/v2.md" in listing
    assert "shared/" in await call(tool_ctx, "fs_list")
    found = await call(tool_ctx, "fs_find", pattern="*.md", path=d)
    assert f"{d}/notes.md" in found and f"{d}/dup.txt" not in found
    meta = await call(tool_ctx, "fs_metadata", path=f"{d}/notes.md")
    assert meta["lines"] == 4 and len(meta["sha256"]) == 64 and meta["is_text"]

    await call(tool_ctx, "fs_mkdir", path=f"{d}/sub")
    await call(tool_ctx, "fs_copy", source=f"{d}/notes.md", destination=f"{d}/sub")
    await call(tool_ctx, "fs_move", source=f"{d}/sub/notes.md", destination=f"{d}/sub/renamed.md")
    assert "BETA" in await call(tool_ctx, "fs_read", path=f"{d}/sub/renamed.md")
    with pytest.raises(ToolError, match="not empty"):
        await call(tool_ctx, "fs_delete", path=f"{d}/sub")
    assert "Deleted folder" in await call(tool_ctx, "fs_delete", path=f"{d}/sub", recursive=True)
    with pytest.raises(ToolError):
        await call(tool_ctx, "fs_read", path=f"{d}/sub/renamed.md")


async def test_fs_write_to_output_registers_artifact(tool_ctx):
    name = f"{uniq('summary')}.md"
    msg = await call(tool_ctx, "fs_write", path=f"output/{name}", content="# Summary\n")
    assert "artifact" in msg
    from sqlalchemy import select

    from app.core.db import session_scope
    from app.models import Artifact

    async with session_scope() as db:
        rows = (await db.execute(select(Artifact).where(Artifact.path == f"output/{name}"))).scalars().all()
    assert len(rows) == 1 and rows[0].mime == "text/markdown"


async def test_fs_rejects_path_traversal(tool_ctx):
    for bad in ("../../etc/passwd", "workspace/../../../secret.txt", "shared/../users/someone/workspace/x.txt"):
        with pytest.raises(ToolError):
            await call(tool_ctx, "fs_read", path=bad)
    with pytest.raises(ToolError):
        await call(tool_ctx, "fs_write", path="../escape.txt", content="nope")
    with pytest.raises(ToolError):
        await call(tool_ctx, "fs_copy", source="workspace", destination="../../outside")
    assert not (tool_ctx.workspace.user_root.parent / "escape.txt").exists()


async def test_fs_viewer_cannot_write(tool_ctx):
    d = f"workspace/{uniq('viewer')}"
    await call(tool_ctx, "fs_write", path=f"{d}/a.txt", content="hello")
    viewer = ToolContext(org_id=tool_ctx.org_id, user_id=tool_ctx.user_id, role="viewer")
    assert await call(viewer, "fs_read", path=f"{d}/a.txt") == "hello"
    attempts = [
        ("fs_write", {"path": f"{d}/b.txt", "content": "x"}),
        ("fs_append", {"path": f"{d}/a.txt", "content": "x"}),
        ("fs_edit", {"path": f"{d}/a.txt", "old_text": "hello", "new_text": "bye"}),
        ("fs_mkdir", {"path": f"{d}/new"}),
        ("fs_move", {"source": f"{d}/a.txt", "destination": f"{d}/c.txt"}),
        ("fs_copy", {"source": f"{d}/a.txt", "destination": f"{d}/c.txt"}),
        ("fs_delete", {"path": f"{d}/a.txt"}),
        ("fs_archive", {"path": d}),
        ("shell_exec", {"command": "echo hi"}),
        ("code_python", {"code": "print(1)"}),
    ]
    for name, args in attempts:
        with pytest.raises(ToolError, match="read-only"):
            await call(viewer, name, **args)
    assert await call(tool_ctx, "fs_read", path=f"{d}/a.txt") == "hello"


async def test_fs_shared_delete_requires_admin(tool_ctx):
    member = ToolContext(org_id=tool_ctx.org_id, user_id=tool_ctx.user_id, role="member")
    path = f"shared/{uniq('team')}/doc.txt"
    await call(member, "fs_write", path=path, content="team notes")  # members may write to shared/
    with pytest.raises(ToolError, match="admin"):
        await call(member, "fs_delete", path=path)
    with pytest.raises(ToolError, match="admin"):
        await call(member, "fs_move", source=path, destination="workspace/doc.txt")
    for root in ("shared", "shared/", "./workspace/.", "output//"):
        with pytest.raises(ToolError, match="top-level"):
            await call(tool_ctx, "fs_delete", path=root, recursive=True)
    assert "Deleted" in await call(tool_ctx, "fs_delete", path=path)


async def test_fs_refuses_git_internals(tool_ctx):
    repo = f"projects/{uniq('r')}"
    with pytest.raises(ToolError, match=r"\.git"):
        await call(tool_ctx, "fs_write", path=f"{repo}/.git/config", content="[core]\n\tfsmonitor = evil\n")
    with pytest.raises(ToolError, match=r"\.git"):
        await call(tool_ctx, "fs_mkdir", path=f"{repo}/.git/hooks")


async def test_fs_archive_extract_roundtrip(tool_ctx):
    d = f"workspace/{uniq('zip')}"
    await call(tool_ctx, "fs_write", path=f"{d}/src/a.txt", content="A")
    await call(tool_ctx, "fs_write", path=f"{d}/src/sub/b.txt", content="B")
    msg = await call(tool_ctx, "fs_archive", path=f"{d}/src", destination=f"{d}/bundle.zip")
    assert "2 files" in msg
    msg = await call(tool_ctx, "fs_extract", path=f"{d}/bundle.zip", destination=f"{d}/out")
    assert "Extracted 2 files" in msg
    assert await call(tool_ctx, "fs_read", path=f"{d}/out/src/sub/b.txt") == "B"
    with pytest.raises(ToolError, match="already exist"):
        await call(tool_ctx, "fs_extract", path=f"{d}/bundle.zip", destination=f"{d}/out")
    assert "Extracted" in await call(tool_ctx, "fs_extract", path=f"{d}/bundle.zip", destination=f"{d}/out",
                                     overwrite=True)


async def test_fs_extract_blocks_zip_slip_and_git_metadata(tool_ctx):
    d = f"workspace/{uniq('slip')}"
    base = tool_ctx.workspace.resolve(d, for_write=True)
    base.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(base / "evil.zip", "w") as zf:
        zf.writestr("innocent.txt", "fine")
        zf.writestr("../../escaped.txt", "pwned")
    with pytest.raises(ToolError, match="escapes"):
        await call(tool_ctx, "fs_extract", path=f"{d}/evil.zip")
    assert not (base / "evil" / "innocent.txt").exists()  # validated before anything is written
    assert not (tool_ctx.workspace.user_root / "escaped.txt").exists()
    with zipfile.ZipFile(base / "abs.zip", "w") as zf:
        zf.writestr("/etc/cron.d/evil", "x")
    with pytest.raises(ToolError, match="absolute"):
        await call(tool_ctx, "fs_extract", path=f"{d}/abs.zip")
    with zipfile.ZipFile(base / "repo.zip", "w") as zf:
        zf.writestr("repo/README.md", "hi")
        zf.writestr("repo/.git/config", "[core]\n\tfsmonitor = evil\n")
    msg = await call(tool_ctx, "fs_extract", path=f"{d}/repo.zip")
    assert "Skipped 1" in msg
    assert (base / "repo" / "repo" / "README.md").exists()
    assert not (base / "repo" / "repo" / ".git").exists()


async def test_fs_extract_text_documents(tool_ctx):
    import docx
    import openpyxl
    from fpdf import FPDF
    from pptx import Presentation

    d = f"workspace/{uniq('docs')}"
    base = tool_ctx.workspace.resolve(d, for_write=True)
    base.mkdir(parents=True, exist_ok=True)
    document = docx.Document()
    document.add_heading("Quarterly Report", 1)
    document.add_paragraph("Revenue grew 12%.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Region", "Amount"
    table.cell(1, 0).text, table.cell(1, 1).text = "EMEA", "1200"
    document.save(base / "report.docx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Amount"])
    ws.append(["APAC", 950.5])
    wb.save(base / "sales.xlsx")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Roadmap"
    slide.placeholders[1].text = "Ship v2 in Q4"
    prs.save(base / "deck.pptx")
    (base / "page.html").write_text("<html><head><title>Intro</title><style>p{color:red}</style></head><body>"
                                    "<h1>Hello</h1><script>steal()</script><p>World</p></body></html>",
                                    encoding="utf-8")
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="Invoice total 42 EUR")
    pdf.output(str(base / "invoice.pdf"))

    text = await call(tool_ctx, "fs_extract_text", path=f"{d}/report.docx")
    assert "# Quarterly Report" in text and "Revenue grew 12%." in text and "EMEA | 1200" in text
    text = await call(tool_ctx, "fs_extract_text", path=f"{d}/sales.xlsx")
    assert "## Sheet: Sales" in text and "APAC\t950.5" in text
    text = await call(tool_ctx, "fs_extract_text", path=f"{d}/deck.pptx")
    assert "Roadmap" in text and "Ship v2 in Q4" in text
    text = await call(tool_ctx, "fs_extract_text", path=f"{d}/page.html")
    assert "Hello" in text and "World" in text and "steal()" not in text and "color:red" not in text
    text = await call(tool_ctx, "fs_extract_text", path=f"{d}/invoice.pdf")
    assert "Invoice total 42 EUR" in text and "page 1" in text
    read = await call(tool_ctx, "fs_read", path=f"{d}/report.docx")
    assert "Quarterly Report" in read and "text extracted" in read
    meta = await call(tool_ctx, "fs_metadata", path=f"{d}/deck.pptx")
    assert meta["slides"] == 1
    (base / "blob.bin").write_bytes(b"\x00\x01\x02binary")
    assert "binary file" in await call(tool_ctx, "fs_read", path=f"{d}/blob.bin")


# =========================================================================== shell
async def test_shell_exec_local(tool_ctx, monkeypatch):
    monkeypatch.setattr(settings, "sandbox_url", None)
    out = await call(tool_ctx, "shell_exec", command="echo hello-ren")
    assert out.startswith("exit code 0") and "hello-ren" in out
    out = await call(tool_ctx, "shell_exec", command="exit 3")
    assert out.startswith("exit code 3")
    out = await call(tool_ctx, "shell_exec", command="sleep 5", timeout=1)
    assert "timed out" in out


async def test_shell_env_does_not_leak_backend_secrets(tool_ctx, monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "sk-should-not-leak")
    monkeypatch.setenv("REN_DB_PASSWORD", "hunter2-should-not-leak")
    probe = "import json, os; print(json.dumps(dict(os.environ)))"
    res = await _exec.run_command(tool_ctx, [sys.executable, "-c", probe], timeout=60)
    assert res["exit_code"] == 0
    child_env = json.loads(res["stdout"])
    assert "should-not-leak" not in res["stdout"]
    assert "OPENAI_API_KEY" not in child_env and "REN_DB_PASSWORD" not in child_env
    assert child_env["REN_WORKSPACE"] == str(tool_ctx.workspace.user_root)
    assert child_env["TMPDIR"] == str(tool_ctx.workspace.user_root / "temp")
    info = await call(tool_ctx, "shell_env")
    assert "OPENAI_API_KEY" not in info["variables"] and "REN_WORKSPACE" in info["variables"]
    assert all("=" not in name for name in info["variables"])
    found = await call(tool_ctx, "shell_which", names=["git" if HAS_GIT else "python", "no-such-program-xyz"])
    assert found["programs"]["no-such-program-xyz"] is None


async def test_run_command_caps_output_and_redacts(tool_ctx):
    res = await _exec.run_command(tool_ctx, [sys.executable, "-c", "print('x' * 100000)"], timeout=60)
    assert res["truncated"] and len(res["stdout"]) < 32_000 and "chars omitted" in res["stdout"]
    res = await _exec.run_command(tool_ctx, [sys.executable, "-c", "print('token=abcd1234secret')"], timeout=60,
                                  redact=["abcd1234secret"])
    assert "abcd1234secret" not in res["stdout"] and "token=***" in res["stdout"]
    res = await _exec.run_command(tool_ctx, ["definitely-not-a-real-binary-xyz"], timeout=10)
    assert res["exit_code"] in (1, 127) and res["stderr"]


async def test_run_command_cancellation_kills_process(tool_ctx):
    d = f"workspace/{uniq('cancel')}"
    tool_ctx.workspace.resolve(d, for_write=True).mkdir(parents=True, exist_ok=True)
    code = "import time; time.sleep(2); open('marker.txt', 'w').write('still alive')"
    task = asyncio.create_task(_exec.run_command(tool_ctx, [sys.executable, "-c", code], cwd_virtual=d, timeout=60))
    await asyncio.sleep(0.7)
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    await asyncio.sleep(3)
    assert not tool_ctx.workspace.resolve(f"{d}/marker.txt").exists()


async def test_shell_background_jobs(tool_ctx, monkeypatch):
    monkeypatch.setattr(settings, "sandbox_url", None)
    job = await call(tool_ctx, "shell_exec_background", command="echo started; sleep 1; echo finished")
    job_id = job["job_id"]
    assert job["status"] == "running" and job["cwd"] == "workspace"
    collected, status = "", {}
    for _ in range(150):
        status = await call(tool_ctx, "shell_poll", job_id=job_id)
        collected += status["output"]
        if status["status"] != "running":
            break
        await asyncio.sleep(0.1)
    collected += (await call(tool_ctx, "shell_poll", job_id=job_id))["output"]
    assert status["status"] == "exited" and status["exit_code"] == 0
    assert "started" in collected and "finished" in collected
    assert "finished" in (await call(tool_ctx, "shell_poll", job_id=job_id, all_output=True))["output"]

    other = ToolContext(org_id=tool_ctx.org_id, user_id=f"other-{uuid.uuid4().hex[:6]}", role="member")
    with pytest.raises(ToolError, match="No background job"):
        await call(other, "shell_poll", job_id=job_id)
    with pytest.raises(ToolError, match="No background job"):
        await call(other, "shell_kill", job_id=job_id)

    long_job = await call(tool_ctx, "shell_exec_background", command="sleep 30")
    killed = await call(tool_ctx, "shell_kill", job_id=long_job["job_id"])
    assert killed["status"] == "killed"
    procs = await call(tool_ctx, "shell_processes")
    ids = {j["job_id"] for j in procs["jobs"]}
    assert {job_id, long_job["job_id"]} <= ids and procs["running"] == 0


async def test_install_package_validates_specifiers(tool_ctx):
    for bad in (["--index-url=http://evil.example/simple"], ["git+https://github.com/x/y"], ["../local.whl"],
                ["pkg; rm -rf /"], ["requests", "-r/etc/passwd"]):
        with pytest.raises(ToolError, match="Invalid pip"):
            await call(tool_ctx, "shell_install_package", packages=bad)
    with pytest.raises(ToolError, match="Invalid npm"):
        await call(tool_ctx, "shell_install_package", packages=["--global", "lodash"], manager="npm")
    preview = registry.get("shell_install_package").render_preview({"packages": ["pandas>=2"], "manager": "pip"})
    assert "pandas>=2" in preview and "third-party code" in preview


async def test_containers_validate_and_report_missing_cli(tool_ctx, monkeypatch):
    from app.tools.builtin import containers

    with pytest.raises(ToolError, match="Secret"):
        await call(tool_ctx, "k8s_get", resource="secrets", output="yaml")
    with pytest.raises(ToolError, match="Invalid"):
        await call(tool_ctx, "docker_logs", container="--help")
    with pytest.raises(ToolError, match="either"):
        await call(tool_ctx, "k8s_apply")
    with pytest.raises(ToolError, match="Invalid"):
        await call(tool_ctx, "docker_run", image="nginx", ports=["80;rm -rf /"])
    monkeypatch.setattr(containers, "find_program", lambda *_a, **_k: None)
    with pytest.raises(ToolError, match="not installed"):
        await call(tool_ctx, "docker_list")
    with pytest.raises(ToolError, match="not installed"):
        await call(tool_ctx, "k8s_get", resource="pods")
    klog_line = "E1005 13:39:13.268980   32232 memcache.go:265] couldn't get API group list: refused"
    noisy = {"stderr": "\n".join([klog_line] * 5 + ["Unable to connect to the server: refused"]), "stdout": ""}
    assert containers._error_text(noisy).splitlines() == [
        "couldn't get API group list: refused", "Unable to connect to the server: refused"]
    preview = registry.get("docker_run").render_preview({"image": "nginx:1.27", "ports": ["8080:80"]})
    assert "docker run" in preview and "8080:80" in preview and "nginx:1.27" in preview
    assert "replicas" in registry.get("k8s_scale").render_preview({"name": "api", "replicas": 3})


# =========================================================================== code
async def test_code_python_runs_and_registers_artifacts(tool_ctx, monkeypatch):
    monkeypatch.setattr(settings, "sandbox_url", None)
    name = f"{uniq('result')}.csv"
    code = ("import os\n"
            "print('sum', 1 + 2)\n"
            f"open(os.path.join(os.environ['REN_OUTPUT'], '{name}'), 'w').write('a,b\\n1,2\\n')\n"
            f"open('../output/rel-{name}', 'w').write('x')\n")
    out = await call(tool_ctx, "code_python", code=code)
    assert out.startswith("exit code 0") and "sum 3" in out
    assert f"output/{name}" in out and f"output/rel-{name}" in out
    from sqlalchemy import select

    from app.core.db import session_scope
    from app.models import Artifact

    async with session_scope() as db:
        rows = (await db.execute(select(Artifact).where(Artifact.path == f"output/{name}"))).scalars().all()
    assert len(rows) == 1 and rows[0].size > 0 and rows[0].name == name

    out = await call(tool_ctx, "code_python", code="raise ValueError('boom')")
    assert out.startswith("exit code 1") and "ValueError: boom" in out

    module = uniq("helper").replace("-", "_")
    await call(tool_ctx, "fs_write", path=f"workspace/{module}.py", content="VALUE = 41\n")
    out = await call(tool_ctx, "code_python", code=f"import {module}\nprint({module}.VALUE + 1)")
    assert "42" in out


@pytest.mark.skipif(not HAS_NODE, reason="node is not installed")
async def test_code_javascript(tool_ctx, monkeypatch):
    monkeypatch.setattr(settings, "sandbox_url", None)
    out = await call(tool_ctx, "code_javascript", code="const os = require('os'); console.log(6 * 7, typeof os)")
    assert out.startswith("exit code 0") and "42 object" in out


async def test_code_test_lint_and_format(tool_ctx, monkeypatch):
    monkeypatch.setattr(settings, "sandbox_url", None)
    d = f"projects/{uniq('pyproj')}"
    await call(tool_ctx, "fs_write", path=f"{d}/calc.py", content="def add(a, b):\n    return a + b\n")
    await call(tool_ctx, "fs_write", path=f"{d}/test_calc.py",
               content="from calc import add\n\ndef test_add():\n    assert add(2, 3) == 5\n")
    out = await call(tool_ctx, "code_test", path=d)
    assert "1 passed" in out
    await call(tool_ctx, "fs_write", path=f"{d}/messy.py", content="import os\nx=1\n")
    lint = await call(tool_ctx, "code_lint", path=f"{d}/messy.py")
    assert "F401" in lint or "imported but unused" in lint
    check = await call(tool_ctx, "code_format", path=f"{d}/messy.py", check=True)
    assert "x = 1" in check
    await call(tool_ctx, "code_format", path=f"{d}/messy.py")
    assert "x = 1" in await call(tool_ctx, "fs_read", path=f"{d}/messy.py")


# =========================================================================== git
@pytest.mark.skipif(not HAS_GIT, reason="git is not installed")
async def test_git_init_add_commit_log_branch_merge(tool_ctx, monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "sandbox_url", None)
    # isolate from the developer's global git config so the Ren user's identity is used
    monkeypatch.setenv("HOME", str(tmp_path))
    monkeypatch.setenv("USERPROFILE", str(tmp_path))
    name = uniq("repo")
    repo = f"projects/{name}"
    assert repo in await call(tool_ctx, "git_init", directory=name)
    await call(tool_ctx, "fs_write", path=f"{repo}/README.md", content="# Demo\n")
    assert "README.md" in await call(tool_ctx, "git_status", repo=repo)
    assert "README.md" in await call(tool_ctx, "git_add", repo=repo)
    assert "Initial commit" in await call(tool_ctx, "git_commit", repo=repo, message="Initial commit")
    log = await call(tool_ctx, "git_log", repo=repo)
    assert "Initial commit" in log and "Test Owner" in log
    assert "working tree clean" in await call(tool_ctx, "git_status", repo=f"{repo}/")
    with pytest.raises(ToolError, match="Nothing to commit"):
        await call(tool_ctx, "git_commit", repo=repo, message="empty")

    await call(tool_ctx, "git_create_branch", repo=repo, name="feature/more")
    await call(tool_ctx, "fs_append", path=f"{repo}/README.md", content="More text.\n")
    assert "+More text." in await call(tool_ctx, "git_diff", repo=repo)
    await call(tool_ctx, "git_commit", repo=repo, message="Add more", add_all=True)
    branches = await call(tool_ctx, "git_branch", repo=repo)
    assert "* feature/more" in branches and "main" in branches
    assert "README.md" in await call(tool_ctx, "git_show", repo=repo, ref="HEAD", stat=True)
    assert "More text." in await call(tool_ctx, "git_show", repo=repo, ref="HEAD", path="README.md")
    await call(tool_ctx, "git_checkout", repo=repo, ref="main")
    merged = await call(tool_ctx, "git_merge", repo=repo, branch="feature/more")
    assert "Fast-forward" in merged
    assert "Add more" in await call(tool_ctx, "git_log", repo=repo, max_count=1)


@pytest.mark.skipif(not HAS_GIT, reason="git is not installed")
async def test_git_requires_repo_inside_workspace(tool_ctx):
    plain = f"projects/{uniq('plain')}"
    await call(tool_ctx, "fs_mkdir", path=plain)
    with pytest.raises(ToolError, match="not inside a git repository"):
        await call(tool_ctx, "git_status", repo=plain)
    with pytest.raises(ToolError):
        await call(tool_ctx, "git_log", repo="../../..")


async def test_git_clone_rejects_unsafe_urls(tool_ctx):
    for url in ("file:///etc", "ext::sh -c touch% /tmp/pwned", "/srv/repo", "git@github.com:org/repo.git",
                "https://user:secret@github.com/org/repo.git", "http://127.0.0.1/repo.git"):
        with pytest.raises(ToolError):
            await call(tool_ctx, "git_clone", url=url)


@pytest.mark.skipif(not HAS_GIT, reason="git is not installed")
async def test_git_github_token_flows_through_credential_helper_only(tool_ctx):
    from app.tools.builtin import git as git_tools

    token = "ghp_TestToken1234567890"
    tool_ctx.extra["connectors"] = {"github": {"token": token}}
    config, env, secrets = await git_tools._auth(tool_ctx, "https://github.com/org/repo.git")
    assert env == {"REN_GIT_TOKEN": token} and secrets == [token]
    assert token not in " ".join(config)
    assert await git_tools._auth(tool_ctx, "https://gitlab.com/org/repo.git") == ([], {}, [])
    # the helper really answers git's credential protocol (exercises the shell helper on this platform)
    res = await _exec.run_command(tool_ctx, ["git", *config, "credential", "fill"], timeout=30,
                                  stdin="protocol=https\nhost=github.com\n\n", env={**env, "GIT_TERMINAL_PROMPT": "0"})
    assert res["exit_code"] == 0, res["stderr"]
    assert "username=x-access-token" in res["stdout"] and f"password={token}" in res["stdout"]
