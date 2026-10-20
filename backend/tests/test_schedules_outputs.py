"""Schedules (time zones, CRUD, pause/resume, run now, permissions) and outputs (artifact listing, previews)."""
import re
import time
from datetime import UTC, datetime

import pytest
from fastapi.testclient import TestClient

TZ = {"X-Timezone": "Asia/Kolkata"}


# ------------------------------------------------------------------------------------------- time zone maths
def test_timezone_and_cron_helpers():
    from app.services.scheduler import (
        describe_cron,
        describe_schedule,
        next_cron_run,
        parse_when,
        upcoming_runs,
        valid_timezone,
        validate_cron,
    )

    assert valid_timezone("Asia/Kolkata") == "Asia/Kolkata"
    assert valid_timezone("utc") == "UTC" and valid_timezone("Mars/Olympus") is None and valid_timezone("") is None
    # browsers report CLDR legacy IDs (Chrome: Asia/Calcutta); Debian's tz database lacks those links
    assert valid_timezone("Asia/Calcutta") == "Asia/Kolkata" and valid_timezone("Europe/Kiev") == "Europe/Kyiv"
    # 09:00 in India is 03:30 UTC; the next one after that is the following day
    first = next_cron_run("0 9 * * *", datetime(2026, 10, 5, 0, 0, tzinfo=UTC), "Asia/Kolkata")
    assert first == datetime(2026, 10, 5, 3, 30, tzinfo=UTC)
    assert next_cron_run("0 9 * * *", first, "Asia/Kolkata") == datetime(2026, 10, 6, 3, 30, tzinfo=UTC)
    # wall-clock time survives the DST change (Berlin leaves summer time on 25 Oct 2026)
    before = next_cron_run("0 9 * * *", datetime(2026, 10, 24, 0, 0, tzinfo=UTC), "Europe/Berlin")
    after = next_cron_run("0 9 * * *", datetime(2026, 10, 26, 0, 0, tzinfo=UTC), "Europe/Berlin")
    assert (before.hour, after.hour) == (7, 8)
    assert len(upcoming_runs("*/15 * * * *", None, "UTC", 4)) == 4

    assert parse_when("2026-10-06T09:00", "Asia/Kolkata") == datetime(2026, 10, 6, 3, 30, tzinfo=UTC)
    assert parse_when("2026-10-06T09:00Z", "Asia/Kolkata") == datetime(2026, 10, 6, 9, 0, tzinfo=UTC)
    with pytest.raises(ValueError):
        parse_when("next tuesday", "UTC")
    with pytest.raises(ValueError):
        validate_cron("every day")

    cases = {
        "0 11 * * *": "Every day at 11:00", "30 8 * * 1-5": "Weekdays at 08:30", "0 10 * * 0,6": "Weekends at 10:00",
        "0 9 * * MON,WED": "Every Mon, Wed at 09:00", "*/15 * * * *": "Every 15 minutes", "0 */2 * * *": "Every 2 hours",
        "5 * * * *": "Every hour at :05", "0 7 1 * *": "Monthly on day 1 at 07:00", "0 0 1 1 *": "Cron 0 0 1 1 *",
    }
    for cron, text in cases.items():
        assert describe_cron(cron) == text, cron
    assert describe_schedule("0 11 * * *", "UTC", None) == "Every day at 11:00 (UTC)"
    assert describe_schedule(None, "Asia/Kolkata", datetime(2026, 10, 6, 3, 30, tzinfo=UTC)) == \
        "Once · Tue 6 Oct 2026, 09:00 (Asia/Kolkata)"


def test_apply_changes_rules():
    from app.models import ScheduledJob
    from app.services.scheduler import apply_changes

    job = ScheduledJob(kind="task", title="Brief", payload={"agent_id": "a1", "instructions": "x"}, cron="0 9 * * *",
                       timezone="UTC", status="scheduled", run_at=datetime(2099, 1, 1, 9, 0, tzinfo=UTC))
    apply_changes(job, {"timezone": "Asia/Kolkata"})  # recurring: recomputed in the new zone
    assert job.timezone == "Asia/Kolkata" and (job.run_at.hour, job.run_at.minute) == (3, 30)
    with pytest.raises(ValueError, match="needs run_at"):
        apply_changes(job, {"cron": None})
    apply_changes(job, {"cron": None, "run_at": "2099-01-01T10:00", "title": "Kickoff", "instructions": "y"})
    assert job.cron is None and job.run_at == datetime(2099, 1, 1, 4, 30, tzinfo=UTC)
    assert job.title == "Kickoff" and job.payload["instructions"] == "y" and job.payload["title"] == "Kickoff"
    apply_changes(job, {"timezone": "Europe/London"})  # one-off: same wall-clock time (10:00) in the new zone
    assert job.run_at == datetime(2099, 1, 1, 10, 0, tzinfo=UTC)
    with pytest.raises(ValueError):
        apply_changes(job, {"title": "  "})
    with pytest.raises(ValueError, match="past"):
        apply_changes(job, {"run_at": "2001-01-01T10:00"})
    job.status = "done"
    with pytest.raises(ValueError, match="active or paused"):
        apply_changes(job, {"title": "Again"})


# ------------------------------------------------------------------------------------------- API
@pytest.fixture(scope="module")
def client(db_ready):
    from app.main import app

    with TestClient(app) as c:
        yield c


def _register(client, email: str, name: str = "Sched User"):
    r = client.post("/api/auth/register", json={"email": email, "password": "supersecret1", "name": name,
                                                "org_name": "Schedules Co", "mode": "business"})
    assert r.status_code == 200, r.text
    d = r.json()
    h = {"Authorization": f"Bearer {d['access_token']}", "X-Org-Id": d["orgs"][0]["id"], **TZ}
    client.post("/api/providers", headers=h, json={"name": "Mock", "kind": "mock", "default_model": "mock-1"})
    return h, d


@pytest.fixture(scope="module")
def owner(client):
    return _register(client, "sched-owner@ren.dev", "Sched Owner")


def _specialist(client, h) -> dict:
    return next(a for a in client.get("/api/agents", headers=h).json()["items"] if not a["is_orchestrator"])


def _wait(pred, timeout=15.0):
    end = time.time() + timeout
    while time.time() < end:
        v = pred()
        if v:
            return v
        time.sleep(0.2)
    return None


def test_timezone_header_is_remembered(client, owner):
    h, _ = owner
    me = client.get("/api/auth/me", headers=h).json()
    assert me["user"]["timezone"] == "Asia/Kolkata"
    client.get("/api/auth/me", headers={**h, "X-Timezone": "Not/AZone"})  # ignored
    assert client.get("/api/auth/me", headers=h).json()["user"]["timezone"] == "Asia/Kolkata"


def test_schedule_preview(client, owner):
    h, _ = owner
    p = client.get("/api/schedules/preview", headers=h, params={"cron": "0 9 * * 1-5"}).json()
    assert p["valid"] and p["description"] == "Weekdays at 09:00 (Asia/Kolkata)"
    assert len(p["next"]) == 5 and all(n.endswith("T03:30:00Z") for n in p["next"])
    bad = client.get("/api/schedules/preview", headers=h, params={"cron": "nonsense"}).json()
    assert bad["valid"] is False and bad["error"]
    past = client.get("/api/schedules/preview", headers=h, params={"run_at": "2001-01-01T09:00"}).json()
    assert past["valid"] is False and "past" in past["error"]
    once = client.get("/api/schedules/preview", headers=h,
                      params={"run_at": "2099-06-01T09:00", "timezone": "Europe/London"}).json()
    assert once["valid"] and once["next"] == ["2099-06-01T08:00:00Z"] and "Europe/London" in once["description"]
    assert client.get("/api/schedules/preview", headers=h, params={"cron": "0 9 * * *", "timezone": "Nope/Nope"}).json()["valid"] is False


def test_schedule_lifecycle(client, owner):
    h, _ = owner
    agent = _specialist(client, h)
    orch = next(a for a in client.get("/api/agents", headers=h).json()["items"] if a["is_orchestrator"])
    base = {"title": "Daily Tech & AI Newsletter", "agent_id": agent["id"], "instructions": "Compile the newsletter."}
    assert client.post("/api/schedules", headers=h, json={**base, "agent_id": orch["id"], "cron": "0 9 * * *"}).status_code == 400
    assert client.post("/api/schedules", headers=h, json=base).status_code == 400  # neither cron nor run_at
    assert client.post("/api/schedules", headers=h, json={**base, "cron": "0 9 * * *", "timezone": "Mars/Base"}).status_code == 400
    assert client.post("/api/schedules", headers=h, json={**base, "run_at": "2001-01-01T09:00"}).status_code == 400

    job = client.post("/api/schedules", headers=h, json={**base, "cron": "0 9 * * *"}).json()
    assert job["timezone"] == "Asia/Kolkata" and job["description"] == "Every day at 09:00 (Asia/Kolkata)"
    assert job["run_at"].endswith("T03:30:00Z") and job["agent_name"] == agent["name"]
    assert job["status"] == "scheduled" and job["run_count"] == 0 and job["last_task"] is None
    assert job["instructions"] == "Compile the newsletter." and job["kind"] == "task"
    jid = job["id"]
    assert any(j["id"] == jid for j in client.get("/api/schedules", headers=h).json()["items"])

    upd = client.patch(f"/api/schedules/{jid}", headers=h, json={"cron": "30 8 * * 1-5", "title": "Morning brief"}).json()
    assert upd["title"] == "Morning brief" and upd["description"] == "Weekdays at 08:30 (Asia/Kolkata)"
    upd = client.patch(f"/api/schedules/{jid}", headers=h, json={"timezone": "UTC"}).json()
    assert upd["description"] == "Weekdays at 08:30 (UTC)" and upd["run_at"].endswith("T08:30:00Z")
    assert client.patch(f"/api/schedules/{jid}", headers=h, json={"cron": "bad"}).status_code == 400

    paused = client.post(f"/api/schedules/{jid}/pause", headers=h).json()
    assert paused["status"] == "paused"
    assert any(j["id"] == jid for j in client.get("/api/schedules", headers=h, params={"status": "paused"}).json()["items"])
    assert any(j["id"] == jid for j in client.get("/api/schedules", headers=h).json()["items"])  # active includes paused
    resumed = client.post(f"/api/schedules/{jid}/resume", headers=h).json()
    assert resumed["status"] == "scheduled" and resumed["run_at"] > datetime.now(UTC).isoformat()

    run = client.post(f"/api/schedules/{jid}/run", headers=h).json()
    task = run["task"]
    assert task["schedule_id"] == jid and task["title"] == "Morning brief"
    assert client.get(f"/api/tasks/{task['id']}", headers=h).json()["schedule_id"] == jid
    detail = client.get(f"/api/schedules/{jid}", headers=h).json()
    assert detail["job"]["run_count"] == 1 and detail["job"]["last_task"]["id"] == task["id"]
    assert [r["id"] for r in detail["runs"]] == [task["id"]] and "artifacts" in detail["runs"][0]
    assert detail["job"]["status"] == "scheduled"  # running now doesn't change the schedule
    _wait(lambda: client.get(f"/api/tasks/{task['id']}", headers=h).json()["status"] in ("completed", "failed"))

    assert client.delete(f"/api/schedules/{jid}", headers=h).json() == {"ok": True}
    assert not any(j["id"] == jid for j in client.get("/api/schedules", headers=h).json()["items"])
    hist = client.get("/api/schedules", headers=h, params={"status": "history"}).json()["items"]
    assert next(j for j in hist if j["id"] == jid)["status"] == "cancelled"
    assert client.patch(f"/api/schedules/{jid}", headers=h, json={"title": "x"}).status_code == 400
    assert client.post(f"/api/schedules/{jid}/pause", headers=h).status_code == 400
    assert client.get("/api/schedules", headers=h, params={"status": "weird"}).status_code == 400


def test_schedule_permissions(client, owner):
    h, _ = owner
    agent = _specialist(client, h)
    mine = client.post("/api/schedules", headers=h, json={"title": "Owner job", "agent_id": agent["id"],
                                                         "instructions": "x", "cron": "0 6 * * *"}).json()
    members = {}
    for role in ("member", "viewer"):
        inv = client.post("/api/orgs/current/invites", headers=h, json={"email": f"sched-{role}@ren.dev", "role": role}).json()
        acc = client.post("/api/auth/accept-invite", json={"token": inv["token"], "name": role.title(), "password": "rolepass123"})
        members[role] = {"Authorization": f"Bearer {acc.json()['access_token']}", "X-Org-Id": h["X-Org-Id"]}
    m = members["member"]
    assert all(j["id"] != mine["id"] for j in client.get("/api/schedules", headers=m).json()["items"])
    assert client.get(f"/api/schedules/{mine['id']}", headers=m).status_code == 404
    assert client.patch(f"/api/schedules/{mine['id']}", headers=m, json={"title": "hijack"}).status_code == 404
    assert client.post(f"/api/schedules/{mine['id']}/run", headers=m).status_code == 404
    theirs = client.post("/api/schedules", headers=m, json={"title": "Member job", "agent_id": agent["id"],
                                                           "instructions": "y", "run_at": "2099-01-01T09:00"}).json()
    assert theirs["timezone"] == "UTC"  # this member never sent X-Timezone
    assert any(j["id"] == theirs["id"] for j in client.get("/api/schedules", headers=h).json()["items"])  # admins see all
    assert client.post("/api/schedules", headers=members["viewer"], json={"title": "v", "agent_id": agent["id"],
                                                                          "instructions": "z", "cron": "0 9 * * *"}).status_code == 403


# ------------------------------------------------------------------------------------------- outputs
def _png() -> bytes:
    import io

    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (4, 4), (31, 90, 69)).save(buf, "PNG")
    return buf.getvalue()


def test_outputs_listing_and_previews(client, owner):
    import asyncio

    from docx import Document

    from app.core.db import session_scope
    from app.models import Artifact
    from app.services.workspace import Workspace

    h, data = owner
    org_id, user_id = h["X-Org-Id"], data["user"]["id"]
    agent = _specialist(client, h)
    job = client.post("/api/schedules", headers=h, json={"title": "Report run", "agent_id": agent["id"],
                                                        "instructions": "Write it", "cron": "0 7 * * *"}).json()
    task = client.post(f"/api/schedules/{job['id']}/run", headers=h).json()["task"]

    ws = Workspace(org_id, user_id)
    files = {
        "output/q3-report.md": b"# Q3 Report\n\nRevenue grew **12%**.\n",
        "output/sales-data.csv": b"region,amount\nNorth,10\nSouth,7\n",
        "output/config.json": b'{"a":1,"b":[1,2]}',
        "output/chart.png": _png(),
        "output/blob.bin": bytes(range(256)) * 4,
        "uploads/brief.txt": b"Customer brief",
    }
    for rel, content in files.items():
        ws.resolve(rel, for_write=True).write_bytes(content)
    doc = Document()
    doc.add_heading("Board memo", 1)
    doc.add_paragraph("Approve the budget.")
    doc.save(str(ws.resolve("output/memo.docx", for_write=True)))

    async def register():
        async with session_scope() as db:
            for rel in [*files, "output/memo.docx"]:
                generated = rel.startswith("output/")
                db.add(Artifact(org_id=org_id, user_id=user_id, task_id=task["id"] if generated else None,
                                name=rel.rsplit("/", 1)[-1], path=rel, mime="application/octet-stream",
                                size=ws.resolve(rel).stat().st_size))

    asyncio.run(register())
    ids = {a["path"]: a["id"] for a in client.get("/api/artifacts", headers=h).json()["items"]}

    gen = client.get("/api/artifacts", headers=h, params={"source": "generated"}).json()
    assert all(a["source"] == "generated" for a in gen["items"])
    assert {p for p in files if p.startswith("output/")} | {"output/memo.docx"} <= {a["path"] for a in gen["items"]}
    ours = next(a for a in gen["items"] if a["path"] == "output/q3-report.md")
    assert ours["task_title"] == "Report run" and ours["agent_name"] == agent["name"]
    assert ours["schedule_id"] == job["id"] and ours["kind"] == "document"
    up = client.get("/api/artifacts", headers=h, params={"source": "uploaded"}).json()
    assert "brief.txt" in [a["name"] for a in up["items"]] and all(a["source"] == "uploaded" for a in up["items"])
    docs = client.get("/api/artifacts", headers=h, params={"kind": "document"}).json()
    assert {"memo.docx", "q3-report.md"} <= {a["name"] for a in docs["items"]}
    assert all(a["kind"] == "document" for a in docs["items"])
    assert client.get("/api/artifacts", headers=h, params={"q": "sales-data"}).json()["items"][0]["kind"] == "data"
    paged = client.get("/api/artifacts", headers=h, params={"limit": 2, "offset": 0}).json()
    assert len(paged["items"]) == 2 and paged["total"] >= 7
    assert client.get("/api/artifacts", headers=h, params={"kind": "weird"}).status_code == 400

    def preview(path):
        r = client.get(f"/api/artifacts/{ids[path]}/preview", headers=h)
        assert r.status_code == 200, r.text
        return r.json()

    md = preview("output/q3-report.md")
    assert md["render"] == "markdown" and "# Q3 Report" in md["content"] and md["kind"] == "document"
    assert md["download"] == {"path": f"/artifacts/{ids['output/q3-report.md']}/download", "query": {}}
    assert md["task_title"] == "Report run" and md["agent_name"] == agent["name"]
    csv_p = preview("output/sales-data.csv")
    assert csv_p["render"] == "table" and csv_p["table"]["columns"] == ["region", "amount"]
    assert csv_p["table"]["rows"] == [["North", "10"], ["South", "7"]] and csv_p["table"]["total_rows"] == 2
    js = preview("output/config.json")
    assert js["render"] == "code" and js["language"] == "json" and '"b": [' in js["content"]
    assert preview("output/chart.png")["render"] == "image"
    memo = preview("output/memo.docx")
    assert memo["render"] == "markdown" and memo["converted_from"] == "docx" and "Board memo" in memo["content"]
    assert preview("output/blob.bin")["render"] == "none"

    by_path = client.get("/api/files/preview", headers=h, params={"path": "./output/q3-report.md"}).json()
    assert by_path["path"] == "output/q3-report.md" and by_path["artifact_id"] == ids["output/q3-report.md"]
    assert by_path["download"] == {"path": "/files/download", "query": {"path": "output/q3-report.md"}}
    assert client.get("/api/files/preview", headers=h, params={"path": "output/missing.md"}).status_code == 404
    assert client.get("/api/files/download", headers=h, params={"path": "output/q3-report.md"}).content.startswith(b"# Q3")
    raw = client.get(f"/api/artifacts/{ids['output/chart.png']}/download", headers=h)
    assert raw.status_code == 200 and raw.content == files["output/chart.png"]

    # an admin can open a member's output by its path; other members can't
    inv = client.post("/api/orgs/current/invites", headers=h, json={"email": "sched-writer@ren.dev", "role": "member"}).json()
    acc = client.post("/api/auth/accept-invite", json={"token": inv["token"], "name": "Writer", "password": "rolepass123"}).json()
    mh = {"Authorization": f"Bearer {acc['access_token']}", "X-Org-Id": org_id}
    mws = Workspace(org_id, acc["user"]["id"])
    mws.resolve("output/member-notes.md", for_write=True).write_text("# Member notes\n", encoding="utf-8")

    async def register_member():
        async with session_scope() as db:
            db.add(Artifact(org_id=org_id, user_id=acc["user"]["id"], task_id=task["id"], name="member-notes.md",
                            path="output/member-notes.md", mime="text/markdown", size=15))

    asyncio.run(register_member())
    adm = client.get("/api/files/preview", headers=h, params={"path": "output/member-notes.md"})
    assert adm.status_code == 200 and "Member notes" in adm.json()["content"]
    assert client.get("/api/files/download", headers=h, params={"path": "output/member-notes.md"}).status_code == 200
    other = _register(client, "sched-outsider@ren.dev")[0]  # a different workspace entirely
    assert client.get("/api/files/preview", headers=other, params={"path": "output/member-notes.md"}).status_code == 404
    assert client.get(f"/api/artifacts/{ids['output/q3-report.md']}/preview", headers=mh).status_code == 404


# ------------------------------------------------------------------------------------------- Ren's schedule tools
async def test_orchestrator_schedule_tools(org_ctx):
    from app.agents.orchestrator import ConversationOrchestrator
    from app.core.db import session_scope
    from app.models import Agent

    async with session_scope() as db:
        agent = Agent(org_id=org_ctx["org_id"], name="Newsletter Writer", slug="newsletter-writer", role="Writer")
        db.add(agent)
        await db.flush()
        agents = [agent]

    orch = object.__new__(ConversationOrchestrator)  # tool handlers only — no inbox loop
    orch.org_id, orch.user_id, orch.conv_id, orch.user_tz = org_ctx["org_id"], org_ctx["user_id"], None, "Asia/Kolkata"

    out = await orch._schedule({"agent": "Newsletter Writer", "title": "Daily Tech & AI Newsletter",
                                "instructions": "Compile it.", "cron": "0 11 * * *"}, agents)
    assert "Every day at 11:00 (Asia/Kolkata)" in out, out
    jid = re.search(r"\[([0-9a-f]{8})\]", out).group(1)
    assert "Error" in await orch._schedule({"agent": "Newsletter Writer", "title": "x", "instructions": "y",
                                            "cron": "0 9 * * *", "timezone": "Moon/Base"}, agents)
    async with session_scope() as db:
        board = await orch._schedules(db)
    assert jid in board and "Daily Tech & AI Newsletter" in board and "(Asia/Kolkata)" in board

    out = await orch._manage_schedule({"schedule": jid, "action": "update", "cron": "0 9 * * 1-5"})
    assert out.startswith("Updated") and "Weekdays at 09:00 (Asia/Kolkata)" in out, out
    assert (await orch._manage_schedule({"schedule": "Daily Tech & AI Newsletter", "action": "pause"})).startswith("Paused")
    async with session_scope() as db:
        assert "PAUSED" in await orch._schedules(db)
    assert (await orch._manage_schedule({"schedule": jid, "action": "resume"})).startswith("Resumed")
    assert "Error" in await orch._manage_schedule({"schedule": jid, "action": "update", "cron": "nope"})
    assert "Error" in await orch._manage_schedule({"schedule": jid, "action": "explode"})
    assert "no active schedule" in await orch._manage_schedule({"schedule": "deadbeef", "action": "pause"})
    assert (await orch._manage_schedule({"schedule": jid, "action": "cancel"})).startswith("Cancelled")
    assert "no active schedule" in await orch._manage_schedule({"schedule": jid, "action": "resume"})
