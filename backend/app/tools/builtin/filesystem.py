"""Filesystem tools: browse, read, write, search and organize files in the user's sandboxed workspace.

Virtual layout (see app.services.workspace): workspace/ projects/ downloads/ output/ uploads/ temp/ are private to the
user, shared/ is the organization's shared drive. Every path is resolved through ``ctx.workspace`` (no escapes);
mutating tools refuse viewers, refuse to touch .git internals, and deleting/moving out of shared/ needs an admin.
"""
from __future__ import annotations

import asyncio
import codecs
import collections
import difflib
import hashlib
import mimetypes
import os
import re
import shutil
import stat
import sys
import tarfile
import threading
import time
import uuid
import zipfile
from datetime import UTC, datetime
from pathlib import Path, PurePosixPath

from app.services.workspace import TOP_DIRS
from app.tools.base import Risk, ToolContext, ToolError, ToolGroup
from app.tools.builtin._exec import require_write, ws_path

try:  # the `regex` module supports match timeouts (protects against catastrophic patterns)
    import regex as _regex
except ImportError:  # pragma: no cover
    _regex = None

fs = ToolGroup(
    "filesystem", "Files",
    "Browse, read, write, search and organize files in your private workspace and the organization's shared drive",
    prefix="fs", icon="folder", category="core",
)

READ_BUDGET = 15_000  # characters returned by read-style tools
MAX_LINE_CHARS = 2_000
MAX_WRITE_BYTES = 20 * 1024 * 1024
MAX_SEARCH_FILE_BYTES = 10 * 1024 * 1024
MAX_DIFF_FILE_BYTES = 5 * 1024 * 1024
MAX_ARCHIVE_BYTES = 2 * 1024 ** 3
MAX_ARCHIVE_FILES = 100_000
MAX_EXTRACT_MEMBERS = 20_000
SKIP_DIRS = {".git", ".hg", ".svn", "node_modules", "__pycache__", ".venv", "venv", ".mypy_cache", ".pytest_cache",
             ".ruff_cache", ".tox", ".next", ".cache"}
DOC_EXTS = {".pdf", ".docx", ".xlsx", ".xlsm", ".pptx", ".html", ".htm", ".xhtml"}
BINARY_EXTS = {
    ".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".ico", ".tif", ".tiff", ".psd", ".zip", ".gz", ".tgz", ".bz2",
    ".xz", ".7z", ".rar", ".exe", ".dll", ".so", ".dylib", ".bin", ".mp3", ".mp4", ".mov", ".avi", ".mkv", ".wav",
    ".flac", ".ogg", ".woff", ".woff2", ".ttf", ".otf", ".class", ".jar", ".pyc", ".o", ".a", ".lib", ".sqlite",
    ".db", ".parquet", ".feather", ".npy", ".npz", ".pkl", ".xls", ".doc", ".ppt", ".odt", ".ods", ".odp",
}
MOUNTS = {
    "workspace": "your private scratch area (relative paths default here)",
    "projects": "code repositories and projects",
    "downloads": "files fetched from the web",
    "output": "deliverables: files here are shown to the user",
    "uploads": "files the user uploaded",
    "temp": "temporary files",
    "shared": "the organization's shared drive (visible to all members)",
}
_ARCHIVE_SUFFIXES = (".tar.gz", ".tar.bz2", ".tar.xz", ".tgz", ".tbz2", ".txz", ".tar", ".zip")


# --------------------------------------------------------------------------- helpers
def _norm(ctx: ToolContext, path: str | None) -> str:
    """Canonical virtual path: no empty or '.' segments, no trailing slash ('..' is kept and rejected later)."""
    return "/".join(seg for seg in ctx.workspace.normalize(path).split("/") if seg not in ("", "."))


def _human(n: int) -> str:
    size = float(n)
    for unit in ("B", "KB", "MB", "GB"):
        if size < 1024 or unit == "GB":
            return f"{int(size)} {unit}" if unit == "B" else f"{size:.1f} {unit}"
        size /= 1024
    return f"{n} B"


def _stamp(ts: float) -> str:
    return datetime.fromtimestamp(ts, tz=UTC).strftime("%Y-%m-%d %H:%M UTC")


def _is_mount_root(vpath: str) -> bool:
    return vpath == "" or vpath in TOP_DIRS


def _guard_write(ctx: ToolContext, vpath: str) -> None:
    require_write(ctx)
    if ".git" in PurePosixPath(vpath).parts:
        raise ToolError("Files inside .git/ can't be modified with file tools; use the git tools instead.")


def _guard_remove(ctx: ToolContext, vpath: str, verb: str) -> None:
    _guard_write(ctx, vpath)
    if _is_mount_root(vpath):
        raise ToolError(f"'{vpath or '/'}' is a top-level workspace folder and can't be {verb}d.")
    if vpath.split("/", 1)[0] == "shared" and not ctx.is_admin:
        raise ToolError(f"Only workspace admins can {verb} items in the shared drive (shared/).")


def _lexical(ctx: ToolContext, vpath: str) -> Path:
    """Absolute path of `vpath` without following a symlink in its final component."""
    if "/" not in vpath:
        return ws_path(ctx, vpath)
    parent_v, name = vpath.rsplit("/", 1)
    if name in ("", ".", ".."):
        raise ToolError("Invalid path.")
    return ws_path(ctx, parent_v, must_exist=True) / name


def _exists(p: Path) -> bool:
    return p.exists() or p.is_symlink()


def _roots(ctx: ToolContext, path: str | None) -> list[Path]:
    if _norm(ctx, path) == "":
        return [ws_path(ctx, d) for d in TOP_DIRS]
    return [ws_path(ctx, path, must_exist=True)]


def _walk(root: Path, skip: set[str] = SKIP_DIRS):
    """os.walk without following symlinks, sorted, skipping dependency/VCS folders."""
    for dirpath, dirnames, filenames in os.walk(root, followlinks=False):
        dirnames[:] = sorted((d for d in dirnames if d not in skip), key=str.lower)
        yield Path(dirpath), dirnames, sorted(filenames, key=str.lower)


def _glob_regex(pattern: str) -> re.Pattern:
    pat = pattern.replace("\\", "/").strip()
    while pat.startswith("./"):
        pat = pat[2:]
    pat = pat.lstrip("/")
    out: list[str] = []
    i = 0
    while i < len(pat):
        c = pat[i]
        if c == "*":
            if pat.startswith("**/", i):
                out.append("(?:.*/)?")
                i += 3
                continue
            if pat.startswith("**", i):
                out.append(".*")
                i += 2
                continue
            out.append("[^/]*")
        elif c == "?":
            out.append("[^/]")
        elif c == "[" and pat.find("]", i + 1) != -1:
            j = pat.find("]", i + 1)
            body = pat[i + 1:j]
            if body.startswith("!"):
                body = "^" + body[1:]
            out.append("[" + body.replace("\\", "\\\\") + "]")
            i = j + 1
            continue
        elif c == "{" and pat.find("}", i + 1) != -1:
            j = pat.find("}", i + 1)
            out.append("(?:" + "|".join(re.escape(a) for a in pat[i + 1:j].split(",")) + ")")
            i = j + 1
            continue
        else:
            out.append(re.escape(c))
        i += 1
    try:
        return re.compile("^" + "".join(out) + "$", re.IGNORECASE)
    except re.error as e:
        raise ToolError(f"Invalid glob pattern {pattern!r}: {e}") from e


_BOMS = ((codecs.BOM_UTF8, "utf-8-sig"), (codecs.BOM_UTF32_LE, "utf-32"), (codecs.BOM_UTF32_BE, "utf-32"),
         (codecs.BOM_UTF16_LE, "utf-16"), (codecs.BOM_UTF16_BE, "utf-16"))


def _sniff(p: Path) -> tuple[bool, str]:
    """(is_text, encoding) from the first 8 KB."""
    with open(p, "rb") as f:
        head = f.read(8192)
    for bom, enc in _BOMS:
        if head.startswith(bom):
            return True, enc
    if p.suffix.lower() in BINARY_EXTS or b"\x00" in head:
        return False, ""
    return True, "utf-8"


def _errors_for(enc: str) -> str:
    return "surrogateescape" if enc.startswith("utf-8") else "strict"


def _count_lines(p: Path) -> int | None:
    if p.stat().st_size > 20 * 1024 * 1024:
        return None
    n, last = 0, b""
    with open(p, "rb") as f:
        while chunk := f.read(1 << 20):
            n += chunk.count(b"\n")
            last = chunk
    return n + (1 if last and not last.endswith(b"\n") else 0)


def _atomic_write(p: Path, data: bytes) -> None:
    tmp = p.with_name(f".{p.name}.{uuid.uuid4().hex[:8]}.tmp")
    try:
        with open(tmp, "wb") as f:
            f.write(data)
        if p.exists():
            try:
                shutil.copymode(p, tmp)
            except OSError:
                pass
        os.replace(tmp, p)
    finally:
        if tmp.exists():
            tmp.unlink(missing_ok=True)


def _rmtree(p: Path) -> None:
    def onexc(func, path, _exc):  # read-only files (e.g. .git objects on Windows)
        os.chmod(path, stat.S_IWRITE)
        func(path)

    if sys.version_info >= (3, 12):
        shutil.rmtree(p, onexc=onexc)
    else:  # pragma: no cover
        shutil.rmtree(p, onerror=lambda f, path, exc: onexc(f, path, exc))


def _tree_size(root: Path, max_files: int = MAX_ARCHIVE_FILES) -> tuple[int, int]:
    files = size = 0
    for d, _, names in _walk(root, skip=set()):
        for n in names:
            fp = d / n
            if fp.is_symlink():
                continue
            files += 1
            size += fp.stat().st_size
            if files > max_files:
                return files, size
    return files, size


async def _maybe_artifact(ctx: ToolContext, vpath: str, p: Path) -> str:
    if vpath.startswith("output/") and p.is_file():
        await ctx.add_artifact(vpath)
        return " Attached to the conversation as an artifact."
    return ""


# --------------------------------------------------------------------------- text extraction
def _cell(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(round(v, 10))
    if isinstance(v, datetime):
        return v.isoformat(sep=" ")
    return str(v).replace("\t", " ").replace("\n", " ")


def _pdf_text(p: Path, max_chars: int) -> tuple[str, str]:
    parts: list[str] = []
    total = 0
    try:
        import pdfplumber

        with pdfplumber.open(str(p)) as pdf:
            n = len(pdf.pages)
            for i, page in enumerate(pdf.pages, 1):
                t = (page.extract_text() or "").strip()
                parts.append(f"--- page {i} ---\n{t}")
                total += len(t)
                if total >= max_chars:
                    break
        return "\n\n".join(parts), f"PDF, {n} pages"
    except ImportError:
        pass
    except Exception:  # noqa: BLE001  — fall back to pypdf
        parts, total = [], 0
    from pypdf import PdfReader

    reader = PdfReader(str(p))
    n = len(reader.pages)
    for i, page in enumerate(reader.pages, 1):
        t = (page.extract_text() or "").strip()
        parts.append(f"--- page {i} ---\n{t}")
        total += len(t)
        if total >= max_chars:
            break
    return "\n\n".join(parts), f"PDF, {n} pages"


def _docx_text(p: Path, max_chars: int) -> tuple[str, str]:
    import docx

    d = docx.Document(str(p))
    out: list[str] = []
    total = 0
    for block in d.iter_inner_content():  # paragraphs and tables in document order
        if hasattr(block, "rows"):
            lines = []
            for row in block.rows:
                cells: list[str] = []
                for c in row.cells:
                    t = c.text.strip().replace("\n", " ")
                    if not cells or cells[-1] != t:  # merged cells repeat their text
                        cells.append(t)
                lines.append(" | ".join(cells))
            line = "\n".join(lines)
        else:
            text = block.text
            style = (block.style.name if block.style is not None else "") or ""
            m = re.match(r"heading\s*(\d)", style, re.IGNORECASE)
            if m and text.strip():
                line = f"{'#' * int(m.group(1))} {text}"
            elif style.lower().startswith("title") and text.strip():
                line = f"# {text}"
            elif style.lower().startswith("list") and text.strip():
                line = f"- {text}"
            else:
                line = text
        out.append(line)
        total += len(line) + 1
        if total >= max_chars:
            break
    return "\n".join(out), "Word document"


def _xlsx_text(p: Path, max_chars: int) -> tuple[str, str]:
    import openpyxl

    wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    try:
        names = list(wb.sheetnames)
        out: list[str] = []
        total = 0
        for ws in wb.worksheets:
            out.append(f"## Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                vals = list(row or ())
                while vals and vals[-1] in (None, ""):
                    vals.pop()
                if not vals:
                    continue
                line = "\t".join(_cell(v) for v in vals)
                out.append(line)
                total += len(line) + 1
                if total >= max_chars:
                    break
            if total >= max_chars:
                break
        return "\n".join(out), f"Excel workbook, {len(names)} sheet(s): {', '.join(names[:12])}"
    finally:
        wb.close()


def _pptx_shape_text(shape) -> list[str]:  # noqa: ANN001
    out: list[str] = []
    if hasattr(shape, "shapes"):  # group shape
        for s in shape.shapes:
            out.extend(_pptx_shape_text(s))
    if getattr(shape, "has_text_frame", False):
        t = shape.text_frame.text.strip()
        if t:
            out.append(t)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            out.append(" | ".join(c.text.strip().replace("\n", " ") for c in row.cells))
    return out


def _pptx_text(p: Path, max_chars: int) -> tuple[str, str]:
    from pptx import Presentation

    prs = Presentation(str(p))
    out: list[str] = []
    total = 0
    n = len(prs.slides)
    for i, slide in enumerate(prs.slides, 1):
        lines = [f"--- slide {i} ---"]
        for shape in slide.shapes:
            lines.extend(_pptx_shape_text(shape))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"[speaker notes] {notes}")
        block = "\n".join(lines)
        out.append(block)
        total += len(block)
        if total >= max_chars:
            break
    return "\n\n".join(out), f"PowerPoint, {n} slides"


def _html_text(p: Path, max_chars: int) -> tuple[str, str]:
    from bs4 import BeautifulSoup

    with open(p, "rb") as f:
        raw = f.read(20 * 1024 * 1024)
    try:
        soup = BeautifulSoup(raw, "lxml")
    except Exception:  # noqa: BLE001
        soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style", "noscript", "template", "svg"]):
        tag.decompose()
    title = soup.title.get_text(strip=True) if soup.title else ""
    lines = (re.sub(r"[ \t ]+", " ", ln).strip() for ln in soup.get_text("\n").splitlines())
    text = "\n".join(ln for ln in lines if ln)
    return text[: max_chars + 1], f"HTML page{': ' + title if title else ''}"


_EXTRACT_CACHE: collections.OrderedDict[tuple, tuple[str, str]] = collections.OrderedDict()
_EXTRACT_CACHE_LOCK = threading.Lock()


def _extract_text(p: Path, max_chars: int) -> tuple[str, str]:
    """(text, description). Documents are converted (cached by path/mtime/size, so paging is cheap); plain text is
    read; binaries raise ToolError."""
    ext = p.suffix.lower()
    handlers = {".pdf": _pdf_text, ".docx": _docx_text, ".xlsx": _xlsx_text, ".xlsm": _xlsx_text,
                ".pptx": _pptx_text, ".html": _html_text, ".htm": _html_text, ".xhtml": _html_text}
    if ext in handlers:
        st = p.stat()
        key = (str(p), st.st_mtime_ns, st.st_size, max_chars)
        with _EXTRACT_CACHE_LOCK:
            if key in _EXTRACT_CACHE:
                _EXTRACT_CACHE.move_to_end(key)
                return _EXTRACT_CACHE[key]
        try:
            result = handlers[ext](p, max_chars)
        except ToolError:
            raise
        except ImportError as e:
            raise ToolError(f"Extracting text from {ext} files needs the '{e.name}' package on the server.") from e
        except Exception as e:  # noqa: BLE001
            raise ToolError(f"Could not extract text from {p.name} ({type(e).__name__}: {e}). "
                            "The file may be corrupt, encrypted or not really a " + ext + " file.") from e
        with _EXTRACT_CACHE_LOCK:
            _EXTRACT_CACHE[key] = result
            while len(_EXTRACT_CACHE) > 16:
                _EXTRACT_CACHE.popitem(last=False)
        return result
    if ext in (".doc", ".xls", ".ppt"):
        raise ToolError(f"Legacy {ext} files are not supported; convert them to {ext}x first.")
    is_text, enc = _sniff(p)
    if not is_text:
        mime = mimetypes.guess_type(p.name)[0] or "unknown type"
        raise ToolError(f"{p.name} is a binary file ({mime}). Text extraction supports PDF, DOCX, XLSX, PPTX, HTML "
                        "and plain-text files.")
    with open(p, encoding=enc, errors="replace") as f:
        return f.read(max_chars + 1), "text file"


# --------------------------------------------------------------------------- browse
@fs.tool(name="list", risk=Risk.LOW, scope="read")
async def list_(ctx: ToolContext, path: str = "", recursive: bool = False, limit: int = 200) -> str:
    """List files and folders in the workspace.

    With no path, shows the top-level folders: workspace/ (your scratch area), projects/, downloads/, output/
    (deliverables shown to the user), uploads/ (files the user uploaded), temp/ and shared/ (the organization's
    shared drive). Folders end with "/".

    Args:
        path: Folder to list, e.g. "projects/app" (empty = top level).
        recursive: Include nested files and folders (skips .git, node_modules and similar).
        limit: Maximum number of entries to return (1-1000).
    """
    limit = max(1, min(limit, 1000))
    vp = _norm(ctx, path)
    if vp == "":
        lines = [f"{name}/  — {desc}" for name, desc in MOUNTS.items()]
        return "Top-level folders:\n" + "\n".join(lines)
    root = ws_path(ctx, vp, must_exist=True)

    def run() -> str:
        if root.is_file():
            st = root.stat()
            return f"{vp}  ({_human(st.st_size)}, modified {_stamp(st.st_mtime)})"
        entries: list[str] = []
        more = False

        def add(p: Path) -> bool:
            nonlocal more
            if len(entries) >= limit:
                more = True
                return False
            rel = ctx.workspace.virtual(p) if not p.is_symlink() else f"{vp}/{p.relative_to(root).as_posix()}"
            try:
                if p.is_symlink():
                    entries.append(f"{rel}  (symlink)")
                elif p.is_dir():
                    entries.append(f"{rel}/")
                else:
                    st = p.stat()
                    entries.append(f"{rel}  ({_human(st.st_size)}, {_stamp(st.st_mtime)})")
            except OSError:
                entries.append(f"{rel}  (unreadable)")
            return True

        if recursive:
            for d, dirnames, filenames in _walk(root):
                for name in dirnames + filenames:
                    if not add(d / name):
                        break
                if more:
                    break
        else:
            for p in sorted(root.iterdir(), key=lambda x: (not x.is_dir(), x.name.lower())):
                if not add(p):
                    break
        if not entries:
            return f"{vp}/ is empty."
        header = f"{vp}/ ({len(entries)}{'+' if more else ''} entries)"
        footer = f"\n[limit of {limit} reached; narrow the path or raise limit]" if more else ""
        return header + "\n" + "\n".join(entries) + footer

    return await asyncio.to_thread(run)


@fs.tool(risk=Risk.LOW, scope="read")
async def find(ctx: ToolContext, pattern: str, path: str = "", limit: int = 200) -> str:
    """Find files and folders by name or glob pattern (recursive).

    A pattern without "/" matches names anywhere below `path` (e.g. "*.py", "report*.docx"); a pattern with "/"
    matches paths relative to `path` (e.g. "src/**/*.ts"). Matching is case-insensitive; .git, node_modules and
    similar folders are skipped.

    Args:
        pattern: Glob pattern; supports *, ?, **, [abc] and {a,b}.
        path: Folder to search in (empty = all workspace folders, including shared/).
        limit: Maximum number of results (1-1000).
    """
    if not pattern.strip():
        raise ToolError("pattern is required, e.g. '*.py'.")
    limit = max(1, min(limit, 1000))
    rx = _glob_regex(pattern)
    by_name = "/" not in pattern.replace("\\", "/").strip("/")
    roots = _roots(ctx, path)
    root_v = {root: ctx.workspace.virtual(root) for root in roots}

    def run() -> str:
        hits: list[str] = []
        scanned = 0
        deadline = time.monotonic() + 20
        for root in roots:
            for d, dirnames, filenames in _walk(root):
                for name in dirnames + filenames:
                    scanned += 1
                    p = d / name
                    rel = p.relative_to(root).as_posix()
                    if rx.match(name if by_name else rel):
                        rv = root_v[root]
                        v = f"{rv}/{rel}" if rv else rel
                        hits.append(v + ("/" if p.is_dir() and not p.is_symlink() else ""))
                        if len(hits) >= limit:
                            return "\n".join(hits) + f"\n[limit of {limit} results reached]"
                if time.monotonic() > deadline:
                    return "\n".join(hits) + f"\n[search stopped after 20s / {scanned} entries; narrow the path]"
        if not hits:
            return f"No files match {pattern!r}."
        return "\n".join(hits)

    return await asyncio.to_thread(run)


@fs.tool(risk=Risk.LOW, scope="read")
async def search(ctx: ToolContext, query: str, path: str = "", regex: bool = False, case_sensitive: bool = False,
                 glob: str | None = None, max_results: int = 100) -> str:
    """Search inside text files for a word, phrase or regular expression (like grep -n).

    Returns matching lines as "path:line: text". Binary files, files over 10 MB and folders like .git or node_modules
    are skipped. To search PDFs or Office documents, extract their text with fs_extract_text first.

    Args:
        query: Text to look for (a regular expression when regex=true).
        path: File or folder to search (empty = all workspace folders, including shared/).
        regex: Treat query as a regular expression.
        case_sensitive: Match upper/lower case exactly.
        glob: Only search files whose name matches this glob, e.g. "*.py" or "*.{md,txt}".
        max_results: Maximum number of matching lines to return (1-500).
    """
    if not query:
        raise ToolError("query is required.")
    max_results = max(1, min(max_results, 500))
    flags = 0 if case_sensitive else re.IGNORECASE
    engine = _regex if (regex and _regex is not None) else re
    try:
        rx = engine.compile(query if regex else re.escape(query), flags)
    except Exception as e:  # noqa: BLE001  (re.error / regex.error)
        raise ToolError(f"Invalid regular expression: {e}") from e
    name_rx = _glob_regex(glob) if glob else None
    vp = _norm(ctx, path)
    targets = _roots(ctx, path)

    def matches(line: str) -> bool:
        if engine is _regex:
            try:
                return rx.search(line, timeout=0.5) is not None
            except TimeoutError:
                return False
        return rx.search(line) is not None

    def files():
        for root in targets:
            if root.is_file():
                yield root
                continue
            for d, _, names in _walk(root):
                for n in names:
                    if name_rx is None or name_rx.match(n):
                        yield d / n

    def run() -> str:
        hits: list[str] = []
        files_hit = scanned = 0
        deadline = time.monotonic() + 25
        for fp in files():
            if fp.is_symlink():
                continue
            try:
                if fp.stat().st_size > MAX_SEARCH_FILE_BYTES:
                    continue
                is_text, enc = _sniff(fp)
                if not is_text:
                    continue
                scanned += 1
                found = False
                with open(fp, encoding=enc, errors="replace") as f:
                    for lineno, line in enumerate(f, 1):
                        if matches(line):
                            if not found:
                                files_hit += 1
                                found = True
                            text = line.strip()
                            if len(text) > 300:
                                text = text[:300] + "…"
                            hits.append(f"{ctx.workspace.virtual(fp)}:{lineno}: {text}")
                            if len(hits) >= max_results:
                                return _render(hits, files_hit, True)
            except (OSError, UnicodeError):
                continue
            if time.monotonic() > deadline:
                return _render(hits, files_hit, False) + f"\n[search stopped after 25s / {scanned} files]"
        if not hits:
            return f"No matches for {query!r} in {vp or 'the workspace'} ({scanned} text files searched)."
        return _render(hits, files_hit, False)

    def _render(hits: list[str], nfiles: int, capped: bool) -> str:
        head = f"{len(hits)} matching line(s) in {nfiles} file(s)"
        if capped:
            head += f" (stopped at max_results={max_results})"
        return head + "\n" + "\n".join(hits)

    return await asyncio.to_thread(run)


# --------------------------------------------------------------------------- read
@fs.tool(risk=Risk.LOW, scope="read")
async def read(ctx: ToolContext, path: str, offset: int = 1, limit: int = 400, line_numbers: bool = False) -> str:
    """Read a file from the workspace: text files, or the extracted text of PDF/DOCX/XLSX/PPTX/HTML documents.

    Returns up to `limit` lines starting at line `offset` and says how to continue when the file is longer.
    Binary files (images, archives…) are not returned as text; use fs_metadata for their details.

    Args:
        path: Workspace path, e.g. "workspace/notes.md", "projects/app/src/main.py" or "shared/policies/hr.pdf".
        offset: 1-based line number to start from.
        limit: Maximum number of lines to return (1-2000).
        line_numbers: Prefix each line with its line number (handy before fs_edit on code).
    """
    vp = _norm(ctx, path)
    p = ws_path(ctx, vp, must_exist=True)
    if p.is_dir():
        raise ToolError(f"{vp} is a folder; use fs_list to see its contents.")
    offset, limit = max(1, offset), max(1, min(limit, 2000))

    def render(lines: list[str], start: int, total: int | None, note: str) -> str:
        out: list[str] = []
        used = 0
        last = start - 1
        for i, line in enumerate(lines, start):
            if len(line) > MAX_LINE_CHARS:
                line = line[:MAX_LINE_CHARS] + f"…[line truncated, {len(line)} chars]"
            entry = f"{i:>6}\t{line}" if line_numbers else line
            if out and used + len(entry) + 1 > READ_BUDGET:
                break
            out.append(entry)
            used += len(entry) + 1
            last = i
        text = "\n".join(out)
        if total is not None:
            more = last < total
        else:  # very large file: line count unknown
            more = len(out) < len(lines) or len(lines) >= limit
        if more or note:
            tail = f"lines {start}-{last} of {total if total is not None else '?'}"
            hint = f"; continue with offset={last + 1}" if more else ""
            text += f"\n[{note}{'; ' if note else ''}{tail}{hint}]"
        return text

    def run() -> str:
        if p.suffix.lower() in DOC_EXTS:
            text, kind = _extract_text(p, 400_000)
            lines = text.splitlines()
            if offset > max(1, len(lines)):
                return f"[{vp} has only {len(lines)} lines of extracted text]"
            return render(lines[offset - 1: offset - 1 + limit], offset, len(lines), f"{kind}, text extracted")
        is_text, enc = _sniff(p)
        if not is_text:
            mime = mimetypes.guess_type(p.name)[0] or "unknown type"
            return (f"{vp} is a binary file ({mime}, {_human(p.stat().st_size)}) and can't be shown as text. "
                    "Use fs_metadata for details.")
        lines: list[str] = []
        with open(p, encoding=enc, errors="replace") as f:
            for i, line in enumerate(f, 1):
                if i < offset:
                    continue
                if len(lines) >= limit:
                    break
                lines.append(line.rstrip("\r\n"))
        total = _count_lines(p)
        if not lines:
            if offset == 1:
                return f"[{vp} is empty]"
            return f"[{vp} has only {total} lines]"
        return render(lines, offset, total, "")

    return await asyncio.to_thread(run)


@fs.tool(risk=Risk.LOW, scope="read")
async def extract_text(ctx: ToolContext, path: str, max_chars: int = 15000) -> str:
    """Extract the plain text of a document: PDF, Word (.docx), Excel (.xlsx), PowerPoint (.pptx), HTML or text.

    PDFs are split by page, slides by slide (with speaker notes) and spreadsheets by sheet (tab-separated rows).

    Args:
        path: Workspace path of the document, e.g. "uploads/contract.pdf".
        max_chars: Maximum number of characters to return (500-100000).
    """
    vp = _norm(ctx, path)
    p = ws_path(ctx, vp, must_exist=True)
    if p.is_dir():
        raise ToolError(f"{vp} is a folder.")
    max_chars = max(500, min(max_chars, 100_000))
    text, kind = await asyncio.to_thread(_extract_text, p, max_chars)
    text = text.strip()
    if not text:
        return f"[{vp}: {kind}; no extractable text (it may be scanned images)]"
    if len(text) > max_chars:
        return f"[{vp}: {kind}]\n{text[:max_chars]}\n[truncated at {max_chars} characters]"
    return f"[{vp}: {kind}]\n{text}"


@fs.tool(risk=Risk.LOW, scope="read")
async def metadata(ctx: ToolContext, path: str) -> dict:
    """Get details about a file or folder: size, dates, type, line count, SHA-256, page/sheet/slide counts.

    Args:
        path: Workspace path of the file or folder.
    """
    vp = _norm(ctx, path)
    lex = _lexical(ctx, vp) if vp else ws_path(ctx, vp)
    if not _exists(lex):
        raise ToolError(f"Not found: {vp}")

    def run() -> dict:
        if lex.is_symlink():
            return {"path": vp, "type": "symlink", "target_inside_workspace": False}
        st = lex.stat()
        info: dict = {"path": vp or "/", "name": lex.name, "type": "folder" if lex.is_dir() else "file",
                      "modified": _stamp(st.st_mtime), "created": _stamp(getattr(st, "st_birthtime", st.st_ctime))}
        if lex.is_dir():
            entries = sum(1 for _ in lex.iterdir())
            files, size = _tree_size(lex, max_files=20_000)
            info.update(entries=entries, total_files=files if files <= 20_000 else "20000+",
                        total_size=_human(size))
            return info
        info.update(size=_human(st.st_size), size_bytes=st.st_size,
                    mime=mimetypes.guess_type(lex.name)[0] or "application/octet-stream")
        is_text, enc = _sniff(lex)
        info["is_text"] = is_text
        if is_text:
            info["encoding"] = enc
            lines = _count_lines(lex)
            if lines is not None:
                info["lines"] = lines
        if st.st_size <= 50 * 1024 * 1024:
            h = hashlib.sha256()
            with open(lex, "rb") as f:
                while chunk := f.read(1 << 20):
                    h.update(chunk)
            info["sha256"] = h.hexdigest()
        ext = lex.suffix.lower()
        try:
            if ext == ".pdf":
                from pypdf import PdfReader

                info["pages"] = len(PdfReader(str(lex)).pages)
            elif ext in (".xlsx", ".xlsm"):
                import openpyxl

                wb = openpyxl.load_workbook(str(lex), read_only=True)
                info["sheets"] = list(wb.sheetnames)
                wb.close()
            elif ext == ".pptx":
                from pptx import Presentation

                info["slides"] = len(Presentation(str(lex)).slides)
        except Exception:  # noqa: BLE001
            info["note"] = "document structure could not be read"
        return info

    return await asyncio.to_thread(run)


@fs.tool(risk=Risk.LOW, scope="read")
async def diff(ctx: ToolContext, path_a: str, path_b: str, context: int = 3) -> str:
    """Show a unified diff between two text files.

    Args:
        path_a: Original file, e.g. "workspace/draft_v1.md".
        path_b: Changed file, e.g. "workspace/draft_v2.md".
        context: Number of unchanged context lines around each change (0-20).
    """
    va, vb = _norm(ctx, path_a), _norm(ctx, path_b)
    pa, pb = ws_path(ctx, va, must_exist=True), ws_path(ctx, vb, must_exist=True)
    context = max(0, min(context, 20))

    def load(p: Path, v: str) -> list[str]:
        if p.is_dir():
            raise ToolError(f"{v} is a folder; fs_diff compares files.")
        if p.stat().st_size > MAX_DIFF_FILE_BYTES:
            raise ToolError(f"{v} is too large to diff (max {_human(MAX_DIFF_FILE_BYTES)}).")
        is_text, enc = _sniff(p)
        if not is_text:
            raise ToolError(f"{v} is a binary file.")
        with open(p, encoding=enc, errors="replace") as f:
            return f.read().splitlines()

    def run() -> str:
        a, b = load(pa, va), load(pb, vb)
        lines = list(difflib.unified_diff(a, b, fromfile=va, tofile=vb, n=context, lineterm=""))
        if not lines:
            return "Files are identical."
        added = sum(1 for ln in lines if ln.startswith("+") and not ln.startswith("+++"))
        removed = sum(1 for ln in lines if ln.startswith("-") and not ln.startswith("---"))
        text = "\n".join(lines)
        if len(text) > READ_BUDGET:
            text = text[:READ_BUDGET] + "\n…[diff truncated]"
        return f"{added} line(s) added, {removed} removed\n{text}"

    return await asyncio.to_thread(run)


# --------------------------------------------------------------------------- write
@fs.tool(risk=Risk.MEDIUM, scope="write")
async def write(ctx: ToolContext, path: str, content: str, overwrite: bool = True) -> str:
    """Create or overwrite a text file in the workspace (missing parent folders are created).

    Files written under "output/" are deliverables and are attached to the conversation automatically.

    Args:
        path: Workspace path of the file, e.g. "workspace/notes.md" or "output/summary.md".
        content: The complete file content (UTF-8 text).
        overwrite: Replace the file if it already exists; false fails instead.
    """
    vp = _norm(ctx, path)
    _guard_write(ctx, vp)
    if _is_mount_root(vp):
        raise ToolError("Give a file path inside a folder, e.g. 'workspace/notes.md'.")
    data = content.encode("utf-8", errors="surrogateescape")
    if len(data) > MAX_WRITE_BYTES:
        raise ToolError(f"Content is too large ({_human(len(data))}; max {_human(MAX_WRITE_BYTES)}).")
    p = ws_path(ctx, vp, for_write=True)
    if p.is_dir():
        raise ToolError(f"{vp} is a folder.")
    existed = p.exists()
    if existed and not overwrite:
        raise ToolError(f"{vp} already exists; pass overwrite=true to replace it.")
    await asyncio.to_thread(_atomic_write, p, data)
    msg = f"{'Overwrote' if existed else 'Created'} {vp} ({_human(len(data))})."
    return msg + await _maybe_artifact(ctx, vp, p)


@fs.tool(risk=Risk.MEDIUM, scope="write")
async def append(ctx: ToolContext, path: str, content: str) -> str:
    """Append text to the end of a file (the file is created if it doesn't exist).

    Args:
        path: Workspace path of the file, e.g. "workspace/log.md".
        content: Text to append; include a leading newline if the file doesn't end with one.
    """
    vp = _norm(ctx, path)
    _guard_write(ctx, vp)
    if _is_mount_root(vp):
        raise ToolError("Give a file path inside a folder, e.g. 'workspace/log.md'.")
    p = ws_path(ctx, vp, for_write=True)
    if p.is_dir():
        raise ToolError(f"{vp} is a folder.")
    data = content.encode("utf-8", errors="surrogateescape")
    if (p.stat().st_size if p.exists() else 0) + len(data) > MAX_WRITE_BYTES * 5:
        raise ToolError("The file would become too large.")

    def run() -> int:
        if p.exists() and not _sniff(p)[0]:
            raise ToolError(f"{vp} is a binary file.")
        with open(p, "ab") as f:
            f.write(data)
        return p.stat().st_size

    size = await asyncio.to_thread(run)
    return f"Appended {_human(len(data))} to {vp} (now {_human(size)})."


@fs.tool(risk=Risk.MEDIUM, scope="write")
async def edit(ctx: ToolContext, path: str, old_text: str, new_text: str, count: int = 1) -> str:
    """Replace exact text in a file (safe find-and-replace).

    `old_text` must match the file exactly, including whitespace and indentation. The edit is refused, and the file
    left untouched, unless the number of matches equals `count`, so include enough surrounding lines to make the
    match unique.

    Args:
        path: Workspace path of a text file.
        old_text: Exact text to find (copy it from fs_read output, without line-number prefixes).
        new_text: Replacement text (empty string deletes the match).
        count: Number of occurrences you expect to replace; -1 replaces every occurrence (at least one must exist).
    """
    vp = _norm(ctx, path)
    _guard_write(ctx, vp)
    if not old_text:
        raise ToolError("old_text must not be empty (use fs_write to create or replace a whole file).")
    if count == 0 or count < -1:
        raise ToolError("count must be 1 or more, or -1 to replace all occurrences.")
    p = ws_path(ctx, vp, must_exist=True)
    if p.is_dir():
        raise ToolError(f"{vp} is a folder.")

    def run() -> str:
        is_text, enc = _sniff(p)
        if not is_text:
            raise ToolError(f"{vp} is a binary file and can't be edited as text.")
        errors = _errors_for(enc)
        try:
            with open(p, encoding=enc, errors=errors, newline="") as f:
                data = f.read()
        except UnicodeError as e:
            raise ToolError(f"{vp} could not be decoded as {enc}.") from e
        old, new = old_text, new_text
        n = data.count(old)
        if n == 0 and "\r\n" in data and "\n" in old and "\r\n" not in old:  # CRLF file, LF arguments
            crlf_old = old.replace("\n", "\r\n")
            if data.count(crlf_old):
                old, new = crlf_old, new.replace("\r\n", "\n").replace("\n", "\r\n")
                n = data.count(old)
        if n == 0:
            raise ToolError(f"old_text was not found in {vp}. Re-read the file with fs_read and copy the exact text "
                            "(whitespace and indentation matter).")
        if count != -1 and n != count:
            raise ToolError(f"Found {n} occurrences of old_text in {vp}, expected {count}. Include more surrounding "
                            f"context to make it unique, or pass count={n} (or -1) to replace them all.")
        line = data.count("\n", 0, data.find(old)) + 1
        _atomic_write(p, data.replace(old, new).encode(enc, errors=errors))
        return f"Edited {vp}: replaced {n} occurrence(s) (first at line {line})."

    return await asyncio.to_thread(run)


@fs.tool(risk=Risk.MEDIUM, scope="write")
async def mkdir(ctx: ToolContext, path: str) -> str:
    """Create a folder (and any missing parent folders).

    Args:
        path: Workspace path of the folder, e.g. "projects/new-app/src".
    """
    vp = _norm(ctx, path)
    _guard_write(ctx, vp)
    p = ws_path(ctx, vp)
    if p.exists() and not p.is_dir():
        raise ToolError(f"A file already exists at {vp}.")
    if p.is_dir():
        return f"Folder {vp}/ already exists."
    p.mkdir(parents=True, exist_ok=True)
    return f"Created folder {vp}/."


@fs.tool(risk=Risk.MEDIUM, scope="write")
async def move(ctx: ToolContext, source: str, destination: str, overwrite: bool = False) -> str:
    """Move or rename a file or folder within the workspace.

    If `destination` is an existing folder, the source is moved into it. Moving or renaming items that are in
    shared/ requires an admin.

    Args:
        source: Current workspace path, e.g. "workspace/draft.md".
        destination: New path or target folder, e.g. "output/final.md" or "projects/archive".
        overwrite: Replace an existing destination file.
    """
    sv, dv = _norm(ctx, source), _norm(ctx, destination)
    _guard_remove(ctx, sv, "move")
    _guard_write(ctx, dv)
    if not dv:
        raise ToolError("Choose a destination inside a folder, e.g. 'workspace/notes.md' or 'projects'.")
    src = _lexical(ctx, sv)
    if not _exists(src):
        raise ToolError(f"Not found: {sv}")
    dst = ws_path(ctx, dv)
    if dst.is_dir():
        dst, dv = dst / src.name, f"{dv}/{src.name}" if dv else src.name
        _guard_write(ctx, dv)
    if src.is_dir() and not src.is_symlink() and (dst == src.resolve() or src.resolve() in dst.parents):
        raise ToolError("A folder can't be moved into itself.")
    if _exists(dst):
        if not overwrite:
            raise ToolError(f"{dv} already exists; pass overwrite=true to replace it.")
        if dst.is_dir():
            raise ToolError(f"{dv} is an existing folder; it won't be overwritten.")

    def run() -> None:
        dst.parent.mkdir(parents=True, exist_ok=True)
        if _exists(dst):
            dst.unlink()
        shutil.move(str(src), str(dst))

    await asyncio.to_thread(run)
    return f"Moved {sv} → {dv}." + await _maybe_artifact(ctx, dv, dst)


@fs.tool(risk=Risk.MEDIUM, scope="write")
async def copy(ctx: ToolContext, source: str, destination: str, overwrite: bool = False) -> str:
    """Copy a file or folder within the workspace (e.g. from shared/ into your workspace).

    If `destination` is an existing folder, the source is copied into it. Symbolic links are copied as links.

    Args:
        source: Workspace path to copy, e.g. "shared/templates/report.docx".
        destination: New path or target folder, e.g. "workspace/report.docx".
        overwrite: Replace existing files at the destination.
    """
    sv, dv = _norm(ctx, source), _norm(ctx, destination)
    _guard_write(ctx, dv)
    if not dv:
        raise ToolError("Choose a destination inside a folder, e.g. 'workspace/report.docx' or 'workspace'.")
    src = ws_path(ctx, sv, must_exist=True)
    dst = ws_path(ctx, dv)
    if dst.is_dir():
        dst, dv = dst / src.name, f"{dv}/{src.name}" if dv else src.name
        _guard_write(ctx, dv)
    if src.is_dir() and (dst == src or src in dst.parents):
        raise ToolError("A folder can't be copied into itself.")
    if _exists(dst) and not overwrite:
        raise ToolError(f"{dv} already exists; pass overwrite=true to replace it.")

    def run() -> str:
        dst.parent.mkdir(parents=True, exist_ok=True)
        if src.is_dir():
            files, size = _tree_size(src)
            if files > MAX_ARCHIVE_FILES or size > MAX_ARCHIVE_BYTES:
                raise ToolError("The folder is too large to copy with fs_copy.")
            if dst.exists() and not dst.is_dir():
                raise ToolError(f"{dv} is a file; can't copy a folder over it.")
            shutil.copytree(src, dst, symlinks=True, dirs_exist_ok=overwrite)
            return f"Copied folder {sv}/ → {dv}/ ({files} files, {_human(size)})."
        if dst.is_dir():
            raise ToolError(f"{dv} is a folder.")
        shutil.copy2(src, dst)
        return f"Copied {sv} → {dv} ({_human(dst.stat().st_size)})."

    msg = await asyncio.to_thread(run)
    return msg + await _maybe_artifact(ctx, dv, dst)


@fs.tool(risk=Risk.MEDIUM, scope="write")
async def delete(ctx: ToolContext, path: str, recursive: bool = False) -> str:
    """Delete a file or folder from the workspace. Deleting in shared/ requires an admin.

    Args:
        path: Workspace path to delete, e.g. "temp/scratch.txt".
        recursive: Required to delete a folder that isn't empty (deletes everything inside it).
    """
    vp = _norm(ctx, path)
    _guard_remove(ctx, vp, "delete")
    p = _lexical(ctx, vp)
    if not _exists(p):
        raise ToolError(f"Not found: {vp}")

    def run() -> str:
        if p.is_symlink() or p.is_file():
            p.unlink()
            return f"Deleted {vp}."
        if any(p.iterdir()):
            if not recursive:
                raise ToolError(f"{vp}/ is not empty; pass recursive=true to delete it and everything inside.")
            files, _ = _tree_size(p)
            _rmtree(p)
            return f"Deleted folder {vp}/ ({files} files)."
        p.rmdir()
        return f"Deleted empty folder {vp}/."

    return await asyncio.to_thread(run)


# --------------------------------------------------------------------------- archives
def _archive_kind(name: str) -> str | None:
    low = name.lower()
    for suffix in _ARCHIVE_SUFFIXES:
        if low.endswith(suffix):
            return "zip" if suffix == ".zip" else "tar"
    return None


def _strip_archive_suffix(name: str) -> str:
    low = name.lower()
    for suffix in _ARCHIVE_SUFFIXES:
        if low.endswith(suffix):
            return name[: -len(suffix)] or "archive"
    return Path(name).stem or "archive"


def _safe_member(name: str) -> PurePosixPath | None:
    """Validate an archive member name (zip-slip protection); None for entries that are just '.' or '/'."""
    norm = name.replace("\\", "/")
    if norm.startswith("/") or re.match(r"^[A-Za-z]:", norm):
        raise ToolError(f"Refusing to extract: the archive contains an absolute path ({name!r}).")
    parts = [part for part in norm.split("/") if part not in ("", ".")]
    if not parts:
        return None
    if any(part == ".." for part in parts):
        raise ToolError(f"Refusing to extract: the archive contains a path that escapes the target ({name!r}).")
    if any("\x00" in part or (os.name == "nt" and ":" in part) for part in parts):  # ':' = NTFS alternate stream
        raise ToolError(f"Refusing to extract: the archive contains an invalid file name ({name!r}).")
    return PurePosixPath(*parts)


@fs.tool(risk=Risk.MEDIUM, scope="write", timeout=600)
async def archive(ctx: ToolContext, path: str, destination: str | None = None, overwrite: bool = False) -> str:
    """Compress a folder (or a single file) into a .zip archive.

    Archives written under "output/" are attached to the conversation automatically.

    Args:
        path: Folder or file to compress, e.g. "projects/site".
        destination: Path of the .zip to create (default: "output/<name>.zip").
        overwrite: Replace the destination if it already exists.
    """
    sv = _norm(ctx, path)
    src = ws_path(ctx, sv, must_exist=True)
    if not sv:
        raise ToolError("Choose a folder to archive, e.g. 'projects/site'.")
    dv = _norm(ctx, destination) if destination else f"output/{src.name or 'archive'}.zip"
    if not dv.lower().endswith(".zip"):
        dv += ".zip"
    _guard_write(ctx, dv)
    dst = ws_path(ctx, dv, for_write=True)
    if dst.is_dir():
        raise ToolError(f"{dv} is a folder.")
    if dst.exists() and not overwrite:
        raise ToolError(f"{dv} already exists; pass overwrite=true to replace it.")

    def run() -> str:
        items: list[tuple[Path, str]] = []
        total = 0
        if src.is_file():
            items.append((src, src.name))
            total = src.stat().st_size
        else:
            for d, _, names in _walk(src, skip=set()):
                for n in names:
                    fp = d / n
                    if fp.is_symlink() or fp.resolve() == dst.resolve():
                        continue
                    items.append((fp, f"{src.name}/{fp.relative_to(src).as_posix()}"))
                    total += fp.stat().st_size
                    if len(items) > MAX_ARCHIVE_FILES or total > MAX_ARCHIVE_BYTES:
                        raise ToolError(f"Too much data to archive (max {MAX_ARCHIVE_FILES} files / "
                                        f"{_human(MAX_ARCHIVE_BYTES)}).")
        tmp = dst.with_name(f".{dst.name}.{uuid.uuid4().hex[:8]}.tmp")
        try:
            with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as zf:
                if src.is_dir() and not items:
                    zf.writestr(f"{src.name}/", "")
                for fp, arc in items:
                    zf.write(fp, arc)
            os.replace(tmp, dst)
        finally:
            tmp.unlink(missing_ok=True)
        return f"Created {dv} ({len(items)} files, {_human(total)} → {_human(dst.stat().st_size)})."

    msg = await asyncio.to_thread(run)
    return msg + await _maybe_artifact(ctx, dv, dst)


@fs.tool(risk=Risk.MEDIUM, scope="write", timeout=600)
async def extract(ctx: ToolContext, path: str, destination: str | None = None, overwrite: bool = False) -> str:
    """Extract a .zip or .tar(.gz/.bz2/.xz) archive into a folder.

    Entries with absolute paths or ".." are rejected (nothing is extracted), git metadata (.git/) is skipped, and
    archives over 2 GB uncompressed are refused.

    Args:
        path: Workspace path of the archive, e.g. "uploads/site.zip".
        destination: Folder to extract into (default: a folder named after the archive, next to it).
        overwrite: Allow replacing files that already exist in the destination.
    """
    sv = _norm(ctx, path)
    src = ws_path(ctx, sv, must_exist=True)
    kind = _archive_kind(src.name)
    if src.is_dir() or kind is None:
        raise ToolError("fs_extract supports .zip, .tar, .tar.gz/.tgz, .tar.bz2 and .tar.xz files.")
    parent_v = sv.rsplit("/", 1)[0] if "/" in sv else "workspace"
    dv = _norm(ctx, destination) if destination else f"{parent_v}/{_strip_archive_suffix(src.name)}"
    _guard_write(ctx, dv)
    if not dv:
        raise ToolError("Choose a destination folder, e.g. 'workspace/site'.")
    dest = ws_path(ctx, dv)
    if dest.exists() and not dest.is_dir():
        raise ToolError(f"{dv} is a file.")

    def plan(names: list[tuple[str, bool]]) -> tuple[list[tuple[int, Path]], int]:
        out: list[tuple[int, Path]] = []
        skipped_git = 0
        conflicts: list[str] = []
        base = dest.resolve()
        for idx, (name, is_dir) in enumerate(names):
            rel = _safe_member(name)
            if rel is None:
                continue
            if ".git" in rel.parts:
                skipped_git += 1
                continue
            target = (base / rel).resolve()
            if target != base and base not in target.parents:
                raise ToolError(f"Refusing to extract: entry {name!r} would land outside the destination.")
            if not is_dir and target.exists() and not overwrite:
                conflicts.append(rel.as_posix())
            out.append((idx, target))
        if conflicts:
            raise ToolError(f"{len(conflicts)} file(s) already exist in {dv} (e.g. {conflicts[0]}); pass "
                            "overwrite=true to replace them.")
        return out, skipped_git

    def run_zip() -> tuple[int, int, int]:
        with zipfile.ZipFile(src) as zf:
            infos = zf.infolist()
            if len(infos) > MAX_EXTRACT_MEMBERS:
                raise ToolError(f"The archive has too many entries ({len(infos)}; max {MAX_EXTRACT_MEMBERS}).")
            if sum(i.file_size for i in infos) > MAX_ARCHIVE_BYTES:
                raise ToolError(f"The archive is too large when extracted (max {_human(MAX_ARCHIVE_BYTES)}).")
            todo, skipped = plan([(i.filename, i.is_dir()) for i in infos])
            dest.mkdir(parents=True, exist_ok=True)
            written = files = 0
            for idx, target in todo:
                info = infos[idx]
                if info.is_dir():
                    target.mkdir(parents=True, exist_ok=True)
                    continue
                target.parent.mkdir(parents=True, exist_ok=True)
                with zf.open(info) as fin, open(target, "wb") as fout:
                    while chunk := fin.read(1 << 20):
                        written += len(chunk)
                        if written > MAX_ARCHIVE_BYTES:
                            raise ToolError("Extraction aborted: the archive expands beyond the size limit.")
                        fout.write(chunk)
                files += 1
            return files, written, skipped

    def run_tar() -> tuple[int, int, int]:
        if not hasattr(tarfile, "data_filter"):
            raise ToolError("Extracting tar archives safely needs Python 3.11.4 or newer.")
        with tarfile.open(src) as tf:
            members = tf.getmembers()
            if len(members) > MAX_EXTRACT_MEMBERS:
                raise ToolError(f"The archive has too many entries ({len(members)}; max {MAX_EXTRACT_MEMBERS}).")
            if sum(m.size for m in members if m.isfile()) > MAX_ARCHIVE_BYTES:
                raise ToolError(f"The archive is too large when extracted (max {_human(MAX_ARCHIVE_BYTES)}).")
            for m in members:
                if not (m.isfile() or m.isdir() or m.issym() or m.islnk()):
                    raise ToolError(f"Refusing to extract: special file {m.name!r} in archive.")
            todo, skipped = plan([(m.name, m.isdir()) for m in members])
            selected = [members[idx] for idx, _ in todo]
            dest.mkdir(parents=True, exist_ok=True)
            try:
                tf.extractall(dest, members=selected, filter="data")
            except tarfile.FilterError as e:
                raise ToolError(f"Refusing to extract: {e}") from e
            return (sum(1 for m in selected if m.isfile()), sum(m.size for m in selected if m.isfile()), skipped)

    try:
        files, size, skipped = await asyncio.to_thread(run_zip if kind == "zip" else run_tar)
    except (zipfile.BadZipFile, tarfile.TarError, RuntimeError, EOFError) as e:
        raise ToolError(f"Could not extract {sv}: {e} (the archive may be corrupt or encrypted).") from e
    msg = f"Extracted {files} files ({_human(size)}) from {sv} into {dv}/."
    if skipped:
        msg += f" Skipped {skipped} entries under .git/ (git metadata is never extracted; use git_clone for repos)."
    return msg
